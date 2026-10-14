#!/usr/bin/env python3
"""Shared helpers for R-index indicator review workflows."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

try:
    import pandas as pd
except ModuleNotFoundError:
    pd = None  # type: ignore[assignment]

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
except ModuleNotFoundError:
    load_workbook = None  # type: ignore[assignment]
    Font = None  # type: ignore[assignment]


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PDF_DIR = ROOT / "data/raw_pdfs"
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "outputs"
DEFAULT_TEMPLATE = ROOT / "assets/r3_results_template.html"
SUPPORTED_INDICATORS = {"R3", "R5", "R7"}
BASE_COLUMNS = ["REFNO", "RESULTNO", "TARGET"]
PAYLOAD_PLACEHOLDER = "__R3_RESULTS_PAYLOAD__"


def require_pandas() -> Any:
    if pd is None:
        raise SystemExit("This command requires pandas. Install dependencies with: pip install -r requirements.txt")
    return pd


def require_openpyxl() -> tuple[Any, Any]:
    if load_workbook is None or Font is None:
        raise SystemExit("This command requires openpyxl. Install dependencies with: pip install -r requirements.txt")
    return load_workbook, Font


def normalize_indicator(value: str) -> str:
    indicator = value.strip().upper()
    if indicator not in SUPPORTED_INDICATORS:
        raise argparse.ArgumentTypeError(
            f"indicator must be one of: {', '.join(sorted(SUPPORTED_INDICATORS))}"
        )
    return indicator


def indicator_from_script_name(script_name: str, default: str = "R3") -> str:
    match = re.search(r"_(r[357])_", script_name.casefold())
    if match:
        return match.group(1).upper()
    return default


def final_columns(indicator: str) -> list[str]:
    return [*BASE_COLUMNS, indicator, f"{indicator}_RESULT_JSON"]


def output_pattern(indicator: str) -> re.Pattern[str]:
    return re.compile(rf"^{indicator.lower()}_\d{{8}}_(?:\d{{4,5}}|unlabeled)\.xlsx$")


def default_predictions(indicator: str) -> Path:
    return DEFAULT_OUTPUT_DIR / f"{indicator.lower()}_working_predictions.xlsx"


def default_manifest(indicator: str) -> Path:
    return DEFAULT_OUTPUT_DIR / f"{indicator.lower()}_chat_batch_manifest.json"


def default_comparison(indicator: str) -> Path:
    return DEFAULT_OUTPUT_DIR / f"{indicator.lower()}_accuracy_vs_standard.xlsx"


def _xlsx_column(cell_reference: str) -> str:
    return "".join(character for character in cell_reference if character.isalpha())


def _xlsx_cell_value(cell: ElementTree.Element, shared_strings: list[str], namespace: dict[str, str]) -> str:
    value = cell.findtext("main:v", default="", namespaces=namespace)
    if cell.get("t") == "s" and value:
        return shared_strings[int(value)]
    if cell.get("t") == "inlineStr":
        return cell.findtext("main:is/main:t", default="", namespaces=namespace)
    return value


def load_targets_for_queue(target_xlsx: Path) -> dict[str, list[dict[str, str]]]:
    """Load REFNO -> target rows using only the Python standard library."""
    if not target_xlsx.exists():
        return {}

    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(target_xlsx) as archive:
        shared_strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            strings_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared_strings = [
                "".join(item.itertext()) for item in strings_root.findall("main:si", namespace)
            ]
        sheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))

    rows = sheet.findall("main:sheetData/main:row", namespace)
    if not rows:
        return {}
    header = {
        _xlsx_column(cell.get("r", "")): _xlsx_cell_value(cell, shared_strings, namespace).strip()
        for cell in rows[0].findall("main:c", namespace)
    }
    try:
        refno_column = next(column for column, name in header.items() if name == "REFNO")
        resultno_column = next(column for column, name in header.items() if name == "RESULTNO")
        rockname_column = next(column for column, name in header.items() if name == "ROCKNAME")
    except StopIteration as exc:
        raise SystemExit(f"{target_xlsx} must contain REFNO, RESULTNO, and ROCKNAME columns.") from exc

    targets: dict[str, list[dict[str, str]]] = {}
    for row in rows[1:]:
        values = {
            _xlsx_column(cell.get("r", "")): _xlsx_cell_value(cell, shared_strings, namespace).strip()
            for cell in row.findall("main:c", namespace)
        }
        if values.get(refno_column) and values.get(resultno_column) and values.get(rockname_column):
            targets.setdefault(values[refno_column], []).append(
                {
                    "resultno": values[resultno_column],
                    "target": values[rockname_column],
                }
            )
    return targets


def ordered_pdfs(pdf_dir: Path) -> list[Path]:
    if not pdf_dir.is_dir():
        raise SystemExit(f"PDF directory not found: {pdf_dir}")
    return sorted(
        (path for path in pdf_dir.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"),
        key=lambda path: path.name.casefold(),
    )


def build_manifest(
    pdf_paths: list[Path],
    targets: dict[str, list[dict[str, str]]],
    indicator: str,
) -> list[dict[str, Any]]:
    documents: list[dict[str, Any]] = []
    sequence = 1
    pdf_stems = {pdf_path.stem for pdf_path in pdf_paths}
    for pdf_path in pdf_paths:
        target_rows = targets.get(pdf_path.stem)
        if not target_rows:
            documents.append(
                {
                    "sequence": sequence,
                    "filename": pdf_path.name,
                    "refno": pdf_path.stem,
                    "pdf_path": str(pdf_path),
                    "pdf_found": True,
                    "resultno": None,
                    "target": None,
                    "target_found": False,
                    "indicator": indicator,
                    "instruction": "Open this PDF in the Codex conversation and record a target-mapping manual-review row before continuing.",
                }
            )
            sequence += 1
            continue

        for target_index, target_row in enumerate(target_rows, start=1):
            documents.append(
                {
                    "sequence": sequence,
                    "target_sequence_for_refno": target_index,
                    "filename": pdf_path.name,
                    "refno": pdf_path.stem,
                    "pdf_path": str(pdf_path),
                    "pdf_found": True,
                    "resultno": target_row["resultno"],
                    "target": target_row["target"],
                    "target_found": True,
                    "indicator": indicator,
                    "instruction": f"Open this PDF in the Codex conversation and evaluate {indicator} for this RESULTNO/target only before continuing.",
                }
            )
            sequence += 1
    for refno, target_rows in targets.items():
        if refno in pdf_stems:
            continue
        for target_index, target_row in enumerate(target_rows, start=1):
            documents.append(
                {
                    "sequence": sequence,
                    "target_sequence_for_refno": target_index,
                    "filename": None,
                    "refno": refno,
                    "pdf_path": None,
                    "pdf_found": False,
                    "resultno": target_row["resultno"],
                    "target": target_row["target"],
                    "target_found": True,
                    "indicator": indicator,
                    "instruction": f"No local PDF was found for this REFNO. Record a target-specific manual-review row for {indicator} before continuing.",
                }
            )
            sequence += 1
    return documents


def run_prepare_queue(indicator: str, pdf_dir: Path, target_xlsx: Path, output_json: Path) -> int:
    pdf_paths = ordered_pdfs(pdf_dir)
    targets = load_targets_for_queue(target_xlsx)
    documents = build_manifest(pdf_paths, targets, indicator)
    unmatched_pdfs = [item["filename"] for item in documents if item.get("pdf_found") and not item["target_found"]]
    missing_pdf_targets = [
        f"{item['refno']}/{item['resultno']}" for item in documents if item["target_found"] and not item.get("pdf_found", True)
    ]

    payload = {
        "workflow": f"Codex chat-model direct PDF review for {indicator}; no external API call",
        "indicator": indicator,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "order": "ascending filename (case-insensitive), then workbook order for multiple targets under the same REFNO; target rows without local PDFs are appended in workbook order",
        "selection": "all PDF files in --pdf-dir plus every target row in --target-xlsx; no limit, no resume-based skipping, and no collapsing of multiple targets per paper",
        "pdf_count": len(pdf_paths),
        "target_row_count": sum(len(rows) for rows in targets.values()),
        "unmatched_pdf_count": len(unmatched_pdfs),
        "missing_pdf_target_count": len(missing_pdf_targets),
        "review_count": len(documents),
        "documents": documents,
    }
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Prepared {len(documents)} {indicator} review item(s) from {len(pdf_paths)} PDF(s).")
    print(f"Manifest: {output_json}")
    if unmatched_pdfs:
        print(f"WARNING: {len(unmatched_pdfs)} PDF(s) have no target in {target_xlsx}: {', '.join(unmatched_pdfs)}")
    if missing_pdf_targets:
        print(
            f"WARNING: {len(missing_pdf_targets)} target row(s) have no local PDF in {pdf_dir}: "
            f"{', '.join(missing_pdf_targets)}"
        )
    return 0


def build_final_frame(predictions: Path, targets: Path, indicator: str) -> pd.DataFrame:
    pandas = require_pandas()
    columns = final_columns(indicator)
    score_column = indicator
    json_column = f"{indicator}_RESULT_JSON"
    pred = pandas.read_excel(predictions, dtype={"REFNO": str, "RESULTNO": str})
    missing = sorted(set(columns) - set(pred.columns))
    if missing:
        raise SystemExit(f"{predictions} is missing required column(s): {', '.join(missing)}")

    known_predictions = pred["RESULTNO"].dropna()
    known_predictions = known_predictions[known_predictions.astype(str).str.upper() != "UNKNOWN"]
    if known_predictions.duplicated().any():
        duplicated = ", ".join(known_predictions[known_predictions.duplicated()].astype(str).head(10))
        raise SystemExit(f"{predictions} contains duplicate RESULTNO values: {duplicated}")

    target_columns = set(pandas.read_excel(targets, nrows=0).columns)
    if {"RESULTNO", score_column}.issubset(target_columns):
        standard = pandas.read_excel(targets, usecols=["RESULTNO", score_column], dtype={"RESULTNO": str})
        if standard["RESULTNO"].duplicated().any():
            duplicated = ", ".join(standard.loc[standard["RESULTNO"].duplicated(), "RESULTNO"].astype(str).head(10))
            raise SystemExit(f"{targets} contains duplicate RESULTNO values: {duplicated}")

        standard = standard.rename(columns={score_column: f"STANDARD_{indicator}"})
        merged = pred.merge(standard, on="RESULTNO", how="left", validate="many_to_one")
    else:
        merged = pred.copy()
        merged[f"STANDARD_{indicator}"] = pandas.NA
    merged["RESULTNO"] = merged["RESULTNO"].where(merged["RESULTNO"].notna(), "UNKNOWN")
    compared = merged[f"STANDARD_{indicator}"].isin([0, 1])
    matches = pandas.Series(False, index=merged.index)
    matches.loc[compared] = merged.loc[compared, score_column].eq(merged.loc[compared, f"STANDARD_{indicator}"])
    merged["MISMATCH"] = compared & ~matches
    merged[json_column] = merged[json_column].fillna("")
    return merged


def write_final_workbook(frame: pd.DataFrame, output: Path, indicator: str) -> None:
    workbook_loader, font_class = require_openpyxl()
    output.parent.mkdir(parents=True, exist_ok=True)
    columns = final_columns(indicator)
    frame[columns].to_excel(output, index=False)

    workbook = workbook_loader(output)
    worksheet = workbook.active
    worksheet.title = f"{indicator} Results"

    for cell in worksheet[1]:
        cell.font = font_class(bold=True)

    red_font = font_class(color="FFFF0000")
    for excel_row, mismatch in enumerate(frame["MISMATCH"].tolist(), start=2):
        if mismatch:
            for cell in worksheet[excel_row]:
                cell.font = red_font

    widths = {"A": 12, "B": 12, "C": 34, "D": 8, "E": 120}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(output)


def accuracy_suffix(frame: pd.DataFrame, indicator: str) -> tuple[str, int, int]:
    standard_column = f"STANDARD_{indicator}"
    compared = frame[standard_column].isin([0, 1])
    total = int(compared.sum())
    if total == 0:
        return "unlabeled", 0, 0
    correct = int((compared & frame[indicator].eq(frame[standard_column])).sum())
    suffix = f"{round(correct / total * 10000):04d}"
    return suffix, correct, total


def validate_timestamp(value: str) -> str:
    if not re.fullmatch(r"\d{8}", value):
        raise argparse.ArgumentTypeError("timestamp must use MMDDHHMM, for example 06281543")
    return value


def validate_output_name(path: Path, indicator: str) -> Path:
    if not output_pattern(indicator).fullmatch(path.name):
        lower = indicator.lower()
        raise SystemExit(
            f"Final workbook name must match {lower}_<MMDDHHMM>_<accuracy-or-unlabeled>.xlsx, "
            f"for example {lower}_06281543_9333.xlsx or {lower}_06281543_unlabeled.xlsx"
        )
    return path


def accuracy_part_from_name(path: Path, indicator: str) -> str:
    validate_output_name(path, indicator)
    return path.stem.split("_")[-1]


def resolve_final_output_path(
    frame: pd.DataFrame,
    indicator: str,
    output: Path | None,
    output_dir: Path,
    timestamp: str | None,
) -> tuple[Path, str, int, int]:
    suffix, correct, total = accuracy_suffix(frame, indicator)
    if output is not None:
        validate_output_name(output, indicator)
        requested_suffix = accuracy_part_from_name(output, indicator)
        if requested_suffix != suffix:
            raise SystemExit(
                f"Output filename accuracy suffix {requested_suffix} does not match computed suffix {suffix}."
            )
        return output, suffix, correct, total
    stamp = timestamp or datetime.now().strftime("%m%d%H%M")
    return output_dir / f"{indicator.lower()}_{stamp}_{suffix}.xlsx", suffix, correct, total


def run_create_final_workbook(
    indicator: str,
    predictions: Path,
    targets: Path,
    output: Path | None,
    output_dir: Path,
    timestamp: str | None,
) -> None:
    frame = build_final_frame(predictions, targets, indicator)
    output_path, suffix, correct, total = resolve_final_output_path(frame, indicator, output, output_dir, timestamp)
    write_final_workbook(frame, output_path, indicator)

    mismatch_ids = (
        frame.loc[frame["MISMATCH"], ["REFNO", "RESULTNO"]]
        .astype(str)
        .apply(lambda row: f"{row['REFNO']}/{row['RESULTNO']}", axis=1)
        .tolist()
    )
    print(f"Saved: {output_path}")
    print(f"Rows: {len(frame)}")
    print(f"Columns: {', '.join(final_columns(indicator))}")
    if total:
        print(f"Accuracy: {correct}/{total} = {correct / total:.2%}")
    else:
        print(f"Accuracy: not computed (no standard {indicator} labels in {targets})")
    print(f"Accuracy suffix: {suffix}")
    print(f"Mismatches highlighted red: {len(mismatch_ids)}")
    if mismatch_ids:
        print(f"Mismatch REFNO/RESULTNOs: {', '.join(mismatch_ids)}")


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if pd is not None:
        try:
            if pd.isna(value):
                return ""
        except (TypeError, ValueError):
            pass
    return str(value)


def short_text(value: Any, limit: int = 650) -> str:
    text = re.sub(r"\s+", " ", as_text(value)).strip()
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "..."


def load_results(path: Path, indicator: str) -> pd.DataFrame:
    pandas = require_pandas()
    columns = final_columns(indicator)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        frame = pandas.read_excel(path, dtype={"REFNO": str, "RESULTNO": str})
    elif suffix == ".csv":
        frame = pandas.read_csv(path, dtype={"REFNO": str, "RESULTNO": str})
    elif suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            for key in ("results", "rows", "data"):
                if isinstance(payload.get(key), list):
                    payload = payload[key]
                    break
        frame = pandas.DataFrame(payload)
    else:
        raise SystemExit(f"Unsupported input format: {path.suffix}")

    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise SystemExit(f"{path} is missing required column(s): {', '.join(missing)}")
    return frame[columns].copy()


def parse_json_cell(value: Any) -> tuple[dict[str, Any], str | None]:
    text = as_text(value).strip()
    if not text:
        return {}, "empty JSON cell"
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        return {}, f"invalid JSON: {exc.msg}"
    if not isinstance(parsed, dict):
        return {}, "JSON cell is not an object"
    return parsed, None


def normalize_score(value: Any) -> str:
    text = as_text(value).strip()
    if text in {"1", "1.0", "True", "true"}:
        return "1"
    if text in {"0", "0.0", "False", "false"}:
        return "0"
    return text or "UNKNOWN"


def get_path(data: dict[str, Any], *keys: str) -> Any:
    current: Any = data
    for key in keys:
        if not isinstance(current, dict):
            return None
        current = current.get(key)
    return current


def evidence_quote(block: dict[str, Any]) -> str:
    quotes = block.get("evidence_quotes")
    if not isinstance(quotes, list) or not quotes:
        return ""
    first = quotes[0]
    if not isinstance(first, dict):
        return ""
    quote = as_text(first.get("quote")).strip()
    page = as_text(first.get("section_or_page")).strip()
    if not quote or quote.upper() == "ABSENT":
        return ""
    quote = short_text(quote, 550)
    return f"{quote} ({page})" if page else quote


def status_satisfied(status: Any) -> bool | None:
    upper = as_text(status).strip().upper()
    if "NOT_SATISFIED" in upper or upper in {"ABSENT", "INVALID", "NO", "NOT_MET", "FALSE"}:
        return False
    if "SATISFIED" in upper or upper in {"PRESENT", "VALID", "YES", "MET", "TRUE"}:
        return True
    return None


def split_status_reason(value: Any) -> tuple[str, str]:
    text = short_text(value, 900)
    if " - " in text:
        status, reasoning = text.split(" - ", 1)
        return status.strip() or "UNKNOWN", reasoning.strip()
    return text.strip() or "UNKNOWN", ""


def evidence_status(block: dict[str, Any]) -> tuple[str, bool | None]:
    status = as_text(block.get("status")).strip() or "UNKNOWN"
    satisfied = block.get("satisfied")
    if isinstance(satisfied, bool):
        return status, satisfied
    return status, status_satisfied(status)


def condition_label(key: str, block: dict[str, Any]) -> str:
    label = as_text(block.get("label") or block.get("criterion_question") or block.get("condition")).strip()
    if label:
        return label
    return key.replace("Evidence_", "").replace("_", " ")


def r3_condition_blocks(parsed: dict[str, Any]) -> list[dict[str, Any]]:
    mapping = [
        ("A", "Evidence_A_Carrier_Investigation", "Carrier investigation"),
        ("B", "Evidence_B_Carrier_Identity", "Carrier identity"),
        ("C", "Evidence_C_Carrier_Significance", "Carrier significance"),
    ]
    blocks = []
    for short_key, json_key, fallback_label in mapping:
        block = parsed.get(json_key)
        if isinstance(block, dict):
            status, satisfied = evidence_status(block)
            reasoning = short_text(block.get("reasoning") or block.get("reason"), 650)
            quote = evidence_quote(block)
        else:
            status, reasoning = split_status_reason(block)
            satisfied = status_satisfied(status)
            quote = ""
        if status == "UNKNOWN" and not reasoning and not quote:
            continue
        blocks.append(
            {
                "key": short_key,
                "label": fallback_label,
                "status": status,
                "satisfied": satisfied,
                "quote": quote,
                "reasoning": reasoning,
            }
        )

    metadata = parsed.get("metadata_rockmag")
    if isinstance(metadata, dict):
        quote = short_text(metadata.get("quote"), 550)
        if quote.upper() == "ABSENT":
            quote = ""
        reason = short_text(metadata.get("reason"), 650)
        status = as_text(metadata.get("status")).strip() or ("valid" if metadata.get("valid_r3_evidence") else "absent")
        satisfied = metadata.get("valid_r3_evidence")
        if not isinstance(satisfied, bool):
            satisfied = status_satisfied(status)
        if quote or reason:
            blocks.append(
                {
                    "key": "metadata",
                    "label": "PMAGRESULT rock-magnetic evidence",
                    "status": status,
                    "satisfied": satisfied,
                    "quote": quote,
                    "reasoning": reason,
                }
            )
    return blocks


def generic_condition_blocks(parsed: dict[str, Any], indicator: str) -> list[dict[str, Any]]:
    explicit = parsed.get("condition_summary")
    if isinstance(explicit, list):
        blocks = []
        for index, item in enumerate(explicit, start=1):
            if not isinstance(item, dict):
                continue
            if indicator == "R5" and as_text(item.get("key")) == "source_default":
                continue
            status, satisfied = evidence_status(item)
            quote = evidence_quote(item)
            if indicator == "R7" and satisfied is False:
                quote = ""
            blocks.append(
                {
                    "key": as_text(item.get("key") or index),
                    "label": condition_label(f"Condition {index}", item),
                    "status": status,
                    "satisfied": satisfied,
                    "quote": quote,
                    "reasoning": short_text(item.get("reasoning"), 650),
                }
            )
        if blocks:
            return blocks

    blocks = []
    for key, value in parsed.items():
        if not isinstance(value, dict):
            continue
        if "status" not in value and "satisfied" not in value:
            continue
        status, satisfied = evidence_status(value)
        blocks.append(
            {
                "key": key.replace("_", " "),
                "label": condition_label(key, value),
                "status": status,
                "satisfied": satisfied,
                "quote": evidence_quote(value),
                "reasoning": short_text(value.get("reasoning"), 650),
            }
        )
    return blocks


def standard_lookup(standard: Path, indicator: str) -> dict[str, str]:
    pandas = require_pandas()
    if not standard.exists():
        return {}
    columns = set(pandas.read_excel(standard, nrows=0).columns)
    if not {"RESULTNO", indicator}.issubset(columns):
        return {}
    frame = pandas.read_excel(standard, usecols=["RESULTNO", indicator], dtype={"RESULTNO": str})
    lookup: dict[str, str] = {}
    for _, row in frame.iterrows():
        score = normalize_score(row.get(indicator))
        if score in {"0", "1"}:
            lookup[as_text(row.get("RESULTNO"))] = score
    return lookup


def mismatch_metadata(resultno: Any, score: str, standards: dict[str, str]) -> tuple[bool, str, str]:
    standard_score = standards.get(as_text(resultno))
    if standard_score is None:
        return False, "UNKNOWN", "no_standard"
    if score == standard_score:
        return False, standard_score, "correct"
    if score == "1" and standard_score == "0":
        return True, standard_score, "false_positive"
    if score == "0" and standard_score == "1":
        return True, standard_score, "false_negative"
    return True, standard_score, "mismatch"


def summarize_rows(
    frame: pd.DataFrame,
    indicator: str,
    standards: dict[str, str] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    standards = standards or {}
    stats: dict[str, Any] = {
        "total": 0,
        "score_yes": 0,
        "score_no": 0,
        "outcome_unknown": 0,
        "manual_review": 0,
        "json_errors": 0,
        "mismatches": 0,
        "false_positive": 0,
        "false_negative": 0,
        "no_standard": 0,
        "conditions": {},
    }
    score_column = indicator
    json_column = f"{indicator}_RESULT_JSON"

    for _, source in frame.iterrows():
        parsed, error = parse_json_cell(source[json_column])
        score = normalize_score(source[score_column])
        conditions = r3_condition_blocks(parsed) if indicator == "R3" else generic_condition_blocks(parsed, indicator)
        manual = bool(parsed.get("manual_review")) if parsed else False
        outcome = as_text(parsed.get("outcome") or "DECIDED").strip().upper()
        summary = as_text(parsed.get("summary") or parsed.get("final_decision_reasoning")).strip()
        mismatch, standard_score, error_type = mismatch_metadata(source["RESULTNO"], score, standards)

        stats["total"] += 1
        if score == "1":
            stats["score_yes"] += 1
        elif score == "0":
            stats["score_no"] += 1
        if outcome == "UNKNOWN":
            stats["outcome_unknown"] += 1
        if manual:
            stats["manual_review"] += 1
        if error:
            stats["json_errors"] += 1
        if mismatch:
            stats["mismatches"] += 1
        if error_type in {"false_positive", "false_negative", "no_standard"}:
            stats[error_type] += 1
        for condition in conditions:
            key = as_text(condition["key"])
            entry = stats["conditions"].setdefault(
                key,
                {"key": key, "label": as_text(condition["label"]), "satisfied": 0, "total": 0},
            )
            entry["total"] += 1
            if condition["satisfied"] is True:
                entry["satisfied"] += 1

        rows.append(
            {
                "refno": as_text(source["REFNO"]),
                "resultno": as_text(source["RESULTNO"]),
                "target": as_text(source["TARGET"]),
                "score": score,
                "outcome": outcome,
                "summary": summary or (error or "No summary available."),
                "mismatch": mismatch,
                "standard_score": standard_score,
                "error_type": error_type,
                "conditions": conditions,
            }
        )

    stats["conditions"] = list(stats["conditions"].values())
    return rows, stats


def safe_json_for_script(payload: dict[str, Any]) -> str:
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return text.replace("</", "<\\/")


def build_payload(
    frame: pd.DataFrame,
    input_path: Path,
    indicator: str,
    title: str,
    standards: dict[str, str] | None = None,
) -> dict[str, Any]:
    rows, stats = summarize_rows(frame, indicator, standards)
    public_stats = {key: value for key, value in stats.items() if key not in {"manual_review", "json_errors"}}
    return {
        "title": title,
        "indicator": indicator,
        "score_label": indicator,
        "source": str(input_path),
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "standard_mode": "labelled_rows" if standards else "unlabelled_extension_rows",
        "stats": public_stats,
        "rows": rows,
    }


def render_template(template_path: Path, payload: dict[str, Any]) -> str:
    template = template_path.read_text(encoding="utf-8")
    if PAYLOAD_PLACEHOLDER not in template:
        raise SystemExit(f"{template_path} is missing {PAYLOAD_PLACEHOLDER}")
    return template.replace(PAYLOAD_PLACEHOLDER, safe_json_for_script(payload))


def latest_named_workbook(input_dir: Path, indicator: str) -> Path:
    candidates = sorted(
        (
            path
            for path in input_dir.glob(f"{indicator.lower()}_*.xlsx")
            if output_pattern(indicator).fullmatch(path.name)
        ),
        key=lambda path: (path.stat().st_mtime, path.name),
        reverse=True,
    )
    if not candidates:
        lower = indicator.lower()
        raise SystemExit(
            f"No final {indicator} workbook matching {lower}_<MMDDHHMM>_<accuracy-or-unlabeled>.xlsx found in {input_dir}"
        )
    return candidates[0]


def resolve_html_output_path(input_path: Path, output: Path | None, indicator: str) -> Path:
    if input_path.suffix.lower() in {".xlsx", ".xlsm", ".xls"} and not output_pattern(indicator).fullmatch(input_path.name):
        lower = indicator.lower()
        raise SystemExit(
            f"Final {indicator} workbook name must match {lower}_<MMDDHHMM>_<accuracy-or-unlabeled>.xlsx "
            "before creating the matching HTML page."
        )
    resolved = output or input_path.with_suffix(".html")
    if resolved.stem != input_path.stem:
        raise SystemExit(
            "HTML output basename must match the workbook basename, "
            f"expected {input_path.with_suffix('.html').name}"
        )
    return resolved


def run_create_results_page(
    indicator: str,
    input_path: Path | None,
    output: Path | None,
    template: Path,
    title: str | None,
    standard: Path = DEFAULT_TARGETS,
) -> None:
    resolved_input = input_path or latest_named_workbook(DEFAULT_OUTPUT_DIR, indicator)
    frame = load_results(resolved_input, indicator)
    page_title = title or f"{indicator} Results Visual Review"
    payload = build_payload(frame, resolved_input, indicator, page_title, standard_lookup(standard, indicator))
    output_path = resolve_html_output_path(resolved_input, output, indicator)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(render_template(template, payload), encoding="utf-8")

    stats = payload["stats"]
    print(f"Saved: {output_path}")
    print(f"Template: {template}")
    print(f"Rows: {stats['total']}")
    print(f"{indicator}=1: {stats['score_yes']} | {indicator}=0: {stats['score_no']}")
    print(f"Mismatches: {stats['mismatches']}")


def metrics(frame: pd.DataFrame, indicator: str) -> dict[str, object]:
    n = len(frame)
    standard_column = f"STANDARD_{indicator}"
    tp = int(((frame[indicator] == 1) & (frame[standard_column] == 1)).sum())
    tn = int(((frame[indicator] == 0) & (frame[standard_column] == 0)).sum())
    fp = int(((frame[indicator] == 1) & (frame[standard_column] == 0)).sum())
    fn = int(((frame[indicator] == 0) & (frame[standard_column] == 1)).sum())
    return {
        "n_compared": n,
        "correct": tp + tn,
        "accuracy": (tp + tn) / n if n else None,
        "true_positive": tp,
        "true_negative": tn,
        "false_positive": fp,
        "false_negative": fn,
        f"precision_{indicator}_1": tp / (tp + fp) if tp + fp else None,
        f"recall_{indicator}_1": tp / (tp + fn) if tp + fn else None,
    }


def defaulted_mask(frame: pd.DataFrame, indicator: str) -> pd.Series:
    pandas = require_pandas()
    column = f"{indicator}_DEFAULTED"
    if column not in frame.columns:
        return pandas.Series(False, index=frame.index)

    values = frame[column].fillna(False)
    if pandas.api.types.is_bool_dtype(values):
        return values.astype(bool)
    if pandas.api.types.is_numeric_dtype(values):
        return values.ne(0)
    return values.astype(str).str.strip().str.casefold().isin({"true", "1", "yes", "y"})


def run_compare_with_standard(indicator: str, predictions: Path | None, standard: Path, output: Path) -> None:
    pandas = require_pandas()
    prediction_path = predictions or latest_named_workbook(DEFAULT_OUTPUT_DIR, indicator)
    pred = pandas.read_excel(prediction_path, dtype={"REFNO": str, "RESULTNO": str})
    if "RESULTNO" not in pred.columns:
        raise SystemExit(
            f"{prediction_path} must contain RESULTNO because the standard workbook can contain multiple targets per REFNO."
        )
    std = pandas.read_excel(
        standard,
        usecols=["RESULTNO", "REFNO", "ROCKNAME", indicator],
        dtype={"RESULTNO": str, "REFNO": str},
    )
    std = std.rename(columns={"ROCKNAME": "STANDARD_TARGET", indicator: f"STANDARD_{indicator}"})

    comparison = pred.merge(std, on="RESULTNO", how="left", validate="many_to_one", suffixes=("", "_STANDARD"))
    if "REFNO_STANDARD" in comparison.columns:
        comparison["REFNO"] = comparison["REFNO"].where(
            comparison["REFNO"].notna(), comparison["REFNO_STANDARD"]
        )
    standard_column = f"STANDARD_{indicator}"
    comparison["AGREE"] = comparison[indicator].eq(comparison[standard_column])
    comparison["ERROR_TYPE"] = "correct"
    comparison.loc[(comparison[indicator] == 1) & (comparison[standard_column] == 0), "ERROR_TYPE"] = "false_positive"
    comparison.loc[(comparison[indicator] == 0) & (comparison[standard_column] == 1), "ERROR_TYPE"] = "false_negative"
    comparison.loc[comparison[standard_column].isna(), "ERROR_TYPE"] = "no_standard"

    comparable = comparison[comparison[standard_column].isin([0, 1])].copy()
    nondefaulted = comparable[~defaulted_mask(comparable, indicator)].copy()
    summary = pandas.DataFrame(
        [
            {"scope": "all matched targets (includes defaults)", **metrics(comparable, indicator)},
            {f"scope": f"non-defaulted {indicator} results", **metrics(nondefaulted, indicator)},
        ]
    )
    disagreements = comparable[~comparable.AGREE].sort_values(["ERROR_TYPE", "REFNO"])

    assert len(comparable) == len(pred), "Every prediction must have one standard label."
    output.parent.mkdir(parents=True, exist_ok=True)
    with pandas.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        comparison.to_excel(writer, sheet_name="Comparison", index=False)
        disagreements.to_excel(writer, sheet_name="Disagreements", index=False)

    all_result = summary.iloc[0]
    valid_result = summary.iloc[1]
    print(f"All: {int(all_result.correct)}/{int(all_result.n_compared)} = {all_result.accuracy:.2%}")
    print(f"Non-defaulted: {int(valid_result.correct)}/{int(valid_result.n_compared)} = {valid_result.accuracy:.2%}")
    print(f"Predictions: {prediction_path}")
    print(f"Saved: {output}")
