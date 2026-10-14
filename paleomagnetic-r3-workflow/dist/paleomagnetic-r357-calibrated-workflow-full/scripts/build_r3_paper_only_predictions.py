#!/usr/bin/env python3
"""Build R3 predictions from target identities and original-paper text only."""

from __future__ import annotations

import argparse
from collections import Counter, OrderedDict
import json
from pathlib import Path
import sys
from typing import Any

import pandas as pd
from openpyxl import load_workbook

try:
    from .paper_text_retrieval import PaperTextRetriever, canonical_target_identity
    from .r3_carrier_claim_status import carrier_claim_is_affirmed
    from .expert_conflict_engine import (
        DEFAULT_KNOWLEDGE_BASE,
        ExpertConflictEngine,
        compact_analysis,
    )
except ImportError:
    from paper_text_retrieval import PaperTextRetriever, canonical_target_identity
    from r3_carrier_claim_status import carrier_claim_is_affirmed
    from expert_conflict_engine import (
        DEFAULT_KNOWLEDGE_BASE,
        ExpertConflictEngine,
        compact_analysis,
    )


ROOT = Path(__file__).resolve().parents[1]
IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _clean_identity(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_targets(path: Path, limit: int | None = None) -> pd.DataFrame:
    """Load a label-free identity workbook with an exact three-column schema."""

    if path.suffix.casefold() != ".xlsx":
        raise SystemExit(f"Identity manifest must be an XLSX file: {path}")
    if not path.is_file():
        raise SystemExit(f"Identity manifest does not exist: {path}")
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise SystemExit(f"Could not open identity manifest {path}: {exc}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise SystemExit("Strict identity manifest must contain exactly one worksheet")
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=False)
        header_cells = next(rows, ())
        if any(cell.data_type == "f" for cell in header_cells):
            raise SystemExit("Identity manifest header cannot contain formulas")
        header = tuple(_clean_identity(cell.value) for cell in header_cells)
        if header != IDENTITY_COLUMNS:
            raise SystemExit(
                "Strict identity columns must be exactly REFNO, RESULTNO, ROCKNAME"
            )

        identities: list[dict[str, str]] = []
        for excel_row, cells in enumerate(rows, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise SystemExit(
                    f"Identity manifest row {excel_row} contains a formula"
                )
            values = tuple(_clean_identity(cell.value) for cell in cells)
            if not any(values):
                continue
            if len(values) != len(IDENTITY_COLUMNS) or any(not value for value in values):
                raise SystemExit(
                    f"Identity manifest row {excel_row} has an empty or invalid identity value"
                )
            identities.append(dict(zip(IDENTITY_COLUMNS, values)))
    finally:
        workbook.close()

    if not identities:
        raise SystemExit(f"Identity manifest contains no data rows: {path}")
    frame = pd.DataFrame.from_records(identities, columns=IDENTITY_COLUMNS)
    if frame["RESULTNO"].duplicated().any():
        examples = ", ".join(frame.loc[frame["RESULTNO"].duplicated(), "RESULTNO"].head(10))
        raise SystemExit(f"{path} contains duplicate RESULTNO values: {examples}")
    if limit is not None:
        frame = frame.head(limit).copy()
    return frame


def load_page_cache(path: Path) -> dict[str, dict[str, Any]]:
    cache: dict[str, dict[str, Any]] = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise SystemExit(f"{path}:{line_number}: invalid JSON: {exc}") from exc
            refno = clean(row.get("REFNO"))
            if refno:
                cache[refno] = row
    return cache


def evidence_quote(hit: dict[str, Any] | None, why: str) -> list[dict[str, Any]]:
    if not hit or not clean(hit.get("quote")):
        return []
    location = f"page {hit.get('page')}"
    section = clean(hit.get("section"))
    if section and section.casefold() != "unknown":
        location += f" | {section}"
    return [
        {
            "quote": clean(hit.get("quote")),
            "section_or_page": location,
            "why_relevant": why,
            "target_binding": clean(hit.get("target_anchor")),
            "target_relevance": hit.get("target_relevance"),
        }
    ]


def compact_hit(hit: dict[str, Any] | None, *, quote_limit: int = 360) -> dict[str, Any]:
    """Keep the evidence needed for review without duplicating full excerpts."""

    if not hit:
        return {}
    output = {
        key: hit.get(key)
        for key in (
            "category",
            "subtype",
            "label",
            "matched_text",
            "page",
            "section",
            "target_relevance",
            "binding",
            "target_anchor",
            "local_target",
            "scope_id",
            "scope_kind",
            "scope_target_exclusive",
            "scope_competing_targets",
            "reference_noise",
            "source",
        )
        if hit.get(key) is not None and hit.get(key) != ""
    }
    quote = clean(hit.get("quote"))
    if quote:
        output["quote"] = (
            quote
            if len(quote) <= quote_limit
            else quote[: quote_limit - 3].rstrip() + "..."
        )
    return output


def compact_retrieval(retrieval: dict[str, Any], direct_verdict: dict[str, Any]) -> dict[str, Any]:
    chain = dict(retrieval.get("best_chain") or {})
    chain["hits"] = {
        category: compact_hit(hit)
        for category, hit in (chain.get("hits") or {}).items()
        if hit
    }
    alternatives = {
        category: [compact_hit(dict(hit)) for hit in list(hits or [])[:2]]
        for category, hits in (retrieval.get("categories") or {}).items()
        if category != "verdict" and hits
    }
    ignored = list(direct_verdict.get("ignored_existing_scores") or [])
    return {
        "target_anchor": retrieval.get("target_anchor"),
        "scope_binding": retrieval.get("scope_binding"),
        "best_chain": chain,
        "direct_verdict": {
            "status": direct_verdict.get("status", "absent"),
            "ignored_existing_scores_count": len(ignored),
        },
        "alternatives": alternatives,
        "provenance": retrieval.get("provenance"),
    }


def selected_chain_hits(
    retrieval: dict[str, Any],
    *,
    require_local: bool = False,
) -> dict[str, dict[str, Any]]:
    chain = retrieval.get("best_chain") or {}
    raw = chain.get("hits") or {}
    selected: dict[str, dict[str, Any]] = {}
    for category in ("method", "carrier", "bridge"):
        hit = dict(raw.get(category) or {})
        selected[category] = (
            hit if hit and (not require_local or hit.get("local_target") is True) else {}
        )
    return selected


def selected_independent_method(
    retrieval: dict[str, Any],
    *,
    require_local: bool = False,
) -> dict[str, Any]:
    """Return the strongest target-bound carrier-investigation method."""

    candidates = [
        dict(hit)
        for hit in retrieval.get("categories", {}).get("method", []) or []
        if not hit.get("reference_noise")
        and hit.get("subtype") != "demagnetization"
        and float(hit.get("target_relevance") or 0.0) >= 0.2
        and (not require_local or hit.get("local_target") is True)
    ]
    return max(
        candidates,
        key=lambda hit: (
            bool(hit.get("local_target")),
            float(hit.get("target_relevance") or 0.0),
        ),
        default={},
    )


def selected_category_hit(
    retrieval: dict[str, Any],
    category: str,
    *,
    require_local: bool = False,
) -> dict[str, Any]:
    candidates = [
        dict(hit)
        for hit in retrieval.get("categories", {}).get(category, []) or []
        if not hit.get("reference_noise")
        and float(hit.get("target_relevance") or 0.0) >= 0.2
        and (not require_local or hit.get("local_target") is True)
    ]
    return max(
        candidates,
        key=lambda hit: (
            bool(hit.get("local_target")),
            float(hit.get("target_relevance") or 0.0),
        ),
        default={},
    )


def selected_direct_verdict(retrieval: dict[str, Any]) -> dict[str, Any]:
    """Ignore embedded R3 scores so existing evaluations cannot become evidence."""

    ignored = [
        dict(hit)
        for hit in retrieval.get("categories", {}).get("verdict", []) or []
    ]
    return {
        "status": "absent",
        "hit": {},
        "positive_hit": {},
        "negative_hit": {},
        "ignored_existing_scores": ignored,
    }


PROMOTED_SCOPE_ANCHORS = {
    "target_same_paragraph",
    "target_same_table_row",
    "target_same_discourse",
}
CARRIER_IDENTIFYING_BRIDGE_SUBTYPES = {
    "explicit_carrier_bridge",
    "mineral_remanence_bridge",
    "carrier_significance",
}


def carrier_identifying_bridge(hit: dict[str, Any]) -> bool:
    return bool(
        hit
        and hit.get("subtype") in CARRIER_IDENTIFYING_BRIDGE_SUBTYPES
        and carrier_claim_is_affirmed(hit)
    )


def compatible_promoted_scope(
    hits: list[dict[str, Any]],
    *,
    multi_target: bool,
) -> bool:
    if not multi_target or not any(
        hit.get("target_anchor") in PROMOTED_SCOPE_ANCHORS for hit in hits
    ):
        return True
    scope_ids = [clean(hit.get("scope_id")) for hit in hits]
    return bool(scope_ids) and all(scope_id == scope_ids[0] for scope_id in scope_ids)


def selected_compatible_criterion_chain(
    retrieval: dict[str, Any],
    *,
    multi_target: bool,
) -> dict[str, dict[str, Any]]:
    """Select the strongest complete chain that satisfies locality constraints."""

    best_chain_hits = dict((retrieval.get("best_chain") or {}).get("hits") or {})

    def candidates(category: str) -> list[dict[str, Any]]:
        members = [
            dict(hit)
            for hit in retrieval.get("categories", {}).get(category, []) or []
        ]
        best_hit = dict(best_chain_hits.get(category) or {})
        if best_hit and best_hit not in members:
            members.append(best_hit)
        return [
            hit
            for hit in members
            if not hit.get("reference_noise")
            and float(hit.get("target_relevance") or 0.0) >= 0.2
            and (not multi_target or hit.get("local_target") is True)
            and not (category == "method" and hit.get("subtype") == "demagnetization")
            and not (
                category in {"carrier", "bridge"}
                and not carrier_claim_is_affirmed(hit)
            )
        ]

    methods = candidates("method")
    carriers = candidates("carrier")
    bridges = candidates("bridge")
    best: tuple[tuple[float, ...], dict[str, dict[str, Any]]] | None = None
    for method in methods:
        for bridge in bridges:
            carrier_options: list[dict[str, Any] | None] = list(carriers)
            if carrier_identifying_bridge(bridge):
                carrier_options.append(None)
            for carrier in carrier_options:
                hits = [hit for hit in (method, carrier, bridge) if hit]
                if not carrier and not carrier_identifying_bridge(bridge):
                    continue
                pages = [int(hit.get("page") or 0) for hit in hits]
                if not pages or any(page <= 0 for page in pages):
                    continue
                page_span = max(pages) - min(pages)
                if page_span > 1 or not compatible_promoted_scope(
                    hits,
                    multi_target=multi_target,
                ):
                    continue
                relevance = [float(hit.get("target_relevance") or 0.0) for hit in hits]
                rank = (
                    float(-page_span),
                    min(relevance),
                    sum(relevance) / len(relevance),
                    float(len(hits)),
                )
                selected = {
                    "method": method,
                    "carrier": carrier or {},
                    "bridge": bridge,
                }
                if best is None or rank > best[0]:
                    best = (rank, selected)
    return best[1] if best else {}


def judge_r3(
    retrieval: dict[str, Any],
    *,
    multi_target: bool,
    identity_ambiguous: bool,
    source_flags: list[str],
) -> dict[str, Any]:
    chain = retrieval.get("best_chain") or {}
    chain_hits = selected_chain_hits(retrieval, require_local=multi_target)
    independent_method = selected_independent_method(
        retrieval,
        require_local=multi_target,
    )
    compatible_chain = selected_compatible_criterion_chain(
        retrieval,
        multi_target=multi_target,
    )
    if compatible_chain:
        hits = compatible_chain
    else:
        hits = dict(chain_hits)
        if not hits["carrier"]:
            hits["carrier"] = selected_category_hit(
                retrieval,
                "carrier",
                require_local=multi_target,
            )
        if not hits["bridge"]:
            hits["bridge"] = selected_category_hit(
                retrieval,
                "bridge",
                require_local=multi_target,
            )
        method = chain_hits["method"]
        if not method or method.get("subtype") == "demagnetization":
            method = independent_method
        hits["method"] = method
    method = hits["method"]
    carrier = hits["carrier"]
    bridge = hits["bridge"]
    direct_verdict = selected_direct_verdict(retrieval)
    method_ok = bool(method) and float(method.get("target_relevance") or 0.0) >= 0.2
    independent_method_ok = bool(independent_method)
    carrier_ok = bool(
        carrier
        and float(carrier.get("target_relevance") or 0.0) >= 0.2
        and carrier_claim_is_affirmed(carrier)
    )
    bridge_ok = bool(
        bridge
        and float(bridge.get("target_relevance") or 0.0) >= 0.2
        and carrier_claim_is_affirmed(bridge)
    )
    criterion_hits = [hit for hit in (method, carrier, bridge) if hit]
    criterion_pages = [int(hit.get("page") or 0) for hit in criterion_hits if hit.get("page")]
    page_span = max(criterion_pages) - min(criterion_pages) if criterion_pages else None
    connected = bool(page_span is not None and page_span <= 1)
    scope_compatible = compatible_promoted_scope(
        criterion_hits,
        multi_target=multi_target,
    )
    bridge_identifies_carrier = carrier_identifying_bridge(bridge)
    carrier_identified = bool(carrier_ok or bridge_identifies_carrier)
    significance_ok = bridge_ok
    target_anchor = retrieval.get("target_anchor") or {}
    target_ok = not multi_target or target_anchor.get("strength") in {"strong", "moderate"}
    if multi_target and target_anchor.get("status") == "generic_only":
        chain_hits = [hit for hit in hits.values() if hit]
        if direct_verdict["status"] in {"positive", "negative", "conflict"}:
            target_ok = True
        else:
            target_ok = bool(chain_hits) and all(
                bool(hit.get("local_target")) for hit in chain_hits
            )
    source_available = bool(retrieval.get("source_available"))
    relationship_chain_ok = bool(
        source_available
        and method_ok
        and carrier_identified
        and significance_ok
        and connected
        and scope_compatible
        and target_ok
    )
    expert_investigation_ok = bool(
        source_available
        and target_ok
        and relationship_chain_ok
    )
    if not source_available:
        score = 0
        outcome = "UNKNOWN"
    elif direct_verdict["status"] == "conflict":
        score = 0
        outcome = "UNKNOWN"
    elif direct_verdict["status"] == "positive" and target_ok:
        score = 1
        outcome = "POSITIVE"
    elif direct_verdict["status"] == "negative" and target_ok:
        score = 0
        outcome = "NEGATIVE"
    else:
        score = int(expert_investigation_ok)
        outcome = "POSITIVE" if score else "NEGATIVE"
    quality_risk = bool(
        set(source_flags)
        & {"short_text", "glyph_or_cid_text", "possibly_incomplete_extraction", "no_extractable_text"}
    )
    manual_review = bool(
        not source_available
        or identity_ambiguous
        or quality_risk
        or not target_ok
        or (method_ok and not carrier_identified)
        or (method_ok and not significance_ok)
        or (len(criterion_hits) >= 2 and not connected)
        or (len(criterion_hits) >= 2 and not scope_compatible)
        or outcome == "UNKNOWN"
    )
    conditions = {
        "method": method_ok,
        "independent_method": independent_method_ok,
        "carrier": carrier_identified,
        "bridge": significance_ok,
        "connected": connected,
        "scope_compatible": scope_compatible,
        "target": target_ok,
    }
    failed = [name for name, passed in conditions.items() if not passed]
    if not source_available:
        reason = (
            "R3 binary compatibility value is 0, but outcome is UNKNOWN because no "
            "original-paper text was available."
        )
    elif direct_verdict["status"] == "conflict":
        reason = (
            "R3 binary compatibility value is 0, but outcome is UNKNOWN because "
            "conflicting target-bound R3 verdicts were retrieved."
        )
    elif direct_verdict["status"] == "positive":
        reason = "R3=1 because the original paper explicitly reports target-bound R3=1."
    elif direct_verdict["status"] == "negative":
        reason = "R3=0 because the original paper explicitly reports target-bound R3=0."
    elif score:
        reason = (
            "R3=1 because a target-bound diagnostic investigation, carrier identification, "
            "and an interpretation of carrier significance form a connected original-paper "
            "evidence chain."
        )
    else:
        reason = (
            "R3=0 in paper-only mode because the original-paper gate failed: "
            + ", ".join(failed or ["source unavailable"])
            + "."
        )
    return {
        "score": score,
        "outcome": outcome,
        "manual_review": manual_review,
        "reason": reason,
        "conditions": conditions,
        "hits": hits,
        "direct_verdict": direct_verdict,
    }


def evidence_block(
    hit: dict[str, Any],
    satisfied: bool,
    label: str,
) -> dict[str, Any]:
    return {
        "status": "SATISFIED" if satisfied else "NOT_SATISFIED",
        "satisfied": satisfied,
        "evidence_quotes": evidence_quote(hit, label),
        "reasoning": label if hit else f"No target-bound {label.casefold()} description was retrieved.",
        "provenance": "original_paper_text",
    }


def build_rows(
    targets: pd.DataFrame,
    cache: dict[str, dict[str, Any]],
    page_cache_path: Path,
    knowledge_engine: ExpertConflictEngine | None = None,
) -> list[dict[str, Any]]:
    knowledge_engine = knowledge_engine or ExpertConflictEngine.from_path()
    canonical_targets_by_refno: dict[str, set[str]] = {}
    target_names_by_refno: dict[str, dict[str, str]] = {}
    result_count_by_target: Counter[tuple[str, str]] = Counter()
    for identity in targets.to_dict(orient="records"):
        refno = clean(identity["REFNO"])
        canonical_target = canonical_target_identity(clean(identity["ROCKNAME"]))
        canonical_targets_by_refno.setdefault(refno, set()).add(canonical_target)
        target_names_by_refno.setdefault(refno, {}).setdefault(
            canonical_target,
            clean(identity["ROCKNAME"]),
        )
        result_count_by_target[(refno, canonical_target)] += 1
    rows: list[dict[str, Any]] = []
    retrievers: OrderedDict[str, PaperTextRetriever] = OrderedDict()
    analysis_cache: dict[
        tuple[str, str, bool],
        tuple[dict[str, Any], dict[str, Any], dict[str, Any], list[str], dict[str, Any]],
    ] = {}
    for source_row in targets.to_dict(orient="records"):
        refno = clean(source_row["REFNO"])
        resultno = clean(source_row["RESULTNO"])
        target = clean(source_row["ROCKNAME"])
        canonical_target = canonical_target_identity(target)
        cache_row = cache.get(
            refno,
            {"REFNO": refno, "source": {"source_quality_flags": ["no_extractable_text"]}, "pages": []},
        )
        multi_target = len(canonical_targets_by_refno.get(refno, ())) > 1
        identity_ambiguous = result_count_by_target[(refno, canonical_target)] > 1
        if refno not in retrievers:
            retrievers[refno] = PaperTextRetriever(cache_row)
            if len(retrievers) > 8:
                retrievers.popitem(last=False)
        else:
            retrievers.move_to_end(refno)
        source = cache_row.get("source") or {}
        flags = [str(value) for value in source.get("source_quality_flags", [])]
        analysis_key = (refno, canonical_target, multi_target)
        cached = analysis_cache.get(analysis_key)
        if cached is None:
            retrieval = retrievers[refno].retrieve(
                target,
                multi_target,
                "R3",
                competing_targets=[
                    name
                    for sibling_target, name in target_names_by_refno[refno].items()
                    if sibling_target != canonical_target
                ],
                allow_structured_scope=True,
                limit=12,
                per_category_limit=4,
            )
            judgement = judge_r3(
                retrieval,
                multi_target=multi_target,
                identity_ambiguous=identity_ambiguous,
                source_flags=flags,
            )
            conflict = knowledge_engine.evaluate(
                "R3",
                target,
                retrieval=retrieval,
                pages=cache_row,
                multi_target=multi_target,
            )
            cached = (retrieval, judgement, conflict, flags, source)
            analysis_cache[analysis_key] = cached
        retrieval, judgement, conflict, flags, source = cached
        conflict_decision = conflict["decision"]
        recommended = conflict_decision.get("recommended_value")
        if recommended == 1 and judgement["score"] == 1:
            score = 1
            outcome = str(conflict_decision.get("outcome") or "POSITIVE")
        elif recommended == 1:
            score = 0
            outcome = "UNKNOWN"
            conflict_decision = {
                **conflict_decision,
                "reason": (
                    "Knowledge groups were detected, but the formal R3 method/carrier/"
                    "significance evidence was not connected on the same or an adjacent page."
                ),
            }
        elif recommended == 0:
            score = int(recommended)
            outcome = str(conflict_decision.get("outcome") or "UNKNOWN")
        else:
            score = 0
            outcome = "UNKNOWN"
        final_reason = str(conflict_decision.get("reason") or judgement["reason"])
        if identity_ambiguous:
            score = 0
            outcome = "UNKNOWN"
            final_reason = (
                "UNKNOWN: multiple result rows share this REFNO and ROCKNAME, while the "
                "strict identity manifest contains no component locator. "
                + final_reason
            )
        if outcome == "UNKNOWN" and "UNKNOWN" not in final_reason.upper():
            final_reason = "UNKNOWN: " + final_reason
        manual_review = bool(
            judgement["manual_review"]
            or identity_ambiguous
            or outcome == "UNKNOWN"
            or conflict.get("relation") == "MIXED"
        )
        selected = judgement["hits"]
        payload = {
            "resultno": resultno,
            "target": target,
            "R3": score,
            "outcome": outcome,
            "evaluator": "Label-free expert-knowledge conflict workflow",
            "workflow_profile": "knowledge_strict_v2",
            "Evidence_A_Carrier_Investigation": evidence_block(
                selected["method"],
                judgement["conditions"]["method"],
                "Original paper reports a carrier-relevant investigation.",
            ),
            "Evidence_B_Carrier_Identity": evidence_block(
                selected["carrier"] or selected["bridge"],
                judgement["conditions"]["carrier"],
                "Original paper identifies a magnetic carrier.",
            ),
            "Evidence_C_Carrier_Significance": evidence_block(
                selected["bridge"],
                judgement["conditions"]["bridge"],
                "Original paper links carrier evidence to remanence interpretation.",
            ),
            "original_source_descriptions": compact_retrieval(
                retrieval,
                judgement["direct_verdict"],
            ),
            "knowledge_conflict_analysis": compact_analysis(
                conflict,
                relation_limit=4,
                quote_limit=360,
            ),
            "decision_basis": "expert_knowledge_base_and_original_paper_text",
            "runtime_scoring_sources": [
                "identity manifest: REFNO/RESULTNO/ROCKNAME",
                "original paper page text",
                "frozen expert knowledge base",
            ],
            "forbidden_runtime_scoring_sources": [
                "PMAGRESULT",
                "ROCKUNIT",
                "R3/R5/R7 labels",
                "previous prediction workbooks",
            ],
            "source_audit": {
                "page_cache": str(page_cache_path),
                "source": source,
                "multi_target_refno": multi_target,
                "identity_ambiguous": identity_ambiguous,
            },
            "logic_gate": (
                "diagnostic investigation + carrier identification + significance interpretation -> "
                "target/locality binding -> SUPPORT/CONFLICT/UNKNOWN -> conflict-preserving aggregation"
            ),
            "final_decision_reasoning": final_reason,
            "summary": (
                f"Knowledge-strict outcome={outcome}; R3={score}; "
                f"relation={conflict['relation']}; confidence={conflict['confidence']}."
            ),
            "manual_review": manual_review,
        }
        rows.append(
            {
                "REFNO": refno,
                "RESULTNO": resultno,
                "TARGET": target,
                "R3": score,
                "R3_RESULT_JSON": payload,
            }
        )
    return rows


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--identity-manifest",
        "--targets",
        dest="identity_manifest",
        type=Path,
        required=True,
        help="Strict XLSX containing exactly REFNO, RESULTNO, and ROCKNAME.",
    )
    parser.add_argument("--page-cache", type=Path, required=True)
    parser.add_argument("--knowledge-base", type=Path, default=DEFAULT_KNOWLEDGE_BASE)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary-json", type=Path, default=None)
    parser.add_argument("--limit", type=int, default=None)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    targets = load_targets(args.identity_manifest, args.limit)
    rows = build_rows(
        targets,
        load_page_cache(args.page_cache),
        args.page_cache,
        ExpertConflictEngine.from_path(args.knowledge_base),
    )
    write_jsonl(args.output, rows)
    counts = Counter(int(row["R3"]) for row in rows)
    outcome_counts = Counter(str(row["R3_RESULT_JSON"].get("outcome", "UNKNOWN")) for row in rows)
    review_count = sum(bool(row["R3_RESULT_JSON"].get("manual_review")) for row in rows)
    summary = {
        "workflow_profile": "knowledge_strict_v2",
        "identity_manifest": str(args.identity_manifest),
        "page_cache": str(args.page_cache),
        "output": str(args.output),
        "rows": len(rows),
        "score_counts": {str(key): value for key, value in sorted(counts.items())},
        "outcome_counts": dict(sorted(outcome_counts.items())),
        "manual_review_rows": review_count,
        "runtime_input_columns": ["REFNO", "RESULTNO", "ROCKNAME"],
        "knowledge_base": str(args.knowledge_base),
        "runtime_scoring_sources": [
            "identity manifest",
            "original paper page text",
            "frozen expert knowledge base",
        ],
        "forbidden_runtime_scoring_sources": ["PMAGRESULT", "ROCKUNIT", "R3/R5/R7 labels"],
    }
    if args.summary_json:
        args.summary_json.parent.mkdir(parents=True, exist_ok=True)
        args.summary_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
