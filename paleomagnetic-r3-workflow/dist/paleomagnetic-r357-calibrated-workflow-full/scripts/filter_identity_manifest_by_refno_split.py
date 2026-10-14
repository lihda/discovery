#!/usr/bin/env python3
"""Partition a strict identity manifest by a deterministic REFNO hash."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook


IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
DEFAULT_SALT = "r3-rockmag-distillation-v1"
DEFAULT_HOLDOUT_FRACTION = 0.20


class IdentitySplitError(ValueError):
    """Raised when an identity manifest cannot be partitioned safely."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def holdout_bucket(refno: str, salt: str = DEFAULT_SALT) -> float:
    digest = hashlib.sha256(f"{salt}\0{refno}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _partition_rows(
    rows: Iterable[tuple[str, str, str]],
    *,
    partition: str,
    salt: str,
    holdout_fraction: float,
) -> list[tuple[str, str, str]]:
    want_holdout = partition == "holdout"
    return [
        row
        for row in rows
        if (holdout_bucket(row[0], salt) < holdout_fraction) == want_holdout
    ]


def split_identity_manifest(
    input_path: Path,
    output_path: Path,
    *,
    partition: str = "holdout",
    salt: str = DEFAULT_SALT,
    holdout_fraction: float = DEFAULT_HOLDOUT_FRACTION,
) -> dict[str, Any]:
    input_path = input_path.resolve()
    output_path = output_path.resolve(strict=False)
    if input_path == output_path:
        raise IdentitySplitError("input and output identity manifests must be distinct")
    if partition not in {"training", "holdout"}:
        raise IdentitySplitError("partition must be training or holdout")
    if not 0.0 < holdout_fraction < 1.0:
        raise IdentitySplitError("holdout_fraction must be between zero and one")
    if not salt:
        raise IdentitySplitError("salt must be non-empty")
    if not input_path.is_file():
        raise IdentitySplitError(f"identity manifest does not exist: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise IdentitySplitError("identity manifest must contain one worksheet")
        worksheet = workbook.worksheets[0]
        row_iterator = worksheet.iter_rows(values_only=False)
        header = tuple(_clean(cell.value) for cell in next(row_iterator, ()))
        if header != IDENTITY_COLUMNS:
            raise IdentitySplitError(
                "identity columns must be exactly REFNO, RESULTNO, ROCKNAME"
            )
        rows: list[tuple[str, str, str]] = []
        for row_number, cells in enumerate(row_iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise IdentitySplitError(
                    f"identity formulas are forbidden at Excel row {row_number}"
                )
            values = tuple(_clean(cell.value) for cell in cells[:3])
            if not any(values):
                continue
            if len(values) != 3 or any(not value for value in values):
                raise IdentitySplitError(
                    f"empty identity value at Excel row {row_number}"
                )
            rows.append(values)
    finally:
        workbook.close()

    if not rows:
        raise IdentitySplitError("identity manifest contains no data rows")
    if len({row[1] for row in rows}) != len(rows):
        raise IdentitySplitError("identity manifest contains duplicate RESULTNO values")
    selected = _partition_rows(
        rows,
        partition=partition,
        salt=salt,
        holdout_fraction=holdout_fraction,
    )
    if not selected:
        raise IdentitySplitError(f"{partition} partition contains no rows")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    output = Workbook(write_only=True)
    sheet = output.create_sheet("identity")
    sheet.append(IDENTITY_COLUMNS)
    for row in selected:
        sheet.append(row)
    output.save(temporary)
    temporary.replace(output_path)

    input_refnos = {row[0] for row in rows}
    selected_refnos = {row[0] for row in selected}
    return {
        "schema_version": "strict_identity_refno_partition_v1",
        "status": "PASS",
        "partition": partition,
        "split": {
            "method": "sha256_threshold_by_refno_v1",
            "salt": salt,
            "holdout_fraction": holdout_fraction,
            "assignment_rule": (
                "uint64_be(SHA256(UTF8(salt + NUL + REFNO))[0:8]) / 2^64 "
                "< holdout_fraction"
            ),
        },
        "input": {
            "path": str(input_path),
            "sha256": file_sha256(input_path),
            "rows": len(rows),
            "unique_refnos": len(input_refnos),
        },
        "output": {
            "path": str(output_path),
            "sha256": file_sha256(output_path),
            "rows": len(selected),
            "unique_refnos": len(selected_refnos),
        },
        "labels_read": False,
        "source_tables_read": False,
        "output_columns": list(IDENTITY_COLUMNS),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit-json", type=Path, required=True)
    parser.add_argument(
        "--partition", choices=("training", "holdout"), default="holdout"
    )
    parser.add_argument("--salt", default=DEFAULT_SALT)
    parser.add_argument(
        "--holdout-fraction", type=float, default=DEFAULT_HOLDOUT_FRACTION
    )
    args = parser.parse_args()
    try:
        audit = split_identity_manifest(
            args.input,
            args.output,
            partition=args.partition,
            salt=args.salt,
            holdout_fraction=args.holdout_fraction,
        )
    except IdentitySplitError as exc:
        print(f"Identity split error: {exc}")
        return 2
    args.audit_json.parent.mkdir(parents=True, exist_ok=True)
    args.audit_json.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
