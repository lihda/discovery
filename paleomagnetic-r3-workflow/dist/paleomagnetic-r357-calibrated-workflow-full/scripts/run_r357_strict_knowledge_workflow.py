#!/usr/bin/env python3
"""Run physically isolated R3/R5/R7 expert-knowledge scoring.

The scoring process accepts exactly four data inputs: an identity-only XLSX
manifest, a complete original-paper page cache, a frozen expert knowledge
base, and a frozen normative criterion specification. It never accepts a label
workbook or source database table. Evaluation is intentionally absent; use
``evaluate_r357_strict_predictions.py`` only after this command has finished
and the prediction hashes have been recorded.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

try:
    from .audit_r357_scoring_boundary import audit_scoring_boundary
    from .criterion_spec import (
        DEFAULT_CRITERION_SPEC,
        CriterionSpecValidationError,
        load_criterion_spec,
    )
    from .expert_conflict_engine import (
        KnowledgeBaseValidationError,
        load_knowledge_base,
    )
except ImportError:
    from audit_r357_scoring_boundary import audit_scoring_boundary
    from criterion_spec import (
        DEFAULT_CRITERION_SPEC,
        CriterionSpecValidationError,
        load_criterion_spec,
    )
    from expert_conflict_engine import KnowledgeBaseValidationError, load_knowledge_base


ROOT = Path(__file__).resolve().parents[1]
WORKER = ROOT / "scripts/strict_scoring_worker.py"
R3_BUILDER = ROOT / "scripts/build_r3_paper_only_predictions.py"
R5_R7_BUILDER = ROOT / "scripts/build_r5_r7_source_recovered_predictions.py"
IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
DEFAULT_KNOWLEDGE_BASE = ROOT / "data/knowledge/r357_expert_kb_v1.json"
ALLOWED_CACHE_TOP_LEVEL = {"REFNO", "source", "pages"}
ALLOWED_CACHE_SOURCE_FIELDS = {
    "refno",
    "source_kind",
    "parsed_path",
    "pdf_path",
    "page_count",
    "text_length",
    "cache_sha256",
    "source_quality_flags",
}
ALLOWED_CACHE_PAGE_FIELDS = {
    "page",
    "source",
    "text",
    "text_sha256",
    "section_hints",
}


class StrictInputError(ValueError):
    """Raised when an input does not satisfy the strict scoring contract."""


def validate_knowledge_criterion_alignment(
    knowledge: dict[str, Any],
    criterion: dict[str, Any],
) -> None:
    """Fail closed when the rules do not declare the validated criterion source."""

    provenance = knowledge.get("provenance") or {}
    declared = provenance.get("criterion_spec") or {}
    expected_version = str(criterion.get("spec_version") or "")
    expected_source = str((criterion.get("source_artifact") or {}).get("sha256") or "")
    declared_version = str(declared.get("spec_version") or "")
    declared_source = str(declared.get("source_sha256") or "")
    if declared_version != expected_version or declared_source != expected_source:
        raise StrictInputError(
            "Knowledge base criterion provenance does not match the validated criterion "
            f"spec: declared=({declared_version}, {declared_source}), "
            f"validated=({expected_version}, {expected_source})"
        )


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def validate_identity_manifest(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise StrictInputError(f"Identity manifest does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise StrictInputError("Identity manifest must contain exactly one worksheet")
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=False)
        header_cells = next(rows, ())
        header = tuple(_clean(cell.value) for cell in header_cells)
        if header != IDENTITY_COLUMNS:
            raise StrictInputError(
                "Identity manifest columns must be exactly "
                + ", ".join(IDENTITY_COLUMNS)
                + f" in that order; found {list(header)}"
            )
        identities: list[tuple[str, str, str]] = []
        for row_number, cells in enumerate(rows, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise StrictInputError(
                    f"Identity manifest formulas are forbidden (worksheet row {row_number})"
                )
            values = tuple(_clean(cell.value) for cell in cells[:3])
            if not any(values):
                continue
            if len(values) != 3 or any(not value for value in values):
                raise StrictInputError(
                    f"Identity manifest row {row_number} has an empty identity value"
                )
            identities.append(values)
    finally:
        workbook.close()
    if not identities:
        raise StrictInputError("Identity manifest contains no target rows")
    resultnos = [row[1] for row in identities]
    if len(resultnos) != len(set(resultnos)):
        raise StrictInputError("Identity manifest RESULTNO values must be unique")
    if len(identities) != len(set(identities)):
        raise StrictInputError("Identity manifest identity tuples must be unique")
    return {
        "rows": len(identities),
        "refnos": sorted({row[0] for row in identities}),
        "resultnos": set(resultnos),
    }


def _unexpected(actual: Iterable[str], allowed: set[str]) -> list[str]:
    return sorted(set(actual) - allowed)


def validate_page_cache(path: Path, expected_refnos: Iterable[str]) -> dict[str, Any]:
    if not path.is_file():
        raise StrictInputError(f"Page cache does not exist: {path}")
    seen_refnos: set[str] = set()
    page_count = 0
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise StrictInputError(
                    f"Page cache line {line_number} is invalid JSON: {exc}"
                ) from exc
            if not isinstance(row, dict):
                raise StrictInputError(f"Page cache line {line_number} must be an object")
            extra = _unexpected(row, ALLOWED_CACHE_TOP_LEVEL)
            if extra:
                raise StrictInputError(
                    f"Page cache line {line_number} has forbidden top-level fields: {extra}"
                )
            refno = _clean(row.get("REFNO"))
            if not refno:
                raise StrictInputError(f"Page cache line {line_number} has no REFNO")
            if refno in seen_refnos:
                raise StrictInputError(f"Page cache contains duplicate REFNO {refno}")
            source = row.get("source")
            pages = row.get("pages")
            if not isinstance(source, dict) or not isinstance(pages, list):
                raise StrictInputError(
                    f"Page cache REFNO {refno} requires object source and list pages"
                )
            extra_source = _unexpected(source, ALLOWED_CACHE_SOURCE_FIELDS)
            if extra_source:
                raise StrictInputError(
                    f"Page cache REFNO {refno} has forbidden source fields: {extra_source}"
                )
            source_refno = _clean(source.get("refno"))
            if source_refno and source_refno != refno:
                raise StrictInputError(
                    f"Page cache REFNO {refno} disagrees with source.refno={source_refno}"
                )
            for page_index, page in enumerate(pages, start=1):
                if not isinstance(page, dict):
                    raise StrictInputError(
                        f"Page cache REFNO {refno} page {page_index} must be an object"
                    )
                extra_page = _unexpected(page, ALLOWED_CACHE_PAGE_FIELDS)
                if extra_page:
                    raise StrictInputError(
                        f"Page cache REFNO {refno} page {page_index} has forbidden fields: {extra_page}"
                    )
                if not isinstance(page.get("text", ""), str):
                    raise StrictInputError(
                        f"Page cache REFNO {refno} page {page_index} text must be a string"
                    )
            page_count += len(pages)
            seen_refnos.add(refno)

    expected = set(expected_refnos)
    actual = seen_refnos
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    if missing or extra:
        raise StrictInputError(
            "Page cache REFNO set must exactly match the identity manifest; "
            f"missing={missing[:20]}, extra={extra[:20]}"
        )
    return {"refnos": len(seen_refnos), "pages": page_count}


def _ensure_distinct(paths: Sequence[Path]) -> None:
    resolved = [path.resolve(strict=False) for path in paths]
    if len(resolved) != len(set(resolved)):
        raise StrictInputError(
            "Identity manifest, page cache, knowledge base, criterion spec, "
            "and output directory must be distinct"
        )


def _run_guarded(script: Path, script_args: Sequence[str], inputs: Sequence[Path], output_dir: Path) -> None:
    command = [
        sys.executable,
        str(WORKER),
        "--script",
        str(script),
    ]
    for path in inputs:
        command.extend(["--allow-read", str(path)])
    command.extend(["--write-root", str(output_dir), "--", *script_args])
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    environment["PYTHONHASHSEED"] = "0"
    # openpyxl creates worksheet fragments in the process temp directory.
    # Keep those transient writes inside the same empty, guarded output root.
    environment["TMPDIR"] = str(output_dir)
    environment["TMP"] = str(output_dir)
    environment["TEMP"] = str(output_dir)
    subprocess.run(command, cwd=ROOT, env=environment, check=True)


def _write_json(path: Path, value: dict[str, Any]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity-manifest", type=Path, required=True)
    parser.add_argument("--page-cache", type=Path, required=True)
    parser.add_argument("--knowledge-base", type=Path, default=DEFAULT_KNOWLEDGE_BASE)
    parser.add_argument("--criterion-spec", type=Path, default=DEFAULT_CRITERION_SPEC)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args(argv)


def run_workflow(args: argparse.Namespace) -> dict[str, Any]:
    identity = args.identity_manifest.resolve(strict=False)
    cache = args.page_cache.resolve(strict=False)
    knowledge = args.knowledge_base.resolve(strict=False)
    criterion_spec = args.criterion_spec.resolve(strict=False)
    output_dir = args.output_dir.resolve(strict=False)
    _ensure_distinct([identity, cache, knowledge, criterion_spec, output_dir])

    manifest_info = validate_identity_manifest(identity)
    cache_info = validate_page_cache(cache, manifest_info["refnos"])
    if knowledge.suffix.casefold() != ".json" or not knowledge.is_file():
        raise StrictInputError(f"Frozen knowledge base must be an existing JSON file: {knowledge}")
    try:
        criterion = load_criterion_spec(criterion_spec)
    except CriterionSpecValidationError as exc:
        raise StrictInputError(f"Criterion spec failed validation: {exc}") from exc
    try:
        knowledge_document = load_knowledge_base(knowledge)
    except KnowledgeBaseValidationError as exc:
        raise StrictInputError(f"Knowledge base failed validation: {exc}") from exc
    validate_knowledge_criterion_alignment(knowledge_document, criterion)
    knowledge_audit = audit_scoring_boundary([], [knowledge, criterion_spec])
    if not knowledge_audit["passed"]:
        codes = sorted({item["code"] for item in knowledge_audit["violations"]})
        raise StrictInputError(f"Frozen knowledge inputs failed strict boundary audit: {codes}")

    if output_dir.exists() and any(output_dir.iterdir()):
        raise StrictInputError(
            f"Strict output directory must be new or empty: {output_dir}"
        )
    output_dir.mkdir(parents=True, exist_ok=True)

    r3 = output_dir / "r3_predictions.jsonl"
    r5 = output_dir / "r5_predictions.xlsx"
    r7 = output_dir / "r7_predictions.xlsx"
    r3_summary = output_dir / "r3_summary.json"
    boundary_report = output_dir / "scoring_boundary_audit.json"
    run_manifest = output_dir / "strict_run_manifest.json"
    inputs = [identity, cache, knowledge, criterion_spec]
    worker_inputs = [identity, cache, knowledge]
    input_hashes_before = {str(path): file_sha256(path) for path in inputs}

    r3_stage = output_dir / ".r3_scoring"
    r5_stage = output_dir / ".r5_scoring"
    r7_stage = output_dir / ".r7_scoring"
    for stage in (r3_stage, r5_stage, r7_stage):
        stage.mkdir()
    staged_r3 = r3_stage / r3.name
    staged_r3_summary = r3_stage / r3_summary.name
    staged_r5 = r5_stage / r5.name
    staged_r7 = r7_stage / r7.name

    _run_guarded(
        R3_BUILDER,
        [
            "--targets",
            str(identity),
            "--page-cache",
            str(cache),
            "--knowledge-base",
            str(knowledge),
            "--output",
            str(staged_r3),
            "--summary-json",
            str(staged_r3_summary),
        ],
        worker_inputs,
        r3_stage,
    )
    _run_guarded(
        R5_R7_BUILDER,
        [
            "--indicator",
            "R5",
            "--targets",
            str(identity),
            "--page-cache",
            str(cache),
            "--knowledge-base",
            str(knowledge),
            "--r5-profile",
            "knowledge_strict",
            "--r7-profile",
            "knowledge_strict",
            "--r5-output",
            str(staged_r5),
        ],
        worker_inputs,
        r5_stage,
    )
    _run_guarded(
        R5_R7_BUILDER,
        [
            "--indicator",
            "R7",
            "--targets",
            str(identity),
            "--page-cache",
            str(cache),
            "--knowledge-base",
            str(knowledge),
            "--r5-profile",
            "knowledge_strict",
            "--r7-profile",
            "knowledge_strict",
            "--r7-output",
            str(staged_r7),
        ],
        worker_inputs,
        r7_stage,
    )
    for staged, final in (
        (staged_r3, r3),
        (staged_r3_summary, r3_summary),
        (staged_r5, r5),
        (staged_r7, r7),
    ):
        staged.replace(final)
    for stage in (r3_stage, r5_stage, r7_stage):
        stage.rmdir()

    expected_outputs = [r3, r5, r7]
    missing_outputs = [path for path in expected_outputs if not path.is_file()]
    if missing_outputs:
        raise RuntimeError(
            "Scoring did not produce expected prediction file(s): "
            + ", ".join(str(path) for path in missing_outputs)
        )
    input_hashes_after = {str(path): file_sha256(path) for path in inputs}
    if input_hashes_after != input_hashes_before:
        raise RuntimeError("A strict scoring input changed while predictions were being built")

    boundary = audit_scoring_boundary(expected_outputs, [knowledge, criterion_spec])
    _write_json(boundary_report, boundary)
    if not boundary["passed"]:
        codes = sorted({item["code"] for item in boundary["violations"]})
        raise RuntimeError(f"Predictions failed strict scoring boundary audit: {codes}")

    summary = {
        "schema_version": "r357_strict_knowledge_workflow_v2",
        "status": "PASS",
        "workflow_profile": "knowledge_strict",
        "rows": manifest_info["rows"],
        "inputs": {
            "identity_manifest": {"path": str(identity), "sha256": input_hashes_before[str(identity)]},
            "page_cache": {
                "path": str(cache),
                "sha256": input_hashes_before[str(cache)],
                **cache_info,
            },
            "knowledge_base": {"path": str(knowledge), "sha256": input_hashes_before[str(knowledge)]},
            "criterion_spec": {
                "path": str(criterion_spec),
                "sha256": input_hashes_before[str(criterion_spec)],
                "spec_version": criterion["spec_version"],
                "source_sha256": criterion["source_artifact"]["sha256"],
            },
        },
        "predictions": {
            "R3": {"path": str(r3), "sha256": file_sha256(r3)},
            "R5": {"path": str(r5), "sha256": file_sha256(r5)},
            "R7": {"path": str(r7), "sha256": file_sha256(r7)},
        },
        "boundary_audit": {"path": str(boundary_report), "status": boundary["status"]},
        "scoring_contract": {
            "allowed_runtime_data_inputs": [
                "identity manifest: REFNO/RESULTNO/ROCKNAME only",
                "original-paper page cache",
                "frozen expert knowledge base",
                "frozen normative criterion specification",
            ],
            "criterion_spec_required": True,
            "criterion_spec_validator": "scripts.criterion_spec.load_criterion_spec",
            "filesystem_read_guard": "enabled",
            "indicator_process_isolation": "separate empty write roots",
            "network_access": "blocked",
            "child_processes_inside_scoring": "blocked",
            "labels_accepted_by_this_command": False,
            "source_tables_accepted_by_this_command": False,
            "previous_predictions_accepted_by_this_command": False,
        },
        "evaluation": {
            "status": "NOT_RUN",
            "reason": "Evaluation is a separate post-prediction command and cannot feed values back into this scoring run.",
            "command": "scripts/evaluate_r357_strict_predictions.py",
        },
    }
    _write_json(run_manifest, summary)
    return summary


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        summary = run_workflow(args)
    except StrictInputError as exc:
        print(f"Strict input error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
