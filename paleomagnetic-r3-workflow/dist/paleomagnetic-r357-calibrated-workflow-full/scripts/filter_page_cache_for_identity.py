#!/usr/bin/env python3
"""Filter a page-cache JSONL to the exact REFNO set in a strict manifest."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")


class CacheFilterError(ValueError):
    """Raised when filtering would produce an incomplete or ambiguous cache."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def manifest_refnos(path: Path) -> set[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise CacheFilterError("strict identity manifest must contain one worksheet")
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = tuple(str(value or "").strip() for value in next(rows, ()))
        if header != IDENTITY_COLUMNS:
            raise CacheFilterError(
                "strict identity columns must be exactly REFNO, RESULTNO, ROCKNAME"
            )
        refnos: set[str] = set()
        row_count = 0
        for excel_row, values in enumerate(rows, start=2):
            cleaned = tuple(str(value or "").strip() for value in values[:3])
            if len(cleaned) != 3 or any(not value for value in cleaned):
                raise CacheFilterError(f"empty identity value at Excel row {excel_row}")
            refnos.add(cleaned[0])
            row_count += 1
    finally:
        workbook.close()
    if not row_count:
        raise CacheFilterError("strict identity manifest contains no data rows")
    return refnos


def filter_cache(
    identity_manifest: Path,
    input_cache: Path,
    output_cache: Path,
) -> dict[str, Any]:
    identity_manifest = identity_manifest.resolve()
    input_cache = input_cache.resolve()
    output_cache = output_cache.resolve(strict=False)
    if len({identity_manifest, input_cache, output_cache}) != 3:
        raise CacheFilterError("identity, input cache, and output cache must be distinct")
    if not identity_manifest.is_file() or not input_cache.is_file():
        raise CacheFilterError("identity manifest and input cache must exist")

    required = manifest_refnos(identity_manifest)
    found: set[str] = set()
    input_refnos: set[str] = set()
    selected_lines: list[bytes] = []
    with input_cache.open("rb") as handle:
        for line_number, raw_line in enumerate(handle, start=1):
            if not raw_line.strip():
                continue
            try:
                row = json.loads(raw_line)
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise CacheFilterError(
                    f"invalid page-cache JSON at line {line_number}: {exc}"
                ) from exc
            if not isinstance(row, dict):
                raise CacheFilterError(f"page-cache line {line_number} must be an object")
            refno = str(row.get("REFNO") or "").strip()
            if not refno:
                raise CacheFilterError(f"page-cache line {line_number} has no REFNO")
            if refno in input_refnos:
                raise CacheFilterError(f"duplicate page-cache REFNO: {refno}")
            input_refnos.add(refno)
            if refno in required:
                found.add(refno)
                selected_lines.append(raw_line if raw_line.endswith(b"\n") else raw_line + b"\n")

    missing = sorted(required - found)
    if missing:
        raise CacheFilterError(
            "input page cache is missing manifest REFNO values: " + ", ".join(missing[:20])
        )
    output_cache.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_cache.with_suffix(output_cache.suffix + ".tmp")
    with temporary.open("wb") as handle:
        handle.writelines(selected_lines)
    temporary.replace(output_cache)
    return {
        "schema_version": "strict_page_cache_filter_audit_v1",
        "status": "PASS",
        "identity_manifest": {
            "path": str(identity_manifest),
            "sha256": file_sha256(identity_manifest),
            "unique_refnos": len(required),
        },
        "input_cache": {
            "path": str(input_cache),
            "sha256": file_sha256(input_cache),
            "unique_refnos": len(input_refnos),
        },
        "output_cache": {
            "path": str(output_cache),
            "sha256": file_sha256(output_cache),
            "unique_refnos": len(found),
        },
        "ignored_extra_refnos": len(input_refnos - required),
        "raw_cache_rows_preserved": True,
        "labels_read": False,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity-manifest", type=Path, required=True)
    parser.add_argument("--input-cache", type=Path, required=True)
    parser.add_argument("--output-cache", type=Path, required=True)
    parser.add_argument("--audit-json", type=Path, required=True)
    args = parser.parse_args()
    try:
        audit = filter_cache(args.identity_manifest, args.input_cache, args.output_cache)
    except CacheFilterError as exc:
        print(f"Page-cache filter error: {exc}")
        return 2
    args.audit_json.parent.mkdir(parents=True, exist_ok=True)
    args.audit_json.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
