#!/usr/bin/env python3
"""Calculate table-derived R1/R2/R4/R6 reliability indicators.

These indicators are deterministic joins over PMAGRESULT and ROCKUNIT, keyed by
RESULTNO and ROCKUNITNO. Benchmark labels in the target workbook are used only
for optional post-calculation comparison.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_ROCKUNIT = ROOT / "data/source_metadata/ROCKUNIT.xlsx"
DEFAULT_PMAGRESULT = ROOT / "data/source_metadata/PMAGRESULT.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "outputs"
INDICATORS = ("R1", "R2", "R4", "R6")
SOURCE_COLUMNS = [
    "HIMAGAGE",
    "LOMAGAGE",
    "HIGHAGE",
    "LOWAGE",
    "B",
    "N",
    "KD",
    "ED95",
    "DEMAGCODE",
    "TESTS",
]


def as_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def validate_timestamp(value: str) -> str:
    if not re.fullmatch(r"\d{8}", value):
        raise argparse.ArgumentTypeError("timestamp must use MMDDHHMM, for example 07010632")
    return value


def clean_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.fillna("").copy()
    for column in frame.columns:
        frame[column] = frame[column].map(as_text)
    return frame


def numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series.map(as_text).replace("", pd.NA), errors="coerce")


def test_tokens(value: Any) -> set[str]:
    text = as_text(value)
    if not text:
        return set()
    return {token.strip() for token in text.split(",") if token.strip()}


def load_unique_table(path: Path, key: str, required: set[str]) -> pd.DataFrame:
    frame = clean_frame(pd.read_excel(path, dtype=str))
    missing = sorted(required - set(frame.columns))
    if missing:
        raise SystemExit(f"{path} is missing required column(s): {', '.join(missing)}")
    if frame[key].duplicated().any():
        duplicated = ", ".join(frame.loc[frame[key].duplicated(), key].head(10))
        raise SystemExit(f"{path} contains duplicate {key} value(s): {duplicated}")
    return frame


def load_targets(path: Path) -> pd.DataFrame:
    targets = clean_frame(pd.read_excel(path, dtype=str))
    required = {"RESULTNO", "REFNO", "ROCKNAME"}
    missing = sorted(required - set(targets.columns))
    if missing:
        raise SystemExit(f"{path} is missing required column(s): {', '.join(missing)}")
    if targets["RESULTNO"].duplicated().any():
        duplicated = ", ".join(targets.loc[targets["RESULTNO"].duplicated(), "RESULTNO"].head(10))
        raise SystemExit(f"{path} contains duplicate RESULTNO value(s): {duplicated}")
    return targets


def build_source_frame(
    targets: pd.DataFrame | None,
    pmag: pd.DataFrame,
    rockunit: pd.DataFrame,
) -> pd.DataFrame:
    pmag_columns = [
        "RESULTNO",
        "ROCKUNITNO",
        "HIMAGAGE",
        "LOMAGAGE",
        "B",
        "N",
        "KD",
        "ED95",
        "DEMAGCODE",
        "TESTS",
    ]
    rockunit_columns = ["ROCKUNITNO", "REFNO", "ROCKNAME", "HIGHAGE", "LOWAGE"]

    if targets is None:
        base = pmag[["RESULTNO"]].copy()
    else:
        base = targets[["RESULTNO", "REFNO", "ROCKNAME"]].rename(columns={"ROCKNAME": "TARGET"})

    frame = base.merge(pmag[pmag_columns], on="RESULTNO", how="left", validate="one_to_one")
    frame = frame.merge(
        rockunit[rockunit_columns].rename(
            columns={
                "REFNO": "ROCKUNIT_REFNO",
                "ROCKNAME": "ROCKUNIT_ROCKNAME",
            }
        ),
        on="ROCKUNITNO",
        how="left",
        validate="many_to_one",
    )

    if targets is None:
        frame["REFNO"] = frame["ROCKUNIT_REFNO"]
        frame["TARGET"] = frame["ROCKUNIT_ROCKNAME"]

    frame["REFNO"] = frame["REFNO"].fillna("")
    frame["TARGET"] = frame["TARGET"].fillna(frame["ROCKUNIT_ROCKNAME"]).fillna("")
    return frame


def calculate_indicators(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    himagage = numeric(result["HIMAGAGE"])
    lomagage = numeric(result["LOMAGAGE"])
    highage = numeric(result["HIGHAGE"])
    lowage = numeric(result["LOWAGE"])
    b_value = numeric(result["B"])
    n_value = numeric(result["N"])
    kd_value = numeric(result["KD"])
    ed95_value = numeric(result["ED95"])
    demagcode = numeric(result["DEMAGCODE"])
    tokens = result["TESTS"].map(test_tokens)

    result["R1"] = ((himagage - lomagage).sub(highage - lowage).abs() <= 15).fillna(False).astype(int)
    result["R2"] = (
        (b_value >= 8)
        & (n_value >= 25)
        & (kd_value >= 10)
        & (kd_value <= 70)
        & (ed95_value <= 16)
        & (demagcode >= 3)
    ).fillna(False).astype(int)
    result["R4"] = tokens.map(lambda items: any("+" in item and item != "R+" for item in items)).astype(int)
    result["R6"] = tokens.map(lambda items: bool(items.intersection({"R+", "Ra", "Rb", "Rc"}))).astype(int)
    return result


def add_standard_comparison(frame: pd.DataFrame, targets: pd.DataFrame | None) -> pd.DataFrame:
    if targets is None:
        return frame
    available = [indicator for indicator in INDICATORS if indicator in targets.columns]
    if not available:
        return frame
    standard = targets[["RESULTNO", *available]].rename(
        columns={indicator: f"STANDARD_{indicator}" for indicator in available}
    )
    compared = frame.merge(standard, on="RESULTNO", how="left", validate="one_to_one")
    for indicator in available:
        standard_col = f"STANDARD_{indicator}"
        standard_score = numeric(compared[standard_col])
        compared[standard_col] = standard_score
        compared[f"MISMATCH_{indicator}"] = standard_score.isin([0, 1]) & ~compared[indicator].eq(standard_score)
    return compared


def summary(frame: pd.DataFrame) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "rows": int(len(frame)),
        "unique_resultnos": int(frame["RESULTNO"].nunique()),
        "score_counts": {
            indicator: {str(key): int(value) for key, value in frame[indicator].value_counts().sort_index().items()}
            for indicator in INDICATORS
        },
        "missing_source_rows": {
            "pmagresult": int(frame["ROCKUNITNO"].map(as_text).eq("").sum()),
            "rockunit": int((frame["HIGHAGE"].map(as_text).eq("") | frame["LOWAGE"].map(as_text).eq("")).sum()),
        },
        "comparison": {},
    }
    for indicator in INDICATORS:
        standard_col = f"STANDARD_{indicator}"
        if standard_col not in frame.columns:
            continue
        standard_score = numeric(frame[standard_col])
        labelled = standard_score.isin([0, 1])
        total = int(labelled.sum())
        correct = int((labelled & frame[indicator].eq(standard_score)).sum())
        payload["comparison"][indicator] = {
            "labelled_rows": total,
            "correct": correct,
            "accuracy": correct / total if total else None,
            "mismatches": int(frame.get(f"MISMATCH_{indicator}", pd.Series(False, index=frame.index)).sum()),
        }
    return payload


def ordered_columns(frame: pd.DataFrame) -> list[str]:
    columns = [
        "REFNO",
        "RESULTNO",
        "TARGET",
        "ROCKUNITNO",
        *INDICATORS,
        *SOURCE_COLUMNS,
    ]
    for indicator in INDICATORS:
        standard_col = f"STANDARD_{indicator}"
        mismatch_col = f"MISMATCH_{indicator}"
        if standard_col in frame.columns:
            columns.append(standard_col)
        if mismatch_col in frame.columns:
            columns.append(mismatch_col)
    audit_columns = ["ROCKUNIT_REFNO", "ROCKUNIT_ROCKNAME"]
    columns.extend(column for column in audit_columns if column in frame.columns)
    return [column for column in columns if column in frame.columns]


def write_workbook(frame: pd.DataFrame, summary_payload: dict[str, Any], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame[ordered_columns(frame)].to_excel(writer, sheet_name="R1246 Results", index=False)
        pd.DataFrame([summary_payload]).to_excel(writer, sheet_name="Summary", index=False)

    workbook = load_workbook(output)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    red_font = Font(color="FFFF0000")
    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, column_cells in enumerate(worksheet.columns, start=1):
            header = as_text(column_cells[0].value)
            width = 16
            if header == "TARGET" or header.endswith("ROCKNAME"):
                width = 34
            elif header in SOURCE_COLUMNS:
                width = 12
            elif header.startswith("MISMATCH_") or header.startswith("STANDARD_"):
                width = 14
            worksheet.column_dimensions[get_column_letter(index)].width = width

    result_sheet = workbook["R1246 Results"]
    mismatch_columns = {
        cell.column
        for cell in result_sheet[1]
        if as_text(cell.value).startswith("MISMATCH_")
    }
    if mismatch_columns:
        for row in result_sheet.iter_rows(min_row=2):
            if any(row[column - 1].value is True for column in mismatch_columns):
                for cell in row:
                    cell.font = red_font
    workbook.save(output)


def resolve_output_path(output: Path | None, output_dir: Path, timestamp: str | None) -> Path:
    if output:
        return output
    stamp = timestamp or datetime.now().strftime("%m%d%H%M")
    return output_dir / f"r1246_{stamp}.xlsx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--pmagresult", type=Path, default=DEFAULT_PMAGRESULT)
    parser.add_argument("--rockunit", type=Path, default=DEFAULT_ROCKUNIT)
    parser.add_argument("--all-pmag-results", action="store_true", help="Calculate every PMAGRESULT row instead of only target rows.")
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--summary-json", type=Path, default=None)
    parser.add_argument("--timestamp", type=validate_timestamp, default=None)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    required_paths = {"pmagresult": args.pmagresult, "rockunit": args.rockunit}
    if not args.all_pmag_results:
        required_paths["targets"] = args.targets
    missing_paths = [f"{name}={path}" for name, path in required_paths.items() if not path.exists()]
    if missing_paths:
        raise SystemExit("Required input(s) missing: " + "; ".join(missing_paths))

    pmag = load_unique_table(
        args.pmagresult,
        "RESULTNO",
        {"RESULTNO", "ROCKUNITNO", "HIMAGAGE", "LOMAGAGE", "B", "N", "KD", "ED95", "DEMAGCODE", "TESTS"},
    )
    rockunit = load_unique_table(
        args.rockunit,
        "ROCKUNITNO",
        {"ROCKUNITNO", "REFNO", "ROCKNAME", "HIGHAGE", "LOWAGE"},
    )
    targets = None if args.all_pmag_results else load_targets(args.targets)

    frame = calculate_indicators(build_source_frame(targets, pmag, rockunit))
    frame = add_standard_comparison(frame, targets)
    summary_payload = summary(frame)

    output = resolve_output_path(args.output, args.output_dir, args.timestamp)
    write_workbook(frame, summary_payload, output)
    if args.summary_json:
        args.summary_json.parent.mkdir(parents=True, exist_ok=True)
        args.summary_json.write_text(json.dumps(summary_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Saved: {output}")
    if args.summary_json:
        print(f"Summary: {args.summary_json}")
    print(json.dumps(summary_payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
