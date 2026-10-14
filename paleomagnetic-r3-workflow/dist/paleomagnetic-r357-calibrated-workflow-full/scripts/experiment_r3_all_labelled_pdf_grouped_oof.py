#!/usr/bin/env python3
"""Evaluate every PDF-backed labelled R3 target with REFNO-grouped OOF."""

from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Mapping, Sequence

import numpy as np
import sklearn
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import experiment_r3_exact300_auxiliary_training as auxiliary  # noqa: E402
import experiment_r3_exact300_nested_router as nested  # noqa: E402
import experiment_r3_exact300_paper_consistency as paper_v1  # noqa: E402
import experiment_r3_pdf_only_strategies as base  # noqa: E402
import experiment_r3_small_screen_curation_style as curation  # noqa: E402


IDENTITY = ROOT / "outputs/r357_strict_eval_07111758/inputs/identity_manifest.xlsx"
IDENTITY_SHA256 = "5bdb7476a5b7d61494939fc5cd7bc8fa3195ef7b032b29306f146892862b970f"
PAGE_CACHE = ROOT / "tmp/r3_all_labelled_pdf_page_cache_v1/page_cache.jsonl"
PAGE_CACHE_SHA256 = "388bd9d405119a2ee19aa08d9b0194c3902bdcb0b52b9d5e8925f7ff2405a076"
PAGE_CACHE_AUDIT = ROOT / "tmp/r3_all_labelled_pdf_page_cache_v1/audit.json"
PAGE_CACHE_AUDIT_SHA256 = "a12ce6a2b930eb4feef49338ecf04c9c5b4e75364f9981d401b704ea963612c2"
OUTPUT_DIR = ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1"
EXPECTED_LABELLED_ROWS = 3173
EXPECTED_LABELLED_PAPERS = 1627
EXPECTED_PDF_ROWS = 2751
EXPECTED_PDF_PAPERS = 1478
EXPECTED_SKIPPED_ROWS = 422
N_SPLITS = 5
PRIMARY_STRATEGY = "mean4"
FOLD_SEED = 35_720_260_716


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("all-labelled identity must contain one worksheet")
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = tuple(clean(value) for value in next(rows, ()))
        if header != ("REFNO", "RESULTNO", "ROCKNAME"):
            raise ValueError("all-labelled identity must contain exactly three columns")
        records = [
            {"REFNO": clean(row[0]), "RESULTNO": clean(row[1]), "ROCKNAME": clean(row[2])}
            for row in rows
            if any(clean(value) for value in row)
        ]
    finally:
        workbook.close()
    if any(not all(record.values()) for record in records):
        raise ValueError("all-labelled identity contains an empty value")
    if len(records) != EXPECTED_LABELLED_ROWS or len(
        {record["REFNO"] for record in records}
    ) != EXPECTED_LABELLED_PAPERS:
        raise ValueError("all-labelled identity membership changed")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise ValueError("all-labelled identity contains duplicate RESULTNO")
    return records


def pdf_only_records(
    records: Sequence[dict[str, str]], pdf_dir: Path
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    pdf_refnos = base.local_pdf_refnos(pdf_dir)
    selected = [record for record in records if record["REFNO"] in pdf_refnos]
    skipped = [record for record in records if record["REFNO"] not in pdf_refnos]
    coverage = {
        "labelled_rows": len(records),
        "labelled_paper_groups": len({record["REFNO"] for record in records}),
        "pdf_rows": len(selected),
        "pdf_paper_groups": len({record["REFNO"] for record in selected}),
        "skipped_without_pdf_rows": len(skipped),
        "skipped_without_pdf_paper_groups": len(
            {record["REFNO"] for record in skipped}
        ),
    }
    if (
        coverage["pdf_rows"] != EXPECTED_PDF_ROWS
        or coverage["pdf_paper_groups"] != EXPECTED_PDF_PAPERS
        or coverage["skipped_without_pdf_rows"] != EXPECTED_SKIPPED_ROWS
    ):
        raise ValueError("all-labelled PDF coverage changed")
    return selected, coverage


def load_labels(path: Path, resultnos: set[str]) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = tuple(clean(value) for value in next(rows, ()))
        columns = {name: header.index(name) for name in ("RESULTNO", "R3")}
        labels: dict[str, int] = {}
        for row in rows:
            resultno = clean(row[columns["RESULTNO"]])
            if resultno not in resultnos:
                continue
            try:
                label = int(float(row[columns["R3"]]))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"non-binary R3 label for RESULTNO {resultno}") from exc
            if label not in {0, 1} or resultno in labels:
                raise ValueError(f"invalid or duplicate R3 label for RESULTNO {resultno}")
            labels[resultno] = label
    finally:
        workbook.close()
    if set(labels) != resultnos:
        raise ValueError("all-labelled R3 labels do not align with PDF identities")
    return labels


def record_sha256(records: Sequence[Mapping[str, str]]) -> str:
    return auxiliary.sha256_records(records)


def write_array_archive(path: Path, arrays: Mapping[str, np.ndarray]) -> None:
    payload = {
        name: np.asarray(value, dtype=np.float32)
        for name, value in sorted(arrays.items())
    }
    if any(
        values.shape != (EXPECTED_PDF_ROWS,) or not np.isfinite(values).all()
        for values in payload.values()
    ):
        raise ValueError("all-labelled OOF arrays must contain 2,751 finite values")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("wb") as handle:
        np.savez_compressed(handle, **payload)
    temporary.replace(path)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=PAGE_CACHE)
    parser.add_argument("--page-cache-audit", type=Path, default=PAGE_CACHE_AUDIT)
    parser.add_argument("--pdf-dir", type=Path, default=base.DEFAULT_PDF_DIR)
    parser.add_argument("--targets", type=Path, default=base.DEFAULT_TARGETS)
    parser.add_argument("--taxonomy", type=Path, default=base.DEFAULT_TAXONOMY)
    parser.add_argument("--positive-kb", type=Path, default=base.DEFAULT_POSITIVE_KB)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args(argv)
    paths = {
        "identity": args.identity.resolve(),
        "page_cache": args.page_cache.resolve(),
        "page_cache_audit": args.page_cache_audit.resolve(),
        "pdf_dir": args.pdf_dir.resolve(),
        "targets": args.targets.resolve(),
        "taxonomy": args.taxonomy.resolve(),
        "positive_kb": args.positive_kb.resolve(),
    }
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    expected_hashes = {
        "identity": IDENTITY_SHA256,
        "page_cache": PAGE_CACHE_SHA256,
        "page_cache_audit": PAGE_CACHE_AUDIT_SHA256,
        "targets": auxiliary.TARGETS_SHA256,
        "taxonomy": auxiliary.TAXONOMY_SHA256,
        "positive_kb": auxiliary.POSITIVE_KB_SHA256,
    }
    for name, digest in expected_hashes.items():
        if base.file_sha256(paths[name]) != digest:
            raise ValueError(f"frozen input hash mismatch: {name}")
    runtime = {
        "python": sys.version.split()[0],
        "numpy": np.__version__,
        "sklearn": sklearn.__version__,
    }
    if runtime != paper_v1.CANONICAL_RUNTIME:
        raise ValueError(
            f"canonical runtime required: expected={paper_v1.CANONICAL_RUNTIME}, "
            f"observed={runtime}"
        )

    labelled_records = load_identity(paths["identity"])
    records, coverage = pdf_only_records(labelled_records, paths["pdf_dir"])
    alignment_sha256 = record_sha256(records)
    refnos = {record["REFNO"] for record in records}
    paper_text = base.load_page_text(paths["page_cache"], refnos)
    pages = curation.load_pages(paths["page_cache"], refnos)
    coverage.update(
        {
            "parsed_pdf_rows": len(records),
            "parsed_pdf_paper_groups": len(refnos),
            "skipped_unparseable_pdf_rows": 0,
            "skipped_unparseable_pdf_paper_groups": 0,
            "evaluated_rows": len(records),
            "evaluated_paper_groups": len(refnos),
        }
    )
    concepts, knowledge_audit = base.load_knowledge_patterns(
        paths["taxonomy"], paths["positive_kb"]
    )
    base_features = base.build_features(records, paper_text, concepts)
    curation_features = curation.build_features(
        records, pages, concepts, paths["pdf_dir"]
    )
    engineered = np.asarray(base_features["engineered_rows"], dtype=np.float32)
    curation_all = np.asarray(curation_features["all"], dtype=np.float32)
    combined = np.concatenate((engineered, curation_all), axis=1).astype(
        np.float32, copy=False
    )
    if combined.shape != (EXPECTED_PDF_ROWS, 351) or not np.isfinite(combined).all():
        raise ValueError(f"unexpected all-labelled combined matrix: {combined.shape}")
    feature_path = output_dir / "combined_features.npz"
    temporary_feature = feature_path.with_name(feature_path.name + ".tmp")
    with temporary_feature.open("wb") as handle:
        np.savez_compressed(handle, combined=combined)
    temporary_feature.replace(feature_path)
    feature_manifest = {
        "schema_version": "r3_all_labelled_pdf_combined_features_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "rows": len(records),
        "features": combined.shape[1],
        "alignment_sha256": alignment_sha256,
        "matrix_sha256": auxiliary.sha256_array(combined),
        "contains_labels": False,
        "contains_identifiers": False,
        "inputs": {
            name: {"path": str(path), "sha256": base.file_sha256(path)}
            for name, path in paths.items()
            if name not in {"pdf_dir", "targets"}
        },
        "knowledge": knowledge_audit,
        "output": {
            "path": str(feature_path),
            "sha256": base.file_sha256(feature_path),
        },
    }
    feature_manifest_path = output_dir / "combined_features.manifest.json"
    base.atomic_json(feature_manifest_path, feature_manifest)

    all_indices = np.arange(len(records), dtype=np.int64)
    groups = [record["REFNO"] for record in records]
    folds = nested.deterministic_group_folds(
        all_indices, groups, n_splits=N_SPLITS, seed=FOLD_SEED
    )
    fold_ids = np.zeros(len(records), dtype=np.int8)
    fold_audit: list[dict[str, Any]] = []
    for fold, (train, test) in enumerate(folds, start=1):
        train_refnos = {records[int(index)]["REFNO"] for index in train}
        test_refnos = {records[int(index)]["REFNO"] for index in test}
        if train_refnos & test_refnos:
            raise AssertionError("all-labelled fold splits a REFNO")
        fold_ids[test] = fold
        fold_audit.append(
            {
                "fold": fold,
                "train_rows": len(train),
                "train_papers": len(train_refnos),
                "test_rows": len(test),
                "test_papers": len(test_refnos),
                "train_test_refno_overlap": 0,
            }
        )
    if Counter(int(value) for value in fold_ids) != Counter(
        {fold: len(test) for fold, (_, test) in enumerate(folds, start=1)}
    ):
        raise AssertionError("all-labelled fold assignment is incomplete")
    fold_manifest = {
        "schema_version": "r3_all_labelled_pdf_refno_folds_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "seed": FOLD_SEED,
        "method": "deterministic row-balanced complete-REFNO assignment",
        "alignment_sha256": alignment_sha256,
        "fold_counts": {
            str(fold): int((fold_ids == fold).sum())
            for fold in range(1, N_SPLITS + 1)
        },
        "fold_id_sha256": hashlib.sha256(fold_ids.tobytes()).hexdigest(),
        "labels_read": False,
        "rows": [
            {
                "row_index": index,
                "REFNO": record["REFNO"],
                "RESULTNO": record["RESULTNO"],
                "fold": int(fold_ids[index]),
            }
            for index, record in enumerate(records)
        ],
    }
    fold_manifest_path = output_dir / "refno_fold_manifest.json"
    base.atomic_json(fold_manifest_path, fold_manifest)

    experts = (
        nested.ExpertSpec("expert_hgb_combined", "combined", "hgb10"),
        nested.ExpertSpec("expert_rf_combined", "combined", "rf8"),
        nested.ExpertSpec("expert_extra_combined", "combined", "extra8"),
        nested.ExpertSpec(
            "expert_logreg_combined", "combined", "logreg001", standardized=True
        ),
    )
    freeze = {
        "schema_version": "r3_all_labelled_pdf_mean4_candidate_freeze_v1",
        "status": "FROZEN_BEFORE_LABEL_ACCESS",
        "primary_strategy": PRIMARY_STRATEGY,
        "experts": [asdict(expert) for expert in experts],
        "fusion": "uniform arithmetic mean of four expert probabilities",
        "threshold": 0.5,
        "coverage": coverage,
        "fold_manifest": {
            "path": str(fold_manifest_path),
            "sha256": base.file_sha256(fold_manifest_path),
        },
        "feature_manifest": {
            "path": str(feature_manifest_path),
            "sha256": base.file_sha256(feature_manifest_path),
        },
        "runtime_boundary": {
            "labels_read_at_freeze": False,
            "pmagresult_or_row_rockmag_used_as_feature": False,
            "prior_oof_or_evaluation_read": False,
            "refno_used_only_for_group_isolation": True,
            "resultno_used_only_for_alignment": True,
        },
        "canonical_runtime": runtime,
        "script": {
            "path": str(Path(__file__).resolve()),
            "sha256": base.file_sha256(Path(__file__).resolve()),
        },
    }
    freeze_path = output_dir / "candidate_freeze.json"
    base.atomic_json(freeze_path, freeze)

    labels_by_resultno = load_labels(
        paths["targets"], {record["RESULTNO"] for record in records}
    )
    labels = np.asarray(
        [labels_by_resultno[record["RESULTNO"]] for record in records],
        dtype=np.int8,
    )
    canonical_targets = [
        base.canonical_target_identity(record["ROCKNAME"]) for record in records
    ]
    probabilities = {
        expert.name: np.full(len(records), np.nan, dtype=np.float64)
        for expert in experts
    }
    for fold, (train, test) in enumerate(folds, start=1):
        weights = base.inverse_target_multiplicity_weights(
            train, records, canonical_targets
        )
        for expert in experts:
            fitted = nested.fit_expert(expert, combined, train, labels, weights)
            probabilities[expert.name][test] = fitted.predict(combined[test])
        print(
            f"fold {fold}: train_rows={len(train)}, test_rows={len(test)}, "
            f"test_papers={fold_audit[fold - 1]['test_papers']}",
            flush=True,
        )
    if any(not np.isfinite(values).all() for values in probabilities.values()):
        raise AssertionError("all-labelled expert OOF is incomplete")
    probabilities[PRIMARY_STRATEGY] = np.column_stack(
        [probabilities[expert.name] for expert in experts]
    ).mean(axis=1)
    oof_path = output_dir / "oof_probabilities.npz"
    write_array_archive(oof_path, probabilities)
    receipt = {
        "schema_version": "r3_all_labelled_pdf_oof_prediction_receipt_v1",
        "status": "OOF_PREDICTIONS_COMPLETE_BEFORE_AGGREGATE_SCORING",
        "primary_strategy": PRIMARY_STRATEGY,
        "candidate_freeze_sha256": base.file_sha256(freeze_path),
        "fold_manifest_sha256": base.file_sha256(fold_manifest_path),
        "oof": {"path": str(oof_path), "sha256": base.file_sha256(oof_path)},
        "each_row_predicted_by_model_excluding_its_refno": True,
        "rows": [
            {
                "row_index": index,
                **record,
                "fold": int(fold_ids[index]),
                "probability": float(probabilities[PRIMARY_STRATEGY][index]),
                "prediction": int(probabilities[PRIMARY_STRATEGY][index] >= 0.5),
            }
            for index, record in enumerate(records)
        ],
    }
    receipt_path = output_dir / "prediction_receipt.json"
    base.atomic_json(receipt_path, receipt)

    results: dict[str, Any] = {}
    for name, values in probabilities.items():
        results[name] = {
            "aggregate": base.metric_record(labels, values),
            "folds": [
                {
                    "fold": fold,
                    **base.metric_record(labels[fold_ids == fold], values[fold_ids == fold]),
                }
                for fold in range(1, N_SPLITS + 1)
            ],
        }
    original_training = np.asarray(
        [base.is_training_refno(record["REFNO"]) for record in records], dtype=bool
    )
    primary = probabilities[PRIMARY_STRATEGY]
    subgroup_metrics = {
        "original_training_partition": base.metric_record(
            labels[original_training], primary[original_training]
        ),
        "original_holdout_partition": base.metric_record(
            labels[~original_training], primary[~original_training]
        ),
    }
    prediction = (primary >= 0.5).astype(np.int8)
    errors = [
        {
            "row_index": int(index),
            **records[int(index)],
            "fold": int(fold_ids[int(index)]),
            "label": int(labels[int(index)]),
            "prediction": int(prediction[int(index)]),
            "probability": float(primary[int(index)]),
            "error_type": "FP" if prediction[int(index)] == 1 else "FN",
        }
        for index in np.flatnonzero(prediction != labels)
    ]
    report = {
        "schema_version": "r3_all_labelled_pdf_mean4_grouped_oof_v1",
        "status": "complete",
        "scope": coverage,
        "method": {
            "primary_strategy": PRIMARY_STRATEGY,
            "outer_split": "five deterministic REFNO-grouped folds",
            "coverage": "every PDF-backed labelled row predicted exactly once",
        },
        "compliance_audit": {
            "feature_and_fold_artifacts_frozen_before_label_access": True,
            "all_test_refno_rows_excluded_from_each_fold_fit": True,
            "each_row_has_exactly_one_oof_prediction": True,
            "prior_oof_or_evaluation_read": False,
            "pmagresult_or_row_rockmag_used_as_feature": False,
            "rows_without_pdf_excluded_from_denominator": True,
        },
        "inputs": {
            name: {"path": str(path), "sha256": base.file_sha256(path)}
            for name, path in paths.items()
            if name != "pdf_dir"
        },
        "feature_manifest": {
            "path": str(feature_manifest_path),
            "sha256": base.file_sha256(feature_manifest_path),
        },
        "fold_manifest": {
            "path": str(fold_manifest_path),
            "sha256": base.file_sha256(fold_manifest_path),
            "audit": fold_audit,
        },
        "candidate_freeze": {
            "path": str(freeze_path),
            "sha256": base.file_sha256(freeze_path),
        },
        "prediction_receipt": {
            "path": str(receipt_path),
            "sha256": base.file_sha256(receipt_path),
        },
        "models": results,
        "primary_result": {
            "strategy": PRIMARY_STRATEGY,
            **results[PRIMARY_STRATEGY]["aggregate"],
        },
        "subgroup_metrics": subgroup_metrics,
        "label_distribution": {
            "negative": int((labels == 0).sum()),
            "positive": int((labels == 1).sum()),
        },
        "errors": errors,
        "artifacts": {
            "oof_probabilities": {
                "path": str(oof_path),
                "sha256": base.file_sha256(oof_path),
            }
        },
    }
    report_path = output_dir / "report.json"
    base.atomic_json(report_path, report)
    print(
        json.dumps(
            {
                "report": str(report_path),
                "scope": report["scope"],
                "primary_result": report["primary_result"],
                "subgroup_metrics": report["subgroup_metrics"],
                "fold_accuracies": [
                    item["accuracy"]
                    for item in results[PRIMARY_STRATEGY]["folds"]
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
