#!/usr/bin/env python3
"""Create consolidated, portable R1-R7 result files.

The detailed HTML remains the visual review surface. This module adds a
compact Excel workbook for ordinary use and a lossless JSONL file for audit or
downstream programs. R3/R5/R7 evidence JSON is kept in JSONL because Excel
cells have a 32,767 character limit.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from create_r1234567_total_review_html import (
    INDICATORS,
    combine,
    parse_json_cell,
    row_payload,
)


TEXT_INDICATORS = ("R3", "R5", "R7")
SOURCE_COLUMNS = (
    "HIMAGAGE",
    "LOMAGAGE",
    "HIGHAGE",
    "LOWAGE",
    "B",
    "N",
    "KD",
    "ED95",
    "DEMAGCODE",
    "TESTS",
)


def as_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def payload_outcome(value: Any) -> str:
    payload = value if isinstance(value, dict) else parse_json_cell(value)
    return as_text(payload.get("outcome") or "DECIDED").upper()


def payload_summary(value: Any) -> str:
    payload = value if isinstance(value, dict) else parse_json_cell(value)
    return as_text(payload.get("final_decision_reasoning") or payload.get("summary"))


def validate_complete(frame: pd.DataFrame) -> None:
    if frame.empty:
        raise SystemExit("Cannot create R1-R7 deliverables from an empty result set.")
    if frame["RESULTNO"].duplicated().any():
        examples = ", ".join(frame.loc[frame["RESULTNO"].duplicated(), "RESULTNO"].head(10))
        raise SystemExit(f"Combined results contain duplicate RESULTNO values: {examples}")
    missing: dict[str, int] = {}
    for indicator in INDICATORS:
        count = int(frame[indicator].map(lambda value: value not in {0, 1}).sum())
        if count:
            missing[indicator] = count
    if missing:
        raise SystemExit(f"Combined results are incomplete; non-binary/missing scores: {missing}")


def result_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame[
        ["REFNO", "RESULTNO", "ROCKUNITNO", "TARGET", *INDICATORS, "R_SUM"]
    ].copy()
    for indicator in TEXT_INDICATORS:
        result[f"{indicator}_OUTCOME"] = frame[f"{indicator}_PAYLOAD"].map(payload_outcome)
        result[f"{indicator}_SUMMARY"] = frame[f"{indicator}_PAYLOAD"].map(payload_summary)
    result["REVIEW_REQUIRED"] = result[
        [f"{indicator}_OUTCOME" for indicator in TEXT_INDICATORS]
    ].apply(lambda row: any(value not in {"DECIDED", "SUPPORTED"} for value in row), axis=1)
    return result


def evidence_frame(frame: pd.DataFrame) -> pd.DataFrame:
    return frame[
        ["REFNO", "RESULTNO", "ROCKUNITNO", "TARGET", *SOURCE_COLUMNS]
    ].copy()


def summary_frame(summary: dict[str, Any], files: dict[str, str]) -> pd.DataFrame:
    rows: list[dict[str, str]] = [
        {"section": "run", "key": "created_at", "value": str(summary["created_at"])},
        {"section": "run", "key": "rows", "value": str(summary["rows"])},
        {"section": "run", "key": "unique_refnos", "value": str(summary["unique_refnos"])},
        {"section": "run", "key": "unique_resultnos", "value": str(summary["unique_resultnos"])},
    ]
    for name, value in summary.get("inputs", {}).items():
        rows.append({"section": "input", "key": name, "value": str(value)})
    for name, value in files.items():
        rows.append({"section": "output", "key": name, "value": value})
    for indicator, counts in summary.get("score_counts", {}).items():
        for score, count in counts.items():
            rows.append(
                {
                    "section": "score_count",
                    "key": f"{indicator}={score}",
                    "value": str(count),
                }
            )
    return pd.DataFrame(rows, columns=["section", "key", "value"])


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for index, cells in enumerate(worksheet.columns, start=1):
            header = as_text(cells[0].value)
            if header == "TARGET":
                width = 38
            elif header.endswith("_SUMMARY"):
                width = 72
            elif header == "value":
                width = 90
            elif header in SOURCE_COLUMNS:
                width = 15
            else:
                width = 16
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_jsonl(frame: pd.DataFrame, path: Path) -> None:
    readable_rows = row_payload(frame)
    by_resultno = {row["RESULTNO"]: row for row in readable_rows}
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix == ".gz":
        handle = gzip.open(path, "wt", encoding="utf-8")
    else:
        handle = path.open("w", encoding="utf-8")
    with handle:
        for _, row in frame.iterrows():
            resultno = as_text(row["RESULTNO"])
            record = dict(by_resultno[resultno])
            record["raw_evidence"] = {
                indicator: row[f"{indicator}_PAYLOAD"] for indicator in TEXT_INDICATORS
            }
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")


def create_deliverables(
    args: argparse.Namespace,
    output_xlsx: Path,
    output_jsonl: Path,
    manifest_json: Path,
) -> dict[str, Any]:
    frame, summary = combine(args)
    validate_complete(frame)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    write_jsonl(frame, output_jsonl)
    files = {"workbook": str(output_xlsx), "jsonl": str(output_jsonl)}
    html_path = getattr(args, "html", None)
    if html_path is not None:
        if not html_path.exists():
            raise SystemExit(f"Detailed review HTML does not exist: {html_path}")
        files["html"] = str(html_path)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        result_frame(frame).to_excel(writer, sheet_name="R1-R7 Results", index=False)
        evidence_frame(frame).to_excel(writer, sheet_name="Table Evidence", index=False)
        summary_frame(summary, files).to_excel(writer, sheet_name="Run Summary", index=False)
    style_workbook(output_xlsx)

    manifest = {
        "schema_version": 1,
        "artifact": "paleomagnetic_r1_r7_results",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "rows": int(len(frame)),
        "inputs": summary["inputs"],
        "score_counts": summary["score_counts"],
        "outputs": {
            "workbook": {"path": str(output_xlsx), "sha256": sha256(output_xlsx)},
            "jsonl": {"path": str(output_jsonl), "sha256": sha256(output_jsonl)},
        },
        "notes": [
            "R1/R2/R4/R6 are deterministic calculations from PMAGRESULT and ROCKUNIT.",
            "R3/R5/R7 are generated by the selected paper/knowledge profiles.",
            "JSONL is the lossless audit format; Excel contains compact summaries.",
        ],
    }
    if html_path is not None:
        manifest["outputs"]["html"] = {
            "path": str(html_path),
            "sha256": sha256(html_path),
        }
    manifest_json.parent.mkdir(parents=True, exist_ok=True)
    manifest_json.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--targets", type=Path, required=True)
    parser.add_argument("--pmagresult", type=Path, required=True)
    parser.add_argument("--rockunit", type=Path, required=True)
    parser.add_argument("--r3", type=Path, required=True)
    parser.add_argument("--r5", type=Path, required=True)
    parser.add_argument("--r7", type=Path, required=True)
    parser.add_argument("--html", type=Path, default=None, help="Detailed R1-R7 review HTML to hash.")
    parser.add_argument("--output-dir", type=Path, default=Path("outputs"))
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument(
        "--output-jsonl",
        type=Path,
        required=True,
        help="Lossless audit JSONL; use a .gz suffix for recommended gzip compression.",
    )
    parser.add_argument("--manifest-json", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    manifest = create_deliverables(
        args,
        args.output_xlsx,
        args.output_jsonl,
        args.manifest_json,
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
