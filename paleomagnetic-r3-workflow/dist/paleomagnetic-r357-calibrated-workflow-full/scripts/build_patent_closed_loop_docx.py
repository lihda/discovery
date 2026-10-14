#!/usr/bin/env python3
"""Build the baseline-evidence closed-loop patent method planning document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

import build_patent_method_docx as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/patent_method_flow_baseline_evidence_closed_loop_20260717.md"
OUTPUT = ROOT / "docs/古地磁可靠性指标判定方法_六步闭环流程详细稿_20260717.docx"
FIGURE = ROOT / "docs/figures/patent-baseline-evidence-closed-loop.png"


def draw_flow_figure(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1700, 2100
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    title_font = ImageFont.truetype(font_path, 50)
    step_font = ImageFont.truetype(font_path, 34)
    small_font = ImageFont.truetype(font_path, 25)
    note_font = ImageFont.truetype(font_path, 23)

    draw.text(
        (width // 2, 62),
        "内部专家基线与源感知证据冲突消解六步闭环",
        font=title_font,
        fill=(31, 45, 61),
        anchor="ma",
    )

    steps = [
        ("S1", "构建古地磁专用知识库", "规则 / 向量 / 案例 / PDF / 元数据"),
        ("S2", "生成内部专家基线", "初值 / 触发规则 / 置信信息 / 记录哈希"),
        ("S3", "检索外部候选证据", "目标 / 分量 / 年龄 / tests / 多路召回"),
        ("S4", "执行源感知知识整合", "绑定度 / 来源质量 / 四类证据 / 关系图"),
        ("S5", "执行保守冲突决策", "保留 / 纠正 / 复核 / 来源拒判"),
        ("S6", "专家反馈与规则沉淀", "案例化 / 离线验证 / 审批 / 新版本"),
    ]
    fills = [
        (232, 241, 231),
        (231, 240, 247),
        (243, 234, 213),
        (220, 234, 247),
        (245, 232, 204),
        (232, 236, 239),
    ]
    left, right = 235, width - 310
    box_h, gap = 205, 72
    y = 185
    for index, ((code, label, detail), fill) in enumerate(zip(steps, fills)):
        draw.rounded_rectangle(
            (left, y, right, y + box_h),
            radius=18,
            fill=fill,
            outline=(80, 105, 128),
            width=4,
        )
        draw.text((left + 70, y + 48), code, font=title_font, fill=(46, 116, 181), anchor="la")
        draw.text((left + 245, y + 46), label, font=step_font, fill=(31, 45, 61), anchor="la")
        draw.text((left + 245, y + 128), detail, font=small_font, fill=(82, 92, 102), anchor="la")
        if index < len(steps) - 1:
            x = (left + right) // 2
            draw.line((x, y + box_h + 8, x, y + box_h + gap - 13), fill=(80, 105, 128), width=6)
            draw.polygon(
                [(x - 16, y + box_h + gap - 30), (x + 16, y + box_h + gap - 30), (x, y + box_h + gap)],
                fill=(80, 105, 128),
            )
        y += box_h + gap

    loop_x = right + 95
    first_y = 185 + box_h // 2
    last_y = 185 + 5 * (box_h + gap) + box_h // 2
    draw.line((right, last_y, loop_x, last_y), fill=(124, 90, 0), width=6)
    draw.line((loop_x, last_y, loop_x, first_y), fill=(124, 90, 0), width=6)
    draw.line((loop_x, first_y, right + 12, first_y), fill=(124, 90, 0), width=6)
    draw.polygon(
        [(right + 30, first_y - 16), (right + 30, first_y + 16), (right, first_y)],
        fill=(124, 90, 0),
    )
    draw.text(
        (loop_x + 30, (first_y + last_y) // 2),
        "仅经验证并发布的新版本\n进入下一轮 S1",
        font=note_font,
        fill=(124, 90, 0),
        anchor="lm",
        spacing=10,
    )

    note_y = height - 185
    draw.rounded_rectangle(
        (235, note_y, width - 235, note_y + 88),
        radius=12,
        fill=(249, 250, 252),
        outline=(180, 188, 198),
        width=3,
    )
    draw.text(
        (width // 2, note_y + 44),
        "当前运行使用冻结知识版本；单条专家反馈不即时改变本次结果",
        font=small_font,
        fill=(124, 90, 0),
        anchor="mm",
    )
    image.save(path, dpi=(200, 200))


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("专利技术交底书 | 方法流程规划")
    base.set_run_font(run, size=11, color=base.BLUE, bold=True)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("古地磁可靠性指标判定方法")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("内部专家基线、源感知证据整合与保守冲突决策六步闭环")

    metadata = [
        ("文件用途", "技术交底、方法流程定稿、附图规划和权利要求展开"),
        ("保护主线", "知识库 -> baseline -> 外部证据 -> 源感知整合 -> 保守决策 -> 专家反馈"),
        ("版本日期", "2026-07-17"),
        ("文档状态", "方法流程详细稿"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{label}：")
        base.set_run_font(run, size=10.5, color=base.INK, bold=True)
        run = p.add_run(value)
        base.set_run_font(run, size=10.5, color=base.INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    base.shade_paragraph(p, base.LIGHT)
    run = p.add_run("核心约束：")
    base.set_run_font(run, size=11, color=base.BLUE_DARK, bold=True)
    run = p.add_run("不因单一关键词或低绑定证据直接翻转 baseline；专家反馈只通过新版本进入下一轮运行。")
    base.set_run_font(run, size=11, color=base.INK)
    doc.add_page_break()


def add_lead_callout(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(12)
    base.shade_paragraph(p, base.LIGHT)
    run = p.add_run("方法定位：")
    base.set_run_font(run, size=11, color=base.BLUE_DARK, bold=True)
    run = p.add_run(
        "本方案保护的是从不可变内部基线到可追溯外部证据、再到保守纠正和版本化专家反馈的连续数据处理闭环。"
    )
    base.set_run_font(run, size=11, color=base.INK)


def add_figure(doc: Document) -> None:
    draw_flow_figure(FIGURE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(FIGURE), width=Inches(5.65))
    cap = doc.add_paragraph(style="Figure Caption")
    cap.add_run("图1  内部专家基线与源感知证据冲突消解六步闭环")


def audit_document(doc: Document) -> None:
    assert doc.sections[0].page_width == Inches(8.5)
    assert doc.sections[0].left_margin == Inches(1)
    assert doc.styles["Normal"].font.size == Pt(11)
    assert doc.styles["Heading 1"].font.size == Pt(16)
    assert doc.styles["Heading 2"].font.size == Pt(13)
    assert doc.styles["Heading 3"].font.size == Pt(12)
    assert len(doc.tables) == 0
    assert any(paragraph.style.name == "List Number" for paragraph in doc.paragraphs)
    assert any("S68" in paragraph.text for paragraph in doc.paragraphs)
    assert any("专家反馈" in paragraph.text for paragraph in doc.paragraphs)


def main() -> None:
    base.SOURCE = SOURCE
    base.OUTPUT = OUTPUT
    base.FIGURE = FIGURE
    base.add_cover = add_cover
    base.add_lead_callout = add_lead_callout
    base.add_figure = add_figure
    base.audit_document = audit_document

    doc = Document()
    base.configure_sections(doc)
    base.configure_styles(doc)
    base.add_running_header_footer(doc.sections[0])
    doc.core_properties.title = "古地磁可靠性指标判定方法：六步闭环流程详细稿"
    doc.core_properties.subject = "专利技术交底方法流程"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "古地磁,R3,R5,R7,baseline,证据冲突,专家反馈,专利"

    add_cover(doc)
    base.parse_source(doc)
    audit_document(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
