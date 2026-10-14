import unittest
from unittest import mock
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import Workbook

from scripts.build_r3_paper_only_predictions import (
    build_rows,
    judge_r3,
    load_targets,
    parse_args,
)
from scripts.expert_conflict_engine import ExpertConflictEngine
from scripts.paper_text_retrieval import retrieve_descriptions


def hit(page=1, *, relevance=0.9, local_target=True, subtype="synthetic"):
    if subtype in {
        "explicit_carrier_bridge",
        "mineral_remanence_bridge",
        "carrier_significance",
    }:
        quote = "Magnetite is the carrier of remanence."
    elif subtype == "origin_bridge":
        quote = "Magnetite supports a primary remanence interpretation."
    elif subtype == "named_carrier":
        quote = "Magnetite was identified."
    elif subtype in {"rock_magnetic_experiment", "demagnetization"}:
        quote = "IRM acquisition curves were measured."
    else:
        quote = "Synthetic source evidence."
    return {
        "page": page,
        "target_relevance": relevance,
        "local_target": local_target,
        "target_anchor": "target_same_sentence" if local_target else "same_target_page",
        "subtype": subtype,
        "quote": quote,
    }


def retrieval(
    *,
    source_available=True,
    anchor_status="exact",
    anchor_strength="strong",
    page_span=0,
    connected=True,
    verdicts=None,
):
    chain_hits = {
        "method": hit(page=1, subtype="rock_magnetic_experiment"),
        "carrier": hit(page=1, subtype="named_carrier"),
        "bridge": hit(page=1 + page_span, subtype="mineral_remanence_bridge"),
    }
    return {
        "source_available": source_available,
        "multi_target": True,
        "target_anchor": {
            "status": anchor_status,
            "strength": anchor_strength,
        },
        "categories": {
            "method": [chain_hits["method"]],
            "carrier": [chain_hits["carrier"]],
            "bridge": [chain_hits["bridge"]],
            "verdict": list(verdicts or []),
        },
        "best_chain": {
            "complete": True,
            "connected": connected,
            "page_span": page_span,
            "hits": chain_hits,
        },
    }


class R3PaperOnlyJudgementTests(unittest.TestCase):
    def test_cli_requires_explicit_identity_manifest(self):
        with self.assertRaises(SystemExit) as raised:
            parse_args(["--page-cache", "cache.jsonl", "--output", "predictions.jsonl"])

        self.assertEqual(raised.exception.code, 2)

    def test_load_targets_accepts_exact_label_free_identity_schema(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "identity.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["REFNO", "RESULTNO", "ROCKNAME"])
            worksheet.append([1141, 1, "Alpha Formation"])
            workbook.save(path)

            frame = load_targets(path)

        self.assertEqual(
            frame.to_dict(orient="records"),
            [{"REFNO": "1141", "RESULTNO": "1", "ROCKNAME": "Alpha Formation"}],
        )

    def test_load_targets_rejects_workbook_with_label_column(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "targets.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["REFNO", "RESULTNO", "ROCKNAME", "R3"])
            worksheet.append([1141, 1, "Alpha Formation", 1])
            workbook.save(path)

            with self.assertRaisesRegex(
                SystemExit,
                "exactly REFNO, RESULTNO, ROCKNAME",
            ):
                load_targets(path)

    def test_load_targets_rejects_formulas_and_multiple_worksheets(self):
        with TemporaryDirectory() as directory:
            formula_path = Path(directory) / "formula.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["REFNO", "RESULTNO", "ROCKNAME"])
            worksheet.append([1141, 1, "=\"Alpha Formation\""])
            workbook.save(formula_path)

            with self.assertRaisesRegex(SystemExit, "contains a formula"):
                load_targets(formula_path)

            multi_sheet_path = Path(directory) / "multi-sheet.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["REFNO", "RESULTNO", "ROCKNAME"])
            worksheet.append([1141, 1, "Alpha Formation"])
            workbook.create_sheet("labels")
            workbook.save(multi_sheet_path)

            with self.assertRaisesRegex(SystemExit, "exactly one worksheet"):
                load_targets(multi_sheet_path)

    def test_multi_target_score_requires_target_anchor(self):
        source = retrieval(anchor_status="missing", anchor_strength="missing")

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertFalse(result["conditions"]["target"])
        self.assertEqual(result["score"], 0)
        self.assertEqual(result["outcome"], "NEGATIVE")
        self.assertTrue(result["manual_review"])

    def test_multi_target_strong_article_anchor_cannot_rescue_nonlocal_chain(self):
        source = retrieval(anchor_status="exact", anchor_strength="strong")
        for category in ("method", "carrier", "bridge"):
            source["best_chain"]["hits"][category]["local_target"] = False
            source["best_chain"]["hits"][category]["target_anchor"] = "same_target_page"

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertEqual(result["outcome"], "NEGATIVE")
        self.assertFalse(result["conditions"]["method"])
        self.assertFalse(result["conditions"]["carrier"])
        self.assertFalse(result["conditions"]["bridge"])

    def test_build_rows_does_not_transfer_beta_rockmag_chain_to_alpha(self):
        targets = pd.DataFrame(
            [
                {"REFNO": "shared", "RESULTNO": "1", "ROCKNAME": "Alpha Formation"},
                {"REFNO": "shared", "RESULTNO": "2", "ROCKNAME": "Beta Formation"},
            ]
        )
        cache = {
            "shared": {
                "REFNO": "shared",
                "source": {"source_kind": "synthetic", "source_quality_flags": []},
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "The Alpha Formation consists of shale.\n\n"
                            "The Beta Formation was tested using IRM acquisition curves; "
                            "magnetite is the main carrier of characteristic remanence."
                        ),
                    }
                ],
            }
        }

        rows = build_rows(targets, cache, Path("synthetic-cache.jsonl"))
        by_result = {row["RESULTNO"]: row for row in rows}

        self.assertEqual(by_result["1"]["R3"], 0)
        self.assertNotEqual(by_result["1"]["R3_RESULT_JSON"]["outcome"], "POSITIVE")
        self.assertEqual(by_result["2"]["R3"], 1)
        self.assertEqual(by_result["2"]["R3_RESULT_JSON"]["outcome"], "POSITIVE")

    def test_connected_chain_accepts_only_same_or_adjacent_pages(self):
        adjacent = judge_r3(
            retrieval(page_span=1, connected=True),
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )
        separated = judge_r3(
            retrieval(page_span=2, connected=False),
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(adjacent["score"], 1)
        self.assertTrue(adjacent["conditions"]["connected"])
        self.assertEqual(separated["score"], 0)
        self.assertFalse(separated["conditions"]["connected"])

    def test_multi_target_promoted_chain_cannot_cross_structured_scopes(self):
        source = retrieval()
        method = source["best_chain"]["hits"]["method"]
        carrier = source["best_chain"]["hits"]["carrier"]
        bridge = source["best_chain"]["hits"]["bridge"]
        method.update(
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:1",
        )
        carrier.update(
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:2",
        )
        bridge.update(
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:2",
        )

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertFalse(result["conditions"]["scope_compatible"])
        self.assertIn("scope_compatible", result["reason"])

    def test_lower_scored_same_scope_complete_chain_is_selected(self):
        row = {
            "pages": [
                {
                    "page": 1,
                    "source": "synthetic",
                    "section_hints": ["Rock magnetism"],
                    "text": (
                        "Alpha Formation samples were selected. IRM acquisition curves "
                        "identified magnetite as the carrier of remanence.\n\n"
                        + "Geological context was documented in detail. " * 8
                        + "\n\nAlpha Formation contains hematite.\n\n"
                        "Beta Formation was thermally demagnetized only."
                    ),
                }
            ]
        }
        source = retrieve_descriptions(
            row,
            "Alpha Formation",
            True,
            "R3",
            competing_targets=["Beta Formation"],
        )

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 1)
        self.assertEqual(
            {
                evidence["scope_id"]
                for evidence in result["hits"].values()
                if evidence
            },
            {"page:1:paragraph:1"},
        )

    def test_carrier_identifying_bridge_can_form_chain_without_incompatible_carrier(self):
        source = retrieval()
        method = source["best_chain"]["hits"]["method"]
        carrier = source["best_chain"]["hits"]["carrier"]
        bridge = source["best_chain"]["hits"]["bridge"]
        method.update(
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:1",
        )
        bridge.update(
            subtype="explicit_carrier_bridge",
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:1",
        )
        carrier.update(
            target_relevance=0.99,
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:2",
        )

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 1)
        self.assertEqual(result["hits"]["carrier"], {})
        self.assertEqual(result["hits"]["bridge"]["scope_id"], "page:1:paragraph:1")

    def test_non_carrier_bridge_still_requires_compatible_carrier(self):
        source = retrieval()
        method = source["best_chain"]["hits"]["method"]
        bridge = source["best_chain"]["hits"]["bridge"]
        method.update(
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:1",
        )
        bridge.update(
            subtype="origin_bridge",
            target_anchor="target_same_paragraph",
            scope_id="page:1:paragraph:1",
        )
        source["best_chain"]["hits"]["carrier"] = {}
        source["categories"]["carrier"] = []

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertFalse(result["conditions"]["carrier"])

    def test_unresolved_or_negated_carrier_claims_do_not_satisfy_r3(self):
        claims = (
            "The magnetic carrier of remanence remains unknown.",
            "The carrier of magnetization could not be identified.",
            "It is unclear which mineral carries the remanence.",
            "Magnetite is not the carrier of remanence.",
            "The results exclude magnetite as the carrier of remanence.",
            "There is no evidence that hematite carries the remanence.",
            "The experiment was intended to identify the remanence carrier, but results were inconclusive.",
            "The remanence carrier has yet to be identified.",
            "The carrier was not established.",
            "The identity of the remanence carrier is ambiguous.",
            "No specific remanence carrier was identified.",
            "The data do not constrain the remanence carrier.",
            "We cannot determine which mineral carries the remanence.",
            "Magnetite was not established as the carrier.",
            "The data do not support magnetite as the remanence carrier.",
            "Magnetite cannot be confirmed as the carrier.",
            "Whether magnetite carries the remanence is ambiguous.",
        )
        for claim in claims:
            with self.subTest(claim=claim):
                source = retrieve_descriptions(
                    {
                        "pages": [
                            {
                                "page": 1,
                                "source": "synthetic",
                                "section_hints": ["Rock magnetism"],
                                "text": (
                                    "Alpha Formation was tested using IRM acquisition curves; "
                                    + claim
                                    + "\n\nBeta Formation was thermally demagnetized only."
                                ),
                            }
                        ]
                    },
                    "Alpha Formation",
                    True,
                    "R3",
                    competing_targets=["Beta Formation"],
                )
                result = judge_r3(
                    source,
                    multi_target=True,
                    identity_ambiguous=False,
                    source_flags=[],
                )
                self.assertEqual(result["score"], 0)

        positive = retrieve_descriptions(
            {
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "Alpha Formation IRM acquisition curves show that magnetite is "
                            "the carrier of remanence, although its acquisition age remains "
                            "uncertain.\n\n"
                            "Beta Formation was thermally demagnetized only."
                        ),
                    }
                ]
            },
            "Alpha Formation",
            True,
            "R3",
            competing_targets=["Beta Formation"],
        )
        positive_result = judge_r3(
            positive,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )
        self.assertEqual(positive_result["score"], 1)

    def test_cross_paragraph_regex_match_cannot_bypass_scope_gate(self):
        source = retrieve_descriptions(
            {
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "Alpha Formation was tested using IRM acquisition curves. "
                            "The remanence is interpreted as primary.\n\n"
                            "Alpha Formation contains magnetite.\n\n"
                            "Alpha Formation contains hematite.\n\n"
                            "Alpha Formation contains maghemite.\n\n"
                            "Alpha Formation contains pyrrhotite.\n\n"
                            "Beta Formation was thermally demagnetized only."
                        ),
                    }
                ]
            },
            "Alpha Formation",
            True,
            "R3",
            competing_targets=["Beta Formation"],
        )

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertTrue(
            all(
                hit["scope_id"] == "page:1:paragraph:1"
                and not any(
                    mineral in hit["matched_text"].casefold()
                    for mineral in ("magnetite", "hematite", "maghemite", "pyrrhotite")
                )
                for hit in source["categories"]["bridge"]
            )
        )

    def test_method_alone_is_not_sufficient_without_carrier_significance(self):
        source = retrieval()
        source["best_chain"] = {
            "complete": False,
            "connected": False,
            "page_span": None,
            "hits": {},
        }
        source["categories"]["carrier"] = []
        source["categories"]["bridge"] = []

        result = judge_r3(
            source,
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertTrue(result["conditions"]["independent_method"])
        self.assertFalse(result["conditions"]["carrier"])
        self.assertFalse(result["conditions"]["bridge"])

    def test_method_and_explicit_carrier_significance_are_sufficient(self):
        source = retrieval()
        source["best_chain"] = {
            "complete": False,
            "connected": False,
            "page_span": None,
            "hits": {},
        }
        source["categories"]["carrier"] = []
        source["categories"]["bridge"] = [
            hit(subtype="mineral_remanence_bridge")
        ]

        result = judge_r3(
            source,
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 1)
        self.assertTrue(result["conditions"]["carrier"])
        self.assertTrue(result["conditions"]["bridge"])

    def test_demagnetization_alone_is_not_independent_carrier_investigation(self):
        source = retrieval()
        source["best_chain"] = {
            "complete": False,
            "connected": False,
            "page_span": None,
            "hits": {},
        }
        source["categories"] = {
            "method": [hit(subtype="demagnetization")],
            "carrier": [],
            "bridge": [],
            "verdict": [],
        }

        result = judge_r3(
            source,
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertFalse(result["conditions"]["independent_method"])

    def test_demagnetization_cannot_complete_chain_with_carrier_and_bridge(self):
        source = retrieval()
        source["best_chain"]["hits"]["method"] = hit(
            subtype="demagnetization"
        )
        source["categories"]["method"] = [
            source["best_chain"]["hits"]["method"]
        ]

        result = judge_r3(
            source,
            multi_target=False,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["score"], 0)
        self.assertFalse(result["conditions"]["method"])
        self.assertFalse(result["conditions"]["independent_method"])

    def test_embedded_positive_verdict_is_ignored(self):
        verdict = hit(subtype="direct_r3_positive")
        source = retrieval(verdicts=[verdict])
        source["best_chain"] = {
            "complete": False,
            "connected": False,
            "page_span": None,
            "hits": {},
        }
        source["categories"]["method"] = []
        source["categories"]["carrier"] = []
        source["categories"]["bridge"] = []

        result = judge_r3(
            source,
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["direct_verdict"]["status"], "absent")
        self.assertEqual(result["score"], 0)
        self.assertEqual(result["outcome"], "NEGATIVE")
        self.assertEqual(len(result["direct_verdict"]["ignored_existing_scores"]), 1)

    def test_embedded_negative_verdict_does_not_override_evidence(self):
        verdict = hit(subtype="direct_r3_negative")

        result = judge_r3(
            retrieval(verdicts=[verdict]),
            multi_target=True,
            identity_ambiguous=False,
            source_flags=[],
        )

        self.assertEqual(result["direct_verdict"]["status"], "absent")
        self.assertEqual(result["score"], 1)
        self.assertEqual(result["outcome"], "POSITIVE")

    def test_unavailable_source_is_unknown_with_binary_zero(self):
        targets = pd.DataFrame(
            [{"REFNO": "missing-ref", "RESULTNO": "1", "ROCKNAME": "Missing Formation"}]
        )

        rows = build_rows(targets, {}, Path("missing-page-cache.jsonl"))
        row = rows[0]
        payload = row["R3_RESULT_JSON"]

        self.assertEqual(row["R3"], 0)
        self.assertEqual(payload["R3"], 0)
        self.assertEqual(payload["outcome"], "UNKNOWN")
        self.assertTrue(payload["manual_review"])
        self.assertIn("UNKNOWN", payload["final_decision_reasoning"])

    def test_duplicate_refno_target_group_shares_analysis_but_abstains(self):
        targets = pd.DataFrame(
            [
                {"REFNO": "same-ref", "RESULTNO": "1", "ROCKNAME": "Alpha Formation"},
                {"REFNO": "same-ref", "RESULTNO": "2", "ROCKNAME": "Alpha Formation"},
            ]
        )
        cache = {
            "same-ref": {
                "REFNO": "same-ref",
                "source": {"source_kind": "synthetic", "source_quality_flags": []},
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "IRM acquisition and hysteresis experiments were performed on "
                            "the Alpha Formation and identified magnetite as the principal "
                            "carrier of characteristic remanence."
                        ),
                    }
                ],
            }
        }

        engine = ExpertConflictEngine.from_path()
        with mock.patch.object(engine, "evaluate", wraps=engine.evaluate) as evaluate:
            rows = build_rows(
                targets,
                cache,
                Path("synthetic-cache.jsonl"),
                knowledge_engine=engine,
            )

        self.assertEqual(evaluate.call_count, 1)
        self.assertEqual({row["R3"] for row in rows}, {0})
        self.assertEqual(
            {row["R3_RESULT_JSON"]["outcome"] for row in rows},
            {"UNKNOWN"},
        )
        self.assertTrue(
            all(
                not row["R3_RESULT_JSON"]["source_audit"]["multi_target_refno"]
                and row["R3_RESULT_JSON"]["source_audit"]["identity_ambiguous"]
                and row["R3_RESULT_JSON"]["manual_review"]
                for row in rows
            )
        )

    def test_normalized_same_target_names_share_analysis_but_abstain(self):
        targets = pd.DataFrame(
            [
                {"REFNO": "same-ref", "RESULTNO": "1", "ROCKNAME": "Sao-Tome Formation"},
                {"REFNO": "same-ref", "RESULTNO": "2", "ROCKNAME": "Sao Tome Formation"},
            ]
        )
        cache = {
            "same-ref": {
                "REFNO": "same-ref",
                "source": {"source_kind": "synthetic", "source_quality_flags": []},
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "IRM acquisition on the Sao Tome Formation identifies magnetite "
                            "as the principal carrier of characteristic remanence."
                        ),
                    }
                ],
            }
        }
        engine = ExpertConflictEngine.from_path()
        with mock.patch.object(engine, "evaluate", wraps=engine.evaluate) as evaluate:
            rows = build_rows(
                targets,
                cache,
                Path("synthetic-cache.jsonl"),
                knowledge_engine=engine,
            )

        self.assertEqual(evaluate.call_count, 1)
        self.assertEqual({row["R3"] for row in rows}, {0})
        self.assertEqual(
            {row["R3_RESULT_JSON"]["outcome"] for row in rows},
            {"UNKNOWN"},
        )
        self.assertTrue(
            all(
                not row["R3_RESULT_JSON"]["source_audit"]["multi_target_refno"]
                and row["R3_RESULT_JSON"]["source_audit"]["identity_ambiguous"]
                and row["R3_RESULT_JSON"]["manual_review"]
                for row in rows
            )
        )

    def test_distinct_targets_remain_multi_target(self):
        targets = pd.DataFrame(
            [
                {"REFNO": "same-ref", "RESULTNO": "1", "ROCKNAME": "Alpha Formation"},
                {"REFNO": "same-ref", "RESULTNO": "2", "ROCKNAME": "Beta Formation"},
            ]
        )
        cache = {
            "same-ref": {
                "REFNO": "same-ref",
                "source": {"source_kind": "synthetic", "source_quality_flags": []},
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic",
                        "section_hints": ["Rock magnetism"],
                        "text": (
                            "IRM acquisition identifies magnetite as the principal carrier "
                            "of characteristic remanence in the Alpha Formation. "
                            "IRM acquisition identifies hematite as the principal carrier "
                            "of characteristic remanence in the Beta Formation."
                        ),
                    }
                ],
            }
        }
        engine = ExpertConflictEngine.from_path()
        with mock.patch.object(engine, "evaluate", wraps=engine.evaluate) as evaluate:
            rows = build_rows(
                targets,
                cache,
                Path("synthetic-cache.jsonl"),
                knowledge_engine=engine,
            )

        self.assertEqual(evaluate.call_count, 2)
        self.assertTrue(
            all(
                row["R3_RESULT_JSON"]["source_audit"]["multi_target_refno"]
                and not row["R3_RESULT_JSON"]["source_audit"]["identity_ambiguous"]
                for row in rows
            )
        )
        scope_bindings = [
            row["R3_RESULT_JSON"]["original_source_descriptions"]["scope_binding"]
            for row in rows
        ]
        self.assertTrue(
            all(
                binding["status"] == "enabled_with_competing_target_exclusion"
                for binding in scope_bindings
            )
        )
        self.assertEqual(
            {tuple(binding["competing_targets"]) for binding in scope_bindings},
            {("alpha formation",), ("beta formation",)},
        )


if __name__ == "__main__":
    unittest.main()
