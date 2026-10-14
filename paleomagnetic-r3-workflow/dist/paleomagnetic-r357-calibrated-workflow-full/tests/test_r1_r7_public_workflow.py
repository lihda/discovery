from __future__ import annotations

import json
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(SCRIPTS))

from create_r1_r7_deliverables import create_deliverables
from run_r1_r7_workflow import (
    ConfigurationError,
    core_command,
    input_preflight,
    load_config,
    validate_r3_stable_runtime,
)


class PublicWorkflowInputTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.pdf_dir = self.root / "pdfs"
        self.parsed_dir = self.root / "parsed"
        self.pdf_dir.mkdir()
        self.parsed_dir.mkdir()
        (self.parsed_dir / "10.txt").write_text("paper text", encoding="utf-8")

        self.targets = self.root / "targets.xlsx"
        self.pmag = self.root / "PMAGRESULT.xlsx"
        self.rockunit = self.root / "ROCKUNIT.xlsx"
        pd.DataFrame(
            [{"RESULTNO": "1", "ROCKUNITNO": "2", "REFNO": "10", "ROCKNAME": "Unit A"}]
        ).to_excel(self.targets, index=False)
        pd.DataFrame(
            [
                {
                    "RESULTNO": "1",
                    "ROCKUNITNO": "2",
                    "HIMAGAGE": "110",
                    "LOMAGAGE": "90",
                    "B": "9",
                    "N": "30",
                    "KD": "20",
                    "ED95": "8",
                    "DEMAGCODE": "3",
                    "TESTS": "F+,R+",
                }
            ]
        ).to_excel(self.pmag, index=False)
        pd.DataFrame(
            [
                {
                    "ROCKUNITNO": "2",
                    "REFNO": "10",
                    "ROCKNAME": "Unit A",
                    "HIGHAGE": "110",
                    "LOWAGE": "90",
                }
            ]
        ).to_excel(self.rockunit, index=False)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def write_config(self, adapter: str = "legacy_excel_v1") -> Path:
        path = self.root / "workflow.json"
        path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "input_adapter": adapter,
                    "inputs": {
                        "targets": str(self.targets),
                        "pmagresult": str(self.pmag),
                        "rockunit": str(self.rockunit),
                        "parsed_dir": str(self.parsed_dir),
                        "pdf_dir": str(self.pdf_dir),
                    },
                    "profiles": {
                        "r3_source_mode": "paper_only",
                        "r5_profile": "paper_only",
                        "r7_profile": "source_text",
                    },
                    "runtime": {"r3_workers": 2},
                    "output": {"root": str(self.root / "runs")},
                }
            ),
            encoding="utf-8",
        )
        return path

    def test_current_adapter_preflight_and_command(self) -> None:
        config = load_config(self.write_config())
        report = input_preflight(config)
        self.assertEqual(report["status"], "ok")
        self.assertEqual(report["targets"]["rows"], 1)
        self.assertEqual(report["paper_sources"]["covered_target_refnos"], 1)

        command = core_command(config, self.root / "run", "07221430")
        self.assertIn("--targets", command)
        self.assertIn(str(self.targets), command)
        self.assertIn("paper_only", command)

    def test_unknown_adapter_fails_closed(self) -> None:
        with self.assertRaises(ConfigurationError):
            load_config(self.write_config(adapter="future_csv_v2"))

    def test_identity_mismatch_is_rejected(self) -> None:
        targets = pd.read_excel(self.targets, dtype=str)
        targets.loc[0, "REFNO"] = "wrong"
        targets.to_excel(self.targets, index=False)
        config = load_config(self.write_config())
        with self.assertRaises(ConfigurationError):
            input_preflight(config)

    def test_current_default_uses_verified_stable_model(self) -> None:
        config = load_config(ROOT / "config/workflow.current.json")
        self.assertEqual(config["profiles"]["r3_source_mode"], "supervised_stable")
        verified = validate_r3_stable_runtime(config["runtime"])
        self.assertEqual(verified["internal_grouped_oof"]["correct"], 2316)
        self.assertEqual(verified["internal_grouped_oof"]["rows"], 2751)

        command = core_command(config, self.root / "run", "07221430")
        self.assertIn("supervised_stable", command)
        self.assertIn("--r3-stable-model", command)
        self.assertIn("--r3-model-snapshot", command)


class ConsolidatedDeliverablesTests(unittest.TestCase):
    def test_writes_complete_workbook_and_gzipped_jsonl(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            targets = root / "targets.xlsx"
            pmag = root / "PMAGRESULT.xlsx"
            rockunit = root / "ROCKUNIT.xlsx"
            pd.DataFrame(
                [{"RESULTNO": "1", "ROCKUNITNO": "2", "REFNO": "10", "ROCKNAME": "Unit A"}]
            ).to_excel(targets, index=False)
            pd.DataFrame(
                [
                    {
                        "RESULTNO": "1",
                        "ROCKUNITNO": "2",
                        "HIMAGAGE": "110",
                        "LOMAGAGE": "90",
                        "B": "9",
                        "N": "30",
                        "KD": "20",
                        "ED95": "8",
                        "DEMAGCODE": "3",
                        "TESTS": "F+,R+",
                    }
                ]
            ).to_excel(pmag, index=False)
            pd.DataFrame(
                [
                    {
                        "ROCKUNITNO": "2",
                        "REFNO": "10",
                        "ROCKNAME": "Unit A",
                        "HIGHAGE": "110",
                        "LOWAGE": "90",
                    }
                ]
            ).to_excel(rockunit, index=False)

            indicator_paths = {}
            for indicator, score, outcome in (("R3", 1, "POSITIVE"), ("R5", 0, "NEGATIVE"), ("R7", 1, "DECIDED")):
                path = root / f"{indicator}.xlsx"
                pd.DataFrame(
                    [
                        {
                            "REFNO": "10",
                            "RESULTNO": "1",
                            "TARGET": "Unit A",
                            indicator: score,
                            f"{indicator}_RESULT_JSON": json.dumps(
                                {
                                    indicator: score,
                                    "outcome": outcome,
                                    "summary": f"{indicator} evidence",
                                }
                            ),
                        }
                    ]
                ).to_excel(path, index=False)
                indicator_paths[indicator] = path

            args = Namespace(
                targets=targets,
                pmagresult=pmag,
                rockunit=rockunit,
                r3=indicator_paths["R3"],
                r5=indicator_paths["R5"],
                r7=indicator_paths["R7"],
                output_dir=root,
            )
            workbook = root / "r1_r7.xlsx"
            jsonl = root / "r1_r7.jsonl.gz"
            manifest = root / "r1_r7.manifest.json"
            result = create_deliverables(args, workbook, jsonl, manifest)

            self.assertEqual(result["rows"], 1)
            self.assertTrue(workbook.exists())
            self.assertTrue(jsonl.exists())
            loaded = load_workbook(workbook, read_only=True, data_only=True)
            self.assertEqual(loaded.sheetnames, ["R1-R7 Results", "Table Evidence", "Run Summary"])
            self.assertEqual(loaded["R1-R7 Results"].max_row, 2)

            import gzip

            with gzip.open(jsonl, "rt", encoding="utf-8") as handle:
                record = json.loads(handle.readline())
            self.assertEqual(record["scores"], {"R1": 1, "R2": 1, "R3": 1, "R4": 1, "R5": 0, "R6": 1, "R7": 1})
            self.assertEqual(record["raw_evidence"]["R5"]["outcome"], "NEGATIVE")


if __name__ == "__main__":
    unittest.main()
