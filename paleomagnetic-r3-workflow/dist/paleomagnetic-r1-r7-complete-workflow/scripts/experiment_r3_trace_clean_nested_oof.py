#!/usr/bin/env python3
"""Evaluate clean TRACE-PMAG features with strict REFNO-nested OOF.

Feature generation is already frozen before this command opens R3 labels.  The
fixed monotonic scorer is fully label-free.  The linear structured/text models
use R3 only inside each outer training partition; they never read PMAGRESULT,
ROCKMAG, prior predictions, OOF arrays, reports, or evaluation metrics.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Mapping, Sequence

import numpy as np
import sklearn
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.ensemble import HistGradientBoostingClassifier, RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
)
from sklearn.preprocessing import StandardScaler


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import build_r3_trace_page_graph_features as freeze  # noqa: E402
from paper_text_retrieval import canonical_target_identity  # noqa: E402


OUTPUT_DIR = ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1"
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
FEATURES = OUTPUT_DIR / freeze.FEATURE_NAME
MANIFEST = OUTPUT_DIR / freeze.MANIFEST_NAME
CONTEXT = OUTPUT_DIR / freeze.CONTEXT_NAME
CANDIDATE_FREEZE = OUTPUT_DIR / freeze.CANDIDATE_FREEZE_NAME
FOLD_MANIFEST = OUTPUT_DIR / freeze.FOLD_MANIFEST_NAME
OOF_OUTPUT = OUTPUT_DIR / "trace_clean_nested_oof_v4.npz"
REPORT_OUTPUT = OUTPUT_DIR / "trace_clean_nested_oof_report_v4.json"
RECEIPT_OUTPUT = OUTPUT_DIR / "trace_clean_nested_oof_receipt_v4.json"
INNER_SEED = 35_720_260_761
RANDOM_STATE = 35_720_260


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_labels(path: Path, resultnos: set[str]) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = tuple(clean(value) for value in next(rows, ()))
        columns = {name: header.index(name) for name in ("RESULTNO", "R3")}
        output: dict[str, int] = {}
        for row in rows:
            resultno = clean(row[columns["RESULTNO"]])
            if resultno not in resultnos:
                continue
            value = int(float(row[columns["R3"]]))
            if value not in {0, 1} or resultno in output:
                raise ValueError(f"invalid R3 label for RESULTNO {resultno}")
            output[resultno] = value
    finally:
        workbook.close()
    if set(output) != resultnos:
        raise ValueError("R3 labels do not align with the frozen TRACE identity")
    return output


def load_frozen_inputs(
    *,
    identity_path: Path,
    pdf_dir: Path,
    feature_path: Path,
    manifest_path: Path,
    context_path: Path,
    candidate_path: Path,
    fold_path: Path,
) -> tuple[
    list[dict[str, str]],
    np.ndarray,
    list[str],
    list[str],
    np.ndarray,
    np.ndarray,
    dict[str, Any],
]:
    all_records = freeze.load_three_column_identity(identity_path)
    records, _, _ = freeze.select_pdf_backed_records(all_records, pdf_dir)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    candidate = json.loads(candidate_path.read_text(encoding="utf-8"))
    fold_manifest = json.loads(fold_path.read_text(encoding="utf-8"))
    alignment = freeze._record_alignment_sha256(records)
    if (
        manifest.get("schema_version") != freeze.FEATURE_SCHEMA_VERSION
        or manifest.get("status") != "FROZEN_WITHOUT_LABEL_ACCESS"
    ):
        raise ValueError("TRACE feature manifest has an invalid frozen schema")
    if (
        candidate.get("schema_version") != freeze.CANDIDATE_SCHEMA_VERSION
        or candidate.get("status") != "FROZEN_BEFORE_LABEL_ACCESS"
    ):
        raise ValueError("TRACE candidate set has an invalid frozen schema")
    if (
        fold_manifest.get("schema_version") != freeze.FOLD_SCHEMA_VERSION
        or fold_manifest.get("status") != "FROZEN_WITHOUT_LABEL_ACCESS"
        or fold_manifest.get("labels_read") is not False
    ):
        raise ValueError("TRACE fold manifest has an invalid frozen schema")
    if manifest.get("alignment_sha256") != alignment:
        raise ValueError("TRACE feature alignment changed")
    if (
        candidate.get("alignment_sha256") != alignment
        or fold_manifest.get("alignment_sha256") != alignment
    ):
        raise ValueError("TRACE frozen artifacts have different row alignments")
    for item in (manifest.get("inputs") or {}).get("code") or []:
        code_path = Path(clean(item.get("path"))).resolve()
        if not code_path.is_file() or freeze.file_sha256(code_path) != clean(
            item.get("sha256")
        ):
            raise ValueError(f"frozen TRACE code hash mismatch: {code_path}")
    boundary = candidate.get("runtime_boundary") or {}
    if any(
        boundary.get(name) is not False
        for name in (
            "labels_read",
            "pmagresult_or_rockmag_read",
            "prior_predictions_read",
            "evaluation_results_read",
        )
    ):
        raise ValueError("TRACE candidate freeze has an invalid source boundary")
    for path, expected in (
        (feature_path, (manifest.get("outputs") or {}).get("features", {}).get("sha256")),
        (context_path, (manifest.get("outputs") or {}).get("target_context", {}).get("sha256")),
        (candidate_path, (manifest.get("outputs") or {}).get("candidate_freeze", {}).get("sha256")),
        (fold_path, (manifest.get("outputs") or {}).get("fold_manifest", {}).get("sha256")),
    ):
        if not expected or freeze.file_sha256(path) != expected:
            raise ValueError(f"frozen TRACE input hash mismatch: {path.name}")
    if candidate.get("target_context_sha256") != freeze.file_sha256(context_path):
        raise ValueError("TRACE candidate/context hash mismatch")
    if candidate.get("fold_manifest_sha256") != freeze.file_sha256(fold_path):
        raise ValueError("TRACE candidate/fold hash mismatch")
    frozen_matrix = np.load(feature_path)["trace"]
    matrix_hash = freeze.array_sha256(frozen_matrix)
    if (
        manifest.get("matrix_sha256") != matrix_hash
        or candidate.get("matrix_sha256") != matrix_hash
    ):
        raise ValueError("TRACE frozen matrix hash mismatch")
    feature_names = list(manifest.get("feature_names") or [])
    if (
        frozen_matrix.ndim != 2
        or frozen_matrix.shape != (len(records), len(feature_names))
        or not np.isfinite(frozen_matrix).all()
    ):
        raise ValueError("TRACE structured features do not align")
    matrix = frozen_matrix.astype(np.float64)
    documents = [
        clean(json.loads(line))
        for line in context_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if len(documents) != len(records) or any(not value for value in documents):
        raise ValueError("TRACE target-context documents do not align")
    fold_rows = fold_manifest.get("rows") or []
    if len(fold_rows) != len(records):
        raise ValueError("TRACE fold manifest does not align")
    fold_ids = np.zeros(len(records), dtype=np.int8)
    document_groups = np.empty(len(records), dtype=object)
    for index, (record, item) in enumerate(zip(records, fold_rows)):
        if (
            int(item.get("row_index", -1)) != index
            or clean(item.get("REFNO")) != record["REFNO"]
            or clean(item.get("RESULTNO")) != record["RESULTNO"]
        ):
            raise ValueError("TRACE fold identity changed")
        fold_ids[index] = int(item["fold"])
        document_groups[index] = clean(item.get("document_sha256"))
        if not document_groups[index]:
            raise ValueError("TRACE fold manifest lacks PDF content groups")
    for refno in {record["REFNO"] for record in records}:
        values = {
            int(fold_ids[index])
            for index, record in enumerate(records)
            if record["REFNO"] == refno
        }
        if len(values) != 1:
            raise AssertionError("TRACE outer folds split a REFNO")
    for document_group in set(document_groups):
        values = set(map(int, fold_ids[document_groups == document_group]))
        if len(values) != 1:
            raise AssertionError("TRACE outer folds split normalized document content")
    if hashlib.sha256(fold_ids.tobytes()).hexdigest() != clean(
        fold_manifest.get("fold_id_sha256")
    ):
        raise ValueError("TRACE frozen fold-id hash mismatch")
    return (
        records,
        matrix,
        feature_names,
        documents,
        fold_ids,
        document_groups,
        candidate,
    )


def inverse_identity_weights(
    indices: np.ndarray, records: Sequence[Mapping[str, str]]
) -> np.ndarray:
    counts = Counter(
        (records[int(index)]["REFNO"], canonical_target_identity(records[int(index)]["ROCKNAME"]))
        for index in indices
    )
    return np.asarray(
        [
            1.0
            / counts[
                (
                    records[int(index)]["REFNO"],
                    canonical_target_identity(records[int(index)]["ROCKNAME"]),
                )
            ]
            for index in indices
        ],
        dtype=np.float64,
    )


def metric(labels: np.ndarray, probabilities: np.ndarray) -> dict[str, Any]:
    prediction = (np.asarray(probabilities) >= 0.5).astype(np.int8)
    tn, fp, fn, tp = confusion_matrix(labels, prediction, labels=[0, 1]).ravel()
    return {
        "rows": int(len(labels)),
        "accuracy": round(float(accuracy_score(labels, prediction)), 8),
        "balanced_accuracy": round(float(balanced_accuracy_score(labels, prediction)), 8),
        "macro_f1": round(float(f1_score(labels, prediction, average="macro", zero_division=0)), 8),
        "precision_1": round(float(precision_score(labels, prediction, zero_division=0)), 8),
        "recall_1": round(float(recall_score(labels, prediction, zero_division=0)), 8),
        "confusion": {"tn": int(tn), "fp": int(fp), "fn": int(fn), "tp": int(tp)},
    }


def better_metric(
    labels: np.ndarray, probabilities: np.ndarray, tie_break: float
) -> tuple[float, float, float]:
    observed = metric(labels, probabilities)
    return (
        float(observed["accuracy"]),
        float(observed["macro_f1"]),
        tie_break,
    )


def fit_structured(
    train: np.ndarray,
    test: np.ndarray,
    labels: np.ndarray,
    weights: np.ndarray,
    c_value: float,
) -> np.ndarray:
    scaler = StandardScaler()
    train_scaled = np.clip(scaler.fit_transform(train), -12.0, 12.0)
    test_scaled = np.clip(scaler.transform(test), -12.0, 12.0)
    model = LogisticRegression(
        C=c_value,
        solver="liblinear",
        max_iter=2_000,
        random_state=RANDOM_STATE,
    )
    model.fit(train_scaled, labels, sample_weight=weights)
    with np.errstate(divide="ignore", over="ignore", invalid="ignore"):
        probability = np.asarray(
            model.predict_proba(test_scaled)[:, 1], dtype=np.float64
        )
    if not np.isfinite(probability).all():
        raise AssertionError("structured TRACE model produced non-finite probabilities")
    return probability


def fit_hgb(
    train: np.ndarray,
    test: np.ndarray,
    labels: np.ndarray,
    weights: np.ndarray,
    config: Mapping[str, Any],
) -> np.ndarray:
    model = HistGradientBoostingClassifier(
        max_iter=int(config["max_iter"]),
        learning_rate=float(config["learning_rate"]),
        max_leaf_nodes=int(config["max_leaf_nodes"]),
        min_samples_leaf=int(config["min_samples_leaf"]),
        l2_regularization=float(config["l2_regularization"]),
        random_state=RANDOM_STATE,
    )
    model.fit(train, labels, sample_weight=weights)
    probability = np.asarray(model.predict_proba(test)[:, 1], dtype=np.float64)
    if not np.isfinite(probability).all():
        raise AssertionError("HGB TRACE model produced non-finite probabilities")
    return probability


def fit_random_forest(
    train: np.ndarray,
    test: np.ndarray,
    labels: np.ndarray,
    weights: np.ndarray,
    config: Mapping[str, Any],
) -> np.ndarray:
    model = RandomForestClassifier(
        n_estimators=int(config["n_estimators"]),
        max_features=float(config["max_features"]),
        min_samples_leaf=int(config["min_samples_leaf"]),
        n_jobs=int(config.get("n_jobs", -1)),
        random_state=RANDOM_STATE,
    )
    model.fit(train, labels, sample_weight=weights)
    probability = np.asarray(model.predict_proba(test)[:, 1], dtype=np.float64)
    if not np.isfinite(probability).all():
        raise AssertionError("random-forest TRACE model produced non-finite probabilities")
    return probability


def _rows_by_refno(
    records: Sequence[Mapping[str, str]],
) -> dict[str, list[int]]:
    output: dict[str, list[int]] = defaultdict(list)
    for index, record in enumerate(records):
        output[clean(record.get("REFNO"))].append(index)
    return dict(output)


def _pairwise_examples(
    matrix: np.ndarray,
    labels: np.ndarray,
    train_indices: np.ndarray,
    records: Sequence[Mapping[str, str]],
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    selected = set(map(int, train_indices))
    differences: list[np.ndarray] = []
    pair_labels: list[int] = []
    weights: list[float] = []
    for indices in _rows_by_refno(records).values():
        current = [index for index in indices if index in selected]
        identity_groups: dict[str, list[int]] = defaultdict(list)
        for index in current:
            identity_groups[
                canonical_target_identity(records[index]["ROCKNAME"])
            ].append(index)
        representatives: dict[int, list[int]] = {0: [], 1: []}
        for group in identity_groups.values():
            group_labels = {int(labels[index]) for index in group}
            if len(group_labels) != 1:
                continue
            representatives[next(iter(group_labels))].append(group[0])
        positive = representatives[1]
        negative = representatives[0]
        pairs = len(positive) * len(negative)
        if not pairs:
            continue
        weight = 0.5 / pairs
        for positive_index in positive:
            for negative_index in negative:
                difference = np.asarray(
                    matrix[positive_index] - matrix[negative_index], dtype=np.float64
                )
                differences.extend((difference, -difference))
                pair_labels.extend((1, 0))
                weights.extend((weight, weight))
    if not differences:
        raise ValueError("TRACE pairwise training fold contains no mixed-label papers")
    return (
        np.asarray(differences, dtype=np.float64),
        np.asarray(pair_labels, dtype=np.int8),
        np.asarray(weights, dtype=np.float64),
    )


def fit_pairwise_scores(
    matrix: np.ndarray,
    labels: np.ndarray,
    train_indices: np.ndarray,
    predict_indices: np.ndarray,
    records: Sequence[Mapping[str, str]],
    c_value: float,
) -> np.ndarray:
    differences, pair_labels, weights = _pairwise_examples(
        matrix, labels, train_indices, records
    )
    scaler = StandardScaler()
    transformed = np.clip(scaler.fit_transform(differences), -12.0, 12.0)
    model = LogisticRegression(
        C=c_value,
        solver="liblinear",
        fit_intercept=False,
        max_iter=3_000,
        random_state=RANDOM_STATE,
    )
    model.fit(transformed, pair_labels, sample_weight=weights * len(weights))
    selected = np.clip(scaler.transform(matrix[predict_indices]), -12.0, 12.0)
    with np.errstate(divide="ignore", over="ignore", invalid="ignore"):
        raw = np.asarray(model.decision_function(selected), dtype=np.float64)
    if not np.isfinite(raw).all():
        raise AssertionError("TRACE pairwise model produced non-finite scores")
    output = np.zeros(len(records), dtype=np.float64)
    output[predict_indices] = raw
    selected_set = set(map(int, predict_indices))
    for indices in _rows_by_refno(records).values():
        group = [index for index in indices if index in selected_set]
        if len(group) < 2:
            continue
        identity_groups: dict[str, list[int]] = defaultdict(list)
        for index in group:
            identity_groups[
                canonical_target_identity(records[index]["ROCKNAME"])
            ].append(index)
        unique_scores = np.asarray(
            [float(output[rows].mean()) for rows in identity_groups.values()],
            dtype=np.float64,
        )
        centered = unique_scores - float(unique_scores.mean())
        scale = float(np.sqrt(np.mean(np.square(centered))))
        if scale > 1e-9:
            normalized = np.clip(centered / scale, -3.0, 3.0)
            for rows, value in zip(identity_groups.values(), normalized):
                output[rows] = value
    return output[predict_indices]


def _logit(values: np.ndarray) -> np.ndarray:
    clipped = np.clip(np.asarray(values, dtype=np.float64), 1e-5, 1.0 - 1e-5)
    return np.log(clipped / (1.0 - clipped))


def _sigmoid(values: np.ndarray) -> np.ndarray:
    return 1.0 / (1.0 + np.exp(-np.clip(values, -40.0, 40.0)))


def pairwise_route_specs(candidate: Mapping[str, Any]) -> list[dict[str, Any]]:
    config = candidate["model_candidates"]["within_paper_pairwise"]
    output: list[dict[str, Any]] = [
        {
            "name": "identity",
            "c_value": None,
            "minimum_distinct_target_count": 10**9,
            "alpha": 0.0,
        }
    ]
    for c_value in config["c_values"]:
        for minimum in config["minimum_distinct_target_counts"]:
            for alpha in config["alphas"]:
                output.append(
                    {
                        "name": f"pair_c{float(c_value):g}_n{int(minimum)}_a{float(alpha):g}",
                        "c_value": float(c_value),
                        "minimum_distinct_target_count": int(minimum),
                        "alpha": float(alpha),
                    }
                )
    return output


def apply_pairwise_route(
    baseline: np.ndarray,
    rank_score: np.ndarray,
    selected_indices: np.ndarray,
    records: Sequence[Mapping[str, str]],
    spec: Mapping[str, Any],
) -> np.ndarray:
    if float(spec["alpha"]) == 0.0:
        return np.asarray(baseline, dtype=np.float64).copy()
    distinct_counts: dict[str, set[str]] = defaultdict(set)
    for record in records:
        distinct_counts[record["REFNO"]].add(
            canonical_target_identity(record["ROCKNAME"])
        )
    enabled = np.asarray(
        [
            len(distinct_counts[records[int(index)]["REFNO"]])
            >= int(spec["minimum_distinct_target_count"])
            for index in selected_indices
        ],
        dtype=bool,
    )
    corrected = _sigmoid(_logit(baseline) + float(spec["alpha"]) * rank_score)
    return np.where(enabled, corrected, baseline)


def select_pairwise_route(
    route_specs: Sequence[Mapping[str, Any]],
    route_probabilities: Mapping[str, np.ndarray],
    labels: np.ndarray,
    config: Mapping[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    identity = next(spec for spec in route_specs if spec["name"] == "identity")
    selected = max(
        route_specs,
        key=lambda spec: better_metric(
            labels,
            route_probabilities[spec["name"]],
            -float(route_specs.index(spec)),
        ),
    )
    identity_metric = metric(labels, route_probabilities["identity"])
    selected_metric = metric(labels, route_probabilities[selected["name"]])
    accuracy_gain = selected_metric["accuracy"] - identity_metric["accuracy"]
    macro_f1_gain = selected_metric["macro_f1"] - identity_metric["macro_f1"]
    accepted = selected["name"] != "identity"
    accepted = accepted and accuracy_gain >= float(
        config["minimum_accuracy_gain_over_identity"]
    )
    if bool(config.get("require_non_decreasing_macro_f1")):
        accepted = accepted and macro_f1_gain >= 0.0
    if not accepted:
        selected = identity
    return dict(selected), {
        "unconstrained_best": selected_metric,
        "identity": identity_metric,
        "accuracy_gain": round(float(accuracy_gain), 8),
        "macro_f1_gain": round(float(macro_f1_gain), 8),
        "accepted_non_identity": bool(accepted),
        "minimum_accuracy_gain": float(
            config["minimum_accuracy_gain_over_identity"]
        ),
    }


def vectorize_text(
    train_documents: Sequence[str],
    test_documents: Sequence[str],
    *,
    analyzer: str,
    ngram_range: tuple[int, int],
    max_features: int,
) -> tuple[Any, Any, int]:
    vectorizer = TfidfVectorizer(
        analyzer=analyzer,
        ngram_range=ngram_range,
        min_df=2,
        max_df=0.997,
        max_features=max_features,
        sublinear_tf=True,
        strip_accents="unicode" if analyzer == "word" else None,
        dtype=np.float32,
    )
    train_matrix = vectorizer.fit_transform(train_documents)
    test_matrix = vectorizer.transform(test_documents)
    return train_matrix, test_matrix, len(vectorizer.vocabulary_)


def fit_text_from_matrices(
    train: Any,
    test: Any,
    labels: np.ndarray,
    weights: np.ndarray,
    c_value: float,
) -> np.ndarray:
    model = LogisticRegression(
        C=c_value,
        solver="liblinear",
        max_iter=2_000,
        random_state=RANDOM_STATE,
    )
    model.fit(train, labels, sample_weight=weights)
    return np.asarray(model.predict_proba(test)[:, 1], dtype=np.float64)


def monotonic_probability(
    matrix: np.ndarray,
    feature_names: Sequence[str],
    formula: Mapping[str, float],
) -> np.ndarray:
    output = np.full(len(matrix), float(formula.get("intercept", 0.0)))
    for name, weight in formula.items():
        if name == "intercept":
            continue
        if name not in feature_names:
            raise ValueError(f"monotonic TRACE formula references missing feature: {name}")
        output += float(weight) * np.clip(matrix[:, feature_names.index(name)], 0.0, 1.0)
    return np.clip(output, 0.0, 1.0)


def fusion_probabilities(
    monotonic: np.ndarray,
    structured: np.ndarray,
    hgb: np.ndarray,
    random_forest: np.ndarray,
    word: np.ndarray,
    char: np.ndarray,
) -> dict[str, np.ndarray]:
    return {
        "monotonic_criterion": monotonic,
        "structured": structured,
        "hgb": hgb,
        "random_forest": random_forest,
        "word": word,
        "char": char,
        "mean_hgb_random_forest": (hgb + random_forest) / 2.0,
        "mean_structured_hgb": (structured + hgb) / 2.0,
        "mean_structured_random_forest": (structured + random_forest) / 2.0,
        "mean_structured_hgb_random_forest": (
            structured + hgb + random_forest
        )
        / 3.0,
        "mean_structured_word": (structured + word) / 2.0,
        "mean_structured_char": (structured + char) / 2.0,
        "mean_hgb_char": (hgb + char) / 2.0,
        "mean_random_forest_char": (random_forest + char) / 2.0,
        "mean_structured_hgb_char": (structured + hgb + char) / 3.0,
        "mean_word_char": (word + char) / 2.0,
        "mean_all": (structured + word + char) / 3.0,
        "weighted_structured_50_word_25_char_25": (
            0.50 * structured + 0.25 * word + 0.25 * char
        ),
    }


def select_fusion(
    allowed_fusions: Sequence[str],
    probabilities: Mapping[str, np.ndarray],
    labels: np.ndarray,
    guard: Mapping[str, Any],
) -> tuple[str, dict[str, Any]]:
    baseline_name = clean(guard["baseline"])
    if baseline_name not in allowed_fusions:
        raise ValueError("fusion guard baseline is not a frozen candidate")
    selected = max(
        allowed_fusions,
        key=lambda name: better_metric(
            labels, probabilities[name], -float(allowed_fusions.index(name))
        ),
    )
    baseline_metric = metric(labels, probabilities[baseline_name])
    selected_metric = metric(labels, probabilities[selected])
    accuracy_gain = selected_metric["accuracy"] - baseline_metric["accuracy"]
    macro_f1_gain = selected_metric["macro_f1"] - baseline_metric["macro_f1"]
    accepted = selected != baseline_name and accuracy_gain >= float(
        guard["minimum_accuracy_gain_over_baseline"]
    )
    if bool(guard.get("require_non_decreasing_macro_f1")):
        accepted = accepted and macro_f1_gain >= 0.0
    unconstrained = selected
    if not accepted:
        selected = baseline_name
    return selected, {
        "baseline": baseline_name,
        "unconstrained_best": unconstrained,
        "baseline_metric": baseline_metric,
        "unconstrained_best_metric": selected_metric,
        "accuracy_gain": round(float(accuracy_gain), 8),
        "macro_f1_gain": round(float(macro_f1_gain), 8),
        "accepted_non_baseline": bool(accepted),
        "minimum_accuracy_gain": float(
            guard["minimum_accuracy_gain_over_baseline"]
        ),
    }


def nested_family_selection(
    outer_train: np.ndarray,
    records: Sequence[Mapping[str, str]],
    labels: np.ndarray,
    matrix: np.ndarray,
    documents: Sequence[str],
    monotonic: np.ndarray,
    candidate: Mapping[str, Any],
    document_groups: np.ndarray,
    outer_fold: int,
) -> tuple[dict[str, float], str, dict[str, Any], dict[str, Any]]:
    inner_records = [records[int(index)] for index in outer_train]
    group_key_by_refno = {
        records[int(index)]["REFNO"]: clean(document_groups[int(index)])
        for index in outer_train
    }
    inner_fold_ids = freeze.deterministic_refno_fold_ids(
        inner_records,
        n_splits=int(candidate["inner_folds"]),
        seed=INNER_SEED + outer_fold,
        group_key_by_refno=group_key_by_refno,
    )
    specs = candidate["model_candidates"]
    c_values = {
        "structured": [float(value) for value in specs["structured_logistic"]["c_values"]],
        "word": [float(value) for value in specs["target_context_word"]["c_values"]],
        "char": [float(value) for value in specs["target_context_char"]["c_values"]],
    }
    predictions = {
        family: {value: np.full(len(outer_train), np.nan) for value in values}
        for family, values in c_values.items()
    }
    nonlinear_predictions = {
        family: np.full(len(outer_train), np.nan)
        for family in ("hgb", "random_forest")
    }
    pair_c_values = [
        float(value)
        for value in specs["within_paper_pairwise"]["c_values"]
    ]
    pair_predictions = {
        value: np.full(len(outer_train), np.nan) for value in pair_c_values
    }
    vocabulary_audit = []
    for inner_fold in range(1, int(candidate["inner_folds"]) + 1):
        fit_positions = np.flatnonzero(inner_fold_ids != inner_fold)
        validation_positions = np.flatnonzero(inner_fold_ids == inner_fold)
        fit_indices = outer_train[fit_positions]
        validation_indices = outer_train[validation_positions]
        fit_labels = labels[fit_indices]
        weights = inverse_identity_weights(fit_indices, records)
        for value in c_values["structured"]:
            predictions["structured"][value][validation_positions] = fit_structured(
                matrix[fit_indices],
                matrix[validation_indices],
                fit_labels,
                weights,
                value,
            )
        nonlinear_predictions["hgb"][validation_positions] = fit_hgb(
            matrix[fit_indices],
            matrix[validation_indices],
            fit_labels,
            weights,
            specs["structured_hgb"],
        )
        nonlinear_predictions["random_forest"][validation_positions] = (
            fit_random_forest(
                matrix[fit_indices],
                matrix[validation_indices],
                fit_labels,
                weights,
                specs["structured_random_forest"],
            )
        )
        text_matrices: dict[str, tuple[Any, Any, int]] = {}
        for family, analyzer, ngrams, cap in (
            ("word", "word", (1, 2), 80_000),
            ("char", "char_wb", (3, 5), 100_000),
        ):
            text_matrices[family] = vectorize_text(
                [documents[int(index)] for index in fit_indices],
                [documents[int(index)] for index in validation_indices],
                analyzer=analyzer,
                ngram_range=ngrams,
                max_features=cap,
            )
            train_text, validation_text, vocabulary = text_matrices[family]
            vocabulary_audit.append(
                {"inner_fold": inner_fold, "family": family, "vocabulary": vocabulary}
            )
            for value in c_values[family]:
                predictions[family][value][validation_positions] = fit_text_from_matrices(
                    train_text,
                    validation_text,
                    fit_labels,
                    weights,
                    value,
                )
        for value in pair_c_values:
            pair_predictions[value][validation_positions] = fit_pairwise_scores(
                matrix,
                labels,
                fit_indices,
                validation_indices,
                records,
                value,
            )
    selected_c: dict[str, float] = {}
    c_metrics: dict[str, dict[str, Any]] = {}
    inner_labels = labels[outer_train]
    for family, values in c_values.items():
        if any(not np.isfinite(predictions[family][value]).all() for value in values):
            raise AssertionError(f"inner {family} OOF predictions are incomplete")
        selected = max(
            values,
            key=lambda value: better_metric(
                inner_labels,
                predictions[family][value],
                -abs(np.log10(value)),
            ),
        )
        selected_c[family] = selected
        c_metrics[family] = {
            str(value): metric(inner_labels, predictions[family][value]) for value in values
        }
    fusions = fusion_probabilities(
        monotonic[outer_train],
        predictions["structured"][selected_c["structured"]],
        nonlinear_predictions["hgb"],
        nonlinear_predictions["random_forest"],
        predictions["word"][selected_c["word"]],
        predictions["char"][selected_c["char"]],
    )
    allowed_fusions = list(candidate["fusion_candidates"])
    if set(allowed_fusions) != set(fusions):
        raise ValueError("implemented TRACE fusion set differs from frozen candidates")
    if any(not np.isfinite(values).all() for values in nonlinear_predictions.values()):
        raise AssertionError("inner nonlinear OOF predictions are incomplete")
    selected_fusion, fusion_guard = select_fusion(
        allowed_fusions,
        fusions,
        inner_labels,
        candidate["fusion_selection_guard"],
    )
    if any(not np.isfinite(values).all() for values in pair_predictions.values()):
        raise AssertionError("inner pairwise OOF scores are incomplete")
    base_probability = fusions[selected_fusion]
    route_specs = pairwise_route_specs(candidate)
    route_probabilities = {}
    for spec in route_specs:
        score = (
            np.zeros(len(outer_train), dtype=np.float64)
            if spec["c_value"] is None
            else pair_predictions[float(spec["c_value"])]
        )
        route_probabilities[spec["name"]] = apply_pairwise_route(
            base_probability, score, outer_train, records, spec
        )
    selected_route, route_guard = select_pairwise_route(
        route_specs,
        route_probabilities,
        inner_labels,
        specs["within_paper_pairwise"],
    )
    audit = {
        "selected_c": selected_c,
        "selected_fusion": selected_fusion,
        "fusion_selection_guard": fusion_guard,
        "inner_family_metrics": c_metrics,
        "inner_fusion_metrics": {
            name: metric(inner_labels, values) for name, values in fusions.items()
        },
        "selected_pairwise_route": selected_route,
        "pairwise_route_selection_guard": route_guard,
        "inner_pairwise_route_metrics": {
            name: metric(inner_labels, values)
            for name, values in route_probabilities.items()
        },
        "vocabulary_audit": vocabulary_audit,
    }
    return selected_c, selected_fusion, selected_route, audit


def fit_outer_families(
    train: np.ndarray,
    test: np.ndarray,
    records: Sequence[Mapping[str, str]],
    labels: np.ndarray,
    matrix: np.ndarray,
    documents: Sequence[str],
    selected_c: Mapping[str, float],
    candidate: Mapping[str, Any],
) -> tuple[dict[str, np.ndarray], dict[str, int]]:
    weights = inverse_identity_weights(train, records)
    specs = candidate["model_candidates"]
    output = {
        "structured": fit_structured(
            matrix[train], matrix[test], labels[train], weights, selected_c["structured"]
        ),
        "hgb": fit_hgb(
            matrix[train],
            matrix[test],
            labels[train],
            weights,
            specs["structured_hgb"],
        ),
        "random_forest": fit_random_forest(
            matrix[train],
            matrix[test],
            labels[train],
            weights,
            specs["structured_random_forest"],
        ),
    }
    vocabulary: dict[str, int] = {}
    for family, analyzer, ngrams, cap in (
        ("word", "word", (1, 2), 80_000),
        ("char", "char_wb", (3, 5), 100_000),
    ):
        train_text, test_text, size = vectorize_text(
            [documents[int(index)] for index in train],
            [documents[int(index)] for index in test],
            analyzer=analyzer,
            ngram_range=ngrams,
            max_features=cap,
        )
        output[family] = fit_text_from_matrices(
            train_text,
            test_text,
            labels[train],
            weights,
            selected_c[family],
        )
        vocabulary[family] = size
    return output, vocabulary


def masks(
    records: Sequence[Mapping[str, str]], labels: np.ndarray, matrix: np.ndarray, names: Sequence[str]
) -> dict[str, np.ndarray]:
    paper_counts = Counter(record["REFNO"] for record in records)
    by_refno: dict[str, list[int]] = defaultdict(list)
    identity_counts = Counter(
        (record["REFNO"], canonical_target_identity(record["ROCKNAME"]))
        for record in records
    )
    for index, record in enumerate(records):
        by_refno[record["REFNO"]].append(index)
    mixed = {
        refno: len({int(labels[index]) for index in indices}) > 1
        for refno, indices in by_refno.items()
    }
    return {
        "overall": np.ones(len(records), dtype=bool),
        "single_target_paper": np.asarray(
            [paper_counts[record["REFNO"]] == 1 for record in records], dtype=bool
        ),
        "multi_target_paper": np.asarray(
            [paper_counts[record["REFNO"]] > 1 for record in records], dtype=bool
        ),
        "mixed_label_paper": np.asarray(
            [mixed[record["REFNO"]] for record in records], dtype=bool
        ),
        "qualified_identity": matrix[:, names.index("identity_qualifier_present")] > 0.5,
        "duplicate_identity": np.asarray(
            [
                identity_counts[(record["REFNO"], canonical_target_identity(record["ROCKNAME"]))]
                > 1
                for record in records
            ],
            dtype=bool,
        ),
    }


def indistinguishable_identity_ceiling(
    records: Sequence[Mapping[str, str]], labels: np.ndarray
) -> dict[str, Any]:
    groups: dict[tuple[str, str], list[int]] = defaultdict(list)
    for index, record in enumerate(records):
        groups[(record["REFNO"], canonical_target_identity(record["ROCKNAME"]))].append(index)
    correct = 0
    mixed_rows = mixed_groups = duplicate_rows = 0
    for indices in groups.values():
        counts = Counter(int(labels[index]) for index in indices)
        correct += max(counts.values())
        if len(indices) > 1:
            duplicate_rows += len(indices)
        if len(counts) > 1:
            mixed_groups += 1
            mixed_rows += len(indices)
    return {
        "assumption": "rows with identical REFNO and canonical ROCKNAME have no distinguishable permitted input",
        "oracle_majority_correct_rows": correct,
        "oracle_majority_accuracy_ceiling": round(correct / len(records), 8),
        "duplicate_identity_rows": duplicate_rows,
        "mixed_label_indistinguishable_groups": mixed_groups,
        "mixed_label_indistinguishable_rows": mixed_rows,
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=freeze.DEFAULT_IDENTITY)
    parser.add_argument("--pdf-dir", type=Path, default=freeze.DEFAULT_PDF_DIR)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--features", type=Path, default=FEATURES)
    parser.add_argument("--manifest", type=Path, default=MANIFEST)
    parser.add_argument("--target-context", type=Path, default=CONTEXT)
    parser.add_argument("--candidate-freeze", type=Path, default=CANDIDATE_FREEZE)
    parser.add_argument("--fold-manifest", type=Path, default=FOLD_MANIFEST)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args(argv)

    (
        records,
        matrix,
        feature_names,
        documents,
        fold_ids,
        document_groups,
        candidate,
    ) = load_frozen_inputs(
        identity_path=args.identity.resolve(),
        pdf_dir=args.pdf_dir.resolve(),
        feature_path=args.features.resolve(),
        manifest_path=args.manifest.resolve(),
        context_path=args.target_context.resolve(),
        candidate_path=args.candidate_freeze.resolve(),
        fold_path=args.fold_manifest.resolve(),
    )
    # This is the first label access in this process.
    label_map = load_labels(
        args.targets.resolve(), {record["RESULTNO"] for record in records}
    )
    labels = np.asarray([label_map[record["RESULTNO"]] for record in records], dtype=np.int8)
    monotonic = monotonic_probability(
        matrix, feature_names, candidate["monotonic_criterion_formula"]
    )
    family_oof = {
        name: np.full(len(records), np.nan, dtype=np.float64)
        for name in ("structured", "hgb", "random_forest", "word", "char")
    }
    primary = np.full(len(records), np.nan, dtype=np.float64)
    base_fusion_oof = np.full(len(records), np.nan, dtype=np.float64)
    fold_audit = []
    for outer_fold in range(1, int(candidate["outer_folds"]) + 1):
        train = np.flatnonzero(fold_ids != outer_fold)
        test = np.flatnonzero(fold_ids == outer_fold)
        if {records[int(index)]["REFNO"] for index in train} & {
            records[int(index)]["REFNO"] for index in test
        }:
            raise AssertionError("TRACE outer fold split a REFNO")
        if set(document_groups[train]) & set(document_groups[test]):
            raise AssertionError("TRACE outer fold split normalized document content")
        (
            selected_c,
            selected_fusion,
            selected_pairwise_route,
            inner_audit,
        ) = nested_family_selection(
            train,
            records,
            labels,
            matrix,
            documents,
            monotonic,
            candidate,
            document_groups,
            outer_fold,
        )
        predicted, vocabulary = fit_outer_families(
            train,
            test,
            records,
            labels,
            matrix,
            documents,
            selected_c,
            candidate,
        )
        for name in family_oof:
            family_oof[name][test] = predicted[name]
        fusions = fusion_probabilities(
            monotonic[test],
            predicted["structured"],
            predicted["hgb"],
            predicted["random_forest"],
            predicted["word"],
            predicted["char"],
        )
        base_fusion_oof[test] = fusions[selected_fusion]
        rank_score = (
            np.zeros(len(test), dtype=np.float64)
            if selected_pairwise_route["c_value"] is None
            else fit_pairwise_scores(
                matrix,
                labels,
                train,
                test,
                records,
                float(selected_pairwise_route["c_value"]),
            )
        )
        primary[test] = apply_pairwise_route(
            fusions[selected_fusion],
            rank_score,
            test,
            records,
            selected_pairwise_route,
        )
        fold_audit.append(
            {
                "outer_fold": outer_fold,
                "train_rows": int(len(train)),
                "test_rows": int(len(test)),
                "test_papers": len({records[int(index)]["REFNO"] for index in test}),
                "outer_vocabulary": vocabulary,
                **inner_audit,
                "outer_test_primary_metric": metric(labels[test], primary[test]),
            }
        )
        print(
            f"TRACE clean nested OOF: fold={outer_fold}/{candidate['outer_folds']} "
            f"fusion={selected_fusion} route={selected_pairwise_route['name']} "
            f"accuracy={metric(labels[test], primary[test])['accuracy']:.4f}",
            flush=True,
        )
    probabilities = {
        "monotonic_criterion": monotonic,
        **family_oof,
        "trace_nested_nonlinear_fusion": base_fusion_oof,
        "trace_guarded_nonlinear_fusion": primary,
    }
    if any(not np.isfinite(values).all() for values in probabilities.values()):
        raise AssertionError("TRACE OOF predictions are incomplete")
    subset_masks = masks(records, labels, matrix, feature_names)
    metrics = {
        name: {
            subset: metric(labels[mask], values[mask])
            for subset, mask in subset_masks.items()
            if mask.any()
        }
        for name, values in probabilities.items()
    }
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    oof_path = output_dir / OOF_OUTPUT.name
    report_path = output_dir / REPORT_OUTPUT.name
    receipt_path = output_dir / RECEIPT_OUTPUT.name
    temporary = oof_path.with_name(oof_path.name + ".tmp")
    with temporary.open("wb") as handle:
        np.savez_compressed(
            handle,
            labels=labels,
            fold_ids=fold_ids,
            **{name: values.astype(np.float32) for name, values in probabilities.items()},
        )
    temporary.replace(oof_path)
    report = {
        "schema_version": "r3_trace_clean_nested_oof_report_v4",
        "status": "POSTHOC_EVALUATION_OF_FROZEN_FEATURES_AND_CANDIDATES",
        "rows": len(records),
        "papers": len({record["REFNO"] for record in records}),
        "primary_strategy": "trace_guarded_nonlinear_fusion",
        "threshold": 0.5,
        "metrics": metrics,
        "fold_audit": fold_audit,
        "identifiability_analysis": indistinguishable_identity_ceiling(records, labels),
        "boundary": {
            "feature_generation_used_labels": False,
            "fixed_monotonic_scorer_used_labels": False,
            "nested_linear_models_used_r3_in_outer_training_folds_only": True,
            "nested_nonlinear_models_used_r3_in_outer_training_folds_only": True,
            "pairwise_ranker_used_r3_in_outer_training_folds_only": True,
            "pmagresult_or_rockmag_used": False,
            "prior_predictions_or_evaluation_reports_used": False,
            "outer_test_refno_seen_in_training": False,
            "outer_test_pdf_content_seen_in_training": False,
            "outer_metrics_used_for_selection": False,
        },
        "runtime": {
            "python": sys.version.split()[0],
            "numpy": np.__version__,
            "sklearn": sklearn.__version__,
        },
        "inputs": {
            name: {"path": str(path.resolve()), "sha256": freeze.file_sha256(path.resolve())}
            for name, path in (
                ("features", args.features),
                ("manifest", args.manifest),
                ("target_context", args.target_context),
                ("candidate_freeze", args.candidate_freeze),
                ("fold_manifest", args.fold_manifest),
            )
        },
        "label_source": {
            "path": str(args.targets.resolve()),
            "sha256": freeze.file_sha256(args.targets.resolve()),
            "columns": ["RESULTNO", "R3"],
            "labels_array_sha256": freeze.array_sha256(labels),
            "used_only_after_feature_and_candidate_freeze": True,
        },
        "output": {"path": str(oof_path), "sha256": freeze.file_sha256(oof_path)},
    }
    freeze.atomic_json(report_path, report)
    receipt = {
        "schema_version": "r3_trace_clean_nested_oof_receipt_v4",
        "status": "POSTHOC_EVALUATION",
        "rows": [
            {
                "row_index": index,
                "REFNO": record["REFNO"],
                "RESULTNO": record["RESULTNO"],
                "fold": int(fold_ids[index]),
                "label": int(labels[index]),
                "probability": round(float(primary[index]), 8),
                "prediction": int(primary[index] >= 0.5),
            }
            for index, record in enumerate(records)
        ],
        "report": {"path": str(report_path), "sha256": freeze.file_sha256(report_path)},
        "label_source": report["label_source"],
    }
    freeze.atomic_json(receipt_path, receipt)
    print(
        f"TRACE clean nested OOF complete: accuracy={metrics['trace_guarded_nonlinear_fusion']['overall']['accuracy']:.4f}, "
        f"output={report_path}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "fusion_probabilities",
    "indistinguishable_identity_ceiling",
    "metric",
    "monotonic_probability",
]
