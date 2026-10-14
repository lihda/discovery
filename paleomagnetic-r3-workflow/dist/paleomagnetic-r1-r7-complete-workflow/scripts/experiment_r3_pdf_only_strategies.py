#!/usr/bin/env python3
"""PDF-only R3 target-attribution experiments on the frozen training split.

This development script deliberately lives outside the runtime scorer.  It
opens R3 only for rows already proven to belong to the frozen training split,
requires a local PDF for every evaluated REFNO, and evaluates every learned
model out of fold with REFNO GroupKFold.  REFNO and RESULTNO are alignment
keys only; neither is exposed to a classifier.

Allowed runtime-recomputable inputs are ROCKNAME, cached original-paper text,
the frozen ROCKMAG-derived taxonomy, and (when present) a separately frozen
positive ROCKMAG expert-knowledge artifact.  PMAGRESULT, held-out labels,
prior predictions, and prior evaluation files are never opened.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook
try:
    import numpy as np
    from sklearn.ensemble import (
        ExtraTreesClassifier,
        GradientBoostingClassifier,
        HistGradientBoostingClassifier,
    )
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.linear_model import LogisticRegression
    from sklearn.metrics import (
        accuracy_score,
        balanced_accuracy_score,
        confusion_matrix,
        f1_score,
    )
    from sklearn.model_selection import GroupKFold
except ModuleNotFoundError as exc:  # Optional development dependencies.
    np = None  # type: ignore[assignment]
    _OPTIONAL_DEPENDENCY_ERROR: ModuleNotFoundError | None = exc
else:
    _OPTIONAL_DEPENDENCY_ERROR = None


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from paper_text_retrieval import canonical_target_identity  # noqa: E402


DEFAULT_IDENTITY = ROOT / "tmp/r3_distilled_training_preflight/identity.xlsx"
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_PAGE_CACHE = ROOT / "tmp/mineru_cache_all_hard_risk.jsonl"
DEFAULT_PDF_DIR = ROOT / "data/raw_pdfs"
DEFAULT_TAXONOMY = ROOT / "data/calibration/r3_rockmag_text_distillation_v1.json"
DEFAULT_POSITIVE_KB = ROOT / "data/calibration/r3_positive_rockmag_expert_kb_v2.json"
DEFAULT_OUTPUT = ROOT / "tmp/r3_pdf_only_strategies/aggregate_result.json"

N_SPLITS = 5
RANDOM_STATE = 357
SPLIT_SALT = "r3-rockmag-distillation-v1"
HOLDOUT_FRACTION = 0.20
SPACE_RE = re.compile(r"\s+")
TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z'-]{1,}|\d+(?:\.\d+)?")
COLLECTIVE_RE = re.compile(
    r"\b(?:all|both|these|those|the)\s+(?:samples?|specimens?|sites?|units?|"
    r"rocks?|formations?|lavas?|d[yi]kes?|intrusions?)\b|"
    r"\b(?:samples?|specimens?)\s+from\s+(?:all|both|each)\b",
    re.I,
)

GENERIC_TARGET_TERMS = {
    "and", "area", "basalt", "basalts", "bed", "beds", "body", "complex",
    "dike", "dikes", "dyke", "dykes", "formation", "formations", "group",
    "intrusion", "intrusions", "lava", "lavas", "member", "members", "of",
    "province", "rock", "rocks", "sample", "samples", "section", "sections",
    "site", "sites", "suite", "the", "unit", "units",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def holdout_bucket(refno: str) -> float:
    digest = hashlib.sha256(f"{SPLIT_SALT}\0{refno}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def is_training_refno(refno: str) -> bool:
    return holdout_bucket(refno) >= HOLDOUT_FRACTION


def _header_indices(header: Sequence[Any], required: Sequence[str]) -> dict[str, int]:
    names = [clean(cell.value if hasattr(cell, "value") else cell) for cell in header]
    output: dict[str, int] = {}
    for name in required:
        matches = [index for index, value in enumerate(names) if value == name]
        if len(matches) != 1:
            raise ValueError(f"expected exactly one {name} column; found {len(matches)}")
        output[name] = matches[0]
    return output


def load_training_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("training identity workbook must contain one worksheet")
        rows = workbook.worksheets[0].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "ROCKNAME"))
        records: list[dict[str, str]] = []
        for row in rows:
            record = {name: clean(row[index].value) for name, index in columns.items()}
            if not any(record.values()):
                continue
            if any(not value for value in record.values()):
                raise ValueError("training identity contains an empty value")
            if not is_training_refno(record["REFNO"]):
                raise ValueError(
                    f"identity contains held-out REFNO {record['REFNO']}; refusing label access"
                )
            records.append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError("training identity is empty")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise ValueError("training identity contains duplicate RESULTNO values")
    return records


def load_training_labels(path: Path, training_resultnos: set[str]) -> dict[str, int]:
    """Dereference R3 only after REFNO and RESULTNO pass training gates."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[workbook.sheetnames[0]].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "R3"))
        labels: dict[str, int] = {}
        for row in rows:
            refno = clean(row[columns["REFNO"]].value)
            if not refno or not is_training_refno(refno):
                continue
            resultno = clean(row[columns["RESULTNO"]].value)
            if resultno not in training_resultnos:
                continue
            raw_label = row[columns["R3"]].value
            try:
                label = int(float(raw_label))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"non-binary training R3 for RESULTNO {resultno}") from exc
            if label not in {0, 1} or resultno in labels:
                raise ValueError(f"invalid or duplicate training R3 for RESULTNO {resultno}")
            labels[resultno] = label
    finally:
        workbook.close()
    if set(labels) != training_resultnos:
        raise ValueError(
            "training label/identity mismatch: "
            f"missing={sorted(training_resultnos - labels.keys())[:10]}, "
            f"extra={sorted(labels.keys() - training_resultnos)[:10]}"
        )
    return labels


def require_experiment_dependencies() -> None:
    if _OPTIONAL_DEPENDENCY_ERROR is not None:
        raise SystemExit(
            "Grouped R3 PDF-only experiments require optional development dependencies. "
            "Install requirements-r3-text-model.txt in an isolated environment. "
            f"Missing import: {_OPTIONAL_DEPENDENCY_ERROR.name}"
        )


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def local_pdf_refnos(pdf_dir: Path) -> set[str]:
    if not pdf_dir.is_dir():
        raise ValueError(f"PDF directory does not exist: {pdf_dir}")
    return {
        path.stem
        for path in pdf_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".pdf"
    }


def load_pdf_only_identity(
    identity_path: Path,
    pdf_dir: Path,
) -> tuple[list[dict[str, str]], dict[str, int]]:
    identity = load_training_identity(identity_path)
    pdf_refnos = local_pdf_refnos(pdf_dir)
    selected = [record for record in identity if record["REFNO"] in pdf_refnos]
    skipped = [record for record in identity if record["REFNO"] not in pdf_refnos]
    if not selected:
        raise ValueError("no frozen-training identities have a local PDF")
    return selected, {
        "fixed_training_rows": len(identity),
        "fixed_training_paper_groups": len({row["REFNO"] for row in identity}),
        "pdf_only_rows": len(selected),
        "pdf_only_paper_groups": len({row["REFNO"] for row in selected}),
        "rows_skipped_without_local_pdf": len(skipped),
        "paper_groups_skipped_without_local_pdf": len(
            {row["REFNO"] for row in skipped}
        ),
    }


def load_page_text(path: Path, refnos: set[str]) -> dict[str, str]:
    papers: dict[str, str] = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            row = json.loads(line)
            refno = clean(row.get("REFNO"))
            if refno not in refnos:
                continue
            if refno in papers:
                raise ValueError(
                    f"duplicate page-cache REFNO {refno} at line {line_number}"
                )
            pages: list[str] = []
            for page in row.get("pages", []):
                if not isinstance(page, Mapping):
                    continue
                text = clean(page.get("text"))
                if text:
                    pages.append(
                        f"[[PAGE {clean(page.get('page')) or len(pages) + 1}]]\n{text}"
                    )
            papers[refno] = "\n\n".join(pages)
    missing = sorted(refnos - papers.keys())
    empty = sorted(refno for refno in refnos if not papers.get(refno, "").strip())
    if missing or empty:
        raise ValueError(
            "every PDF-only training paper must have nonempty cached text: "
            f"missing={missing[:10]}, empty={empty[:10]}"
        )
    return papers


def _iter_runtime_pattern_entries(
    value: Any,
    path: tuple[str, ...] = (),
    *,
    pattern_key: str = "runtime_original_text_patterns",
) -> Iterable[tuple[str, str, list[str]]]:
    """Yield explicit runtime pattern lists without interpreting statistics."""

    if isinstance(value, Mapping):
        raw_patterns = value.get(pattern_key)
        if isinstance(raw_patterns, list) and all(
            isinstance(item, str) for item in raw_patterns
        ):
            concept_id = clean(
                value.get("concept_id")
                or value.get("knowledge_id")
                or value.get("expert_concept_id")
                or value.get("semantic_class")
                or ":".join(path)
            )
            yield ":".join(path[:-1]) or "root", concept_id, list(raw_patterns)
        for key, child in value.items():
            if key == pattern_key:
                continue
            yield from _iter_runtime_pattern_entries(
                child, (*path, str(key)), pattern_key=pattern_key
            )
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _iter_runtime_pattern_entries(
                child, (*path, str(index)), pattern_key=pattern_key
            )


def positive_kb_original_text_entries(
    artifact: Mapping[str, Any],
) -> list[tuple[str, str, list[str]]]:
    """Validate v2 and expose only its explicit runtime paper-text patterns."""

    if artifact.get("schema_version") != "r3_positive_rockmag_expert_kb_v2":
        raise ValueError("positive expert KB has an unexpected schema_version")
    runtime_policy = artifact.get("runtime_policy")
    provenance = artifact.get("provenance")
    if not isinstance(runtime_policy, Mapping) or not isinstance(provenance, Mapping):
        raise ValueError("positive expert KB is missing runtime/provenance audits")
    if runtime_policy.get("contains_record_answer_mapping") is not False:
        raise ValueError("positive expert KB must not contain record answers")
    if runtime_policy.get("source_table_access") is not False:
        raise ValueError("positive expert KB must prohibit runtime source-table access")
    if provenance.get("held_out_field_cells_opened") is not False:
        raise ValueError("positive expert KB reports held-out field access")
    if provenance.get("held_out_labels_opened") is not False:
        raise ValueError("positive expert KB reports held-out label access")
    if provenance.get("prior_predictions_or_evaluations_used") is not False:
        raise ValueError("positive expert KB reports prior prediction/evaluation access")
    entries: list[tuple[str, str, list[str]]] = []
    for taxonomy_name in (
        "method_taxonomy",
        "diagnostic_taxonomy",
        "carrier_taxonomy",
    ):
        taxonomy = artifact.get(taxonomy_name)
        if not isinstance(taxonomy, list):
            raise ValueError(f"positive expert KB is missing {taxonomy_name}")
        entries.extend(
            _iter_runtime_pattern_entries(
                taxonomy,
                (taxonomy_name,),
                pattern_key="runtime_original_text_patterns",
            )
        )
    if not entries:
        raise ValueError("positive expert KB contains no runtime_original_text_patterns")
    return entries


def load_knowledge_patterns(
    taxonomy_path: Path,
    positive_kb_path: Path | None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    inputs: list[tuple[str, Path]] = [("frozen_taxonomy", taxonomy_path)]
    positive_status = "not_present"
    if positive_kb_path is not None and positive_kb_path.is_file():
        inputs.append(("positive_rockmag_expert_kb", positive_kb_path))
        positive_status = "loaded"

    concepts: list[dict[str, Any]] = []
    seen_patterns: set[str] = set()
    per_source = Counter()
    invalid_patterns: list[str] = []
    for source, path in inputs:
        artifact = json.loads(path.read_text(encoding="utf-8"))
        if source == "positive_rockmag_expert_kb":
            if not isinstance(artifact, Mapping):
                raise ValueError("positive expert KB must be a JSON object")
            entries = positive_kb_original_text_entries(artifact)
        else:
            entries = list(_iter_runtime_pattern_entries(artifact))
        if not entries:
            raise ValueError(f"knowledge artifact has no runtime patterns: {path}")
        for category, concept_id, expressions in entries:
            retained: list[re.Pattern[str]] = []
            retained_text: list[str] = []
            for expression in expressions:
                if expression in seen_patterns:
                    continue
                try:
                    compiled = re.compile(expression, re.I | re.S)
                except re.error:
                    invalid_patterns.append(expression)
                    continue
                seen_patterns.add(expression)
                retained.append(compiled)
                retained_text.append(expression)
            if retained:
                concepts.append(
                    {
                        "source": source,
                        "category": category,
                        "concept_id": concept_id,
                        "patterns": retained,
                        "pattern_text": retained_text,
                    }
                )
                per_source[source] += len(retained)
    if not concepts:
        raise ValueError("no valid, unique runtime original-text patterns were loaded")
    return concepts, {
        "positive_kb_status": positive_status,
        "positive_kb_pattern_field": (
            "runtime_original_text_patterns"
            if positive_status == "loaded"
            else None
        ),
        "concept_count": len(concepts),
        "runtime_pattern_count": sum(len(item["patterns"]) for item in concepts),
        "runtime_pattern_counts_by_source": dict(sorted(per_source.items())),
        "invalid_pattern_count": len(invalid_patterns),
        "inputs": {
            source: {"path": str(path), "sha256": file_sha256(path)}
            for source, path in inputs
        },
    }


def target_patterns(rockname: str) -> list[re.Pattern[str]]:
    tokens = [token.casefold() for token in TOKEN_RE.findall(rockname)]
    informative = [
        token
        for token in tokens
        if token not in GENERIC_TARGET_TERMS
        and (len(token) >= 3 or token.replace(".", "").isdigit())
    ]
    patterns: list[re.Pattern[str]] = []
    if len(tokens) >= 2 and len(" ".join(tokens)) >= 7:
        patterns.append(
            re.compile(r"\W+".join(re.escape(token) for token in tokens), re.I)
        )
    for token in sorted(set(informative), key=lambda item: (-len(item), item))[:8]:
        if token.replace(".", "").isdigit():
            patterns.append(re.compile(rf"(?<!\d){re.escape(token)}(?!\d)", re.I))
        else:
            patterns.append(re.compile(rf"\b{re.escape(token)}\b", re.I))
    return patterns


def bounded_spans(
    text: str,
    patterns: Iterable[re.Pattern[str]],
    *,
    radius: int,
    cap: int,
) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    for pattern in patterns:
        for match in pattern.finditer(text):
            spans.append((max(0, match.start() - radius), min(len(text), match.end() + radius)))
            if len(spans) >= cap:
                return spans
    return spans


def merge_spans(
    spans: Iterable[tuple[int, int]],
    *,
    gap: int = 80,
) -> list[tuple[int, int]]:
    ordered = sorted((start, end) for start, end in spans if end > start)
    merged: list[list[int]] = []
    for start, end in ordered:
        if not merged or start > merged[-1][1] + gap:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    return [(start, end) for start, end in merged]


def span_document(text: str, spans: Sequence[tuple[int, int]], cap: int = 60000) -> str:
    pieces: list[str] = []
    characters = 0
    for start, end in merge_spans(spans):
        piece = SPACE_RE.sub(" ", text[start:end]).strip()
        if not piece:
            continue
        remaining = cap - characters
        if remaining <= 0:
            break
        pieces.append(piece[:remaining])
        characters += min(len(piece), remaining)
    return "\n".join(pieces)


def count_pattern(pattern: re.Pattern[str], text: str, cap: int = 50) -> int:
    count = 0
    for _ in pattern.finditer(text):
        count += 1
        if count >= cap:
            break
    return count


def concept_counts(text: str, concepts: Sequence[Mapping[str, Any]]) -> list[int]:
    values: list[int] = []
    for concept in concepts:
        count = 0
        for pattern in concept["patterns"]:
            count = min(50, count + count_pattern(pattern, text))
        values.append(count)
    return values


def target_hit_spans(text: str, rockname: str) -> list[tuple[int, int]]:
    return bounded_spans(text, target_patterns(rockname), radius=0, cap=100)


def article_knowledge_spans(
    text: str,
    concepts: Sequence[Mapping[str, Any]],
) -> list[tuple[int, int]]:
    patterns = [
        pattern
        for concept in concepts
        for pattern in concept["patterns"]
    ]
    return bounded_spans(text, patterns, radius=360, cap=220)


def local_target_document(
    text: str,
    rockname: str,
) -> tuple[str, list[tuple[int, int]]]:
    raw_hits = target_hit_spans(text, rockname)
    spans = [
        (max(0, start - 950), min(len(text), end + 950))
        for start, end in raw_hits
    ]
    # The title/abstract fallback is bounded but still labelled as unanchored.
    body = span_document(text, spans, cap=45000) if spans else SPACE_RE.sub(" ", text[:2600])
    return f"TARGET {rockname}\n{body}", raw_hits


def shared_evidence_counts(
    text: str,
    knowledge_spans: Sequence[tuple[int, int]],
    target_names: Sequence[str],
) -> tuple[int, int]:
    profiles = [target_patterns(name) for name in target_names]
    direct = 0
    collective = 0
    for start, end in merge_spans(knowledge_spans, gap=0):
        context = text[start:end]
        target_mentions = sum(
            any(pattern.search(context) for pattern in patterns)
            for patterns in profiles
        )
        direct += int(target_mentions >= 2)
        collective += int(bool(COLLECTIVE_RE.search(context)))
    return direct, collective


def log_features(values: Sequence[float]) -> list[float]:
    return [math.log1p(max(0.0, float(value))) for value in values]


def build_features(
    identity: Sequence[dict[str, str]],
    paper_text: Mapping[str, str],
    concepts: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    raw_names_by_refno: dict[str, dict[str, str]] = defaultdict(dict)
    rows_per_target: Counter[tuple[str, str]] = Counter()
    for record in identity:
        target_key = canonical_target_identity(record["ROCKNAME"])
        raw_names_by_refno[record["REFNO"]].setdefault(target_key, record["ROCKNAME"])
        rows_per_target[(record["REFNO"], target_key)] += 1

    article_docs: dict[str, str] = {}
    article_counts: dict[str, list[int]] = {}
    article_features: dict[str, list[float]] = {}
    article_meta: dict[str, dict[str, Any]] = {}
    for refno, text in paper_text.items():
        knowledge_spans = article_knowledge_spans(text, concepts)
        document = span_document(text, knowledge_spans, cap=60000)
        if not document:
            document = SPACE_RE.sub(" ", text[:2600])
        names = list(raw_names_by_refno[refno].values())
        direct, collective = shared_evidence_counts(text, knowledge_spans, names)
        counts = concept_counts(text, concepts)
        article_docs[refno] = document
        article_counts[refno] = counts
        article_meta[refno] = {
            "character_count": len(text),
            "knowledge_character_count": len(document),
            "knowledge_hit_count": sum(counts),
            "knowledge_distinct_count": sum(value > 0 for value in counts),
            "canonical_target_count": len(names),
            "shared_direct_contexts": direct,
            "shared_collective_contexts": collective,
            "shared_evidence": bool(direct > 0 or collective >= 2),
        }
        article_features[refno] = [
            *log_features(counts),
            *[float(value > 0) for value in counts],
            math.log1p(len(text)),
            math.log1p(len(document)),
            math.log1p(sum(counts)),
            math.log1p(sum(value > 0 for value in counts)),
            math.log1p(len(names)),
            math.log1p(direct),
            math.log1p(collective),
            float(direct > 0 or collective >= 2),
        ]

    target_cache: dict[tuple[str, str], tuple[str, list[int], int]] = {}
    local_documents: list[str] = []
    rockname_documents: list[str] = []
    engineered_rows: list[list[float]] = []
    local_concept_count_rows: list[list[int]] = []
    article_concept_count_rows: list[list[int]] = []
    row_meta: list[dict[str, Any]] = []
    for record in identity:
        refno = record["REFNO"]
        target_key = canonical_target_identity(record["ROCKNAME"])
        cache_key = (refno, target_key)
        if cache_key not in target_cache:
            local_doc, target_spans = local_target_document(
                paper_text[refno], record["ROCKNAME"]
            )
            counts = concept_counts(local_doc, concepts)
            target_cache[cache_key] = (local_doc, counts, len(target_spans))
        local_doc, local_counts, target_hits = target_cache[cache_key]
        article_values = article_counts[refno]
        meta = article_meta[refno]
        local_documents.append(local_doc)
        rockname_documents.append(record["ROCKNAME"])
        local_concept_count_rows.append(list(local_counts))
        article_concept_count_rows.append(list(article_values))
        engineered_rows.append(
            [
                *log_features(local_counts),
                *[float(value > 0) for value in local_counts],
                *log_features(article_values),
                *[float(value > 0) for value in article_values],
                math.log1p(target_hits),
                math.log1p(meta["character_count"]),
                math.log1p(meta["knowledge_character_count"]),
                math.log1p(meta["knowledge_hit_count"]),
                math.log1p(meta["knowledge_distinct_count"]),
                math.log1p(meta["canonical_target_count"]),
                math.log1p(meta["shared_direct_contexts"]),
                math.log1p(meta["shared_collective_contexts"]),
                float(meta["shared_evidence"]),
                math.log1p(rows_per_target[(refno, target_key)]),
            ]
        )
        row_meta.append(
            {
                "canonical_target": target_key,
                "target_hit_count": target_hits,
                "local_knowledge_hit_count": sum(local_counts),
                "local_knowledge_distinct_count": sum(value > 0 for value in local_counts),
                **meta,
            }
        )
    return {
        "local_documents": local_documents,
        "rockname_documents": rockname_documents,
        "article_documents": article_docs,
        "article_features": article_features,
        "engineered_rows": np.asarray(engineered_rows, dtype=np.float64),
        "local_concept_counts": np.asarray(local_concept_count_rows, dtype=np.int16),
        "article_concept_counts": np.asarray(article_concept_count_rows, dtype=np.int16),
        "article_meta": article_meta,
        "row_meta": row_meta,
        "concept_feature_count": len(concepts),
    }


def metric_record(
    labels: np.ndarray,
    probabilities: np.ndarray,
    *,
    threshold: float = 0.5,
) -> dict[str, Any]:
    prediction = (probabilities >= threshold).astype(np.int8)
    tn, fp, fn, tp = confusion_matrix(labels, prediction, labels=[0, 1]).ravel()
    return {
        "rows": int(len(labels)),
        "coverage": 1.0,
        "accuracy": round(float(accuracy_score(labels, prediction)), 8),
        "macro_f1": round(float(f1_score(labels, prediction, average="macro")), 8),
        "balanced_accuracy": round(float(balanced_accuracy_score(labels, prediction)), 8),
        "confusion": {"tn": int(tn), "fp": int(fp), "fn": int(fn), "tp": int(tp)},
        "predicted_positive": int(prediction.sum()),
        "threshold": round(float(threshold), 5),
    }


def selective_metric(
    labels: np.ndarray,
    probabilities: np.ndarray,
    negative_threshold: float,
    positive_threshold: float,
) -> dict[str, Any]:
    known = (probabilities <= negative_threshold) | (probabilities >= positive_threshold)
    prediction = (probabilities >= positive_threshold).astype(np.int8)
    base = {
        "negative_threshold": round(float(negative_threshold), 5),
        "positive_threshold": round(float(positive_threshold), 5),
        "known_rows": int(known.sum()),
        "abstained_rows": int((~known).sum()),
        "coverage": round(float(known.mean()), 8),
    }
    if not known.any():
        return {
            **base,
            "known_accuracy": None,
            "macro_f1": None,
            "balanced_accuracy": None,
            "f1_0": None,
            "f1_1": None,
            "confusion": {"tn": 0, "fp": 0, "fn": 0, "tp": 0},
        }
    true_known = labels[known]
    pred_known = prediction[known]
    tn, fp, fn, tp = confusion_matrix(true_known, pred_known, labels=[0, 1]).ravel()
    f1_0, f1_1 = f1_score(
        true_known, pred_known, labels=[0, 1], average=None, zero_division=0
    )
    return {
        **base,
        "known_accuracy": round(float(accuracy_score(true_known, pred_known)), 8),
        "macro_f1": round(float((f1_0 + f1_1) / 2.0), 8),
        "balanced_accuracy": round(
            float(balanced_accuracy_score(true_known, pred_known)), 8
        ),
        "f1_0": round(float(f1_0), 8),
        "f1_1": round(float(f1_1), 8),
        "confusion": {"tn": int(tn), "fp": int(fp), "fn": int(fn), "tp": int(tp)},
    }


def fit_row_text_model(
    train_documents: Sequence[str],
    test_documents: Sequence[str],
    train_labels: np.ndarray,
    sample_weight: np.ndarray,
    *,
    analyzer: str,
    ngram_range: tuple[int, int],
    max_features: int,
    c_value: float,
) -> tuple[np.ndarray, int]:
    vectorizer = TfidfVectorizer(
        analyzer=analyzer,
        ngram_range=ngram_range,
        min_df=2,
        max_df=0.997,
        max_features=max_features,
        sublinear_tf=True,
        strip_accents="unicode" if analyzer == "word" else None,
        dtype=np.float32,
    )
    train_matrix = vectorizer.fit_transform(train_documents)
    test_matrix = vectorizer.transform(test_documents)
    classifier = LogisticRegression(
        C=c_value,
        solver="liblinear",
        max_iter=1200,
        random_state=RANDOM_STATE,
    )
    classifier.fit(train_matrix, train_labels, sample_weight=sample_weight)
    return classifier.predict_proba(test_matrix)[:, 1], len(vectorizer.vocabulary_)


def inverse_target_multiplicity_weights(
    indices: np.ndarray,
    records: Sequence[dict[str, str]],
    canonical_targets: Sequence[str],
) -> np.ndarray:
    counts = Counter(
        (records[index]["REFNO"], canonical_targets[index]) for index in indices
    )
    return np.asarray(
        [
            1.0 / counts[(records[index]["REFNO"], canonical_targets[index])]
            for index in indices
        ],
        dtype=np.float64,
    )


def make_row_gradient_boosting() -> GradientBoostingClassifier:
    return GradientBoostingClassifier(
        n_estimators=220,
        learning_rate=0.035,
        max_depth=1,
        min_samples_leaf=12,
        random_state=RANDOM_STATE,
    )


def make_hist_gradient_boosting() -> HistGradientBoostingClassifier:
    return HistGradientBoostingClassifier(
        max_iter=240,
        learning_rate=0.04,
        max_leaf_nodes=7,
        min_samples_leaf=16,
        l2_regularization=0.5,
        random_state=RANDOM_STATE,
    )


def make_extra_trees() -> ExtraTreesClassifier:
    return ExtraTreesClassifier(
        n_estimators=350,
        max_features="sqrt",
        min_samples_leaf=5,
        class_weight=None,
        n_jobs=1,
        random_state=RANDOM_STATE,
    )


def best_accuracy_threshold(
    labels: np.ndarray,
    probabilities: np.ndarray,
) -> tuple[float, dict[str, Any]]:
    candidates = [
        metric_record(labels, probabilities, threshold=float(threshold))
        for threshold in np.arange(0.20, 0.801, 0.025)
    ]
    selected = max(
        candidates,
        key=lambda item: (
            item["accuracy"],
            item["macro_f1"],
            item["balanced_accuracy"],
            -abs(item["threshold"] - 0.5),
        ),
    )
    return float(selected["threshold"]), selected


def nested_row_threshold(
    outer_train_index: np.ndarray,
    records: Sequence[dict[str, str]],
    canonical_targets: Sequence[str],
    labels: np.ndarray,
    groups: np.ndarray,
    engineered_rows: np.ndarray,
) -> tuple[float, dict[str, Any]]:
    """Select a row-GB threshold from inner grouped OOF training predictions."""

    inner_probabilities = np.full(len(outer_train_index), np.nan, dtype=np.float64)
    inner_groups = groups[outer_train_index]
    splitter = GroupKFold(n_splits=4)
    for inner_train_position, inner_test_position in splitter.split(
        np.zeros(len(outer_train_index)),
        labels[outer_train_index],
        inner_groups,
    ):
        fit_index = outer_train_index[inner_train_position]
        validation_index = outer_train_index[inner_test_position]
        weights = inverse_target_multiplicity_weights(
            fit_index, records, canonical_targets
        )
        model = make_row_gradient_boosting()
        model.fit(
            engineered_rows[fit_index],
            labels[fit_index],
            sample_weight=weights,
        )
        inner_probabilities[inner_test_position] = model.predict_proba(
            engineered_rows[validation_index]
        )[:, 1]
    if not np.isfinite(inner_probabilities).all():
        raise AssertionError("nested row threshold has incomplete inner OOF predictions")
    threshold, selected = best_accuracy_threshold(
        labels[outer_train_index], inner_probabilities
    )
    return threshold, {
        "inner_rows": int(len(outer_train_index)),
        "inner_groups": int(len(set(inner_groups))),
        "inner_n_splits": 4,
        "selected_inner_oof_metric": selected,
    }


def group_representatives(
    indices: np.ndarray,
    groups: np.ndarray,
    labels: np.ndarray,
) -> tuple[list[str], np.ndarray]:
    members: dict[str, list[int]] = defaultdict(list)
    for index in indices:
        members[str(groups[index])].append(int(index))
    ordered = sorted(members)
    majority = np.asarray(
        [int(labels[members[group]].mean() >= 0.5) for group in ordered],
        dtype=np.int8,
    )
    return ordered, majority


def run_grouped_oof(
    records: Sequence[dict[str, str]],
    labels: np.ndarray,
    features: Mapping[str, Any],
) -> tuple[
    dict[str, np.ndarray],
    list[np.ndarray],
    dict[str, list[dict[str, Any]]],
    list[dict[str, Any]],
]:
    groups = np.asarray([record["REFNO"] for record in records], dtype=object)
    canonical_targets = [meta["canonical_target"] for meta in features["row_meta"]]
    probabilities = {
        name: np.full(len(records), np.nan, dtype=np.float64)
        for name in (
            "target_local_char",
            "target_local_word",
            "rockname_char",
            "row_knowledge_gradient_boosting",
            "row_knowledge_hist_gradient_boosting",
            "row_knowledge_extra_trees",
            "row_knowledge_gb_nested_threshold_decision",
            "article_knowledge_word_majority",
            "article_knowledge_gradient_boosting_majority",
            "article_knowledge_hist_gradient_boosting_majority",
            "article_knowledge_extra_trees_majority",
        )
    }
    fold_indices: list[np.ndarray] = []
    fold_metrics: dict[str, list[dict[str, Any]]] = defaultdict(list)
    nested_threshold_audit: list[dict[str, Any]] = []
    splitter = GroupKFold(n_splits=N_SPLITS)
    for fold, (train_index, test_index) in enumerate(
        splitter.split(np.zeros(len(records)), labels, groups), start=1
    ):
        fold_indices.append(test_index)
        weights = inverse_target_multiplicity_weights(
            train_index, records, canonical_targets
        )
        local_train = [features["local_documents"][index] for index in train_index]
        local_test = [features["local_documents"][index] for index in test_index]
        rock_train = [features["rockname_documents"][index] for index in train_index]
        rock_test = [features["rockname_documents"][index] for index in test_index]
        specs = (
            (
                "target_local_char", local_train, local_test, "char_wb", (3, 5), 100000, 2.0,
            ),
            (
                "target_local_word", local_train, local_test, "word", (1, 2), 80000, 2.0,
            ),
            (
                "rockname_char", rock_train, rock_test, "char_wb", (3, 5), 30000, 1.0,
            ),
        )
        for name, train_docs, test_docs, analyzer, ngrams, cap, c_value in specs:
            predicted, vocabulary_size = fit_row_text_model(
                train_docs,
                test_docs,
                labels[train_index],
                weights,
                analyzer=analyzer,
                ngram_range=ngrams,
                max_features=cap,
                c_value=c_value,
            )
            probabilities[name][test_index] = predicted
            record = metric_record(labels[test_index], predicted)
            record.update({"fold": fold, "vocabulary_size": vocabulary_size})
            fold_metrics[name].append(record)

        row_model = make_row_gradient_boosting()
        row_model.fit(
            features["engineered_rows"][train_index],
            labels[train_index],
            sample_weight=weights,
        )
        predicted = row_model.predict_proba(
            features["engineered_rows"][test_index]
        )[:, 1]
        name = "row_knowledge_gradient_boosting"
        probabilities[name][test_index] = predicted
        fold_metrics[name].append(
            {"fold": fold, **metric_record(labels[test_index], predicted)}
        )
        nested_threshold, nested_audit = nested_row_threshold(
            train_index,
            records,
            canonical_targets,
            labels,
            groups,
            features["engineered_rows"],
        )
        nested_decision = (predicted >= nested_threshold).astype(np.float64)
        nested_name = "row_knowledge_gb_nested_threshold_decision"
        probabilities[nested_name][test_index] = nested_decision
        nested_fold_metric = metric_record(labels[test_index], nested_decision)
        nested_fold_metric["applied_probability_threshold"] = nested_threshold
        fold_metrics[nested_name].append({"fold": fold, **nested_fold_metric})
        nested_threshold_audit.append(
            {
                "outer_fold": fold,
                "outer_test_rows": int(len(test_index)),
                "outer_test_groups": int(len(set(groups[test_index]))),
                "applied_probability_threshold": nested_threshold,
                **nested_audit,
            }
        )

        for name, row_model in (
            ("row_knowledge_hist_gradient_boosting", make_hist_gradient_boosting()),
            ("row_knowledge_extra_trees", make_extra_trees()),
        ):
            row_model.fit(
                features["engineered_rows"][train_index],
                labels[train_index],
                sample_weight=weights,
            )
            predicted_other = row_model.predict_proba(
                features["engineered_rows"][test_index]
            )[:, 1]
            probabilities[name][test_index] = predicted_other
            fold_metrics[name].append(
                {"fold": fold, **metric_record(labels[test_index], predicted_other)}
            )

        train_group_names, train_group_labels = group_representatives(
            train_index, groups, labels
        )
        test_group_names = sorted(set(str(groups[index]) for index in test_index))
        train_article_docs = [features["article_documents"][group] for group in train_group_names]
        test_article_docs = [features["article_documents"][group] for group in test_group_names]
        group_predictions, vocabulary_size = fit_row_text_model(
            train_article_docs,
            test_article_docs,
            train_group_labels,
            np.ones(len(train_group_labels), dtype=np.float64),
            analyzer="word",
            ngram_range=(1, 2),
            max_features=90000,
            c_value=2.0,
        )
        group_probability = dict(zip(test_group_names, group_predictions))
        name = "article_knowledge_word_majority"
        predicted = np.asarray(
            [group_probability[str(groups[index])] for index in test_index],
            dtype=np.float64,
        )
        probabilities[name][test_index] = predicted
        record = metric_record(labels[test_index], predicted)
        record.update({"fold": fold, "vocabulary_size": vocabulary_size})
        fold_metrics[name].append(record)

        group_feature_train = np.asarray(
            [features["article_features"][group] for group in train_group_names],
            dtype=np.float64,
        )
        group_feature_test = np.asarray(
            [features["article_features"][group] for group in test_group_names],
            dtype=np.float64,
        )
        article_models = (
            (
                "article_knowledge_gradient_boosting_majority",
                GradientBoostingClassifier(
                    n_estimators=200,
                    learning_rate=0.04,
                    max_depth=1,
                    min_samples_leaf=8,
                    random_state=RANDOM_STATE,
                ),
            ),
            (
                "article_knowledge_hist_gradient_boosting_majority",
                make_hist_gradient_boosting(),
            ),
            ("article_knowledge_extra_trees_majority", make_extra_trees()),
        )
        for name, article_model in article_models:
            article_model.fit(group_feature_train, train_group_labels)
            group_predictions = article_model.predict_proba(group_feature_test)[:, 1]
            group_probability = dict(zip(test_group_names, group_predictions))
            predicted_article = np.asarray(
                [group_probability[str(groups[index])] for index in test_index],
                dtype=np.float64,
            )
            probabilities[name][test_index] = predicted_article
            fold_metrics[name].append(
                {"fold": fold, **metric_record(labels[test_index], predicted_article)}
            )
        print(
            f"fold {fold}/{N_SPLITS}: local_char="
            f"{fold_metrics['target_local_char'][-1]['accuracy']:.4f}, "
            f"row_kb={fold_metrics['row_knowledge_gradient_boosting'][-1]['accuracy']:.4f}, "
            f"article_kb={fold_metrics['article_knowledge_gradient_boosting_majority'][-1]['accuracy']:.4f}",
            flush=True,
        )
    for name, values in probabilities.items():
        if not np.isfinite(values).all():
            raise AssertionError(f"incomplete OOF probability array: {name}")
    return probabilities, fold_indices, fold_metrics, nested_threshold_audit


def pool_by_key(
    probabilities: np.ndarray,
    keys: Sequence[tuple[str, ...]],
) -> np.ndarray:
    buckets: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for index, key in enumerate(keys):
        buckets[key].append(index)
    output = probabilities.copy()
    for indices in buckets.values():
        value = float(probabilities[indices].mean())
        output[indices] = value
    return output


def build_strategies(
    records: Sequence[dict[str, str]],
    features: Mapping[str, Any],
    base: Mapping[str, np.ndarray],
) -> dict[str, dict[str, Any]]:
    local = (
        0.10 * base["target_local_char"]
        + 0.05 * base["target_local_word"]
        + 0.35 * base["row_knowledge_gradient_boosting"]
        + 0.25 * base["row_knowledge_hist_gradient_boosting"]
        + 0.20 * base["row_knowledge_extra_trees"]
        + 0.05 * base["rockname_char"]
    )
    article = (
        0.10 * base["article_knowledge_word_majority"]
        + 0.35 * base["article_knowledge_gradient_boosting_majority"]
        + 0.30 * base["article_knowledge_hist_gradient_boosting_majority"]
        + 0.25 * base["article_knowledge_extra_trees_majority"]
    )
    target_keys = [
        (record["REFNO"], features["row_meta"][index]["canonical_target"])
        for index, record in enumerate(records)
    ]
    paper_keys = [(record["REFNO"],) for record in records]
    target_pooled = pool_by_key(local, target_keys)
    paper_mean = pool_by_key(target_pooled, paper_keys)
    target_counts = np.asarray(
        [features["row_meta"][index]["canonical_target_count"] for index in range(len(records))]
    )
    shared = np.asarray(
        [bool(features["row_meta"][index]["shared_evidence"]) for index in range(len(records))]
    )
    single = target_counts == 1
    multi = target_counts > 1

    strategies: dict[str, dict[str, Any]] = {}

    def add(name: str, probability: np.ndarray, family: str, rule: str) -> None:
        strategies[name] = {
            "probabilities": probability,
            "family": family,
            "rule": rule,
        }

    for name, probability in base.items():
        add(name, probability, "base_model", "direct grouped-OOF model output")
    add(
        "target_local_predefined_fusion",
        local,
        "target_local",
        "predefined local-text and three-model expert-count fusion",
    )
    add(
        "target_local_canonical_pool",
        target_pooled,
        "target_local",
        "average only rows sharing REFNO and canonical ROCKNAME target",
    )
    for weight in (0.25, 0.50, 0.75, 1.00):
        probability = target_pooled.copy()
        probability[single] = (
            (1.0 - weight) * target_pooled[single] + weight * article[single]
        )
        add(
            f"single_target_article_{int(weight * 100):02d}",
            probability,
            "single_target_article_propagation",
            f"single-target papers blend article probability at weight {weight:.2f}; multi-target rows remain target-local",
        )
    for weight in (0.25, 0.50, 0.75, 1.00):
        selected = multi & shared
        probability = target_pooled.copy()
        probability[selected] = (
            (1.0 - weight) * target_pooled[selected] + weight * article[selected]
        )
        add(
            f"multi_shared_article_{int(weight * 100):02d}",
            probability,
            "multi_target_shared_evidence",
            f"multi-target rows with shared original-text evidence blend article probability at weight {weight:.2f}",
        )
    for weight in (0.25, 0.50, 0.75, 1.00):
        probability = target_pooled.copy()
        probability[multi] = (
            (1.0 - weight) * target_pooled[multi] + weight * article[multi]
        )
        add(
            f"multi_all_article_{int(weight * 100):02d}",
            probability,
            "multi_target_article_ablation",
            f"ablation: every multi-target row blends article probability at weight {weight:.2f}",
        )
    for weight in (0.25, 0.50, 0.75, 1.00):
        selected = single | (multi & shared)
        probability = target_pooled.copy()
        probability[selected] = (
            (1.0 - weight) * target_pooled[selected] + weight * article[selected]
        )
        add(
            f"scope_gated_article_{int(weight * 100):02d}",
            probability,
            "scope_gated_combination",
            f"single-target or shared-evidence multi-target rows blend article probability at weight {weight:.2f}",
        )
    add(
        "paper_mean_target_local",
        paper_mean,
        "paper_majority_layer",
        "mean target-local probability within REFNO, propagated to all rows",
    )
    add(
        "paper_majority_vote_target_local",
        (paper_mean >= 0.5).astype(np.float64),
        "paper_majority_layer",
        "hard majority/mean-threshold target-local decision within REFNO",
    )
    add(
        "paper_text_majority_predefined_fusion",
        article,
        "paper_majority_layer",
        "predefined article knowledge-text and three-model expert-count majority fusion",
    )
    for weight in (0.25, 0.50, 0.75):
        probability = (1.0 - weight) * target_pooled + weight * paper_mean
        add(
            f"target_plus_paper_mean_{int(weight * 100):02d}",
            probability,
            "paper_majority_layer",
            f"canonical target probability blended with within-paper target mean at weight {weight:.2f}",
        )
    return strategies


def fold_records(
    labels: np.ndarray,
    probabilities: np.ndarray,
    fold_indices: Sequence[np.ndarray],
) -> list[dict[str, Any]]:
    return [
        {"fold": fold, **metric_record(labels[index], probabilities[index])}
        for fold, index in enumerate(fold_indices, start=1)
    ]


def selective_analysis(
    labels: np.ndarray,
    probabilities: np.ndarray,
    fold_indices: Sequence[np.ndarray],
) -> dict[str, Any]:
    preset = [
        selective_metric(labels, probabilities, negative, positive)
        for negative, positive in ((0.20, 0.80), (0.25, 0.75), (0.30, 0.70), (0.35, 0.75), (0.40, 0.80))
    ]
    grid = [
        selective_metric(labels, probabilities, negative, positive)
        for negative in np.arange(0.05, 0.501, 0.025)
        for positive in np.arange(0.50, 0.951, 0.025)
    ]
    eligible = [
        item
        for item in grid
        if item["known_accuracy"] is not None
        and item["known_accuracy"] >= 0.90
        and item["f1_0"] >= 0.50
        and item["f1_1"] >= 0.50
    ]
    selected = max(
        eligible,
        key=lambda item: (item["coverage"], item["known_accuracy"], item["macro_f1"]),
    ) if eligible else None
    folds: list[dict[str, Any]] = []
    if selected:
        folds = [
            {
                "fold": fold,
                **selective_metric(
                    labels[index],
                    probabilities[index],
                    selected["negative_threshold"],
                    selected["positive_threshold"],
                ),
            }
            for fold, index in enumerate(fold_indices, start=1)
        ]
    return {
        "preset_operating_points": preset,
        "posthoc_selected_operating_point": selected,
        "posthoc_selected_folds": folds,
        "selection": {
            "candidate_count": len(grid),
            "rule": "maximum coverage subject to known accuracy >=0.90 and both class F1 >=0.50",
            "warning": "threshold selected on the same grouped-OOF training predictions; external REFNO validation is required",
        },
    }


def posthoc_full_threshold_analysis(
    labels: np.ndarray,
    probabilities: np.ndarray,
) -> dict[str, Any]:
    curve = [
        metric_record(labels, probabilities, threshold=float(threshold))
        for threshold in np.arange(0.20, 0.801, 0.025)
    ]
    selected = max(
        curve,
        key=lambda item: (
            item["accuracy"], item["macro_f1"], item["balanced_accuracy"]
        ),
    )
    return {
        "selected": selected,
        "curve": curve,
        "warning": "full threshold selected post-hoc on all grouped-OOF training predictions; diagnostic upper bound only",
    }


def subset_metric(
    labels: np.ndarray,
    probabilities: np.ndarray,
    mask: np.ndarray,
) -> dict[str, Any] | None:
    if not mask.any():
        return None
    return metric_record(labels[mask], probabilities[mask])


def paper_homogeneity_masks(
    records: Sequence[dict[str, str]],
    labels: np.ndarray,
) -> tuple[np.ndarray, np.ndarray]:
    values_by_refno: dict[str, set[int]] = defaultdict(set)
    for index, record in enumerate(records):
        values_by_refno[record["REFNO"]].add(int(labels[index]))
    homogeneous = np.asarray(
        [len(values_by_refno[record["REFNO"]]) == 1 for record in records]
    )
    return homogeneous, ~homogeneous


def error_source_analysis(
    records: Sequence[dict[str, str]],
    labels: np.ndarray,
    probabilities: np.ndarray,
    features: Mapping[str, Any],
) -> dict[str, Any]:
    predicted = probabilities >= 0.5
    false_positive = predicted & (labels == 0)
    false_negative = (~predicted) & (labels == 1)
    metadata = features["row_meta"]
    masks = {
        "single_target_article": np.asarray([row["canonical_target_count"] == 1 for row in metadata]),
        "multi_target_article": np.asarray([row["canonical_target_count"] > 1 for row in metadata]),
        "shared_evidence_detected": np.asarray([bool(row["shared_evidence"]) for row in metadata]),
        "no_shared_evidence_detected": np.asarray([not row["shared_evidence"] for row in metadata]),
        "target_not_anchored": np.asarray([row["target_hit_count"] == 0 for row in metadata]),
        "target_anchored": np.asarray([row["target_hit_count"] > 0 for row in metadata]),
        "zero_local_expert_hits": np.asarray([row["local_knowledge_hit_count"] == 0 for row in metadata]),
        "article_only_expert_evidence": np.asarray([
            row["local_knowledge_hit_count"] == 0 and row["knowledge_hit_count"] > 0
            for row in metadata
        ]),
        "high_article_low_local_evidence": np.asarray([
            row["knowledge_hit_count"] >= 10 and row["local_knowledge_hit_count"] <= 1
            for row in metadata
        ]),
    }
    output: dict[str, Any] = {}
    for name, mask in masks.items():
        count = int(mask.sum())
        if not count:
            continue
        output[name] = {
            "rows": count,
            "row_fraction": round(count / len(labels), 8),
            "false_positive": int((false_positive & mask).sum()),
            "false_negative": int((false_negative & mask).sum()),
            "false_positive_rate_within_subset": round(
                float((false_positive & mask).sum() / count), 8
            ),
            "error_rate_within_subset": round(
                float(((predicted != labels) & mask).sum() / count), 8
            ),
        }
    groups: dict[str, list[int]] = defaultdict(list)
    for index, record in enumerate(records):
        groups[record["REFNO"]].append(index)
    output["prediction_dispersion"] = {
        "papers_with_mixed_binary_predictions": sum(
            len(set(int(predicted[index]) for index in indices)) > 1
            for indices in groups.values()
        ),
        "papers_with_mixed_labels_diagnostic_only": sum(
            len(set(int(labels[index]) for index in indices)) > 1
            for indices in groups.values()
        ),
    }
    return output


def concept_error_analysis(
    concepts: Sequence[Mapping[str, Any]],
    labels: np.ndarray,
    probabilities: np.ndarray,
    features: Mapping[str, Any],
) -> dict[str, Any]:
    predicted = probabilities >= 0.5
    false_positive = predicted & (labels == 0)
    false_negative = (~predicted) & (labels == 1)

    def records_for(counts: np.ndarray) -> list[dict[str, Any]]:
        output: list[dict[str, Any]] = []
        for index, concept in enumerate(concepts):
            present = counts[:, index] > 0
            rows = int(present.sum())
            if not rows:
                continue
            negatives = int(((labels == 0) & present).sum())
            output.append(
                {
                    "source": concept["source"],
                    "category": concept["category"],
                    "concept_id": concept["concept_id"],
                    "rows_present": rows,
                    "positive_label_rate": round(float(labels[present].mean()), 8),
                    "false_positive": int((false_positive & present).sum()),
                    "false_negative": int((false_negative & present).sum()),
                    "false_positive_share_of_present_negative_rows": (
                        round(float((false_positive & present).sum() / negatives), 8)
                        if negatives
                        else None
                    ),
                    "error_rate": round(float((predicted[present] != labels[present]).mean()), 8),
                }
            )
        return output

    local = records_for(features["local_concept_counts"])
    article = records_for(features["article_concept_counts"])
    return {
        "local_top_by_false_positive_count": sorted(
            local, key=lambda item: (item["false_positive"], item["rows_present"]), reverse=True
        )[:15],
        "article_top_by_false_positive_count": sorted(
            article, key=lambda item: (item["false_positive"], item["rows_present"]), reverse=True
        )[:15],
        "local_all_concepts": local,
        "article_all_concepts": article,
    }


def oracle_diagnostics(
    records: Sequence[dict[str, str]],
    labels: np.ndarray,
    canonical_targets: Sequence[str],
) -> dict[str, Any]:
    by_paper: dict[str, list[int]] = defaultdict(list)
    by_target: dict[tuple[str, str], list[int]] = defaultdict(list)
    for index, record in enumerate(records):
        by_paper[record["REFNO"]].append(int(labels[index]))
        by_target[(record["REFNO"], canonical_targets[index])].append(int(labels[index]))

    def majority_ceiling(groups: Mapping[Any, Sequence[int]]) -> dict[str, Any]:
        correct = sum(max(Counter(values).values()) for values in groups.values())
        return {
            "rows": len(labels),
            "group_count": len(groups),
            "accuracy_ceiling": round(correct / len(labels), 8),
            "mixed_label_groups": sum(len(set(values)) > 1 for values in groups.values()),
            "rows_in_mixed_label_groups": sum(
                len(values) for values in groups.values() if len(set(values)) > 1
            ),
        }
    return {
        "paper_majority_oracle": majority_ceiling(by_paper),
        "canonical_target_majority_oracle": majority_ceiling(by_target),
        "warning": "diagnostic label-consistency ceilings only; not deployable strategies and not counted as model results",
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--positive-kb", type=Path, default=DEFAULT_POSITIVE_KB)
    parser.add_argument("--without-positive-kb", action="store_true")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--probabilities",
        type=Path,
        help="optional NPZ of row-aligned OOF probabilities without labels or IDs",
    )
    args = parser.parse_args(argv)
    require_experiment_dependencies()

    required = {
        "training_identity": args.identity.resolve(),
        "targets": args.targets.resolve(),
        "page_cache": args.page_cache.resolve(),
        "taxonomy": args.taxonomy.resolve(),
    }
    for label, path in required.items():
        if not path.is_file():
            raise SystemExit(f"missing {label}: {path}")
    records, coverage = load_pdf_only_identity(
        required["training_identity"], args.pdf_dir.resolve()
    )
    # The training identity membership gate is complete before R3 cells are read.
    labels_by_resultno = load_training_labels(
        required["targets"], {record["RESULTNO"] for record in records}
    )
    labels = np.asarray(
        [labels_by_resultno[record["RESULTNO"]] for record in records], dtype=np.int8
    )
    refnos = {record["REFNO"] for record in records}
    paper_text = load_page_text(required["page_cache"], refnos)
    positive_path = None if args.without_positive_kb else args.positive_kb.resolve()
    concepts, knowledge_audit = load_knowledge_patterns(
        required["taxonomy"], positive_path
    )
    features = build_features(records, paper_text, concepts)
    (
        base_probabilities,
        fold_indices,
        base_fold_metrics,
        nested_threshold_audit,
    ) = run_grouped_oof(
        records, labels, features
    )
    strategies = build_strategies(records, features, base_probabilities)
    homogeneous_mask, mixed_mask = paper_homogeneity_masks(records, labels)
    strategy_results: dict[str, Any] = {}
    for name, strategy in strategies.items():
        probability = strategy["probabilities"]
        strategy_results[name] = {
            "family": strategy["family"],
            "rule": strategy["rule"],
            "aggregate": metric_record(labels, probability),
            "folds": fold_records(labels, probability, fold_indices),
            "diagnostic_label_homogeneity": {
                "homogeneous_papers": subset_metric(
                    labels, probability, homogeneous_mask
                ),
                "mixed_label_papers": subset_metric(labels, probability, mixed_mask),
            },
        }
        if name in base_fold_metrics:
            strategy_results[name]["model_folds"] = base_fold_metrics[name]

    ranked = sorted(
        strategy_results,
        key=lambda name: (
            strategy_results[name]["aggregate"]["accuracy"],
            strategy_results[name]["aggregate"]["macro_f1"],
            strategy_results[name]["aggregate"]["balanced_accuracy"],
        ),
        reverse=True,
    )
    best_name = ranked[0]
    strict_scope_families = {
        "target_local",
        "single_target_article_propagation",
        "multi_target_shared_evidence",
        "scope_gated_combination",
    }
    strict_ranked = [
        name
        for name in ranked
        if strategy_results[name]["family"] in strict_scope_families
    ]
    best_strict_name = strict_ranked[0]
    best_probability = strategies[best_name]["probabilities"]
    probability_artifact: dict[str, Any] | None = None
    if args.probabilities is not None:
        probability_path = args.probabilities.resolve()
        probability_path.parent.mkdir(parents=True, exist_ok=True)
        np.savez_compressed(
            probability_path,
            **{
                name: np.asarray(strategy["probabilities"], dtype=np.float64)
                for name, strategy in sorted(strategies.items())
            },
        )
        probability_artifact = {
            "path": str(probability_path),
            "sha256": file_sha256(probability_path),
            "rows_per_array": len(records),
            "contains_labels": False,
            "contains_record_identifiers": False,
        }
    metadata = features["row_meta"]
    single_mask = np.asarray([row["canonical_target_count"] == 1 for row in metadata])
    multi_mask = ~single_mask
    shared_mask = np.asarray([bool(row["shared_evidence"]) for row in metadata])
    label_counts = Counter(int(value) for value in labels)
    result = {
        "schema_version": "r3_pdf_only_target_attribution_grouped_oof_v1",
        "status": "complete",
        "scope": "development-only frozen training partition; rows without a local PDF excluded",
        "guardrails": {
            "cross_validation": "5-fold GroupKFold by REFNO",
            "labels": "R3 dereferenced only after fixed-training and local-PDF membership gates",
            "allowed_runtime_feature_sources": [
                "ROCKNAME",
                "original-paper text from page cache for a verified local PDF",
                "frozen taxonomy runtime original-text patterns",
                "optional frozen positive ROCKMAG expert-KB runtime original-text patterns",
            ],
            "forbidden_sources_not_read": [
                "held-out R3 labels",
                "PMAGRESULT",
                "ROCKMAG source cells",
                "prior predictions",
                "prior evaluations",
            ],
            "identifier_feature_exclusion": [
                "REFNO used only for alignment, scope aggregation, and GroupKFold",
                "RESULTNO used only to align a permitted training label",
            ],
            "formal_scorer_modified": False,
        },
        "inputs": {
            **{
                label: {"path": str(path), "sha256": file_sha256(path)}
                for label, path in required.items()
            },
            "pdf_dir": str(args.pdf_dir.resolve()),
        },
        "knowledge": knowledge_audit,
        "data": {
            **coverage,
            "label_counts": {str(key): value for key, value in sorted(label_counts.items())},
            "single_target_papers": len(
                {record["REFNO"] for index, record in enumerate(records) if single_mask[index]}
            ),
            "single_target_rows": int(single_mask.sum()),
            "multi_target_papers": len(
                {record["REFNO"] for index, record in enumerate(records) if multi_mask[index]}
            ),
            "multi_target_rows": int(multi_mask.sum()),
            "shared_evidence_papers": len(
                {record["REFNO"] for index, record in enumerate(records) if shared_mask[index]}
            ),
            "shared_evidence_rows": int(shared_mask.sum()),
            "homogeneous_label_papers_diagnostic_only": len(
                {record["REFNO"] for index, record in enumerate(records) if homogeneous_mask[index]}
            ),
            "homogeneous_label_rows_diagnostic_only": int(homogeneous_mask.sum()),
            "mixed_label_papers_diagnostic_only": len(
                {record["REFNO"] for index, record in enumerate(records) if mixed_mask[index]}
            ),
            "mixed_label_rows_diagnostic_only": int(mixed_mask.sum()),
            "all_pdf_rows_have_nonempty_page_text": True,
            "engineered_feature_count": int(features["engineered_rows"].shape[1]),
        },
        "oracle_diagnostics": oracle_diagnostics(
            records,
            labels,
            [row["canonical_target"] for row in metadata],
        ),
        "strategies": strategy_results,
        "ranking_by_full_coverage_accuracy": ranked,
        "ranking_by_strict_scope_accuracy": strict_ranked,
        "best_full_coverage": {
            "strategy": best_name,
            "deployment_status": (
                "diagnostic_unscoped_ablation"
                if strategy_results[best_name]["family"]
                == "multi_target_article_ablation"
                else "scope_constrained_candidate"
            ),
            **strategy_results[best_name]["aggregate"],
        },
        "best_strict_scope_candidate": {
            "strategy": best_strict_name,
            "family": strategy_results[best_strict_name]["family"],
            **strategy_results[best_strict_name]["aggregate"],
        },
        "best_strategy_subsets": {
            "single_target": subset_metric(labels, best_probability, single_mask),
            "multi_target": subset_metric(labels, best_probability, multi_mask),
            "multi_target_shared": subset_metric(
                labels, best_probability, multi_mask & shared_mask
            ),
            "multi_target_not_shared": subset_metric(
                labels, best_probability, multi_mask & ~shared_mask
            ),
            "homogeneous_papers_diagnostic": subset_metric(
                labels, best_probability, homogeneous_mask
            ),
            "mixed_label_papers_diagnostic": subset_metric(
                labels, best_probability, mixed_mask
            ),
        },
        "nested_threshold": {
            "strategy": "row_knowledge_gb_nested_threshold_decision",
            "aggregate": strategy_results[
                "row_knowledge_gb_nested_threshold_decision"
            ]["aggregate"],
            "outer_fold_audit": nested_threshold_audit,
            "selection_boundary": "each outer test fold threshold selected only from 4-fold REFNO-grouped inner OOF predictions over the outer training papers",
        },
        "posthoc_full_threshold": {
            "probability_source": best_name,
            **posthoc_full_threshold_analysis(labels, best_probability),
        },
        "selective_policy": {
            "probability_source": best_name,
            **selective_analysis(labels, best_probability, fold_indices),
        },
        "error_source_analysis": error_source_analysis(
            records, labels, best_probability, features
        ),
        "concept_error_analysis": concept_error_analysis(
            concepts, labels, best_probability, features
        ),
        "interpretation": {
            "full_coverage_target_accuracy": 0.90,
            "full_coverage_target_reached": bool(
                strategy_results[best_name]["aggregate"]["accuracy"] >= 0.90
            ),
            "rows_without_pdf_excluded_from_every_metric": True,
            "posthoc_strategy_ranking_requires_external_validation": True,
            "selective_threshold_requires_external_validation": True,
        },
        "artifacts": {
            "script": str(Path(__file__).resolve()),
            "script_sha256": file_sha256(Path(__file__).resolve()),
            "output": str(args.output.resolve()),
            "oof_probabilities": probability_artifact,
        },
    }
    atomic_json(args.output.resolve(), result)
    print(json.dumps(result["best_full_coverage"], ensure_ascii=False, indent=2))
    selected = result["selective_policy"]["posthoc_selected_operating_point"]
    print(json.dumps({"selective": selected}, ensure_ascii=False, indent=2))
    print(f"wrote {args.output.resolve()}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
