#!/usr/bin/env python3
"""Audit R3/R5/R7 predictions and knowledge artifacts for scoring leakage.

The audit is intentionally structural: it parses JSON embedded in workbooks,
walks JSON/JSONL recursively, and emits stable violation codes with exact
locations. Any violation makes the command exit non-zero.
"""

from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
import re
from typing import Any, Iterable, Iterator, Mapping, Sequence

from openpyxl import load_workbook


REPORT_SCHEMA_VERSION = "r357_strict_scoring_boundary_audit_v1"
FORBIDDEN_R5_MODEL = "r5_paper_text_model_v1"
PREDICTION_EXTENSIONS = {".xlsx", ".json", ".jsonl"}
KNOWLEDGE_EXTENSIONS = {".json", ".jsonl"}

RUNTIME_SOURCE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("PMAGRESULT", re.compile(r"\bpmagresult\b", re.I)),
    ("ROCKUNIT", re.compile(r"\brockunit\b", re.I)),
    ("evaluation", re.compile(r"\bevaluations?\b", re.I)),
    ("label", re.compile(r"\blabels?\b", re.I)),
    (
        "previous workbook",
        re.compile(
            r"\b(?:previous|prior|existing|earlier|legacy)\s+"
            r"(?:(?:prediction|result|evaluation)\s+)?workbooks?\b",
            re.I,
        ),
    ),
)
RECORD_MAPPING_KEY = re.compile(
    r"(?:^|_)(?:refno|resultno)(?:_|$)|"
    r"(?:^|_)(?:record_level_(?:map|mapping|lookup)|"
    r"answer_lookup|prediction_lookup|label_lookup)(?:_|$)",
    re.I,
)
COMPOSITE_RECORD_ID = re.compile(r"^\d+(?:\s*[:/|,-]\s*\d+)+$")


def _canonical_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def _path_key(path: str, key: Any) -> str:
    text = str(key)
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", text):
        return f"{path}.{text}"
    return f"{path}[{json.dumps(text, ensure_ascii=False)}]"


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, (Mapping, list, tuple, set)):
        return len(value) == 0
    return False


def _display(value: Any, limit: int = 500) -> str:
    try:
        rendered = json.dumps(value, ensure_ascii=False, sort_keys=True)
    except (TypeError, ValueError):
        rendered = repr(value)
    return rendered if len(rendered) <= limit else rendered[: limit - 3] + "..."


def _violation(
    code: str,
    file: Path,
    location: str,
    message: str,
    value: Any = None,
) -> dict[str, Any]:
    result = {
        "code": code,
        "severity": "ERROR",
        "file": str(file.resolve()),
        "location": location,
        "message": message,
    }
    if value is not None:
        result["value"] = _display(value)
    return result


def _walk(value: Any, path: str = "$") -> Iterator[tuple[str, Any]]:
    """Yield every value, including containers, with a JSON-style path."""

    yield path, value
    if isinstance(value, Mapping):
        for key, child in value.items():
            yield from _walk(child, _path_key(path, key))
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _walk(child, f"{path}[{index}]")


def _mapping_items(value: Any, path: str = "$") -> Iterator[tuple[str, Any, Any]]:
    if isinstance(value, Mapping):
        for key, child in value.items():
            child_path = _path_key(path, key)
            yield child_path, key, child
            yield from _mapping_items(child, child_path)
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _mapping_items(child, f"{path}[{index}]")


def _strings(value: Any, path: str) -> Iterator[tuple[str, str]]:
    if isinstance(value, str):
        yield path, value
    elif isinstance(value, Mapping):
        for key, child in value.items():
            yield from _strings(child, _path_key(path, key))
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _strings(child, f"{path}[{index}]")


def _detect_indicators(record: Any) -> set[str]:
    indicators: set[str] = set()
    for _, key, _ in _mapping_items(record):
        normalized = str(key).upper()
        if normalized in {"R3", "R5", "R7"}:
            indicators.add(normalized)
        match = re.fullmatch(r"(R[357])_RESULT_JSON", normalized)
        if match:
            indicators.add(match.group(1))
    return indicators


def _parse_json_cell(
    value: Any,
    *,
    expected: bool,
    file: Path,
    location: str,
    violations: list[dict[str, Any]],
) -> Any:
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    should_parse = expected or stripped.startswith("{") or stripped.startswith("[")
    if not should_parse or not stripped:
        return value
    try:
        return json.loads(stripped)
    except json.JSONDecodeError as exc:
        if expected:
            violations.append(
                _violation(
                    "invalid_prediction_json",
                    file,
                    location,
                    f"Prediction JSON could not be parsed: {exc}",
                    value,
                )
            )
        return value


def _prediction_records_xlsx(
    path: Path, violations: list[dict[str, Any]]
) -> list[tuple[str, dict[str, Any]]]:
    records: list[tuple[str, dict[str, Any]]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            raw_header = next(rows, ())
            header = [str(value).strip() if value is not None else "" for value in raw_header]
            if not header or not any(header):
                violations.append(
                    _violation(
                        "missing_prediction_header",
                        path,
                        worksheet.title,
                        "Worksheet has no header row.",
                    )
                )
                continue
            seen: set[str] = set()
            for column in header:
                folded = column.casefold()
                if column and folded in seen:
                    violations.append(
                        _violation(
                            "duplicate_prediction_column",
                            path,
                            f"{worksheet.title}!row 1",
                            f"Duplicate prediction column: {column}",
                        )
                    )
                seen.add(folded)
            for row_number, values in enumerate(rows, start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                record: dict[str, Any] = {}
                for index, column in enumerate(header):
                    if not column:
                        continue
                    value = values[index] if index < len(values) else None
                    location = f"{worksheet.title}!row {row_number}.{column}"
                    record[column] = _parse_json_cell(
                        value,
                        expected=column.upper().endswith("_RESULT_JSON"),
                        file=path,
                        location=location,
                        violations=violations,
                    )
                records.append((f"{worksheet.title}!row {row_number}", record))
    finally:
        workbook.close()
    return records


def _prediction_records_json(
    path: Path, violations: list[dict[str, Any]]
) -> list[tuple[str, Any]]:
    records: list[tuple[str, Any]] = []
    if path.suffix.casefold() == ".jsonl":
        with path.open(encoding="utf-8") as handle:
            for line_number, line in enumerate(handle, start=1):
                if not line.strip():
                    continue
                try:
                    records.append((f"line {line_number}", json.loads(line)))
                except json.JSONDecodeError as exc:
                    violations.append(
                        _violation(
                            "invalid_prediction_json",
                            path,
                            f"line {line_number}",
                            f"Prediction JSONL row could not be parsed: {exc}",
                            line.strip(),
                        )
                    )
        return records
    try:
        document = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        violations.append(
            _violation(
                "invalid_prediction_json",
                path,
                "$",
                f"Prediction JSON could not be parsed: {exc}",
            )
        )
        return records
    if isinstance(document, list):
        records.extend((f"$[{index}]", row) for index, row in enumerate(document))
    else:
        records.append(("$", document))
    return records


def load_prediction_records(
    path: Path, violations: list[dict[str, Any]]
) -> list[tuple[str, Any]]:
    if path.suffix.casefold() == ".xlsx":
        return _prediction_records_xlsx(path, violations)
    return _prediction_records_json(path, violations)


def inspect_prediction_record(
    record: Any,
    *,
    file: Path,
    location: str,
) -> list[dict[str, Any]]:
    violations: list[dict[str, Any]] = []
    runtime_fields = 0
    indicators = _detect_indicators(record)
    for node_path, key, value in _mapping_items(record, location):
        canonical = _canonical_key(key)
        if canonical == "runtimescoringsources":
            runtime_fields += 1
            if _is_empty(value):
                violations.append(
                    _violation(
                        "empty_runtime_scoring_sources",
                        file,
                        node_path,
                        "runtime_scoring_sources must identify the actual scoring inputs.",
                    )
                )
                continue
            source_strings = list(_strings(value, node_path))
            if not source_strings:
                violations.append(
                    _violation(
                        "invalid_runtime_scoring_sources",
                        file,
                        node_path,
                        "runtime_scoring_sources must contain textual source descriptions.",
                        value,
                    )
                )
            for source_path, source in source_strings:
                for forbidden_name, pattern in RUNTIME_SOURCE_PATTERNS:
                    if pattern.search(source):
                        violations.append(
                            _violation(
                                "forbidden_runtime_scoring_source",
                                file,
                                source_path,
                                f"Runtime scoring source references forbidden input: {forbidden_name}.",
                                source,
                            )
                        )
        elif canonical == "sourcemetadata" and not _is_empty(value):
            violations.append(
                _violation(
                    "nonempty_source_metadata",
                    file,
                    node_path,
                    "Strict predictions require source_metadata to be absent or empty.",
                    value,
                )
            )

    if runtime_fields == 0:
        violations.append(
            _violation(
                "missing_runtime_scoring_sources",
                file,
                location,
                "Prediction record does not declare runtime_scoring_sources.",
            )
        )

    if "R5" in indicators:
        for string_path, text in _strings(record, location):
            if FORBIDDEN_R5_MODEL in text.casefold():
                violations.append(
                    _violation(
                        "forbidden_r5_calibration_model",
                        file,
                        string_path,
                        f"Strict R5 predictions may not use {FORBIDDEN_R5_MODEL}.",
                        text,
                    )
                )
    return violations


def _valid_provenance(value: Any) -> bool:
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, Mapping):
        return bool(value) and any(not _is_empty(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return bool(value) and any(not _is_empty(item) for item in value)
    return False


def _is_true(value: Any) -> bool:
    return value is True or (
        isinstance(value, str) and value.strip().casefold() in {"true", "yes", "1"}
    ) or (isinstance(value, int) and not isinstance(value, bool) and value == 1)


def _looks_like_record_id(key: Any) -> bool:
    text = str(key).strip()
    return bool(text.isdigit() or COMPOSITE_RECORD_ID.fullmatch(text))


def _looks_like_answer(value: Any) -> bool:
    if isinstance(value, bool):
        return True
    if isinstance(value, (int, float)) and value in {0, 1}:
        return True
    if isinstance(value, str) and value.strip().casefold() in {
        "0",
        "1",
        "positive",
        "negative",
        "pass",
        "fail",
    }:
        return True
    if isinstance(value, Mapping):
        keys = {_canonical_key(key) for key in value}
        return bool(keys & {"r3", "r5", "r7", "label", "answer", "outcome", "score"})
    return False


def _knowledge_documents(
    path: Path, violations: list[dict[str, Any]]
) -> list[tuple[str, Any]]:
    documents: list[tuple[str, Any]] = []
    if path.suffix.casefold() == ".jsonl":
        with path.open(encoding="utf-8") as handle:
            for line_number, line in enumerate(handle, start=1):
                if not line.strip():
                    continue
                try:
                    documents.append((f"line {line_number}", json.loads(line)))
                except json.JSONDecodeError as exc:
                    violations.append(
                        _violation(
                            "invalid_knowledge_json",
                            path,
                            f"line {line_number}",
                            f"Knowledge JSONL row could not be parsed: {exc}",
                            line.strip(),
                        )
                    )
        return documents
    try:
        root = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        violations.append(
            _violation(
                "invalid_knowledge_json",
                path,
                "$",
                f"Knowledge artifact could not be parsed: {exc}",
            )
        )
        return documents
    if isinstance(root, list):
        documents.extend((f"$[{index}]", item) for index, item in enumerate(root))
    else:
        documents.append(("$", root))
    return documents


def inspect_knowledge_document(
    document: Any,
    *,
    file: Path,
    location: str,
) -> list[dict[str, Any]]:
    violations: list[dict[str, Any]] = []
    if not isinstance(document, Mapping):
        violations.append(
            _violation(
                "invalid_knowledge_document",
                file,
                location,
                "A knowledge artifact document must be a JSON object.",
                document,
            )
        )
        return violations

    provenance = document.get("provenance")
    provenance_path = _path_key(location, "provenance")
    if provenance is None and isinstance(document.get("metadata"), Mapping):
        provenance = document["metadata"].get("provenance")
        provenance_path = _path_key(_path_key(location, "metadata"), "provenance")
    if not _valid_provenance(provenance):
        violations.append(
            _violation(
                "missing_knowledge_provenance",
                file,
                provenance_path,
                "Knowledge artifact requires non-empty provenance.",
            )
        )

    for node_path, key, value in _mapping_items(document, location):
        canonical = _canonical_key(key)
        if canonical in {"refno", "resultno"}:
            violations.append(
                _violation(
                    "forbidden_identity_key_in_knowledge",
                    file,
                    node_path,
                    f"Knowledge artifacts may not contain record identity key {key!r}.",
                    value,
                )
            )
        if canonical == "labelsupervision" and _is_true(value):
            violations.append(
                _violation(
                    "label_supervision_enabled",
                    file,
                    node_path,
                    "Strict knowledge artifacts require label_supervision=false.",
                    value,
                )
            )
        normalized_key = re.sub(r"[^a-z0-9]+", "_", str(key).casefold()).strip("_")
        if isinstance(value, Mapping) and RECORD_MAPPING_KEY.search(normalized_key):
            violations.append(
                _violation(
                    "record_level_mapping_in_knowledge",
                    file,
                    node_path,
                    "Knowledge artifact contains a record-level lookup/mapping container.",
                    list(value)[:20],
                )
            )

        if isinstance(value, Mapping):
            identifier_items = [
                (map_key, map_value)
                for map_key, map_value in value.items()
                if _looks_like_record_id(map_key)
            ]
            answer_items = [
                (map_key, map_value)
                for map_key, map_value in identifier_items
                if _looks_like_answer(map_value)
            ]
            if len(identifier_items) >= 2 or answer_items:
                violations.append(
                    _violation(
                        "record_level_mapping_in_knowledge",
                        file,
                        node_path,
                        "Knowledge artifact maps record-like identifiers to values.",
                        [map_key for map_key, _ in identifier_items[:20]],
                    )
                )

    for string_path, text in _strings(document, location):
        if FORBIDDEN_R5_MODEL in text.casefold():
            violations.append(
                _violation(
                    "forbidden_r5_calibration_model",
                    file,
                    string_path,
                    f"Strict scoring knowledge may not include {FORBIDDEN_R5_MODEL}.",
                    text,
                )
            )
    return violations


def _expand_paths(paths: Iterable[Path], extensions: set[str]) -> list[Path]:
    expanded: list[Path] = []
    for raw_path in paths:
        path = raw_path.resolve()
        if path.is_dir():
            expanded.extend(
                candidate.resolve()
                for candidate in path.rglob("*")
                if candidate.is_file() and candidate.suffix.casefold() in extensions
            )
        else:
            expanded.append(path)
    return sorted(set(expanded), key=str)


def audit_scoring_boundary(
    prediction_paths: Iterable[Path],
    knowledge_artifact_paths: Iterable[Path],
) -> dict[str, Any]:
    predictions = _expand_paths(prediction_paths, PREDICTION_EXTENSIONS)
    artifacts = _expand_paths(knowledge_artifact_paths, KNOWLEDGE_EXTENSIONS)
    violations: list[dict[str, Any]] = []
    files: list[dict[str, Any]] = []

    if not predictions and not artifacts:
        violations.append(
            _violation(
                "no_audit_inputs",
                Path.cwd(),
                "$",
                "At least one prediction or knowledge artifact is required.",
            )
        )

    for path in predictions:
        before = len(violations)
        if not path.is_file():
            violations.append(
                _violation("missing_input_file", path, "$", "Prediction file does not exist.")
            )
            record_count = 0
        elif path.suffix.casefold() not in PREDICTION_EXTENSIONS:
            violations.append(
                _violation(
                    "unsupported_prediction_format",
                    path,
                    "$",
                    f"Unsupported prediction format: {path.suffix}",
                )
            )
            record_count = 0
        else:
            try:
                records = load_prediction_records(path, violations)
            except (OSError, ValueError) as exc:
                violations.append(
                    _violation("prediction_read_error", path, "$", str(exc))
                )
                records = []
            for location, record in records:
                violations.extend(
                    inspect_prediction_record(record, file=path, location=location)
                )
            record_count = len(records)
        files.append(
            {
                "path": str(path),
                "role": "prediction",
                "records": record_count,
                "violations": len(violations) - before,
            }
        )

    for path in artifacts:
        before = len(violations)
        if not path.is_file():
            violations.append(
                _violation(
                    "missing_input_file", path, "$", "Knowledge artifact does not exist."
                )
            )
            document_count = 0
        elif path.suffix.casefold() not in KNOWLEDGE_EXTENSIONS:
            violations.append(
                _violation(
                    "unsupported_knowledge_format",
                    path,
                    "$",
                    f"Unsupported knowledge format: {path.suffix}",
                )
            )
            document_count = 0
        else:
            try:
                documents = _knowledge_documents(path, violations)
            except (OSError, ValueError) as exc:
                violations.append(_violation("knowledge_read_error", path, "$", str(exc)))
                documents = []
            for location, document in documents:
                violations.extend(
                    inspect_knowledge_document(document, file=path, location=location)
                )
            document_count = len(documents)
        files.append(
            {
                "path": str(path),
                "role": "knowledge_artifact",
                "documents": document_count,
                "violations": len(violations) - before,
            }
        )

    code_counts = Counter(violation["code"] for violation in violations)
    return {
        "schema_version": REPORT_SCHEMA_VERSION,
        "status": "PASS" if not violations else "FAIL",
        "passed": not violations,
        "policy": {
            "allowed_prediction_sources": [
                "strict identity manifest",
                "original paper text",
                "provenanced record-independent knowledge artifacts",
            ],
            "forbidden_runtime_source_terms": [
                "PMAGRESULT",
                "ROCKUNIT",
                "evaluation",
                "label",
                "previous workbook",
            ],
            "forbidden_r5_model": FORBIDDEN_R5_MODEL,
            "source_metadata_must_be_empty": True,
            "knowledge_record_identity_keys_forbidden": ["REFNO", "RESULTNO"],
            "knowledge_label_supervision_must_be_false": True,
            "knowledge_provenance_required": True,
        },
        "inputs": {
            "predictions": [str(path) for path in predictions],
            "knowledge_artifacts": [str(path) for path in artifacts],
        },
        "summary": {
            "files": len(files),
            "prediction_files": len(predictions),
            "knowledge_artifacts": len(artifacts),
            "violations": len(violations),
            "violations_by_code": dict(sorted(code_counts.items())),
        },
        "files": files,
        "violations": violations,
    }


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--prediction",
        action="append",
        type=Path,
        default=[],
        help="Prediction XLSX/JSON/JSONL; repeat for R3, R5, and R7.",
    )
    parser.add_argument(
        "--knowledge-artifact",
        action="append",
        type=Path,
        default=[],
        help="Knowledge JSON/JSONL artifact; repeat as needed.",
    )
    parser.add_argument("--report", type=Path, required=True, help="Machine-readable JSON report.")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    report = audit_scoring_boundary(args.prediction, args.knowledge_artifact)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
