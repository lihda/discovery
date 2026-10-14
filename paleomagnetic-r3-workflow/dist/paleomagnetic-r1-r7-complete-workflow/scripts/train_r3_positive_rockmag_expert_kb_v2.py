#!/usr/bin/env python3
"""Induce an aggregate R3 expert vocabulary from positive training ROCKMAG.

This is an offline, historically supervised knowledge-construction utility.
It opens R3 and PMAGRESULT.ROCKMAG cells only for the fixed training paper
partition. Held-out target labels and held-out ROCKMAG cells are not opened,
hashed, or represented in the artifact. The output contains only curated
semantic classes, regular expressions, and aggregate statistics.

The artifact is a retrieval knowledge base, not an online answer table. Its
weak and non-qualifying classes must not independently trigger an R3 score.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path
import re
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_PMAGRESULT = ROOT / "data/source_metadata/PMAGRESULT.xlsx"
DEFAULT_OUTPUT = ROOT / "data/calibration/r3_positive_rockmag_expert_kb_v2.json"

SPLIT_SALT = "r3-rockmag-distillation-v1"
HOLDOUT_FRACTION = 0.20
WILSON_Z_95 = 1.959963984540054


# Field patterns are used by this offline trainer. Original-text patterns are
# deliberately expanded expressions rather than terse database-only codes.
METHOD_TAXONOMY: tuple[dict[str, Any], ...] = (
    {
        "concept_id": "irm_acquisition",
        "field_patterns": (
            r"\bS?IRM\b",
            r"isothermal\s+remanent\s+magneti[sz]ation",
        ),
        "original_text_patterns": (
            r"\b(?:IRM|SIRM)\b(?=.{0,80}(?:acquisition|back[- ]field|applied\s+field|curve|experiment|demagneti[sz]))",
            r"isothermal\s+remanent\s+magneti[sz]ation",
            r"(?:IRM|remanence)\s+acquisition\s+curves?",
        ),
    },
    {
        "concept_id": "thermomagnetic_curve",
        "field_patterns": (
            r"\bJ(?:s|rs)[-/ ]?T\b",
            r"\bHcs[-/ ]?T\b",
            r"thermo[- ]?magnetic",
        ),
        "original_text_patterns": (
            r"thermo[- ]?magnetic(?:\s+(?:curve|analys[ei]s|experiment|measurement))?",
            r"\bJ(?:s|rs)[- ]?\(?T\)?\b",
            r"saturation\s+magneti[sz]ation.{0,80}(?:temperature|heating|cooling)",
        ),
    },
    {
        "concept_id": "temperature_dependent_susceptibility",
        "field_patterns": (
            r"\b(?:SUSC|SUCS)[-/ ]?T\b",
            r"\bSUSC\s*\(\s*T\s*\)",
            r"\bSUSC\b.{0,15}\b(?:VS\.?|VERSUS)\b.{0,15}\bTEMP\b",
            r"\b(?:kappa|chi|k)[-/ ]?T\b",
            r"temperature[- ]dependent\s+(?:magnetic\s+)?susceptibility",
        ),
        "original_text_patterns": (
            r"temperature[- ]dependent\s+(?:magnetic\s+)?susceptibility",
            r"(?:magnetic\s+)?susceptibility\s+(?:versus|vs\.?)\s+temperature",
            r"susceptibility.{0,100}(?:heating|cooling|temperature)",
            r"(?:heating|cooling|temperature).{0,100}susceptibility",
            r"\b(?:kappa|chi)[- ]?T\b",
        ),
    },
    {
        "concept_id": "three_axis_irm",
        "field_patterns": (
            r"\b3[- ]axis\b",
            r"\bthree[- ]axis\b",
            r"\bLowrie(?:[- ]Fuller)?\b",
        ),
        "original_text_patterns": (
            r"three[- ]axis\s+(?:IRM|isothermal\s+remanent\s+magneti[sz]ation)",
            r"3[- ]axis\s+(?:IRM|isothermal\s+remanent\s+magneti[sz]ation)",
            r"\bLowrie(?:[- ]Fuller)?\s+(?:test|experiment|procedure)",
        ),
    },
    {
        "concept_id": "remanence_temperature_experiment",
        "field_patterns": (
            r"\b(?:SIRM|IRM)[- ]?T\b",
            r"\blow[- ]T\b",
            r"low[- ]temperature.{0,30}\b(?:S?IRM|remanence)\b",
            r"\b(?:S?IRM|remanence)\b.{0,30}low[- ]temperature",
            r"\b(?:Verwey|Morin)\b",
        ),
        "original_text_patterns": (
            r"low[- ]temperature.{0,100}(?:IRM|isothermal\s+remanent\s+magneti[sz]ation|remanence)",
            r"(?:IRM|isothermal\s+remanent\s+magneti[sz]ation|remanence).{0,100}low[- ]temperature",
            r"thermal\s+demagneti[sz]ation.{0,100}(?:SIRM|IRM)",
            r"(?:Verwey|Morin)\s+transition",
        ),
    },
    {
        "concept_id": "hysteresis_forc",
        "field_patterns": (
            r"\bHYST\b",
            r"\bHIST\b",
            r"hysteresis",
            r"\bFORC\b",
            r"\b(?:Day|Henkel)\s+plot",
        ),
        "original_text_patterns": (
            r"hysteresis\s+(?:loops?|properties|parameters|measurement|experiment)",
            r"first[- ]order\s+reversal\s+curves?",
            r"\bFORC\s+(?:diagrams?|distributions?|analys[ei]s)",
            r"(?:Day|Henkel)\s+plots?",
        ),
    },
    {
        "concept_id": "coercivity_backfield",
        "field_patterns": (
            r"\b(?:Hc|Hcr|Bc|Bcr)\b",
            r"coerciv(?:e|ity)",
            r"back[- ]field",
        ),
        "original_text_patterns": (
            r"coerciv(?:ity|e)\s+(?:spectra|spectrum|distribution|measurement|parameter)",
            r"back[- ]field\s+(?:curve|remanence|experiment|measurement)",
            r"\b(?:Hc|Hcr|Bc|Bcr)\b.{0,60}(?:mT|ratio|coerciv|field)",
        ),
    },
    {
        "concept_id": "unblocking_spectrum",
        "field_patterns": (
            r"unblocking(?:[- ]temperature)?\s+(?:spectra|spectrum|distribution)",
        ),
        "original_text_patterns": (
            r"unblocking(?:[- ]temperature)?\s+(?:spectra|spectrum|distribution)",
        ),
    },
    {
        "concept_id": "magnetic_fabric",
        "field_patterns": (
            r"\b(?:AMS|AIR|AAR)\b",
            r"magnetic\s+(?:fabric|anisotrop)",
        ),
        "original_text_patterns": (
            r"anisotropy\s+of\s+magnetic\s+susceptibility",
            r"anisotropy\s+of\s+isothermal\s+remanence",
            r"anisotropy\s+of\s+anhysteretic\s+remanence",
            r"\b(?:AMS|AIR|AAR)\b.{0,60}(?:measurements?|analys[ei]s|fabric)",
        ),
    },
    {
        "concept_id": "petrographic_microscopy",
        "field_patterns": (
            r"\bOP\b",
            r"\bTS\b",
            r"\bSEM\b",
            r"\bTEM\b",
            r"\bEPMA\b",
            r"\bEDS\b",
            r"petrograph",
            r"microprobe",
        ),
        "original_text_patterns": (
            r"reflected[- ]light\s+microscop",
            r"polished\s+thin\s+sections?",
            r"scanning\s+electron\s+microscop",
            r"transmission\s+electron\s+microscop",
            r"electron\s+microprobe",
            r"energy[- ]dispersive\s+(?:spectroscopy|analysis)",
        ),
    },
    {
        "concept_id": "mineralogical_phase_analysis",
        "field_patterns": (
            r"\bX[- ]?RAY\b",
            r"\bXRAY\b",
            r"\bXRD\b",
            r"M(?:o|oe)ssbauer",
            r"mineral\s+separation",
            r"mineralogical\s+analys",
        ),
        "original_text_patterns": (
            r"X[- ]?ray\s+diffraction",
            r"\bXRD\b.{0,60}(?:analys[ei]s|patterns?|measurements?)",
            r"M(?:o|oe)ssbauer\s+(?:spectroscopy|spectra)",
            r"mineral\s+separation",
            r"mineralogical\s+(?:analysis|characteri[sz]ation)",
        ),
    },
)

DIAGNOSTIC_TAXONOMY: tuple[dict[str, Any], ...] = (
    {
        "concept_id": "saturation_behavior",
        "field_patterns": (
            r"\bno\s+sat\b",
            r"\bsat\.?\s*[0-9]",
            r"saturat(?:e|ed|ion)",
        ),
        "original_text_patterns": (
            r"(?:does\s+not|did\s+not|failed\s+to)\s+saturate",
            r"saturat(?:e|ed|ion).{0,60}(?:at|by|above|below)\s*[0-9]",
            r"approach(?:es|ed)?\s+(?:to\s+)?saturation",
        ),
    },
    {
        "concept_id": "curie_temperature",
        "field_patterns": (
            r"\b(?:CT|TC)\s*(?:[=:~<>]\s*)?[0-9]{2,4}",
            r"Curie\s+(?:point|temperature)",
        ),
        "original_text_patterns": (
            r"Curie\s+(?:point|temperature)",
            r"Curie[- ]?temperature.{0,60}[0-9]{2,4}",
        ),
    },
    {
        "concept_id": "domain_state",
        "field_patterns": (
            r"\b(?:SD|PSD|MD|SD[-/]PSD)\b",
            r"domain\s+state",
        ),
        "original_text_patterns": (
            r"(?:single|pseudo[- ]single|multi)[- ]domain",
            r"\b(?:SD|PSD|MD)\b.{0,50}domain",
            r"domain\s+(?:state|size)",
        ),
    },
    {
        "concept_id": "coercivity_parameter",
        "field_patterns": (
            r"\b(?:Hc|Hcr|Bc|Bcr)\b",
            r"\bHcr\s*/\s*Hc\b",
            r"coerciv(?:e|ity)",
        ),
        "original_text_patterns": (
            r"coerciv(?:ity|e)\s+(?:spectra|spectrum|distribution|parameter|ratio)",
            r"\b(?:Hcr\s*/\s*Hc|Bcr\s*/\s*Bc)\b",
        ),
    },
    {
        "concept_id": "mineral_transition",
        "field_patterns": (r"\b(?:Verwey|Morin)\b",),
        "original_text_patterns": (r"(?:Verwey|Morin)\s+transition",),
    },
    {
        "concept_id": "mineral_alteration_texture",
        "field_patterns": (
            r"alteration",
            r"oxidi[sz](?:ation|ed)",
            r"exsol(?:ution|ved)",
        ),
        "original_text_patterns": (
            r"(?:oxidation|alteration|exsolution).{0,100}(?:magnetic\s+mineral|oxide|mineral|grain|texture)",
            r"(?:magnetic\s+mineral|oxide|mineral|grain|texture).{0,100}(?:oxidation|alteration|exsolution)",
        ),
    },
)

CARRIER_TAXONOMY: tuple[dict[str, Any], ...] = (
    {
        "concept_id": "magnetite",
        "field_patterns": (r"\bMAGN\b", r"\bmagnetite\b"),
        "original_text_patterns": (r"\bmagnetite\b",),
    },
    {
        "concept_id": "titanomagnetite",
        "field_patterns": (
            r"\bTM\b",
            r"\bTMAGN\b",
            r"titano[- ]?magnetite",
        ),
        "original_text_patterns": (r"\btitano[- ]?magnetite\b",),
    },
    {
        "concept_id": "hematite",
        "field_patterns": (r"\bHEM\b", r"ha?ematite"),
        "original_text_patterns": (r"\bha?ematite\b",),
    },
    {
        "concept_id": "maghemite",
        "field_patterns": (r"\bMAGH\b", r"maghemite"),
        "original_text_patterns": (r"\bmaghemite\b",),
    },
    {
        "concept_id": "pyrrhotite",
        "field_patterns": (r"\bPYRR\b", r"\bPYRH\b", r"pyrrhotite"),
        "original_text_patterns": (r"\bpyrrhotite\b",),
    },
    {
        "concept_id": "greigite",
        "field_patterns": (r"\bGREIG\b", r"greigite"),
        "original_text_patterns": (r"\bgreigite\b",),
    },
    {
        "concept_id": "goethite",
        "field_patterns": (r"\bGOETH\b", r"\bGOEH\b", r"goethite"),
        "original_text_patterns": (r"\bgoethite\b",),
    },
    {
        "concept_id": "ilmenite",
        "field_patterns": (r"\bILM\b", r"ilmenite"),
        "original_text_patterns": (r"\bilmenite\b",),
    },
    {
        "concept_id": "iron_titanium_oxide",
        "field_patterns": (
            r"\bFe[- ]?Ti\b",
            r"iron[- ]titanium\s+oxide",
        ),
        "original_text_patterns": (
            r"\bFe[- ]?Ti\s+oxides?\b",
            r"iron[- ]titanium\s+oxides?",
        ),
    },
    {
        "concept_id": "iron_sulfide",
        "field_patterns": (r"iron\s+sulphide", r"iron\s+sulfide"),
        "original_text_patterns": (r"iron\s+sulphides?", r"iron\s+sulfides?"),
    },
)


ABBREVIATIONS: tuple[dict[str, str], ...] = (
    {"abbreviation": "IRM", "concept_id": "irm_acquisition", "category": "method", "field_pattern": r"\bIRM\b"},
    {"abbreviation": "SIRM", "concept_id": "irm_acquisition", "category": "method", "field_pattern": r"\bSIRM\b"},
    {"abbreviation": "Js-T", "concept_id": "thermomagnetic_curve", "category": "method", "field_pattern": r"\bJs[-/ ]?T\b"},
    {"abbreviation": "Jrs-T", "concept_id": "thermomagnetic_curve", "category": "method", "field_pattern": r"\bJrs[-/ ]?T\b"},
    {"abbreviation": "SUSC-T", "concept_id": "temperature_dependent_susceptibility", "category": "method", "field_pattern": r"\bSUSC[-/ ]?T\b"},
    {"abbreviation": "SUCS-T", "concept_id": "temperature_dependent_susceptibility", "category": "method", "field_pattern": r"\bSUCS[-/ ]?T\b"},
    {"abbreviation": "HYST", "concept_id": "hysteresis_forc", "category": "method", "field_pattern": r"\bHYST\b"},
    {"abbreviation": "HIST", "concept_id": "hysteresis_forc", "category": "method", "field_pattern": r"\bHIST\b"},
    {"abbreviation": "FORC", "concept_id": "hysteresis_forc", "category": "method", "field_pattern": r"\bFORC\b"},
    {"abbreviation": "OP", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bOP\b"},
    {"abbreviation": "TS", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bTS\b"},
    {"abbreviation": "SEM", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bSEM\b"},
    {"abbreviation": "TEM", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bTEM\b"},
    {"abbreviation": "EPMA", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bEPMA\b"},
    {"abbreviation": "EDS", "concept_id": "petrographic_microscopy", "category": "method", "field_pattern": r"\bEDS\b"},
    {"abbreviation": "XRD", "concept_id": "mineralogical_phase_analysis", "category": "method", "field_pattern": r"\bXRD\b"},
    {"abbreviation": "AMS", "concept_id": "magnetic_fabric", "category": "method", "field_pattern": r"\bAMS\b"},
    {"abbreviation": "AIR", "concept_id": "magnetic_fabric", "category": "method", "field_pattern": r"\bAIR\b"},
    {"abbreviation": "AAR", "concept_id": "magnetic_fabric", "category": "method", "field_pattern": r"\bAAR\b"},
    {"abbreviation": "CT", "concept_id": "curie_temperature", "category": "diagnostic", "field_pattern": r"\bCT\b"},
    {"abbreviation": "Tc", "concept_id": "curie_temperature", "category": "diagnostic", "field_pattern": r"\bTC\b"},
    {"abbreviation": "SD", "concept_id": "domain_state", "category": "diagnostic", "field_pattern": r"\bSD\b"},
    {"abbreviation": "PSD", "concept_id": "domain_state", "category": "diagnostic", "field_pattern": r"\bPSD\b"},
    {"abbreviation": "MD", "concept_id": "domain_state", "category": "diagnostic", "field_pattern": r"\bMD\b"},
    {"abbreviation": "Hc", "concept_id": "coercivity_parameter", "category": "diagnostic", "field_pattern": r"\bHC\b"},
    {"abbreviation": "Hcr", "concept_id": "coercivity_parameter", "category": "diagnostic", "field_pattern": r"\bHCR\b"},
    {"abbreviation": "Bc", "concept_id": "coercivity_parameter", "category": "diagnostic", "field_pattern": r"\bBC\b"},
    {"abbreviation": "Bcr", "concept_id": "coercivity_parameter", "category": "diagnostic", "field_pattern": r"\bBCR\b"},
    {"abbreviation": "MAGN", "concept_id": "magnetite", "category": "carrier", "field_pattern": r"\bMAGN\b"},
    {"abbreviation": "TM", "concept_id": "titanomagnetite", "category": "carrier", "field_pattern": r"\bTM\b"},
    {"abbreviation": "TMAGN", "concept_id": "titanomagnetite", "category": "carrier", "field_pattern": r"\bTMAGN\b"},
    {"abbreviation": "HEM", "concept_id": "hematite", "category": "carrier", "field_pattern": r"\bHEM\b"},
    {"abbreviation": "MAGH", "concept_id": "maghemite", "category": "carrier", "field_pattern": r"\bMAGH\b"},
    {"abbreviation": "PYRR", "concept_id": "pyrrhotite", "category": "carrier", "field_pattern": r"\bPYRR\b"},
    {"abbreviation": "PYRH", "concept_id": "pyrrhotite", "category": "carrier", "field_pattern": r"\bPYRH\b"},
    {"abbreviation": "GREIG", "concept_id": "greigite", "category": "carrier", "field_pattern": r"\bGREIG\b"},
    {"abbreviation": "GOETH", "concept_id": "goethite", "category": "carrier", "field_pattern": r"\bGOETH\b"},
    {"abbreviation": "GOEH", "concept_id": "goethite", "category": "carrier", "field_pattern": r"\bGOEH\b"},
    {"abbreviation": "ILM", "concept_id": "ilmenite", "category": "carrier", "field_pattern": r"\bILM\b"},
)


NONQUALIFYING_CLASSES: tuple[dict[str, Any], ...] = (
    {
        "class_id": "tub_unblocking_interval_without_independent_method",
        "field_patterns": (r"\bT(?:ub|b)\s*(?:=|>|<|:)",),
        "derived_condition": "unblocking interval occurs without an independently eligible method",
        "policy": "Component cleaning or unblocking bounds alone are not a carrier investigation.",
    },
    {
        "class_id": "bare_susceptibility",
        "field_patterns": (r"^\s*SUSC\s*[.;]?\s*$",),
        "derived_condition": "the complete value is only a susceptibility code",
        "policy": "A bare susceptibility code lacks a carrier-oriented temperature or diagnostic experiment.",
    },
    {
        "class_id": "bare_anisotropy",
        "field_patterns": (r"^\s*AN\s*[.;]?\s*$",),
        "derived_condition": "the complete value is only the ambiguous AN code",
        "policy": "AN is not expanded to anisotropy unless the source states AMS, AIR, AAR, or magnetic fabric.",
    },
    {
        "class_id": "bare_susc_an_bundle",
        "field_patterns": (r"^\s*(?:SUSC\s*[,;/+]\s*AN|AN\s*[,;/+]\s*SUSC)\s*[.;]?\s*$",),
        "derived_condition": "only the two ambiguous bare codes occur",
        "policy": "Two weak database codes do not establish a performed carrier experiment in original text.",
    },
    {
        "class_id": "carrier_without_method",
        "field_patterns": (),
        "derived_condition": "one or more carrier concepts occur without an eligible method",
        "policy": "A mineral name is a retrieval cue, not standalone proof of a target-bound investigation.",
    },
    {
        "class_id": "unclassified_nonempty",
        "field_patterns": (),
        "derived_condition": "a nonempty value matches no curated semantic concept",
        "policy": "Unclassified historical text is retained only as an aggregate coverage count.",
    },
)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def holdout_bucket(paper_identity: str, salt: str = SPLIT_SALT) -> float:
    digest = hashlib.sha256(f"{salt}\0{paper_identity}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def is_holdout(paper_identity: str) -> bool:
    return holdout_bucket(paper_identity) < HOLDOUT_FRACTION


def _cell_value(cell: Any) -> Any:
    """Read one cell lazily so boundary tests can reject forbidden access."""

    try:
        return cell.value
    except AttributeError:
        return cell


def _header_map(header: Sequence[Any], required: Sequence[str]) -> dict[str, int]:
    names = [clean(_cell_value(cell)) for cell in header]
    result: dict[str, int] = {}
    for name in required:
        matches = [index for index, value in enumerate(names) if value == name]
        if len(matches) != 1:
            raise ValueError(f"expected exactly one {name} column; found {len(matches)}")
        result[name] = matches[0]
    return result


def partition_target_rows(
    header: Sequence[Any],
    rows: Iterable[Sequence[Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Open supervised target cells only after assignment to training."""

    columns = _header_map(header, ("REFNO", "RESULTNO", "R3"))
    training: list[dict[str, Any]] = []
    training_papers: set[str] = set()
    held_out_papers: set[str] = set()
    training_rows = 0
    held_out_rows = 0
    blank_paper_rows = 0

    for row in rows:
        paper_identity = clean(_cell_value(row[columns["REFNO"]]))
        if not paper_identity:
            blank_paper_rows += 1
            continue
        if is_holdout(paper_identity):
            held_out_rows += 1
            held_out_papers.add(paper_identity)
            continue

        training_rows += 1
        training_papers.add(paper_identity)
        record_identity = clean(_cell_value(row[columns["RESULTNO"]]))
        label_value = _cell_value(row[columns["R3"]])
        if not record_identity:
            raise ValueError("training row has an empty record identity")
        try:
            label = int(float(label_value))
        except (TypeError, ValueError) as exc:
            raise ValueError("training row has a non-binary R3 value") from exc
        if label not in {0, 1}:
            raise ValueError("training row has a non-binary R3 value")
        training.append(
            {
                "paper_identity": paper_identity,
                "record_identity": record_identity,
                "label": label,
            }
        )

    if training_papers & held_out_papers:
        raise AssertionError("paper-group split overlap")
    return training, {
        "training": {
            "rows": training_rows,
            "paper_groups": len(training_papers),
        },
        "holdout": {
            "rows": held_out_rows,
            "paper_groups": len(held_out_papers),
        },
        "blank_paper_identity_rows": blank_paper_rows,
    }


def load_training_targets(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows()
        return partition_target_rows(next(rows, ()), rows)
    finally:
        workbook.close()


def collect_training_rockmag_rows(
    header: Sequence[Any],
    rows: Iterable[Sequence[Any]],
    training_record_identities: set[str],
) -> tuple[dict[str, list[str]], dict[str, int]]:
    """Open ROCKMAG only for identities already admitted to training."""

    columns = _header_map(header, ("RESULTNO", "ROCKMAG"))
    values: dict[str, list[str]] = defaultdict(list)
    source_rows = 0
    training_cells_opened = 0
    nonempty_cells = 0
    for row in rows:
        source_rows += 1
        record_identity = clean(_cell_value(row[columns["RESULTNO"]]))
        if record_identity not in training_record_identities:
            continue
        training_cells_opened += 1
        value = clean(_cell_value(row[columns["ROCKMAG"]]))
        if value:
            values[record_identity].append(value)
            nonempty_cells += 1
    return dict(values), {
        "source_rows_identity_scanned": source_rows,
        "training_field_cells_opened": training_cells_opened,
        "nonempty_training_field_cells": nonempty_cells,
        "nontraining_field_cells_opened": 0,
    }


def load_training_rockmag(
    path: Path,
    training_record_identities: set[str],
) -> tuple[dict[str, list[str]], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows()
        return collect_training_rockmag_rows(
            next(rows, ()), rows, training_record_identities
        )
    finally:
        workbook.close()


def _compile_taxonomy(
    taxonomy: Sequence[Mapping[str, Any]],
) -> tuple[tuple[str, tuple[re.Pattern[str], ...]], ...]:
    return tuple(
        (
            str(item["concept_id"]),
            tuple(re.compile(pattern, re.I) for pattern in item["field_patterns"]),
        )
        for item in taxonomy
    )


_COMPILED_METHODS = _compile_taxonomy(METHOD_TAXONOMY)
_COMPILED_DIAGNOSTICS = _compile_taxonomy(DIAGNOSTIC_TAXONOMY)
_COMPILED_CARRIERS = _compile_taxonomy(CARRIER_TAXONOMY)
_COMPILED_ABBREVIATIONS = tuple(
    (item["abbreviation"], re.compile(item["field_pattern"], re.I))
    for item in ABBREVIATIONS
)
_TUB_INTERVAL = re.compile(r"\bT(?:ub|b)\s*(?:=|>|<|:)", re.I)
_BARE_SUSC = re.compile(r"^\s*SUSC\s*[.;]?\s*$", re.I)
_BARE_AN = re.compile(r"^\s*AN\s*[.;]?\s*$", re.I)
_BARE_SUSC_AN = re.compile(
    r"^\s*(?:SUSC\s*[,;/+]\s*AN|AN\s*[,;/+]\s*SUSC)\s*[.;]?\s*$",
    re.I,
)


def _matches(
    text: str,
    compiled: Sequence[tuple[str, Sequence[re.Pattern[str]]]],
) -> set[str]:
    return {
        concept_id
        for concept_id, patterns in compiled
        if any(pattern.search(text) for pattern in patterns)
    }


def classify_rockmag_value(value: Any) -> dict[str, Any]:
    """Return only normalized semantic labels for one historical value."""

    text = clean(value)
    methods = _matches(text, _COMPILED_METHODS)
    diagnostics = _matches(text, _COMPILED_DIAGNOSTICS)
    carriers = _matches(text, _COMPILED_CARRIERS)
    abbreviations = {
        abbreviation
        for abbreviation, pattern in _COMPILED_ABBREVIATIONS
        if pattern.search(text)
    }
    nonqualifying: set[str] = set()
    if _TUB_INTERVAL.search(text) and not methods:
        nonqualifying.add("tub_unblocking_interval_without_independent_method")
    if _BARE_SUSC.fullmatch(text):
        nonqualifying.add("bare_susceptibility")
    if _BARE_AN.fullmatch(text):
        nonqualifying.add("bare_anisotropy")
    if _BARE_SUSC_AN.fullmatch(text):
        nonqualifying.add("bare_susc_an_bundle")
    if carriers and not methods:
        nonqualifying.add("carrier_without_method")
    if text and not (methods or diagnostics or carriers or nonqualifying):
        nonqualifying.add("unclassified_nonempty")

    categories = {
        category
        for category, concepts in (
            ("method", methods),
            ("diagnostic", diagnostics),
            ("carrier", carriers),
        )
        if concepts
    }
    if not text:
        evidence_class = "missing_field"
    elif methods and diagnostics and carriers:
        evidence_class = "method_diagnostic_carrier"
    elif methods and carriers:
        evidence_class = "method_carrier"
    elif methods and diagnostics:
        evidence_class = "method_diagnostic"
    elif methods:
        evidence_class = "method_only"
    elif diagnostics and carriers:
        evidence_class = "diagnostic_carrier_without_method"
    elif diagnostics:
        evidence_class = "diagnostic_only"
    elif carriers:
        evidence_class = "carrier_only"
    else:
        evidence_class = "unclassified_nonempty"

    return {
        "method_concepts": sorted(methods),
        "diagnostic_concepts": sorted(diagnostics),
        "carrier_concepts": sorted(carriers),
        "abbreviations": sorted(abbreviations),
        "nonqualifying_classes": sorted(nonqualifying),
        "category_signature": "+".join(sorted(categories)) if categories else "none",
        "evidence_class": evidence_class,
        "strong_chain": bool(methods and (diagnostics or carriers)),
        "method_event": bool(methods),
        "known_concept": bool(methods or diagnostics or carriers),
    }


def wilson_lower_bound(successes: int, trials: int) -> float:
    if trials <= 0:
        return 0.0
    proportion = successes / trials
    denominator = 1.0 + WILSON_Z_95 * WILSON_Z_95 / trials
    centre = proportion + WILSON_Z_95 * WILSON_Z_95 / (2.0 * trials)
    margin = WILSON_Z_95 * math.sqrt(
        proportion * (1.0 - proportion) / trials
        + WILSON_Z_95 * WILSON_Z_95 / (4.0 * trials * trials)
    )
    return max(0.0, (centre - margin) / denominator)


def _association(positive: int, negative: int, positive_total: int) -> dict[str, Any]:
    observations = positive + negative
    return {
        "positive_rows": positive,
        "negative_rows": negative,
        "observations": observations,
        "positive_precision": round(positive / observations, 8) if observations else None,
        "positive_precision_wilson_lower_95": round(
            wilson_lower_bound(positive, observations), 8
        ),
        "positive_row_coverage": round(positive / positive_total, 8)
        if positive_total
        else 0.0,
    }


def _paper_association(labels: Mapping[str, set[int]]) -> dict[str, int | float | None]:
    positive_only = sum(value == {1} for value in labels.values())
    negative_only = sum(value == {0} for value in labels.values())
    conflict = sum(value == {0, 1} for value in labels.values())
    observations = positive_only + negative_only
    return {
        "positive_only_papers": positive_only,
        "negative_only_papers": negative_only,
        "label_conflict_papers": conflict,
        "nonconflict_observations": observations,
        "positive_precision": round(positive_only / observations, 8)
        if observations
        else None,
    }


def _row_classification(values: Sequence[str]) -> dict[str, Any]:
    classifications = [classify_rockmag_value(value) for value in values if clean(value)]
    methods: set[str] = set()
    diagnostics: set[str] = set()
    carriers: set[str] = set()
    abbreviations: set[str] = set()
    nonqualifying: set[str] = set()
    for item in classifications:
        methods.update(item["method_concepts"])
        diagnostics.update(item["diagnostic_concepts"])
        carriers.update(item["carrier_concepts"])
        abbreviations.update(item["abbreviations"])
        nonqualifying.update(item["nonqualifying_classes"])

    categories = {
        category
        for category, concepts in (
            ("method", methods),
            ("diagnostic", diagnostics),
            ("carrier", carriers),
        )
        if concepts
    }
    if not classifications:
        evidence_class = "missing_field"
    elif methods and diagnostics and carriers:
        evidence_class = "method_diagnostic_carrier"
    elif methods and carriers:
        evidence_class = "method_carrier"
    elif methods and diagnostics:
        evidence_class = "method_diagnostic"
    elif methods:
        evidence_class = "method_only"
    elif diagnostics and carriers:
        evidence_class = "diagnostic_carrier_without_method"
    elif diagnostics:
        evidence_class = "diagnostic_only"
    elif carriers:
        evidence_class = "carrier_only"
    else:
        evidence_class = "unclassified_nonempty"
        nonqualifying.add("unclassified_nonempty")
    return {
        "methods": methods,
        "diagnostics": diagnostics,
        "carriers": carriers,
        "abbreviations": abbreviations,
        "nonqualifying": nonqualifying,
        "category_signature": "+".join(sorted(categories)) if categories else "none",
        "evidence_class": evidence_class,
        "has_value": bool(classifications),
        "known_concept": bool(methods or diagnostics or carriers),
        "strong_chain": bool(methods and (diagnostics or carriers)),
        "method_event": bool(methods),
    }


def induce_statistics(
    training_records: Sequence[Mapping[str, Any]],
    values_by_record: Mapping[str, Sequence[str]],
    *,
    minimum_cooccurrence: int = 2,
) -> dict[str, Any]:
    """Aggregate concepts without returning any row or paper identities."""

    positive_total = sum(int(record["label"]) == 1 for record in training_records)
    negative_total = len(training_records) - positive_total
    concept_rows: dict[str, Counter[int]] = defaultdict(Counter)
    concept_paper_labels: dict[str, dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    abbreviation_rows: dict[str, Counter[int]] = defaultdict(Counter)
    abbreviation_paper_labels: dict[str, dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    nonqualifying_rows: dict[str, Counter[int]] = defaultdict(Counter)
    nonqualifying_paper_labels: dict[str, dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    positive_evidence_classes: Counter[str] = Counter()
    positive_category_signatures: Counter[str] = Counter()
    evidence_class_papers: dict[str, set[str]] = defaultdict(set)
    category_signature_papers: dict[str, set[str]] = defaultdict(set)
    positive_edge_rows: Counter[tuple[str, str]] = Counter()
    positive_edge_papers: dict[tuple[str, str], set[str]] = defaultdict(set)
    aggregate = Counter()

    for record in training_records:
        label = int(record["label"])
        paper_identity = str(record["paper_identity"])
        record_identity = str(record["record_identity"])
        row = _row_classification(values_by_record.get(record_identity, ()))
        concepts = {
            *(f"method:{value}" for value in row["methods"]),
            *(f"diagnostic:{value}" for value in row["diagnostics"]),
            *(f"carrier:{value}" for value in row["carriers"]),
        }
        for concept in concepts:
            concept_rows[concept][label] += 1
            concept_paper_labels[concept][paper_identity].add(label)
        for abbreviation in row["abbreviations"]:
            abbreviation_rows[abbreviation][label] += 1
            abbreviation_paper_labels[abbreviation][paper_identity].add(label)
        for class_id in row["nonqualifying"]:
            nonqualifying_rows[class_id][label] += 1
            nonqualifying_paper_labels[class_id][paper_identity].add(label)

        aggregate[f"label_{label}_rows_with_nonempty_field"] += int(row["has_value"])
        aggregate[f"label_{label}_rows_matching_known_concept"] += int(
            row["known_concept"]
        )
        aggregate[f"label_{label}_rows_with_method_event"] += int(row["method_event"])
        aggregate[f"label_{label}_rows_with_strong_chain"] += int(row["strong_chain"])

        if label != 1:
            continue
        positive_evidence_classes[row["evidence_class"]] += 1
        positive_category_signatures[row["category_signature"]] += 1
        evidence_class_papers[row["evidence_class"]].add(paper_identity)
        category_signature_papers[row["category_signature"]].add(paper_identity)
        ordered = sorted(concepts)
        for left in range(len(ordered)):
            for right in range(left + 1, len(ordered)):
                edge = (ordered[left], ordered[right])
                positive_edge_rows[edge] += 1
                positive_edge_papers[edge].add(paper_identity)

    def concept_stat(key: str) -> dict[str, Any]:
        rows = concept_rows[key]
        return {
            "row_level": _association(rows[1], rows[0], positive_total),
            "paper_deduplicated": _paper_association(concept_paper_labels[key]),
        }

    concept_statistics = {
        key: concept_stat(key)
        for key in sorted(concept_rows)
    }
    abbreviation_statistics = {
        key: {
            "row_level": _association(
                abbreviation_rows[key][1], abbreviation_rows[key][0], positive_total
            ),
            "paper_deduplicated": _paper_association(
                abbreviation_paper_labels[key]
            ),
        }
        for key in sorted(abbreviation_rows)
    }
    nonqualifying_statistics = {
        key: {
            "row_level": _association(
                nonqualifying_rows[key][1], nonqualifying_rows[key][0], positive_total
            ),
            "paper_deduplicated": _paper_association(
                nonqualifying_paper_labels[key]
            ),
        }
        for key in sorted(nonqualifying_rows)
    }
    cooccurrences = [
        {
            "source_concept": source,
            "target_concept": target,
            "positive_row_count": count,
            "positive_paper_count": len(positive_edge_papers[(source, target)]),
            "positive_row_coverage": round(count / positive_total, 8)
            if positive_total
            else 0.0,
        }
        for (source, target), count in sorted(
            positive_edge_rows.items(), key=lambda item: (-item[1], item[0])
        )
        if count >= minimum_cooccurrence
    ]

    def count_records(
        counts: Counter[str], paper_sets: Mapping[str, set[str]]
    ) -> list[dict[str, Any]]:
        return [
            {
                "class_id": class_id,
                "positive_row_count": count,
                "positive_paper_count": len(paper_sets[class_id]),
                "positive_row_coverage": round(count / positive_total, 8)
                if positive_total
                else 0.0,
            }
            for class_id, count in sorted(counts.items())
        ]

    def coverage(key: str, label: int, total: int) -> dict[str, int | float]:
        count = aggregate[f"label_{label}_{key}"]
        return {
            "rows": count,
            "denominator_rows": total,
            "coverage": round(count / total, 8) if total else 0.0,
        }

    return {
        "training_population": {
            "rows": len(training_records),
            "positive_rows": positive_total,
            "negative_rows": negative_total,
        },
        "coverage": {
            "positive_nonempty_field": coverage(
                "rows_with_nonempty_field", 1, positive_total
            ),
            "positive_known_concept": coverage(
                "rows_matching_known_concept", 1, positive_total
            ),
            "positive_method_event": coverage(
                "rows_with_method_event", 1, positive_total
            ),
            "positive_strong_chain": coverage(
                "rows_with_strong_chain", 1, positive_total
            ),
            "negative_nonempty_field": coverage(
                "rows_with_nonempty_field", 0, negative_total
            ),
            "negative_known_concept": coverage(
                "rows_matching_known_concept", 0, negative_total
            ),
            "positive_known_concept_given_nonempty": {
                "rows": aggregate["label_1_rows_matching_known_concept"],
                "denominator_rows": aggregate["label_1_rows_with_nonempty_field"],
                "coverage": round(
                    aggregate["label_1_rows_matching_known_concept"]
                    / aggregate["label_1_rows_with_nonempty_field"],
                    8,
                )
                if aggregate["label_1_rows_with_nonempty_field"]
                else 0.0,
            },
        },
        "positive_evidence_classes": count_records(
            positive_evidence_classes, evidence_class_papers
        ),
        "positive_category_signatures": count_records(
            positive_category_signatures, category_signature_papers
        ),
        "concept_statistics": concept_statistics,
        "abbreviation_statistics": abbreviation_statistics,
        "nonqualifying_statistics": nonqualifying_statistics,
        "positive_concept_cooccurrences": cooccurrences,
        "minimum_cooccurrence": minimum_cooccurrence,
    }


def _taxonomy_output(
    taxonomy: Sequence[Mapping[str, Any]],
    category: str,
    statistics: Mapping[str, Any],
) -> list[dict[str, Any]]:
    concept_statistics = statistics["concept_statistics"]
    empty = {
        "row_level": _association(
            0,
            0,
            int(statistics["training_population"]["positive_rows"]),
        ),
        "paper_deduplicated": _paper_association({}),
    }
    return [
        {
            "concept_id": item["concept_id"],
            "category": category,
            "field_patterns": list(item["field_patterns"]),
            "original_text_patterns": list(item["original_text_patterns"]),
            "runtime_original_text_patterns": list(
                item["original_text_patterns"]
            ),
            "statistics": concept_statistics.get(
                f"{category}:{item['concept_id']}", empty
            ),
        }
        for item in taxonomy
    ]


def _abbreviation_output(statistics: Mapping[str, Any]) -> list[dict[str, Any]]:
    observed = statistics["abbreviation_statistics"]
    positive_total = int(statistics["training_population"]["positive_rows"])
    empty = {
        "row_level": _association(0, 0, positive_total),
        "paper_deduplicated": _paper_association({}),
    }
    return [
        {
            **dict(item),
            "statistics": observed.get(item["abbreviation"], empty),
        }
        for item in ABBREVIATIONS
    ]


def _nonqualifying_output(statistics: Mapping[str, Any]) -> list[dict[str, Any]]:
    observed = statistics["nonqualifying_statistics"]
    positive_total = int(statistics["training_population"]["positive_rows"])
    empty = {
        "row_level": _association(0, 0, positive_total),
        "paper_deduplicated": _paper_association({}),
    }
    return [
        {
            **dict(item),
            "statistics": observed.get(item["class_id"], empty),
        }
        for item in NONQUALIFYING_CLASSES
    ]


def validate_artifact_boundary(artifact: Mapping[str, Any]) -> None:
    """Fail closed if the aggregate artifact violates its leakage contract."""

    serialized = json.dumps(artifact, ensure_ascii=False, sort_keys=True)
    for forbidden in ("REFNO", "RESULTNO", "ROCKNAME"):
        if forbidden in serialized:
            raise ValueError(f"artifact contains forbidden identity field name: {forbidden}")
    provenance = artifact.get("provenance", {})
    required_false = (
        "held_out_labels_opened",
        "held_out_field_cells_opened",
        "raw_field_values_retained",
        "row_identity_mapping_retained",
        "source_quotes_retained",
        "prior_predictions_or_evaluations_used",
        "whole_input_file_hashes_retained",
    )
    for key in required_false:
        if provenance.get(key) is not False:
            raise ValueError(f"provenance.{key} must be false")
    if provenance.get("supervision_scope") != "fixed_training_partition_only":
        raise ValueError("unexpected supervision scope")
    if provenance.get("label_supervision") is not True:
        raise ValueError("historical label supervision must be disclosed")
    if artifact.get("runtime_policy", {}).get("source_table_access") is not False:
        raise ValueError("runtime source-table access must be false")
    if artifact.get("runtime_policy", {}).get("contains_record_answer_mapping") is not False:
        raise ValueError("runtime answer mapping flag must be false")

    for taxonomy_name in (
        "method_taxonomy",
        "diagnostic_taxonomy",
        "carrier_taxonomy",
    ):
        for concept in artifact.get(taxonomy_name, []):
            for pattern in concept.get("field_patterns", []):
                re.compile(pattern, re.I)
            for pattern in concept.get("original_text_patterns", []):
                re.compile(pattern, re.I | re.S)
            runtime_patterns = concept.get("runtime_original_text_patterns", [])
            if runtime_patterns != concept.get("original_text_patterns", []):
                raise ValueError(
                    f"{taxonomy_name} runtime patterns must equal reviewed original-text patterns"
                )
            for pattern in runtime_patterns:
                re.compile(pattern, re.I | re.S)
    for item in artifact.get("abbreviation_lexicon", []):
        re.compile(str(item["field_pattern"]), re.I)
    for item in artifact.get("nonqualifying_classes", []):
        for pattern in item.get("field_patterns", []):
            re.compile(pattern, re.I)


def build_artifact(
    targets_path: Path,
    pmagresult_path: Path,
    *,
    minimum_cooccurrence: int = 2,
) -> dict[str, Any]:
    training_records, split_counts = load_training_targets(targets_path)
    values_by_record, source_access = load_training_rockmag(
        pmagresult_path,
        {str(record["record_identity"]) for record in training_records},
    )
    statistics = induce_statistics(
        training_records,
        values_by_record,
        minimum_cooccurrence=minimum_cooccurrence,
    )
    artifact = {
        "schema_version": "r3_positive_rockmag_expert_kb_v2",
        "knowledge_version": "r3_positive_rockmag_expert_kb_v2.0",
        "artifact_class": "historically_supervised_aggregate_expert_vocabulary",
        "provenance": {
            "phase": "offline_training_only",
            "source_declaration": {
                "supervised_field": "data/targets.xlsx::R3",
                "historical_semantic_field": (
                    "data/source_metadata/PMAGRESULT.xlsx::ROCKMAG"
                ),
                "positive_selection": "binary R3 equals one",
            },
            "supervision_scope": "fixed_training_partition_only",
            "label_supervision": True,
            "label_supervision_disclosed": True,
            "held_out_labels_opened": False,
            "held_out_field_cells_opened": False,
            "raw_field_values_retained": False,
            "row_identity_mapping_retained": False,
            "source_quotes_retained": False,
            "prior_predictions_or_evaluations_used": False,
            "whole_input_file_hashes_retained": False,
            "generator": "scripts/train_r3_positive_rockmag_expert_kb_v2.py",
            "generator_sha256": file_sha256(Path(__file__)),
        },
        "split": {
            "unit": "paper identity",
            "algorithm": (
                "SHA256(salt + NUL + canonical paper identity), first 8 bytes "
                "as unsigned big-endian divided by 2^64"
            ),
            "salt": SPLIT_SALT,
            "holdout_fraction": HOLDOUT_FRACTION,
            "identity_counts": split_counts,
        },
        "offline_access_audit": source_access,
        "runtime_policy": {
            "allowed_sources": [
                "original paper text",
                "identity-only target manifest",
                "this frozen aggregate vocabulary",
            ],
            "source_table_access": False,
            "label_access": False,
            "previous_evaluation_access": False,
            "contains_record_answer_mapping": False,
            "strong_chain_use": (
                "method plus diagnostic or carrier is a high-specificity retrieval template"
            ),
            "method_only_use": (
                "positive-associated retrieval cue requiring target-bound performed context"
            ),
            "weak_class_use": "retrieval or review cue only; never standalone positive proof",
        },
        "method_taxonomy": _taxonomy_output(
            METHOD_TAXONOMY, "method", statistics
        ),
        "diagnostic_taxonomy": _taxonomy_output(
            DIAGNOSTIC_TAXONOMY, "diagnostic", statistics
        ),
        "carrier_taxonomy": _taxonomy_output(
            CARRIER_TAXONOMY, "carrier", statistics
        ),
        "abbreviation_lexicon": _abbreviation_output(statistics),
        "nonqualifying_classes": _nonqualifying_output(statistics),
        "aggregate_statistics": {
            key: value
            for key, value in statistics.items()
            if key
            not in {
                "concept_statistics",
                "abbreviation_statistics",
                "nonqualifying_statistics",
            }
        },
    }
    validate_artifact_boundary(artifact)
    return artifact


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--pmagresult", type=Path, default=DEFAULT_PMAGRESULT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--minimum-cooccurrence", type=int, default=2)
    args = parser.parse_args()
    if args.minimum_cooccurrence < 1:
        raise SystemExit("--minimum-cooccurrence must be at least one")

    artifact = build_artifact(
        args.targets,
        args.pmagresult,
        minimum_cooccurrence=args.minimum_cooccurrence,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(artifact, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    coverage = artifact["aggregate_statistics"]["coverage"]
    print(
        json.dumps(
            {
                "output": str(args.output),
                "schema_version": artifact["schema_version"],
                "split": artifact["split"]["identity_counts"],
                "positive_nonempty_field": coverage["positive_nonempty_field"],
                "positive_known_concept": coverage["positive_known_concept"],
                "positive_method_event": coverage["positive_method_event"],
                "positive_strong_chain": coverage["positive_strong_chain"],
                "cooccurrence_edges": len(
                    artifact["aggregate_statistics"][
                        "positive_concept_cooccurrences"
                    ]
                ),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
