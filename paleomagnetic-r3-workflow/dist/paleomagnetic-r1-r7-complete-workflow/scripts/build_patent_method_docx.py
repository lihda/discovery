#!/usr/bin/env python3
"""Build the detailed patent method-flow planning document."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/patent_method_flow_detailed_20260717.md"
OUTPUT = ROOT / "docs/古地磁可靠性指标知识冲突检测方法_方法流程详细稿_20260717.docx"
FIGURE = ROOT / "docs/figures/patent-overall-method-flow.png"

FONT_WEST = "Arial Unicode MS"
FONT_EAST = "Arial Unicode MS"
INK = RGBColor(31, 45, 61)
BLUE = RGBColor(46, 116, 181)
BLUE_DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(96, 106, 116)
LIGHT = "F2F4F7"


def set_run_font(run, size=None, bold=None, italic=None, color=None, font=FONT_WEST):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size, color=INK, bold=False):
    style.font.name = FONT_WEST
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_WEST)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_WEST)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (fld_char, instr, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    set_style_font(title, 24, INK, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 14, MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 12, BLUE_DARK, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, 11, INK)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    if "Patent Formula" not in doc.styles:
        formula = doc.styles.add_style("Patent Formula", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(formula, 10.5, BLUE_DARK, bold=False)
        formula.paragraph_format.left_indent = Inches(0.3)
        formula.paragraph_format.right_indent = Inches(0.3)
        formula.paragraph_format.space_before = Pt(4)
        formula.paragraph_format.space_after = Pt(8)
        formula.paragraph_format.line_spacing = 1.0
        formula.paragraph_format.keep_together = True

    if "Figure Caption" not in doc.styles:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(caption, 9.5, MUTED)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(10)
        caption.paragraph_format.keep_with_next = True


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_running_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("专利技术交底 | 方法流程详细稿")
    set_run_font(r, size=9, color=MUTED, bold=True)
    r = p.add_run("\t2026-07-17")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(p)
    r = p.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)


def draw_flow_figure(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1500, 1960
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    title_font = ImageFont.truetype(font_path, 48)
    step_font = ImageFont.truetype(font_path, 34)
    small_font = ImageFont.truetype(font_path, 25)
    draw.text((width // 2, 60), "古地磁可靠性指标知识冲突检测方法总体流程", font=title_font, fill=(31, 45, 61), anchor="ma")

    steps = [
        ("S1", "判据知识构建与冻结", "规范命题 / 规则对象 / 版本哈希"),
        ("S2", "目标身份与页级原文", "身份白名单 / 单页单源 / 来源审计"),
        ("S3", "候选检索与目标绑定", "竞争目标 / scope_id / 绑定强度"),
        ("S4", "证据关系与冲突推理", "R3 / R5 / R7 / 规则级例外"),
        ("S5", "判决、拒判与人工复核", "四值证据态 / 1-0-ABSTAIN"),
        ("S6", "结果固化与隔离评测", "输入输出哈希 / 边界审计 / 回执"),
    ]
    colors = [(231, 240, 247), (232, 241, 231), (243, 234, 213), (220, 234, 247), (245, 232, 204), (232, 236, 239)]
    left, right = 180, width - 180
    box_h, gap = 205, 70
    y = 190
    for index, ((code, label, detail), fill) in enumerate(zip(steps, colors)):
        draw.rounded_rectangle((left, y, right, y + box_h), radius=18, fill=fill, outline=(80, 105, 128), width=4)
        draw.text((left + 75, y + 52), code, font=title_font, fill=(46, 116, 181), anchor="la")
        draw.text((left + 250, y + 48), label, font=step_font, fill=(31, 45, 61), anchor="la")
        draw.text((left + 250, y + 128), detail, font=small_font, fill=(82, 92, 102), anchor="la")
        if index < len(steps) - 1:
            x = width // 2
            draw.line((x, y + box_h + 8, x, y + box_h + gap - 12), fill=(80, 105, 128), width=6)
            draw.polygon([(x - 16, y + box_h + gap - 28), (x + 16, y + box_h + gap - 28), (x, y + box_h + gap)], fill=(80, 105, 128))
        y += box_h + gap

    note_y = height - 190
    draw.rounded_rectangle((180, note_y, width - 180, note_y + 90), radius=12, fill=(249, 250, 252), outline=(180, 188, 198), width=3)
    draw.text((width // 2, note_y + 45), "标签仅在评分进程关闭、预测哈希固定之后进入隔离评测", font=small_font, fill=(124, 90, 0), anchor="mm")
    image.save(path, dpi=(200, 200))


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")


def add_inline(paragraph, text, default_size=11):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=default_size, color=INK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=INK, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size - 0.5, color=BLUE_DARK, font="Courier New")
        else:
            run = paragraph.add_run(token.rstrip(".,;，。；"))
            set_run_font(run, size=default_size - 0.5, color=BLUE_DARK)
            run.font.underline = True
            trailing = token[len(token.rstrip(".,;，。；")):]
            if trailing:
                trail = paragraph.add_run(trailing)
                set_run_font(trail, size=default_size, color=INK)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=default_size, color=INK)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("专利技术交底书")
    set_run_font(r, size=12, color=BLUE, bold=True)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("古地磁可靠性指标知识冲突检测方法")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("方法流程规划与具体实施步骤详细稿")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("建议保护主线")
    set_run_font(r, size=10, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.65)
    p.paragraph_format.right_indent = Inches(0.65)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("冻结判据 + 页级保真 + 目标绑定 + 可撤销冲突 + 拒判复核 + 预测固化")
    set_run_font(r, size=13, color=BLUE_DARK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("版本日期：2026-07-17")
    set_run_font(r, size=10.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("用途：技术交底、方法流程定稿、附图与权利要求展开")
    set_run_font(r, size=10.5, color=MUTED)
    doc.add_page_break()


def add_lead_callout(doc):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.10
    shade_paragraph(p, LIGHT)
    r = p.add_run("核心结论：")
    set_run_font(r, size=11, color=BLUE_DARK, bold=True)
    r = p.add_run("本项目应作为一条可审计的数据变换与冲突判决链进行保护，而不是仅描述为从论文中检索关键词并输出 R3/R5/R7。")
    set_run_font(r, size=11, color=INK)


def add_figure(doc):
    draw_flow_figure(FIGURE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(FIGURE), width=Inches(5.55))
    cap = doc.add_paragraph(style="Figure Caption")
    cap.add_run("图1  古地磁可靠性指标知识冲突检测方法总体流程")


def parse_source(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    body_started = False
    inserted_figure = False
    paragraph_buffer = []

    def flush():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, text)
        paragraph_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if not body_started:
            if line.startswith("## 一、"):
                body_started = True
                add_lead_callout(doc)
            else:
                continue

        if not line:
            flush()
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush()
            level = len(heading.group(1)) - 1
            text = heading.group(2).strip()
            if not inserted_figure and text.startswith("五、"):
                add_figure(doc)
                inserted_figure = True
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline(p, text, default_size={1: 16, 2: 13, 3: 12}[min(level, 3)])
            if level == 1 and any(text.startswith(prefix) for prefix in ("五、", "八、", "十、", "十二、", "十四、")):
                p.paragraph_format.page_break_before = True
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or numbered:
            flush()
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or numbered).group(1))
            continue

        if line.startswith("`") and line.endswith("`") and len(line) > 2:
            flush()
            p = doc.add_paragraph(style="Patent Formula")
            r = p.add_run(line[1:-1])
            set_run_font(r, size=10.5, color=BLUE_DARK, font="Courier New")
            shade_paragraph(p, "F7F9FB")
            continue

        paragraph_buffer.append(line)

    flush()


def audit_document(doc):
    assert doc.sections[0].page_width == Inches(8.5)
    assert doc.sections[0].left_margin == Inches(1)
    assert doc.styles["Normal"].font.size == Pt(11)
    assert doc.styles["Heading 1"].font.size == Pt(16)
    assert doc.styles["Heading 2"].font.size == Pt(13)
    assert doc.styles["Heading 3"].font.size == Pt(12)
    assert len(doc.tables) == 0
    assert any(p.style.name == "List Number" for p in doc.paragraphs)
    assert any("S68" in p.text for p in doc.paragraphs)


def main():
    doc = Document()
    configure_sections(doc)
    configure_styles(doc)
    add_running_header_footer(doc.sections[0])
    doc.core_properties.title = "古地磁可靠性指标知识冲突检测方法：方法流程规划与具体实施步骤详细稿"
    doc.core_properties.subject = "专利技术交底方法流程"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "古地磁,R3,R5,R7,证据关系,知识冲突,专利"

    add_cover(doc)
    parse_source(doc)
    audit_document(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
