#!/usr/bin/env python3
"""Freeze label-free TRACE-PMAG page-graph features for downstream R3 tests.

This command has a deliberately narrow input boundary. It accepts only the
three-column target identity, local-PDF membership, original-paper page text,
the frozen R3 criterion, and an independently curated expert terminology KB.
It has no argument for labels, source-table fields, predictions, or metrics.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from typing import Any, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from r3_trace_page_graph_features import (  # noqa: E402
    build_trace_page_graph_features,
    load_trace_concepts,
)
from paper_text_retrieval import canonical_target_identity  # noqa: E402


DEFAULT_IDENTITY = ROOT / "outputs/r357_strict_eval_07111758/inputs/identity_manifest.xlsx"
DEFAULT_PAGE_CACHE = ROOT / "tmp/r3_all_labelled_pdf_page_cache_v1/page_cache.jsonl"
DEFAULT_PDF_DIR = ROOT / "data/raw_pdfs"
DEFAULT_CRITERION = ROOT / "data/knowledge/r357_criterion_spec_v2.json"
DEFAULT_CONCEPTS = ROOT / "data/knowledge/r3_trace_normative_concepts_v1.json"
DEFAULT_OUTPUT_DIR = ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1"

FEATURE_NAME = "trace_page_graph_features_v4.npz"
MANIFEST_NAME = "trace_page_graph_features_v4.manifest.json"
AUDIT_NAME = "trace_page_graph_features_v4.audit.json"
CONTEXT_NAME = "trace_target_context_documents_v4.jsonl"
CONCEPT_SNAPSHOT_NAME = "trace_page_graph_concepts_v4.json"
CANDIDATE_FREEZE_NAME = "trace_page_graph_candidate_freeze_v4.json"
FOLD_MANIFEST_NAME = "trace_refno_fold_manifest_v4.json"
FEATURE_SCHEMA_VERSION = "r3_trace_page_graph_features_v4"
AUDIT_SCHEMA_VERSION = "r3_trace_page_graph_feature_audit_v4"
CONCEPT_SCHEMA_VERSION = "r3_trace_page_graph_concepts_v4"
CANDIDATE_SCHEMA_VERSION = "r3_trace_page_graph_candidate_freeze_v4"
FOLD_SCHEMA_VERSION = "r3_trace_refno_fold_manifest_v4"
OUTER_FOLDS = 5
OUTER_FOLD_SEED = 35_720_260_759
INNER_FOLDS = 3

IDENTITY_FIELDS = ("REFNO", "RESULTNO", "ROCKNAME")
PAGE_CACHE_ROW_FIELDS = frozenset({"REFNO", "pages"})
PAGE_CACHE_PAGE_FIELDS = frozenset({"page", "text", "section_hints"})
KNOWLEDGE_CATEGORIES = ("method", "diagnostic", "carrier", "bridge")
REQUIRED_R3_PROPOSITIONS = frozenset(
    {
        "R3.P.METHOD.UNBLOCKING_COERCIVITY",
        "R3.P.METHOD.IRM",
        "R3.P.METHOD.TEMPERATURE_SUSCEPTIBILITY",
        "R3.P.METHOD.THREE_AXIS_IRM",
        "R3.P.METHOD.LOW_TEMPERATURE_IRM",
        "R3.P.METHOD.HYSTERESIS",
        "R3.P.METHOD.DAY_PLOT",
        "R3.P.METHOD.FORC",
        "R3.P.METHOD.AMS",
        "R3.P.METHOD.AIR",
        "R3.P.METHOD.AAR",
        "R3.P.METHOD.REFLECTED_LIGHT",
        "R3.P.METHOD.SEM",
        "R3.P.METHOD.TEM",
        "R3.P.METHOD.MINERAL_SEPARATION",
        "R3.P.CARRIER.IDENTIFIED",
        "R3.P.CARRIER.REASONABLE_INCONCLUSIVE_ATTEMPT",
        "R3.P.CARRIER.SIGNIFICANCE_COMMENTED",
    }
)
FORBIDDEN_PROVENANCE_RE = re.compile(r"\bPMAGRESULT\b|\bROCKMAG\b", re.I)
FORBIDDEN_INPUT_FIELD_RE = re.compile(
    r"(?:^|_)(?:R3|R5|R7|PMAGRESULT|ROCKMAG|PREDICTIONS?|OOF|EVALUATION|METRICS?)(?:$|_)",
    re.I,
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_sha256(value: Any) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def array_sha256(value: np.ndarray) -> str:
    array = np.ascontiguousarray(value)
    digest = hashlib.sha256()
    digest.update(str(array.dtype).encode("ascii"))
    digest.update(json.dumps(list(array.shape)).encode("ascii"))
    digest.update(array.tobytes(order="C"))
    return digest.hexdigest()


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def atomic_npz(path: Path, **arrays: np.ndarray) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        with temporary.open("wb") as handle:
            np.savez_compressed(handle, **arrays)
        temporary.replace(path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def atomic_jsonl_strings(path: Path, values: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            for value in values:
                handle.write(json.dumps(clean(value), ensure_ascii=False))
                handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def load_three_column_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("TRACE identity must contain exactly one worksheet")
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = tuple(clean(value) for value in next(rows, ()))
        if header != IDENTITY_FIELDS:
            raise ValueError(
                "TRACE identity must contain exactly REFNO, RESULTNO, ROCKNAME; "
                "labels and source-table fields are forbidden"
            )
        records = []
        for raw in rows:
            values = tuple(clean(value) for value in raw)
            if not any(values):
                continue
            if len(values) != len(IDENTITY_FIELDS) or not all(values):
                raise ValueError("TRACE identity contains an incomplete row")
            records.append(dict(zip(IDENTITY_FIELDS, values)))
    finally:
        workbook.close()
    if not records:
        raise ValueError("TRACE identity is empty")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise ValueError("TRACE identity contains duplicate RESULTNO values")
    return records


def select_pdf_backed_records(
    records: Sequence[Mapping[str, str]], pdf_dir: Path
) -> tuple[list[dict[str, str]], dict[str, int], str]:
    if not pdf_dir.is_dir():
        raise ValueError(f"PDF directory does not exist: {pdf_dir}")
    pdf_names = sorted(
        path.name
        for path in pdf_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".pdf"
    )
    pdf_refnos = {Path(name).stem for name in pdf_names}
    selected = [dict(record) for record in records if record["REFNO"] in pdf_refnos]
    if not selected:
        raise ValueError("no identity rows have a local PDF")
    selected_refnos = {record["REFNO"] for record in selected}
    all_refnos = {record["REFNO"] for record in records}
    coverage = {
        "identity_rows": len(records),
        "identity_papers": len(all_refnos),
        "pdf_backed_rows": len(selected),
        "pdf_backed_papers": len(selected_refnos),
        "skipped_without_pdf_rows": len(records) - len(selected),
        "skipped_without_pdf_papers": len(all_refnos - selected_refnos),
    }
    return selected, coverage, json_sha256(pdf_names)


def pdf_content_sha256_by_refno(
    pdf_dir: Path, wanted_refnos: set[str]
) -> dict[str, str]:
    paths = {
        path.stem: path
        for path in pdf_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".pdf"
    }
    missing = sorted(wanted_refnos - paths.keys())
    if missing:
        raise ValueError(f"PDF content grouping is missing REFNO values: {missing[:20]}")
    return {refno: file_sha256(paths[refno]) for refno in sorted(wanted_refnos)}


def normalized_text_document_groups(
    page_rows: Mapping[str, Mapping[str, Any]],
    pdf_sha_by_refno: Mapping[str, str],
) -> tuple[dict[str, str], dict[str, str], dict[str, int]]:
    """Group byte variants of a paper by normalized extracted text when usable."""

    groups: dict[str, str] = {}
    text_hashes: dict[str, str] = {}
    character_counts: dict[str, int] = {}
    for refno in sorted(pdf_sha_by_refno):
        digest = hashlib.sha256()
        character_count = 0
        for page in (page_rows.get(refno) or {}).get("pages") or []:
            normalized = canonical_target_identity(clean(page.get("text")))
            digest.update(normalized.encode("utf-8"))
            digest.update(b"\n\f\n")
            character_count += len(normalized)
        text_sha = digest.hexdigest()
        source_kind = "text" if character_count >= 256 else "pdf"
        source_hash = text_sha if source_kind == "text" else pdf_sha_by_refno[refno]
        groups[refno] = hashlib.sha256(
            f"{source_kind}|{source_hash}".encode("ascii")
        ).hexdigest()
        text_hashes[refno] = text_sha
        character_counts[refno] = character_count
    return groups, text_hashes, character_counts


def load_sanitized_page_cache(
    path: Path, wanted_refnos: set[str]
) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            raw = json.loads(line)
            if not isinstance(raw, Mapping):
                raise ValueError(f"page cache row is not an object at line {line_number}")
            refno = clean(raw.get("REFNO"))
            if refno not in wanted_refnos:
                continue
            forbidden_row_fields = sorted(
                clean(name) for name in raw if FORBIDDEN_INPUT_FIELD_RE.search(clean(name))
            )
            if forbidden_row_fields:
                raise ValueError(
                    f"page-cache REFNO {refno} contains forbidden fields: {forbidden_row_fields}"
                )
            if refno in rows:
                raise ValueError(f"duplicate page-cache REFNO {refno}")
            pages = []
            for page_number, raw_page in enumerate(raw.get("pages") or [], start=1):
                if not isinstance(raw_page, Mapping):
                    continue
                forbidden_page_fields = sorted(
                    clean(name)
                    for name in raw_page
                    if FORBIDDEN_INPUT_FIELD_RE.search(clean(name))
                )
                if forbidden_page_fields:
                    raise ValueError(
                        f"page-cache REFNO {refno} page {page_number} contains forbidden "
                        f"fields: {forbidden_page_fields}"
                    )
                page = {
                    name: raw_page.get(name)
                    for name in PAGE_CACHE_PAGE_FIELDS
                    if name in raw_page
                }
                page["page"] = int(page.get("page") or page_number)
                page["text"] = clean(page.get("text"))
                page["section_hints"] = [
                    clean(value)
                    for value in page.get("section_hints") or []
                    if clean(value)
                ]
                pages.append(page)
            rows[refno] = {"REFNO": refno, "pages": pages}
    missing = sorted(wanted_refnos - rows.keys())
    if missing:
        raise ValueError(f"page cache is missing REFNO values: {missing[:20]}")
    return rows


def _load_json_object(path: Path, label: str) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a JSON object")
    return value


def _reject_forbidden_provenance(
    payload: Mapping[str, Any], label: str, *, require_explicit_boundary: bool
) -> None:
    provenance = payload.get("provenance") or {}
    boundary_fields = (
        "label_supervision",
        "source_table_access",
        "pmagresult_or_rockmag_used",
        "evaluation_results_used",
        "prior_predictions_used",
        "contains_record_identifiers",
        "contains_answer_mapping",
    )
    for field in boundary_fields:
        if provenance.get(field) is True:
            raise ValueError(f"{label} provenance forbids {field}=true")
    if require_explicit_boundary:
        missing = [field for field in boundary_fields if provenance.get(field) is not False]
        if missing:
            raise ValueError(
                f"{label} must explicitly set its no-leakage provenance flags false: {missing}"
            )
    historical_sources = {
        key: provenance.get(key)
        for key in (
            "knowledge_induction",
            "offline_sources",
            "source_columns_read",
            "training_sources",
            "induction_sources",
        )
        if key in provenance
    }
    serialized_sources = json.dumps(historical_sources, ensure_ascii=False, sort_keys=True)
    match = FORBIDDEN_PROVENANCE_RE.search(serialized_sources)
    if match:
        raise ValueError(
            f"{label} provenance contains a forbidden historical source marker: {match.group(0)}"
        )


def build_trace_concept_snapshot(
    criterion_path: Path, concepts_path: Path
) -> tuple[dict[str, Any], dict[str, Any]]:
    criterion = _load_json_object(criterion_path, "criterion")
    normative = _load_json_object(concepts_path, "normative concepts")
    _reject_forbidden_provenance(
        criterion, "criterion", require_explicit_boundary=False
    )
    _reject_forbidden_provenance(
        normative, "normative concepts", require_explicit_boundary=True
    )
    if criterion.get("freeze_status") != "frozen":
        raise ValueError("criterion must be frozen")
    if normative.get("freeze_status") not in {"frozen", "frozen_before_label_access"}:
        raise ValueError("normative concepts must be frozen before label access")
    r3 = ((criterion.get("criteria") or {}).get("R3") or {})
    proposition_ids = {
        clean(item.get("proposition_id")) for item in r3.get("propositions") or []
    }
    missing = REQUIRED_R3_PROPOSITIONS - proposition_ids
    if missing:
        raise ValueError(f"criterion is missing required R3 propositions: {sorted(missing)}")

    source_ids = {
        clean(item.get("source_id"))
        for item in normative.get("sources") or []
        if isinstance(item, Mapping) and clean(item.get("source_id"))
    }
    if not source_ids:
        raise ValueError("normative concepts must declare auditable sources")
    concepts = []
    concept_ids: set[str] = set()
    category_counts: Counter[str] = Counter()
    for number, raw in enumerate(normative.get("concepts") or [], start=1):
        if not isinstance(raw, Mapping):
            raise ValueError(f"normative concept {number} is not an object")
        concept_id = clean(raw.get("concept_id"))
        category = clean(raw.get("category")).casefold()
        patterns = [clean(value) for value in raw.get("search_patterns") or [] if clean(value)]
        concept_source_refs = [
            clean(value)
            for value in raw.get("source_refs") or []
            if clean(value)
        ]
        if not concept_id or concept_id in concept_ids:
            raise ValueError(f"invalid or duplicate normative concept_id: {concept_id!r}")
        if category not in KNOWLEDGE_CATEGORIES:
            raise ValueError(f"invalid TRACE category for {concept_id}: {category!r}")
        if not patterns or not concept_source_refs:
            raise ValueError(f"TRACE concept {concept_id} lacks patterns or source links")
        unknown_refs = set(concept_source_refs) - source_ids
        if unknown_refs:
            raise ValueError(
                f"TRACE concept {concept_id} cites unknown sources: {sorted(unknown_refs)}"
            )
        for pattern in patterns:
            re.compile(pattern, re.I)
        concepts.append(
            {
                "concept_id": concept_id,
                "category": category,
                "source_refs": concept_source_refs,
                "search_patterns": patterns,
            }
        )
        concept_ids.add(concept_id)
        category_counts[category] += 1
    absent_categories = set(KNOWLEDGE_CATEGORIES) - category_counts.keys()
    if absent_categories:
        raise ValueError(f"normative concepts lack TRACE categories: {sorted(absent_categories)}")
    snapshot = {
        "schema_version": CONCEPT_SCHEMA_VERSION,
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "criterion": {
            "path": str(criterion_path.resolve()),
            "sha256": file_sha256(criterion_path),
            "spec_version": criterion.get("spec_version"),
        },
        "normative_concepts": {
            "path": str(concepts_path.resolve()),
            "sha256": file_sha256(concepts_path),
            "schema_version": normative.get("schema_version"),
        },
        "concepts": concepts,
    }
    audit = {
        "concept_count": len(concepts),
        "category_counts": {
            name: category_counts[name] for name in KNOWLEDGE_CATEGORIES
        },
        "criterion_propositions_available": len(proposition_ids),
        "criterion_propositions_required": len(REQUIRED_R3_PROPOSITIONS),
        "forbidden_historical_source_markers_found": [],
    }
    return snapshot, audit


def _record_alignment_sha256(records: Sequence[Mapping[str, str]]) -> str:
    return json_sha256(
        [{name: clean(record.get(name)) for name in IDENTITY_FIELDS} for record in records]
    )


def deterministic_refno_fold_ids(
    records: Sequence[Mapping[str, str]],
    *,
    n_splits: int,
    seed: int,
    group_key_by_refno: Mapping[str, str] | None = None,
) -> np.ndarray:
    members: dict[str, list[int]] = {}
    for index, record in enumerate(records):
        refno = clean(record.get("REFNO"))
        group = clean((group_key_by_refno or {}).get(refno)) or refno
        members.setdefault(group, []).append(index)
    ordered = sorted(
        members,
        key=lambda group: (
            -len(members[group]),
            hashlib.sha256(f"{seed}|{group}".encode("utf-8")).hexdigest(),
        ),
    )
    fold_rows = [0] * n_splits
    fold_groups = [0] * n_splits
    assignments: dict[str, int] = {}
    for group in ordered:
        fold = min(
            range(n_splits),
            key=lambda value: (fold_rows[value], fold_groups[value], value),
        )
        assignments[group] = fold + 1
        fold_rows[fold] += len(members[group])
        fold_groups[fold] += 1
    output = np.asarray(
        [
            assignments[
                clean((group_key_by_refno or {}).get(clean(record.get("REFNO"))))
                or clean(record.get("REFNO"))
            ]
            for record in records
        ],
        dtype=np.int8,
    )
    expected = set(range(1, min(n_splits, len(members)) + 1))
    if set(map(int, output)) != expected:
        raise AssertionError("TRACE REFNO fold assignment is incomplete")
    return output


def _feature_audit(features: Mapping[str, Any], matrix: np.ndarray) -> dict[str, Any]:
    names = list(features["feature_names"])

    def count(name: str) -> int:
        return int(np.sum(matrix[:, names.index(name)] > 0.5))

    paper_rows = list(features["paper_audit"])
    physical_distribution = Counter(int(row["physical_edges"]) for row in paper_rows)
    semantic_distribution = Counter(int(row["semantic_edges"]) for row in paper_rows)
    return {
        "schema_version": AUDIT_SCHEMA_VERSION,
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "rows": int(matrix.shape[0]),
        "papers": int(features["paper_count"]),
        "feature_count": int(matrix.shape[1]),
        "complete_chain_rows": count("trace_chain_complete"),
        "connected_chain_rows": count("trace_chain_connected"),
        "exclusive_complete_chain_rows": int(
            np.sum(
                (matrix[:, names.index("trace_chain_complete")] > 0.5)
                & (matrix[:, names.index("trace_chain_all_pages_exclusive")] > 0.5)
            )
        ),
        "all_four_query_rows": count("query_all_four_available"),
        "papers_with_semantic_edges": sum(int(row["semantic_edges"] > 0) for row in paper_rows),
        "physical_edge_distribution_by_paper": {
            str(key): value for key, value in sorted(physical_distribution.items())
        },
        "semantic_edge_distribution_by_paper": {
            str(key): value for key, value in sorted(semantic_distribution.items())
        },
        "papers_detail": paper_rows,
    }


def build_frozen_trace_features(
    *,
    identity_path: Path,
    page_cache_path: Path,
    pdf_dir: Path,
    criterion_path: Path,
    concepts_path: Path,
    output_dir: Path,
) -> dict[str, Path]:
    paths = {
        "identity": identity_path.resolve(),
        "page_cache": page_cache_path.resolve(),
        "pdf_dir": pdf_dir.resolve(),
        "criterion": criterion_path.resolve(),
        "concepts": concepts_path.resolve(),
    }
    for name in ("identity", "page_cache", "criterion", "concepts"):
        if not paths[name].is_file():
            raise ValueError(f"missing TRACE input {name}: {paths[name]}")
    records_all = load_three_column_identity(paths["identity"])
    records, coverage, pdf_membership_sha256 = select_pdf_backed_records(
        records_all, paths["pdf_dir"]
    )
    alignment_sha256 = _record_alignment_sha256(records)
    refnos = {record["REFNO"] for record in records}
    pdf_sha_by_refno = pdf_content_sha256_by_refno(paths["pdf_dir"], refnos)
    page_rows = load_sanitized_page_cache(paths["page_cache"], refnos)
    document_groups, text_sha_by_refno, text_characters_by_refno = (
        normalized_text_document_groups(page_rows, pdf_sha_by_refno)
    )
    snapshot, knowledge_audit = build_trace_concept_snapshot(
        paths["criterion"], paths["concepts"]
    )

    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    feature_path = output_dir / FEATURE_NAME
    manifest_path = output_dir / MANIFEST_NAME
    audit_path = output_dir / AUDIT_NAME
    context_path = output_dir / CONTEXT_NAME
    concept_path = output_dir / CONCEPT_SNAPSHOT_NAME
    candidate_path = output_dir / CANDIDATE_FREEZE_NAME
    fold_path = output_dir / FOLD_MANIFEST_NAME
    atomic_json(concept_path, snapshot)
    concepts = load_trace_concepts(concept_path)
    features = build_trace_page_graph_features(records, page_rows, concepts)
    matrix = np.asarray(features["matrix"], dtype=np.float32)
    trace_feature_names = list(features["feature_names"])
    if matrix.shape[0] != len(records) or matrix.ndim != 2 or not np.isfinite(matrix).all():
        raise ValueError("TRACE feature matrix is invalid")
    atomic_npz(feature_path, trace=matrix)
    context_documents = [clean(value) for value in features["target_context_documents"]]
    if len(context_documents) != len(records) or any(not value for value in context_documents):
        raise ValueError("TRACE target-context documents are invalid")
    atomic_jsonl_strings(context_path, context_documents)
    audit = _feature_audit(features, matrix)
    audit["trace_feature_count"] = len(trace_feature_names)
    audit["target_context_rows"] = len(context_documents)
    audit["target_context_rows_with_evidence"] = sum(
        "[METHOD_" in value
        or "[DIAGNOSTIC_" in value
        or "[CARRIER_" in value
        or "[BRIDGE_" in value
        for value in context_documents
    )
    audit["target_context_character_distribution"] = {
        "minimum": min(map(len, context_documents)),
        "median": int(np.median([len(value) for value in context_documents])),
        "maximum": max(map(len, context_documents)),
    }
    atomic_json(audit_path, audit)

    fold_ids = deterministic_refno_fold_ids(
        records,
        n_splits=OUTER_FOLDS,
        seed=OUTER_FOLD_SEED,
        group_key_by_refno=document_groups,
    )
    fold_manifest = {
        "schema_version": FOLD_SCHEMA_VERSION,
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "seed": OUTER_FOLD_SEED,
        "method": "deterministic row-balanced complete-REFNO and normalized-document-content assignment",
        "alignment_sha256": alignment_sha256,
        "fold_id_sha256": hashlib.sha256(fold_ids.tobytes()).hexdigest(),
        "labels_read": False,
        "document_content_groups": len(set(document_groups.values())),
        "duplicate_pdf_content_groups": sum(
            count > 1 for count in Counter(pdf_sha_by_refno.values()).values()
        ),
        "duplicate_normalized_document_groups": sum(
            count > 1 for count in Counter(document_groups.values()).values()
        ),
        "fold_counts": {
            str(fold): int(np.sum(fold_ids == fold))
            for fold in range(1, OUTER_FOLDS + 1)
        },
        "rows": [
            {
                "row_index": index,
                "REFNO": clean(record.get("REFNO")),
                "RESULTNO": clean(record.get("RESULTNO")),
                "document_sha256": document_groups[clean(record.get("REFNO"))],
                "pdf_sha256": pdf_sha_by_refno[clean(record.get("REFNO"))],
                "normalized_text_sha256": text_sha_by_refno[
                    clean(record.get("REFNO"))
                ],
                "normalized_text_characters": text_characters_by_refno[
                    clean(record.get("REFNO"))
                ],
                "fold": int(fold_ids[index]),
            }
            for index, record in enumerate(records)
        ],
    }
    atomic_json(fold_path, fold_manifest)

    candidate_freeze = {
        "schema_version": CANDIDATE_SCHEMA_VERSION,
        "status": "FROZEN_BEFORE_LABEL_ACCESS",
        "primary_candidate": "trace_guarded_nonlinear_fusion",
        "feature_candidates": {
            "trace_structured": "TRACE-PMAG graph over original text with criterion concepts, component-aware attribution, and evidence-chain features",
            "trace_coreference_atomic_events": "conservative same-paragraph paleomagnetic coreference and sentence-level atomic evidence relations",
            "trace_fuzzy_target_anchors": "document-missing-only compound and orthographic target anchors with sibling arbitration",
            "trace_monotonic_criterion": "fixed monotonic method/carrier/bridge completeness score with no fitted threshold",
            "trace_target_context_word": "word TF-IDF fitted inside each training fold on target-bound original-paper evidence contexts",
            "trace_target_context_char": "character TF-IDF fitted inside each training fold on target-bound original-paper evidence contexts",
            "trace_structured_plus_target_context": "late fusion of structured TRACE and the two fold-local target-context text models",
            "trace_nonlinear_structured": "fixed HGB and random-forest learners fitted only inside grouped training folds",
            "trace_joint_pairwise": "optional guarded within-paper positive-versus-negative target ranking"
        },
        "fixed_model_family": "predeclared linear, nonlinear, text, and monotonic models evaluated with complete-document isolation",
        "model_candidates": {
            "structured_logistic": {"c_values": [0.01, 0.1, 1.0], "standardize": True},
            "structured_hgb": {
                "max_iter": 260,
                "learning_rate": 0.035,
                "max_leaf_nodes": 7,
                "min_samples_leaf": 10,
                "l2_regularization": 1.0
            },
            "structured_random_forest": {
                "n_estimators": 600,
                "max_features": 0.5,
                "min_samples_leaf": 8,
                "n_jobs": -1
            },
            "target_context_word": {"ngram_range": [1, 2], "min_df": 2, "max_features": 80000, "c_values": [0.5, 1.0, 2.0]},
            "target_context_char": {"analyzer": "char_wb", "ngram_range": [3, 5], "min_df": 2, "max_features": 100000, "c_values": [0.5, 1.0, 2.0]},
            "within_paper_pairwise": {
                "c_values": [0.0001, 0.001, 0.01],
                "minimum_distinct_target_counts": [2, 4, 8],
                "alphas": [0.25, 0.5, 0.75, 1.0],
                "minimum_accuracy_gain_over_identity": 0.01,
                "require_non_decreasing_macro_f1": True
            }
        },
        "fusion_candidates": [
            "monotonic_criterion",
            "structured",
            "hgb",
            "random_forest",
            "word",
            "char",
            "mean_hgb_random_forest",
            "mean_structured_hgb",
            "mean_structured_random_forest",
            "mean_structured_hgb_random_forest",
            "mean_structured_word",
            "mean_structured_char",
            "mean_hgb_char",
            "mean_random_forest_char",
            "mean_structured_hgb_char",
            "mean_word_char",
            "mean_all",
            "weighted_structured_50_word_25_char_25"
        ],
        "fusion_selection_guard": {
            "baseline": "hgb",
            "minimum_accuracy_gain_over_baseline": 0.005,
            "require_non_decreasing_macro_f1": True
        },
        "monotonic_criterion_formula": {
            "intercept": 0.08,
            "query_method_available": 0.16,
            "query_carrier_available": 0.16,
            "query_bridge_available": 0.18,
            "trace_chain_complete": 0.16,
            "trace_chain_connected": 0.10,
            "trace_chain_all_pages_exclusive": 0.08,
            "query_diagnostic_available": 0.08
        },
        "outer_folds": OUTER_FOLDS,
        "outer_fold_seed": OUTER_FOLD_SEED,
        "inner_folds": INNER_FOLDS,
        "fixed_decision_threshold": 0.5,
        "selection_rule": "select hyperparameters and any fusion or pairwise correction only inside outer training folds with frozen minimum-gain guards; never from outer-fold metrics",
        "alignment_sha256": alignment_sha256,
        "matrix_sha256": array_sha256(matrix),
        "target_context_sha256": file_sha256(context_path),
        "fold_manifest_sha256": file_sha256(fold_path),
        "runtime_boundary": {
            "labels_read": False,
            "pmagresult_or_rockmag_read": False,
            "prior_predictions_read": False,
            "evaluation_results_read": False,
            "resultno_used_only_for_alignment": True,
            "refno_used_only_for_paper_grouping": True,
            "pdf_content_sha256_used_only_for_group_isolation": True,
            "normalized_page_text_sha256_used_only_for_group_isolation": True,
        },
    }
    atomic_json(candidate_path, candidate_freeze)

    code_inputs = [
        Path(__file__).resolve(),
        ROOT / "scripts/r3_trace_page_graph_features.py",
        ROOT / "scripts/r3_trace_discourse_roles.py",
        ROOT / "scripts/r3_trace_coreference_events.py",
        ROOT / "scripts/experiment_r3_trace_clean_nested_oof.py",
        ROOT / "scripts/r3_paleomag_scope_features.py",
        ROOT / "scripts/paper_text_retrieval.py",
    ]
    manifest = {
        "schema_version": FEATURE_SCHEMA_VERSION,
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "rows": len(records),
        "papers": features["paper_count"],
        "features": matrix.shape[1],
        "feature_names": features["feature_names"],
        "alignment_sha256": alignment_sha256,
        "matrix_sha256": array_sha256(matrix),
        "contains_labels": False,
        "feature_archive_contains_identifiers": False,
        "target_context_contains_alignment_identifiers": False,
        "coverage": coverage,
        "runtime_boundary": candidate_freeze["runtime_boundary"],
        "allowed_sources": [
            "REFNO/RESULTNO/ROCKNAME identity",
            "local PDF membership and content hash for duplicate-document isolation",
            "original-paper page text and section hints",
            "frozen uploaded R3 criterion formalization",
            "independently curated frozen paleomagnetic normative concepts",
        ],
        "forbidden_sources": [
            "PMAGRESULT and row ROCKMAG",
            "R3 labels or any existing assessment result",
            "prior predictions, OOF arrays, reports, and evaluation metrics",
        ],
        "knowledge_audit": knowledge_audit,
        "inputs": {
            "identity": {"path": str(paths["identity"]), "sha256": file_sha256(paths["identity"])},
            "page_cache": {"path": str(paths["page_cache"]), "sha256": file_sha256(paths["page_cache"])},
            "pdf_membership": {"path": str(paths["pdf_dir"]), "sha256": pdf_membership_sha256},
            "pdf_content_grouping": {
                "path": str(paths["pdf_dir"]),
                "sha256": json_sha256(pdf_sha_by_refno),
            },
            "normalized_text_content_grouping": {
                "path": str(paths["page_cache"]),
                "sha256": json_sha256(document_groups),
            },
            "criterion": {"path": str(paths["criterion"]), "sha256": file_sha256(paths["criterion"])},
            "normative_concepts": {"path": str(paths["concepts"]), "sha256": file_sha256(paths["concepts"])},
            "code": [
                {"path": str(path), "sha256": file_sha256(path)} for path in code_inputs
            ],
        },
        "outputs": {
            "features": {"path": str(feature_path), "sha256": file_sha256(feature_path)},
            "target_context": {"path": str(context_path), "sha256": file_sha256(context_path)},
            "fold_manifest": {"path": str(fold_path), "sha256": file_sha256(fold_path)},
            "audit": {"path": str(audit_path), "sha256": file_sha256(audit_path)},
            "concept_snapshot": {"path": str(concept_path), "sha256": file_sha256(concept_path)},
            "candidate_freeze": {"path": str(candidate_path), "sha256": file_sha256(candidate_path)},
        },
    }
    atomic_json(manifest_path, manifest)
    return {
        "features": feature_path,
        "target_context": context_path,
        "manifest": manifest_path,
        "audit": audit_path,
        "concept_snapshot": concept_path,
        "candidate_freeze": candidate_path,
        "fold_manifest": fold_path,
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--criterion", type=Path, default=DEFAULT_CRITERION)
    parser.add_argument("--concepts", type=Path, default=DEFAULT_CONCEPTS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args(argv)
    outputs = build_frozen_trace_features(
        identity_path=args.identity,
        page_cache_path=args.page_cache,
        pdf_dir=args.pdf_dir,
        criterion_path=args.criterion,
        concepts_path=args.concepts,
        output_dir=args.output_dir,
    )
    matrix = np.load(outputs["features"])["trace"]
    print(
        f"frozen TRACE-PMAG features: rows={matrix.shape[0]}, "
        f"features={matrix.shape[1]}, output={outputs['features']}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "build_frozen_trace_features",
    "build_trace_concept_snapshot",
    "load_sanitized_page_cache",
    "load_three_column_identity",
]
