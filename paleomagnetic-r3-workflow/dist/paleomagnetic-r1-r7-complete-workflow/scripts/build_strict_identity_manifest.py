#!/usr/bin/env python3
"""Build a scoring-safe identity manifest from an Excel workbook.

Only REFNO, RESULTNO, and ROCKNAME are read from data rows. ROCKUNITNO may be
included explicitly as auxiliary identity metadata, but it is never part of
the minimum scoring identity. Criterion values (R1-R7) are not loaded.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any, Sequence

import pandas as pd
from openpyxl import load_workbook


SCORING_IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
AUXILIARY_IDENTITY_COLUMN = "ROCKUNITNO"


class ManifestValidationError(ValueError):
    """Raised when an input cannot form a strict identity manifest."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def workbook_header(path: Path, sheet_name: str | None = None) -> tuple[str, list[str]]:
    """Read only the header row, returning the selected sheet and its columns."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        selected = sheet_name or workbook.sheetnames[0]
        if selected not in workbook.sheetnames:
            raise ManifestValidationError(
                f"Worksheet {selected!r} is not present in {path}"
            )
        worksheet = workbook[selected]
        values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = [_clean(value) for value in values]
    finally:
        workbook.close()
    if not columns or not any(columns):
        raise ManifestValidationError(f"{path} has no header row")
    return selected, columns


def _resolve_columns(header: Sequence[str], include_rockunitno: bool) -> list[str]:
    normalized: dict[str, str] = {}
    duplicates: list[str] = []
    for raw_column in header:
        key = raw_column.casefold()
        if not key:
            continue
        if key in normalized:
            duplicates.append(raw_column)
        else:
            normalized[key] = raw_column
    if duplicates:
        raise ManifestValidationError(
            "Input header contains duplicate column names (case-insensitive): "
            + ", ".join(duplicates)
        )

    required_missing = [
        column
        for column in SCORING_IDENTITY_COLUMNS
        if column.casefold() not in normalized
    ]
    if required_missing:
        raise ManifestValidationError(
            "Missing required identity columns: " + ", ".join(required_missing)
        )

    selected = [normalized[column.casefold()] for column in SCORING_IDENTITY_COLUMNS]
    if include_rockunitno:
        actual = normalized.get(AUXILIARY_IDENTITY_COLUMN.casefold())
        if actual is None:
            raise ManifestValidationError(
                f"--include-rockunitno was requested, but {AUXILIARY_IDENTITY_COLUMN} "
                "is not present"
            )
        selected.append(actual)
    return selected


def load_identity_rows(
    input_path: Path,
    *,
    sheet_name: str | None = None,
    include_rockunitno: bool = False,
) -> tuple[pd.DataFrame, str, list[str]]:
    """Load identity values using an explicit allowlist at the Excel reader."""

    selected_sheet, header = workbook_header(input_path, sheet_name)
    source_columns = _resolve_columns(header, include_rockunitno)
    frame = pd.read_excel(
        input_path,
        sheet_name=selected_sheet,
        usecols=source_columns,
        dtype=str,
        keep_default_na=False,
    )
    rename = {
        source: canonical
        for source, canonical in zip(
            source_columns,
            (*SCORING_IDENTITY_COLUMNS,)
            + ((AUXILIARY_IDENTITY_COLUMN,) if include_rockunitno else ()),
        )
    }
    frame = frame.rename(columns=rename)
    output_columns = list(rename.values())
    frame = frame[output_columns].copy()
    for column in output_columns:
        frame[column] = frame[column].map(_clean)

    empty_locations: list[str] = []
    for column in output_columns:
        empty_rows = frame.index[frame[column].eq("")].tolist()
        if empty_rows:
            excel_rows = ", ".join(str(index + 2) for index in empty_rows[:10])
            empty_locations.append(f"{column} (Excel row(s) {excel_rows})")
    if empty_locations:
        raise ManifestValidationError(
            "Identity values must be non-empty: " + "; ".join(empty_locations)
        )

    duplicated_resultno = frame["RESULTNO"].duplicated(keep=False)
    if duplicated_resultno.any():
        examples = ", ".join(
            frame.loc[duplicated_resultno, "RESULTNO"].drop_duplicates().head(10)
        )
        raise ManifestValidationError(
            "RESULTNO must be globally unique; duplicate value(s): " + examples
        )
    duplicated_identity = frame.duplicated(list(SCORING_IDENTITY_COLUMNS), keep=False)
    if duplicated_identity.any():
        examples = ", ".join(
            "/".join(frame.loc[index, list(SCORING_IDENTITY_COLUMNS)])
            for index in frame.index[duplicated_identity][:10]
        )
        raise ManifestValidationError(
            "The REFNO/RESULTNO/ROCKNAME identity tuple must be unique: " + examples
        )
    if frame.empty:
        raise ManifestValidationError(f"{input_path} contains no identity rows")
    return frame, selected_sheet, header


def build_manifest(
    input_path: Path,
    output_path: Path,
    audit_path: Path,
    *,
    sheet_name: str | None = None,
    include_rockunitno: bool = False,
) -> dict[str, Any]:
    input_path = input_path.resolve()
    output_path = output_path.resolve()
    audit_path = audit_path.resolve()
    if input_path == output_path:
        raise ManifestValidationError("Input and output workbook paths must differ")
    if not input_path.is_file():
        raise ManifestValidationError(f"Input workbook does not exist: {input_path}")

    input_hash = file_sha256(input_path)
    frame, selected_sheet, header = load_identity_rows(
        input_path,
        sheet_name=sheet_name,
        include_rockunitno=include_rockunitno,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_excel(output_path, index=False, engine="openpyxl")
    output_hash = file_sha256(output_path)

    output_columns = frame.columns.tolist()
    audit = {
        "schema_version": "strict_identity_manifest_audit_v1",
        "status": "PASS",
        "input": {
            "path": str(input_path),
            "sha256": input_hash,
            "sheet": selected_sheet,
            "columns_present": header,
            "columns_read": output_columns,
        },
        "output": {
            "path": str(output_path),
            "sha256": output_hash,
            "columns_written": output_columns,
            "rows": len(frame),
        },
        "scoring_identity_columns": list(SCORING_IDENTITY_COLUMNS),
        "auxiliary_identity_columns": (
            [AUXILIARY_IDENTITY_COLUMN] if include_rockunitno else []
        ),
        "criterion_value_columns_read": [],
        "guarantees": {
            "identity_values_non_empty": True,
            "resultno_globally_unique": True,
            "identity_tuple_unique": True,
            "r1_r7_values_read": False,
            "rockunitno_used_for_scoring": False,
        },
    }
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return audit


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Source XLSX workbook.")
    parser.add_argument("--output", type=Path, required=True, help="Identity-only XLSX output.")
    parser.add_argument(
        "--audit-json",
        type=Path,
        required=True,
        help="Machine-readable provenance and boundary audit.",
    )
    parser.add_argument("--sheet", default=None, help="Worksheet name; defaults to the first sheet.")
    parser.add_argument(
        "--include-rockunitno",
        action="store_true",
        help="Include ROCKUNITNO as auxiliary output identity metadata.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        audit = build_manifest(
            args.input,
            args.output,
            args.audit_json,
            sheet_name=args.sheet,
            include_rockunitno=args.include_rockunitno,
        )
    except (ManifestValidationError, OSError, ValueError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
