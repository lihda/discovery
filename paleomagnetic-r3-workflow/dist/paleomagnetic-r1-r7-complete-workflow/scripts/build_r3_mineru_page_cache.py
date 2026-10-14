#!/usr/bin/env python3
"""Optionally build or patch the strict R3 page cache with MinerU output.

This adapter is deliberately separate from ``build_r3_page_source_cache.py``;
the default PDF extraction path is unchanged.  Runtime inputs are restricted
to an identity-only manifest, local PDFs, an optional existing page cache, and
MinerU's own output.  Labels, PMAGRESULT, prior evaluations, and predictions
are not accepted by the command.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
import hashlib
from html.parser import HTMLParser
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Callable, Iterable, Mapping, Optional, Sequence

from openpyxl import load_workbook

try:
    from .build_r3_distilled_predictions import load_page_cache
except ImportError:
    from build_r3_distilled_predictions import load_page_cache


IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
SECTION_HINT = re.compile(
    r"^(?:\d+(?:\.\d+)*\s+)?"
    r"(abstract|introduction|geological setting|geology|sampling|methods?|"
    r"palaeomagnetic methods?|paleomagnetic methods?|rock magneti(?:c|cs)|"
    r"magnetic mineralogy|results?|discussion|conclusions?|appendix|supplement)",
    re.I,
)
MARKDOWN_PAGE_MARKER = re.compile(
    r"(?im)^(?:\s*<!--\s*)?(?:#{1,6}\s*)?page[_\s:-]*(\d+)"
    r"(?:\s*-->)?\s*$"
)
MARKDOWN_TABLE_SEPARATOR = re.compile(r"^:?-{3,}:?$")
QUALITY_KEYWORDS = (
    "abstract",
    "introduction",
    "discussion",
    "conclusion",
    "magnetic",
)
QUALITY_COVERAGE_PATTERNS = (
    re.compile(r"\babstract\b", re.I),
    re.compile(r"\bintroduction\b", re.I),
    re.compile(r"\bmethods?\b", re.I),
    re.compile(r"\bresults?\b", re.I),
    re.compile(r"\bdiscussion\b", re.I),
    re.compile(r"\bconclusions?\b", re.I),
    re.compile(r"\bmagneti(?:c|zation)\b", re.I),
    re.compile(r"\breman(?:ence|ent)\b", re.I),
    re.compile(r"\bpalaeomagnet\w*|\bpaleomagnet\w*", re.I),
    re.compile(r"\bminer(?:al|alogy|alogical)\w*", re.I),
    re.compile(r"\bhysteresis\b", re.I),
    re.compile(r"\bIRM\b", re.I),
    re.compile(r"\bsusceptibility\b", re.I),
)
UNRELIABLE_PAGE_BOUNDARY_FLAGS = frozenset(
    {
        "mineru_page_number_missing",
        "mineru_page_boundaries_unavailable",
    }
)
R3_CRITICAL_CONCEPT_PATTERNS = (
    ("irm_or_sirm", re.compile(r"\b(?:IRM|SIRM)\b|isothermal remanent magneti[sz]ation", re.I)),
    ("hysteresis_or_forc", re.compile(r"\bhysteresis\b|\bFORC\b|\bDay\s+plot\b", re.I)),
    (
        "thermomagnetic_or_curie",
        re.compile(
            r"thermo[- ]?magnetic|Curie\s+(?:point|temperature)|"
            r"(?:susceptibility|magneti[sz]ation).{0,80}temperature|"
            r"temperature.{0,80}(?:susceptibility|magneti[sz]ation)|"
            r"\bJ(?:s|rs)[- ]?\(?T\)?\b",
            re.I | re.S,
        ),
    ),
    (
        "coercivity_or_backfield",
        re.compile(r"back[- ]field|coerciv(?:e|ity)|\bHc(?:r)?\b|\bBcr\b", re.I),
    ),
    ("lowrie_or_three_axis", re.compile(r"\bLowrie(?:[- ]Fuller)?\b|three[- ]axis|3[- ]axis", re.I)),
    (
        "mineralogical_observation",
        re.compile(
            r"reflected[- ]light|ore\s+microscop|\bSEM\b|\bTEM\b|"
            r"electron\s+microprobe|\bEPMA\b|X[- ]?ray\s+diffraction|\bXRD\b|"
            r"M[o\u00f6]ssbauer|mineralogical\s+analysis",
            re.I,
        ),
    ),
    (
        "magnetic_anisotropy",
        re.compile(
            r"\bAMS\b|anisotropy\s+of\s+magnetic\s+susceptibility|"
            r"\bAIR\b|\bAAR\b|magnetic\s+fabric",
            re.I,
        ),
    ),
    (
        "named_carrier",
        re.compile(
            r"\b(?:titano[- ]?magnetite|titano[- ]?maghemite|magnetite|hematite|"
            r"haematite|maghemite|goethite|pyrrhotite|greigite|ilmenite|"
            r"specularite|martite|Fe[- ]Ti\s+oxide(?:s)?)\b",
            re.I,
        ),
    ),
    (
        "carrier_remanence_link",
        re.compile(
            r"(?:carrier|carried\s+by|carries|carrying).{0,160}"
            r"(?:remanence|remanent\s+magneti[sz]ation|\bChRM\b|component)|"
            r"(?:remanence|remanent\s+magneti[sz]ation|\bChRM\b|component).{0,160}"
            r"(?:carrier|carried\s+by|carries|carrying)",
            re.I | re.S,
        ),
    ),
)
DEFAULT_QUALITY_GATE_FLAGS = frozenset(
    {
        "flattened_layout",
        "no_extractable_text",
        "short_text",
        "glyph_or_cid_text",
        "possibly_incomplete_extraction",
        "poor_ocr",
        "low_ocr_quality",
        "missing_pages",
        "truncated_text",
        *UNRELIABLE_PAGE_BOUNDARY_FLAGS,
    }
)
STRONG_RISK_FLAGS = frozenset(
    {
        "no_extractable_text",
        "glyph_or_cid_text",
        "poor_ocr",
        "low_ocr_quality",
        "missing_pages",
        "truncated_text",
    }
)
ARBITRATION_POLICY = {
    "schema_version": "mineru_text_quality_arbitration_v3",
    "risk_flags": sorted(DEFAULT_QUALITY_GATE_FLAGS),
    "strong_risk_flags": sorted(STRONG_RISK_FLAGS),
    "unreliable_page_boundary_flags": sorted(UNRELIABLE_PAGE_BOUNDARY_FLAGS),
    "requires_r3_concept_preservation": True,
    "minimum_character_ratio": 0.80,
    "minimum_page_ratio": 0.75,
    "character_gain_ratio": 1.15,
    "character_gain_minimum": 300,
    "major_character_gain_ratio": 1.50,
    "major_character_gain_minimum": 500,
    "glyph_noise_absolute_change": 5,
    "glyph_noise_reduction_ratio": 0.50,
    "glyph_noise_increase_ratio": 0.25,
    "keyword_loss_tolerance": 1,
    "strong_risk_minimum_character_ratio": 0.90,
    "strong_risk_minimum_page_ratio": 0.80,
    "multiple_improvements_required": 2,
    "flattened_layout_minimum_pages": 5,
    "flattened_layout_minimum_characters": 5000,
    "layout_recovery_minimum_character_ratio": 0.80,
    "layout_recovery_minimum_page_ratio": 0.90,
    "layout_recovery_minimum_structured_pages": 2,
    "layout_recovery_minimum_structured_page_ratio": 0.25,
    "layout_recovery_minimum_boundary_gain": 4,
}
PAGE_FUSION_POLICY = {
    "schema_version": "mineru_page_source_fusion_v1",
    "minimum_candidate_character_ratio": 0.50,
    "structure_recovery_minimum_character_ratio": 0.70,
    "table_recovery_minimum_character_ratio": 0.60,
    "noise_recovery_minimum_character_ratio": 0.60,
    "minimum_structure_boundary_gain": 2,
    "minimum_table_row_gain": 2,
    "character_gain_ratio": 1.15,
    "character_gain_minimum": 120,
    "glyph_noise_reduction_minimum": 1,
    "glyph_noise_reduction_ratio": 0.50,
    "minimum_block_confidence_mean": 0.80,
    "maximum_low_confidence_block_ratio": 0.35,
    "low_confidence_block_threshold": 0.80,
    "requires_r3_concept_preservation": True,
    "allows_cross_source_page_concatenation": False,
}


class MineruAdapterError(ValueError):
    """Raised when identity, cache, MinerU, or output data violates the contract."""


class MineruUnavailableError(MineruAdapterError):
    """Raised when the configured MinerU executable cannot be found."""


class MineruExecutionError(MineruAdapterError):
    """Raised when MinerU exits unsuccessfully."""


class MineruOutputError(MineruAdapterError):
    """Raised when MinerU output cannot be converted to page-aware text."""


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_text(text: Any) -> str:
    value = "" if text is None else str(text)
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r" +", " ", value)
    value = re.sub(r" *\t *", "\t", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def text_fingerprint(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text).strip()
    return hashlib.sha256(normalized.encode("utf-8", errors="ignore")).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def bounded_stderr_tail(value: Any, limit: int = 1000) -> str:
    text = _clean(value)
    if len(text) <= limit:
        return text
    marker = "...[stderr truncated; showing final output]...\n"
    return marker + text[-(limit - len(marker)) :]


def section_hints(text: str) -> list[str]:
    hints: list[str] = []
    for raw_line in text.splitlines()[:100]:
        line = _clean(raw_line)
        if line and len(line) <= 120 and SECTION_HINT.search(line):
            hints.append(line[:120])
        if len(hints) >= 6:
            break
    return hints


def _glyph_noise(text: str) -> int:
    return text.count("(cid:") + len(re.findall(r"/C\d{2,3}", text))


def _replacement_noise(text: str) -> int:
    """Count explicit decoding-replacement/control artifacts, not scientific glyphs."""

    return text.count("\ufffd") + text.count("\x00")


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"td", "th"}:
            self.parts.append("\t")
        elif tag.casefold() in {"tr", "p", "div", "li", "br"}:
            self.parts.append("\n")

    def text(self) -> str:
        value = " ".join(self.parts)
        value = re.sub(r" *\t *", "\t", value)
        value = re.sub(r" *\n *", "\n", value)
        return normalize_text(value)


def html_to_text(value: str) -> str:
    parser = _HTMLTextExtractor()
    try:
        parser.feed(value)
        parser.close()
    except Exception as exc:
        raise MineruOutputError(f"Could not parse MinerU HTML table: {exc}") from exc
    return parser.text()


def _markdown_row(line: str) -> Optional[list[str]]:
    stripped = line.strip()
    if "|" not in stripped:
        return None
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    if not cells:
        return None
    if all(MARKDOWN_TABLE_SEPARATOR.fullmatch(cell.replace(" ", "")) for cell in cells):
        return []
    return cells


def markdown_to_text(value: str) -> str:
    text = value
    table_pattern = re.compile(r"(?is)<table\b.*?</table>")
    text = table_pattern.sub(lambda match: "\n" + html_to_text(match.group(0)) + "\n", text)
    text = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(?m)^\s*```[^\n]*$", "", text)
    converted: list[str] = []
    for raw_line in text.splitlines():
        row = _markdown_row(raw_line)
        if row is not None:
            if row:
                converted.append("\t".join(row))
            continue
        line = re.sub(r"^\s{0,3}#{1,6}\s*", "", raw_line)
        line = re.sub(r"^\s*[-*+]\s+", "", line)
        line = re.sub(r"[*_~`]", "", line)
        converted.append(line)
    return normalize_text("\n".join(converted))


def _flatten_text(value: Any) -> list[str]:
    if isinstance(value, str):
        cleaned = normalize_text(value)
        return [cleaned] if cleaned else []
    if isinstance(value, list):
        flattened: list[str] = []
        for item in value:
            flattened.extend(_flatten_text(item))
        return flattened
    return []


def _page_number(entry: Mapping[str, Any]) -> tuple[int, bool]:
    for key in ("page_idx", "page_index"):
        value = entry.get(key)
        if isinstance(value, int) and not isinstance(value, bool) and value >= 0:
            return value + 1, False
    for key in ("page_no", "page_number", "page"):
        value = entry.get(key)
        if isinstance(value, int) and not isinstance(value, bool):
            return (value if value >= 1 else value + 1), False
    return 1, True


def _table_block_text(entry: Mapping[str, Any]) -> tuple[str, bool]:
    parts: list[str] = []
    for key in ("table_caption", "caption"):
        parts.extend(_flatten_text(entry.get(key)))
    body_found = False
    for key in (
        "table_body",
        "table_html",
        "html",
        "table_markdown",
        "markdown",
        "text",
        "content",
    ):
        raw = entry.get(key)
        if not isinstance(raw, str) or not raw.strip():
            continue
        body_found = True
        if re.search(r"<\s*(?:table|tr|td|th)\b", raw, re.I):
            converted = html_to_text(raw)
        else:
            converted = markdown_to_text(raw)
        if converted:
            parts.append(converted)
        break
    for key in ("table_footnote", "footnote"):
        parts.extend(_flatten_text(entry.get(key)))
    return normalize_text("\n".join(parts)), body_found


def _content_entries(document: Any) -> list[Mapping[str, Any]]:
    if isinstance(document, list):
        if all(isinstance(item, Mapping) for item in document):
            return list(document)
        raise MineruOutputError("MinerU content list contains a non-object entry")
    if not isinstance(document, Mapping):
        raise MineruOutputError("MinerU content list must be an array or object")
    for key in ("content_list", "contents", "items", "blocks"):
        value = document.get(key)
        if isinstance(value, list) and all(isinstance(item, Mapping) for item in value):
            return list(value)
    pages = document.get("pages")
    if isinstance(pages, list):
        entries: list[Mapping[str, Any]] = []
        for page_position, page in enumerate(pages):
            if not isinstance(page, Mapping):
                continue
            blocks = page.get("blocks") or page.get("items") or page.get("contents")
            if not isinstance(blocks, list):
                continue
            page_number, inferred = _page_number(page)
            if inferred:
                # The enclosing pages array itself provides a reliable boundary.
                page_number = page_position + 1
            for block in blocks:
                if isinstance(block, Mapping):
                    copied = dict(block)
                    copied.setdefault("page_no", page_number)
                    entries.append(copied)
        return entries
    raise MineruOutputError("MinerU JSON does not contain a recognized content list")


def _bounded_confidence(value: Any) -> Optional[float]:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    confidence = float(value)
    if not math.isfinite(confidence) or not 0.0 <= confidence <= 1.0:
        return None
    return confidence


def _confidence_summary(values: Iterable[Any]) -> dict[str, Any]:
    scores = [score for value in values if (score := _bounded_confidence(value)) is not None]
    if not scores:
        return {}
    low_threshold = PAGE_FUSION_POLICY["low_confidence_block_threshold"]
    return {
        "block_confidence_mean": round(sum(scores) / len(scores), 6),
        "block_confidence_min": round(min(scores), 6),
        "block_confidence_count": len(scores),
        "low_confidence_block_ratio": round(
            sum(score < low_threshold for score in scores) / len(scores),
            6,
        ),
    }


def _entry_confidences(entry: Mapping[str, Any]) -> list[float]:
    scores: list[float] = []
    for key in ("score", "confidence", "layout_score", "ocr_score"):
        score = _bounded_confidence(entry.get(key))
        if score is not None:
            scores.append(score)
    return scores


def load_mineru_model_page_confidence(path: Path) -> dict[int, dict[str, Any]]:
    """Aggregate MinerU layout-block scores without using paper labels or outcomes."""

    try:
        document = json.loads(path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise MineruOutputError(f"Invalid MinerU model JSON {path}: {exc}") from exc
    if not isinstance(document, list) or not all(
        isinstance(page, Mapping) for page in document
    ):
        raise MineruOutputError(f"MinerU model JSON must contain a page array: {path}")
    by_page: dict[int, dict[str, Any]] = {}
    for position, page in enumerate(document, start=1):
        page_info = page.get("page_info")
        raw_page_number = (
            page_info.get("page_no") if isinstance(page_info, Mapping) else None
        )
        if (
            isinstance(raw_page_number, bool)
            or not isinstance(raw_page_number, int)
            or raw_page_number < 0
        ):
            page_number = position
        else:
            page_number = raw_page_number + 1
        detections = page.get("layout_dets")
        if not isinstance(detections, list):
            continue
        layout_scores = [
            detection.get("score")
            for detection in detections
            if isinstance(detection, Mapping)
            and _clean(detection.get("label")).casefold() != "ocr_text"
        ]
        summary = _confidence_summary(layout_scores)
        if not summary:
            summary = _confidence_summary(
                detection.get("score")
                for detection in detections
                if isinstance(detection, Mapping)
            )
        if summary:
            by_page[page_number] = summary
    return by_page


def attach_mineru_model_confidence(
    parsed: "ParsedMineruOutput",
    model_path: Path,
) -> None:
    confidence_by_page = load_mineru_model_page_confidence(model_path)
    for page in parsed.pages:
        page_number = page.get("page")
        if isinstance(page_number, int) and page_number in confidence_by_page:
            page.update(confidence_by_page[page_number])


@dataclass
class ParsedMineruOutput:
    pages: list[dict[str, Any]]
    parsed_path: Optional[Path]
    source_kind: str
    quality_flags: list[str]
    table_blocks: int


def parse_content_list(path: Path) -> ParsedMineruOutput:
    try:
        document = json.loads(path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise MineruOutputError(f"Invalid MinerU content list {path}: {exc}") from exc
    entries = _content_entries(document)
    page_parts: dict[int, list[str]] = defaultdict(list)
    page_types: dict[int, set[str]] = defaultdict(set)
    page_confidences: dict[int, list[float]] = defaultdict(list)
    page_table_blocks: Counter[int] = Counter()
    flags: set[str] = set()
    table_blocks = 0
    for entry in entries:
        kind = _clean(entry.get("type") or entry.get("content_type") or "text").casefold()
        page, inferred = _page_number(entry)
        if inferred:
            flags.add("mineru_page_number_missing")
        text = ""
        source_type = "text"
        if "table" in kind:
            table_blocks += 1
            page_table_blocks[page] += 1
            source_type = "table"
            text, body_found = _table_block_text(entry)
            if body_found and not text:
                flags.add("mineru_table_text_missing")
        elif kind in {"text", "title", "heading", "paragraph", "list"} or "text" in kind:
            raw = entry.get("text", entry.get("content", ""))
            if isinstance(raw, str):
                text = markdown_to_text(raw)
        elif kind in {"image", "figure", "chart"}:
            captions: list[str] = []
            for key in (
                "image_caption",
                "figure_caption",
                "chart_caption",
                "caption",
                "image_footnote",
                "chart_footnote",
            ):
                captions.extend(_flatten_text(entry.get(key)))
            text = normalize_text("\n".join(captions))
            source_type = "caption"
        if text:
            page_parts[page].append(text)
            page_types[page].add(source_type)
            page_confidences[page].extend(_entry_confidences(entry))

    raw_pages: list[dict[str, Any]] = []
    for page_number in sorted(page_parts):
        text = normalize_text("\n\n".join(page_parts[page_number]))
        types = page_types[page_number]
        source = "mineru_" + "_".join(sorted(types))
        raw_page = {
            "page": page_number,
            "text": text,
            "source": source,
            "table_blocks": page_table_blocks[page_number],
        }
        raw_page.update(_confidence_summary(page_confidences[page_number]))
        raw_pages.append(raw_page)
    if not raw_pages:
        flags.add("no_extractable_text")
    source_kind = "mineru_content_list"
    if table_blocks:
        source_kind += "_with_tables"
    return ParsedMineruOutput(raw_pages, path, source_kind, sorted(flags), table_blocks)


def parse_markdown(path: Path) -> ParsedMineruOutput:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise MineruOutputError(f"MinerU Markdown is not UTF-8: {path}") from exc
    flags: list[str] = []
    pages: list[dict[str, Any]] = []
    markers = list(MARKDOWN_PAGE_MARKER.finditer(raw))
    if markers:
        for index, marker in enumerate(markers):
            start = marker.end()
            end = markers[index + 1].start() if index + 1 < len(markers) else len(raw)
            text = markdown_to_text(raw[start:end])
            if text:
                pages.append(
                    {
                        "page": int(marker.group(1)),
                        "text": text,
                        "source": "mineru_markdown",
                    }
                )
    elif "\f" in raw:
        for page_number, chunk in enumerate(raw.split("\f"), start=1):
            text = markdown_to_text(chunk)
            if text:
                pages.append(
                    {"page": page_number, "text": text, "source": "mineru_markdown"}
                )
    else:
        text = markdown_to_text(raw)
        if text:
            pages.append({"page": 1, "text": text, "source": "mineru_markdown"})
        flags.append("mineru_page_boundaries_unavailable")
    if not pages:
        flags.append("no_extractable_text")
    table_blocks = len(re.findall(r"(?i)<table\b", raw)) + sum(
        1 for line in raw.splitlines() if _markdown_row(line) == []
    )
    return ParsedMineruOutput(
        pages,
        path,
        "mineru_markdown_with_tables" if table_blocks else "mineru_markdown",
        sorted(set(flags)),
        table_blocks,
    )


def _matching_outputs(root: Path, pdf_stem: str, pattern: str) -> list[Path]:
    candidates = [path for path in root.rglob(pattern) if path.is_file()]
    return sorted(
        candidates,
        key=lambda path: (
            pdf_stem.casefold() not in path.name.casefold(),
            pdf_stem.casefold() not in str(path.parent).casefold(),
            len(path.parts),
            str(path),
        ),
    )


def _attach_available_model_confidence(
    parsed: ParsedMineruOutput,
    root: Path,
    pdf_stem: str,
) -> None:
    model_paths = _matching_outputs(root, pdf_stem, "*model.json")
    if not model_paths:
        return
    try:
        attach_mineru_model_confidence(parsed, model_paths[0])
    except MineruOutputError:
        parsed.quality_flags = sorted(
            {*parsed.quality_flags, "mineru_block_confidence_unavailable"}
        )


def parse_mineru_directory(root: Path, pdf_stem: str) -> ParsedMineruOutput:
    errors: list[str] = []
    for path in _matching_outputs(root, pdf_stem, "*content_list.json"):
        try:
            parsed = parse_content_list(path)
        except MineruOutputError as exc:
            errors.append(str(exc))
            continue
        if parsed.pages:
            _attach_available_model_confidence(parsed, root, pdf_stem)
            return parsed
        errors.append(f"{path} contained no extractable page text")
    markdown_candidates = _matching_outputs(root, pdf_stem, "*.md")
    for path in markdown_candidates:
        try:
            parsed = parse_markdown(path)
        except MineruOutputError as exc:
            errors.append(str(exc))
            continue
        if parsed.pages:
            _attach_available_model_confidence(parsed, root, pdf_stem)
            return parsed
        errors.append(f"{path} contained no extractable page text")
    detail = "; ".join(errors[:5]) if errors else "no content-list JSON or Markdown found"
    raise MineruOutputError(f"No usable MinerU output for {pdf_stem}: {detail}")


class MineruCLI:
    """Small injectable wrapper around the MinerU command-line interface."""

    def __init__(
        self,
        executable: str = "mineru",
        *,
        extra_args: Sequence[str] = (),
        timeout_seconds: int = 900,
        runner: Callable[..., Any] = subprocess.run,
        which: Callable[[str], Optional[str]] = shutil.which,
    ) -> None:
        self.executable = executable
        self.extra_args = tuple(extra_args)
        self.timeout_seconds = timeout_seconds
        self.runner = runner
        self.which = which

    def resolved_executable(self) -> Optional[str]:
        if "/" in self.executable or "\\" in self.executable:
            path = Path(self.executable).expanduser()
            return str(path.resolve()) if path.is_file() else None
        return self.which(self.executable)

    def available(self) -> bool:
        return self.resolved_executable() is not None

    def run(self, pdf_path: Path, output_dir: Path) -> Sequence[str]:
        executable = self.resolved_executable()
        if executable is None:
            raise MineruUnavailableError(
                f"MinerU executable is unavailable: {self.executable!r}"
            )
        command = [
            executable,
            "-p",
            str(pdf_path),
            "-o",
            str(output_dir),
            *self.extra_args,
        ]
        try:
            self.runner(
                command,
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=self.timeout_seconds,
            )
        except FileNotFoundError as exc:
            raise MineruUnavailableError(
                f"MinerU executable disappeared before invocation: {executable}"
            ) from exc
        except subprocess.TimeoutExpired as exc:
            raise MineruExecutionError(
                f"MinerU timed out after {self.timeout_seconds}s for {pdf_path.name}"
            ) from exc
        except subprocess.CalledProcessError as exc:
            stderr = bounded_stderr_tail(getattr(exc, "stderr", ""))
            detail = f": {stderr}" if stderr else ""
            raise MineruExecutionError(
                f"MinerU failed for {pdf_path.name} with exit code {exc.returncode}{detail}"
            ) from exc
        return command


def load_identity_refnos(path: Path) -> list[str]:
    if not path.is_file() or path.suffix.casefold() != ".xlsx":
        raise MineruAdapterError(
            f"Identity manifest must be an existing XLSX file: {path}"
        )
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise MineruAdapterError("Identity manifest must contain exactly one worksheet")
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=False)
        header_cells = next(rows, ())
        if any(cell.data_type == "f" for cell in header_cells):
            raise MineruAdapterError("Identity manifest header formulas are forbidden")
        header = tuple(_clean(cell.value) for cell in header_cells)
        if header != IDENTITY_COLUMNS:
            raise MineruAdapterError(
                "Identity columns must be exactly REFNO, RESULTNO, ROCKNAME; "
                "labels and source-table fields are forbidden"
            )
        refnos: list[str] = []
        resultnos: set[str] = set()
        for row_number, cells in enumerate(rows, start=2):
            if len(cells) != 3 or any(cell.data_type == "f" for cell in cells):
                raise MineruAdapterError(
                    f"Identity row {row_number} has formulas or an invalid column count"
                )
            values = tuple(_clean(cell.value) for cell in cells)
            if not any(values):
                continue
            if any(not value for value in values):
                raise MineruAdapterError(f"Identity row {row_number} has an empty value")
            refno, resultno, _ = values
            if resultno in resultnos:
                raise MineruAdapterError(f"Duplicate RESULTNO in identity manifest: {resultno}")
            resultnos.add(resultno)
            refnos.append(refno)
    finally:
        workbook.close()
    unique_refnos = list(dict.fromkeys(refnos))
    if not unique_refnos:
        raise MineruAdapterError("Identity manifest contains no REFNO values")
    return unique_refnos


def find_pdf(pdf_dir: Path, refno: str) -> Optional[Path]:
    matches = [
        path
        for path in (pdf_dir / f"{refno}.pdf", pdf_dir / f"{refno}.PDF")
        if path.is_file()
    ]
    unique: list[Path] = []
    for path in matches:
        if any(path.samefile(existing) for existing in unique):
            continue
        unique.append(path.resolve())
    if len(unique) > 1:
        raise MineruAdapterError(f"Multiple PDFs match REFNO {refno}: {unique}")
    return unique[0] if unique else None


def _quality_flags(pages: Sequence[Mapping[str, Any]], initial: Iterable[str]) -> list[str]:
    flags = set(initial)
    text = "\n".join(str(page.get("text") or "") for page in pages)
    if not text.strip():
        flags.add("no_extractable_text")
    if len(text.strip()) < 2000:
        flags.add("short_text")
    if _glyph_noise(text) > 20:
        flags.add("glyph_or_cid_text")
    lower = text.casefold()
    if text and not any(keyword in lower for keyword in QUALITY_KEYWORDS):
        flags.add("possibly_incomplete_extraction")
    return sorted(flags)


def make_cache_record(
    refno: str,
    pdf_path: Optional[Path],
    parsed: Optional[ParsedMineruOutput],
    *,
    fallback_kind: str,
    fallback_flags: Sequence[str],
) -> dict[str, Any]:
    parsed_flags = set(parsed.quality_flags) if parsed is not None else set()
    # A flattened document must never masquerade as one physical page: R3 uses
    # same/adjacent-page locality as an evidence-chain requirement.
    raw_pages = (
        parsed.pages
        if parsed is not None
        and not (parsed_flags & UNRELIABLE_PAGE_BOUNDARY_FLAGS)
        else []
    )
    pages: list[dict[str, Any]] = []
    page_hashes: list[str] = []
    seen_pages: set[int] = set()
    for position, raw_page in enumerate(raw_pages, start=1):
        page_number = raw_page.get("page")
        if isinstance(page_number, bool) or not isinstance(page_number, int) or page_number < 1:
            page_number = position
        while page_number in seen_pages:
            page_number += 1
        seen_pages.add(page_number)
        text = normalize_text(raw_page.get("text"))
        if not text:
            continue
        digest = text_fingerprint(text)
        page_hashes.append(digest)
        pages.append(
            {
                "page": page_number,
                "source": _clean(raw_page.get("source")) or "mineru_text",
                "text": text,
                "text_sha256": digest,
                "section_hints": section_hints(text),
            }
        )
    initial_flags = list(fallback_flags)
    if parsed is not None:
        initial_flags.extend(parsed.quality_flags)
    flags = _quality_flags(pages, initial_flags)
    combined_text = "\n".join(page["text"] for page in pages)
    return {
        "REFNO": refno,
        "source": {
            "refno": refno,
            "source_kind": parsed.source_kind if parsed is not None else fallback_kind,
            "parsed_path": (
                str(parsed.parsed_path)
                if parsed is not None and parsed.parsed_path is not None
                else None
            ),
            "pdf_path": str(pdf_path) if pdf_path is not None else None,
            "page_count": len(pages),
            "text_length": len(combined_text),
            "cache_sha256": hashlib.sha256(
                "".join(page_hashes).encode("ascii")
            ).hexdigest(),
            "source_quality_flags": flags,
        },
        "pages": pages,
    }


def _atomic_write_jsonl(path: Path, records: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
    temporary.replace(path)


def _atomic_write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def quality_gate_selection(
    refnos: Sequence[str],
    base_records: Mapping[str, Mapping[str, Any]],
    quality_flags: Iterable[str],
) -> tuple[set[str], dict[str, int]]:
    """Select records using only source quality metadata and cached page text."""

    configured_flags = {_clean(flag) for flag in quality_flags if _clean(flag)}
    if not configured_flags:
        raise MineruAdapterError("Quality gate requires at least one non-empty flag")
    selected: set[str] = set()
    reasons: Counter[str] = Counter()
    for refno in refnos:
        record = base_records[refno]
        source = record.get("source")
        pages = record.get("pages")
        if not isinstance(source, Mapping) or not isinstance(pages, list):
            raise MineruAdapterError(f"Base cache record {refno} has an invalid structure")
        record_flags = {
            _clean(flag)
            for flag in source.get("source_quality_flags", [])
            if _clean(flag)
        }
        if cache_text_quality(record).flattened_layout:
            record_flags.add("flattened_layout")
        matched_flags = sorted(record_flags & configured_flags)
        empty_pages = not pages
        empty_text = not any(
            isinstance(page, Mapping) and _clean(page.get("text")) for page in pages
        )
        if not matched_flags and not empty_pages and not empty_text:
            continue
        selected.add(refno)
        for flag in matched_flags:
            reasons[f"flag:{flag}"] += 1
        if empty_pages:
            reasons["empty_pages"] += 1
        if empty_text:
            reasons["empty_text"] += 1
    return selected, dict(sorted(reasons.items()))


@dataclass(frozen=True)
class CacheTextQuality:
    page_count: int
    character_count: int
    risk_flags: tuple[str, ...]
    glyph_noise: int
    keyword_coverage: int
    r3_concepts: tuple[str, ...]
    paragraph_boundaries: int
    structured_pages: int
    table_rows: int
    flattened_layout: bool


@dataclass(frozen=True)
class ReplacementDecision:
    replace: bool
    reason: str
    improvements: tuple[str, ...]
    regressions: tuple[str, ...]
    base: CacheTextQuality
    candidate: CacheTextQuality


@dataclass(frozen=True)
class PageTextQuality:
    character_count: int
    glyph_noise: int
    r3_concepts: tuple[str, ...]
    paragraph_boundaries: int
    table_rows: int
    block_confidence_mean: Optional[float]
    low_confidence_block_ratio: Optional[float]


@dataclass(frozen=True)
class PageFusionResult:
    record: dict[str, Any]
    changed: bool
    decisions: tuple[dict[str, Any], ...]
    candidate_pages_selected: int
    base_pages_selected: int
    candidate_only_pages: int
    base_only_pages: int
    selected_candidate_table_blocks: int


def cache_text_quality(record: Mapping[str, Any]) -> CacheTextQuality:
    source = record.get("source")
    pages = record.get("pages")
    if not isinstance(source, Mapping) or not isinstance(pages, list):
        raise MineruAdapterError("Cache quality comparison received an invalid record")
    nonempty_texts: list[str] = []
    for page in pages:
        if not isinstance(page, Mapping):
            continue
        text = normalize_text(page.get("text"))
        if text:
            nonempty_texts.append(text)
    text = normalize_text("\n".join(nonempty_texts))
    paragraph_boundaries = 0
    structured_pages = 0
    table_rows = 0
    for page_text in nonempty_texts:
        page_paragraph_boundaries = len(re.findall(r"\n[ \t]*\n", page_text))
        page_table_rows = sum(
            1
            for line in page_text.splitlines()
            if "\t" in line
            and len([cell for cell in line.split("\t") if cell.strip()]) >= 2
        )
        paragraph_boundaries += page_paragraph_boundaries
        table_rows += page_table_rows
        if page_paragraph_boundaries or page_table_rows:
            structured_pages += 1
    flattened_layout = bool(
        len(nonempty_texts) >= ARBITRATION_POLICY["flattened_layout_minimum_pages"]
        and len(text)
        >= ARBITRATION_POLICY["flattened_layout_minimum_characters"]
        and structured_pages == 0
    )
    flag_set = {
        _clean(flag)
        for flag in source.get("source_quality_flags", [])
        if _clean(flag) in DEFAULT_QUALITY_GATE_FLAGS
    }
    if flattened_layout:
        flag_set.add("flattened_layout")
    return CacheTextQuality(
        page_count=len(nonempty_texts),
        character_count=len(text),
        risk_flags=tuple(sorted(flag_set)),
        glyph_noise=_glyph_noise(text),
        keyword_coverage=sum(
            1 for pattern in QUALITY_COVERAGE_PATTERNS if pattern.search(text)
        ),
        r3_concepts=tuple(
            concept
            for concept, pattern in R3_CRITICAL_CONCEPT_PATTERNS
            if pattern.search(text)
        ),
        paragraph_boundaries=paragraph_boundaries,
        structured_pages=structured_pages,
        table_rows=table_rows,
        flattened_layout=flattened_layout,
    )


def page_text_quality(page: Mapping[str, Any]) -> PageTextQuality:
    """Measure one physical page without consulting labels or predictions."""

    text = normalize_text(page.get("text"))
    table_rows = sum(
        1
        for line in text.splitlines()
        if "\t" in line
        and len([cell for cell in line.split("\t") if cell.strip()]) >= 2
    )
    confidence_mean = _bounded_confidence(page.get("block_confidence_mean"))
    low_confidence_ratio = _bounded_confidence(
        page.get("low_confidence_block_ratio")
    )
    return PageTextQuality(
        character_count=len(text),
        glyph_noise=_glyph_noise(text),
        r3_concepts=tuple(
            concept
            for concept, pattern in R3_CRITICAL_CONCEPT_PATTERNS
            if pattern.search(text)
        ),
        paragraph_boundaries=len(re.findall(r"\n[ \t]*\n", text)),
        table_rows=table_rows,
        block_confidence_mean=confidence_mean,
        low_confidence_block_ratio=low_confidence_ratio,
    )


def _quality_audit(quality: PageTextQuality) -> dict[str, Any]:
    return {
        "character_count": quality.character_count,
        "glyph_noise": quality.glyph_noise,
        "r3_concepts": list(quality.r3_concepts),
        "paragraph_boundaries": quality.paragraph_boundaries,
        "table_rows": quality.table_rows,
        "block_confidence_mean": quality.block_confidence_mean,
        "low_confidence_block_ratio": quality.low_confidence_block_ratio,
    }


def _choose_page_source(
    page_number: int,
    base_page: Mapping[str, Any],
    candidate_page: Mapping[str, Any],
) -> tuple[str, str, dict[str, Any]]:
    """Choose a complete page from one source using the frozen quality policy."""

    base = page_text_quality(base_page)
    candidate = page_text_quality(candidate_page)
    lost_concepts = sorted(set(base.r3_concepts) - set(candidate.r3_concepts))
    character_ratio = (
        candidate.character_count / base.character_count
        if base.character_count
        else math.inf
    )
    audit = {
        "page": page_number,
        "base": _quality_audit(base),
        "candidate": _quality_audit(candidate),
        "candidate_character_ratio": (
            round(character_ratio, 6) if math.isfinite(character_ratio) else None
        ),
        "lost_r3_concepts": lost_concepts,
    }
    if candidate.character_count == 0:
        return "base", "candidate_page_empty", audit
    if lost_concepts and PAGE_FUSION_POLICY["requires_r3_concept_preservation"]:
        return "base", "candidate_page_r3_concept_loss", audit
    if character_ratio < PAGE_FUSION_POLICY["minimum_candidate_character_ratio"]:
        return "base", "candidate_page_character_loss", audit
    if (
        candidate.block_confidence_mean is not None
        and candidate.block_confidence_mean
        < PAGE_FUSION_POLICY["minimum_block_confidence_mean"]
    ):
        return "base", "candidate_page_low_block_confidence", audit
    if (
        candidate.low_confidence_block_ratio is not None
        and candidate.low_confidence_block_ratio
        > PAGE_FUSION_POLICY["maximum_low_confidence_block_ratio"]
    ):
        return "base", "candidate_page_many_low_confidence_blocks", audit

    boundary_gain = (
        candidate.paragraph_boundaries
        + candidate.table_rows
        - base.paragraph_boundaries
        - base.table_rows
    )
    table_gain = candidate.table_rows - base.table_rows
    structure_recovered = bool(
        character_ratio
        >= PAGE_FUSION_POLICY["structure_recovery_minimum_character_ratio"]
        and boundary_gain
        >= PAGE_FUSION_POLICY["minimum_structure_boundary_gain"]
    )
    table_recovered = bool(
        character_ratio
        >= PAGE_FUSION_POLICY["table_recovery_minimum_character_ratio"]
        and table_gain >= PAGE_FUSION_POLICY["minimum_table_row_gain"]
    )
    noise_recovered = bool(
        base.glyph_noise
        >= PAGE_FUSION_POLICY["glyph_noise_reduction_minimum"]
        and candidate.glyph_noise
        <= base.glyph_noise * PAGE_FUSION_POLICY["glyph_noise_reduction_ratio"]
        and character_ratio
        >= PAGE_FUSION_POLICY["noise_recovery_minimum_character_ratio"]
    )
    content_recovered = bool(
        character_ratio >= PAGE_FUSION_POLICY["character_gain_ratio"]
        and candidate.character_count - base.character_count
        >= PAGE_FUSION_POLICY["character_gain_minimum"]
    )
    audit.update(
        {
            "structure_boundary_gain": boundary_gain,
            "table_row_gain": table_gain,
            "improvements": [
                name
                for name, enabled in (
                    ("structure_recovery", structure_recovered),
                    ("table_recovery", table_recovered),
                    ("glyph_noise_recovery", noise_recovered),
                    ("content_recovery", content_recovered),
                )
                if enabled
            ],
        }
    )
    if table_recovered:
        return "candidate", "candidate_page_table_recovery", audit
    if structure_recovered:
        return "candidate", "candidate_page_structure_recovery", audit
    if noise_recovered:
        return "candidate", "candidate_page_glyph_recovery", audit
    if content_recovered:
        return "candidate", "candidate_page_content_recovery", audit
    return "base", "candidate_page_insufficient_improvement", audit


def fuse_cache_pages(
    base_record: Mapping[str, Any],
    candidate_record: Mapping[str, Any],
    parsed: ParsedMineruOutput,
    pdf_path: Path,
) -> PageFusionResult:
    """Fuse two page-aware parses while never concatenating competing page text."""

    base_pages = {
        int(page["page"]): page
        for page in base_record.get("pages", [])
        if isinstance(page, Mapping) and isinstance(page.get("page"), int)
    }
    candidate_pages = {
        int(page["page"]): dict(page)
        for page in candidate_record.get("pages", [])
        if isinstance(page, Mapping) and isinstance(page.get("page"), int)
    }
    raw_candidate_pages = {
        int(page["page"]): page
        for page in parsed.pages
        if isinstance(page, Mapping) and isinstance(page.get("page"), int)
    }
    for page_number, candidate_page in candidate_pages.items():
        raw_page = raw_candidate_pages.get(page_number, {})
        for key in (
            "block_confidence_mean",
            "block_confidence_min",
            "block_confidence_count",
            "low_confidence_block_ratio",
            "table_blocks",
        ):
            if key in raw_page:
                candidate_page[key] = raw_page[key]

    fused_pages: list[dict[str, Any]] = []
    decisions: list[dict[str, Any]] = []
    candidate_pages_selected = 0
    base_pages_selected = 0
    candidate_only_pages = 0
    base_only_pages = 0
    selected_candidate_table_blocks = 0
    for page_number in sorted(set(base_pages) | set(candidate_pages)):
        base_page = base_pages.get(page_number)
        candidate_page = candidate_pages.get(page_number)
        if base_page is None and candidate_page is not None:
            candidate_quality = page_text_quality(candidate_page)
            confidence_rejected = bool(
                candidate_quality.block_confidence_mean is not None
                and candidate_quality.block_confidence_mean
                < PAGE_FUSION_POLICY["minimum_block_confidence_mean"]
            ) or bool(
                candidate_quality.low_confidence_block_ratio is not None
                and candidate_quality.low_confidence_block_ratio
                > PAGE_FUSION_POLICY["maximum_low_confidence_block_ratio"]
            )
            if confidence_rejected:
                decisions.append(
                    {
                        "page": page_number,
                        "candidate": _quality_audit(candidate_quality),
                        "selected_source": "none",
                        "reason": "candidate_only_page_low_block_confidence",
                    }
                )
                continue
            selected = candidate_page
            selected_source = "candidate"
            reason = "candidate_only_physical_page"
            candidate_pages_selected += 1
            candidate_only_pages += 1
            selected_candidate_table_blocks += int(
                candidate_page.get("table_blocks") or 0
            )
            audit = {
                "page": page_number,
                "candidate": _quality_audit(candidate_quality),
            }
        elif candidate_page is None and base_page is not None:
            selected = base_page
            selected_source = "base"
            reason = "base_only_physical_page"
            base_pages_selected += 1
            base_only_pages += 1
            audit = {
                "page": page_number,
                "base": _quality_audit(page_text_quality(base_page)),
            }
        elif base_page is not None and candidate_page is not None:
            selected_source, reason, audit = _choose_page_source(
                page_number,
                base_page,
                candidate_page,
            )
            if selected_source == "candidate":
                selected = candidate_page
                candidate_pages_selected += 1
                selected_candidate_table_blocks += int(
                    candidate_page.get("table_blocks") or 0
                )
            else:
                selected = base_page
                base_pages_selected += 1
        else:  # pragma: no cover - impossible after the union above
            continue
        decisions.append(
            {
                **audit,
                "selected_source": selected_source,
                "reason": reason,
            }
        )
        fused_pages.append(
            {
                "page": page_number,
                "source": _clean(selected.get("source")) or selected_source,
                "text": normalize_text(selected.get("text")),
            }
        )

    changed = candidate_pages_selected > 0
    parsed_fusion = ParsedMineruOutput(
        pages=fused_pages,
        parsed_path=parsed.parsed_path,
        source_kind="page_fusion",
        quality_flags=[
            flag
            for flag in parsed.quality_flags
            if flag not in UNRELIABLE_PAGE_BOUNDARY_FLAGS
        ],
        table_blocks=selected_candidate_table_blocks,
    )
    record = make_cache_record(
        _clean(base_record.get("REFNO")),
        pdf_path,
        parsed_fusion,
        fallback_kind="page_fusion",
        fallback_flags=(),
    )
    return PageFusionResult(
        record=record,
        changed=changed,
        decisions=tuple(decisions),
        candidate_pages_selected=candidate_pages_selected,
        base_pages_selected=base_pages_selected,
        candidate_only_pages=candidate_only_pages,
        base_only_pages=base_only_pages,
        selected_candidate_table_blocks=selected_candidate_table_blocks,
    )


def compare_cache_quality(
    base_record: Mapping[str, Any],
    candidate_record: Mapping[str, Any],
) -> ReplacementDecision:
    """Apply a deterministic, label-free MinerU replacement policy."""

    base = cache_text_quality(base_record)
    candidate = cache_text_quality(candidate_record)
    candidate_boundary_risks = sorted(
        set(candidate.risk_flags) & UNRELIABLE_PAGE_BOUNDARY_FLAGS
    )
    if candidate_boundary_risks:
        return ReplacementDecision(
            False,
            "unreliable_page_boundaries",
            (),
            tuple(candidate_boundary_risks),
            base,
            candidate,
        )
    if candidate.character_count == 0:
        return ReplacementDecision(
            False,
            "candidate_empty",
            (),
            ("empty_text",),
            base,
            candidate,
        )
    if base.character_count == 0:
        return ReplacementDecision(
            True,
            "base_empty",
            ("nonempty_text",),
            (),
            base,
            candidate,
        )

    char_ratio = candidate.character_count / base.character_count
    page_ratio = candidate.page_count / base.page_count
    improvements: list[str] = []
    regressions: list[str] = []
    lost_r3_concepts = sorted(set(base.r3_concepts) - set(candidate.r3_concepts))
    required_structured_pages = max(
        ARBITRATION_POLICY["layout_recovery_minimum_structured_pages"],
        math.ceil(
            candidate.page_count
            * ARBITRATION_POLICY["layout_recovery_minimum_structured_page_ratio"]
        ),
    )
    structure_boundary_gain = (
        candidate.paragraph_boundaries
        + candidate.table_rows
        - base.paragraph_boundaries
        - base.table_rows
    )
    significant_layout_recovery = bool(
        base.flattened_layout
        and not candidate.flattened_layout
        and candidate.structured_pages >= required_structured_pages
        and structure_boundary_gain
        >= ARBITRATION_POLICY["layout_recovery_minimum_boundary_gain"]
    )
    layout_structure_recovered = bool(
        significant_layout_recovery
        and char_ratio
        >= ARBITRATION_POLICY["layout_recovery_minimum_character_ratio"]
        and page_ratio >= ARBITRATION_POLICY["layout_recovery_minimum_page_ratio"]
        and not lost_r3_concepts
    )
    if (
        char_ratio >= ARBITRATION_POLICY["character_gain_ratio"]
        and candidate.character_count - base.character_count
        >= ARBITRATION_POLICY["character_gain_minimum"]
    ):
        improvements.append("character_gain")
    if candidate.page_count > base.page_count:
        improvements.append("page_gain")
    if len(candidate.risk_flags) < len(base.risk_flags):
        improvements.append("risk_reduction")
    if (
        candidate.glyph_noise < base.glyph_noise
        and (
            base.glyph_noise - candidate.glyph_noise
            >= ARBITRATION_POLICY["glyph_noise_absolute_change"]
            or candidate.glyph_noise
            <= base.glyph_noise * ARBITRATION_POLICY["glyph_noise_reduction_ratio"]
        )
    ):
        improvements.append("glyph_noise_reduction")
    if candidate.keyword_coverage > base.keyword_coverage:
        improvements.append("keyword_coverage_gain")
    if significant_layout_recovery:
        improvements.append("layout_structure_gain")

    if char_ratio < ARBITRATION_POLICY["minimum_character_ratio"]:
        regressions.append("character_loss")
    if base.page_count >= 2 and page_ratio < ARBITRATION_POLICY["minimum_page_ratio"]:
        regressions.append("page_loss")
    if (
        significant_layout_recovery
        and char_ratio
        < ARBITRATION_POLICY["layout_recovery_minimum_character_ratio"]
    ):
        regressions.append("layout_character_loss")
    if (
        significant_layout_recovery
        and page_ratio < ARBITRATION_POLICY["layout_recovery_minimum_page_ratio"]
    ):
        regressions.append("layout_page_loss")
    if len(candidate.risk_flags) > len(base.risk_flags):
        regressions.append("risk_increase")
    noise_limit = base.glyph_noise + max(
        ARBITRATION_POLICY["glyph_noise_absolute_change"],
        int(base.glyph_noise * ARBITRATION_POLICY["glyph_noise_increase_ratio"]),
    )
    if candidate.glyph_noise > noise_limit:
        regressions.append("glyph_noise_increase")
    if (
        candidate.keyword_coverage + ARBITRATION_POLICY["keyword_loss_tolerance"]
        < base.keyword_coverage
    ):
        regressions.append("keyword_coverage_loss")
    if lost_r3_concepts:
        regressions.extend(f"r3_concept_loss:{concept}" for concept in lost_r3_concepts)

    base_strong = set(base.risk_flags) & STRONG_RISK_FLAGS
    candidate_strong = set(candidate.risk_flags) & STRONG_RISK_FLAGS
    strong_risk_resolved = (
        bool(base_strong - candidate_strong)
        and len(candidate_strong) < len(base_strong)
        and char_ratio >= ARBITRATION_POLICY["strong_risk_minimum_character_ratio"]
        and page_ratio >= ARBITRATION_POLICY["strong_risk_minimum_page_ratio"]
        and candidate.keyword_coverage >= base.keyword_coverage
    )
    major_character_gain = (
        char_ratio >= ARBITRATION_POLICY["major_character_gain_ratio"]
        and candidate.character_count - base.character_count
        >= ARBITRATION_POLICY["major_character_gain_minimum"]
    )
    if regressions:
        reason = (
            "content_regression"
            if set(regressions)
            & {
                "character_loss",
                "page_loss",
                "layout_character_loss",
                "layout_page_loss",
            }
            else "r3_concept_loss"
            if lost_r3_concepts
            else "quality_regression"
        )
        replace = False
    elif layout_structure_recovered:
        reason = "layout_structure_recovered"
        replace = True
    elif strong_risk_resolved:
        reason = "strong_risk_resolved"
        replace = True
    elif major_character_gain:
        reason = "major_content_gain"
        replace = True
    elif len(improvements) >= ARBITRATION_POLICY["multiple_improvements_required"]:
        reason = "multi_metric_improvement"
        replace = True
    else:
        reason = "insufficient_improvement"
        replace = False
    return ReplacementDecision(
        replace,
        reason,
        tuple(improvements),
        tuple(regressions),
        base,
        candidate,
    )


def build_mineru_cache(
    identity_manifest: Path,
    pdf_dir: Path,
    output: Path,
    summary_json: Path,
    *,
    base_cache: Optional[Path] = None,
    selected_refnos: Sequence[str] = (),
    select_all: bool = False,
    use_quality_gate: bool = False,
    quality_gate_flags: Sequence[str] = (),
    mineru_cli: Optional[MineruCLI] = None,
    mineru_work_dir: Optional[Path] = None,
    page_fusion: bool = False,
    on_unavailable: str = "fail",
    on_error: str = "fail",
    on_missing_pdf: str = "skip",
) -> dict[str, Any]:
    identity_manifest = identity_manifest.expanduser().resolve()
    pdf_dir = pdf_dir.expanduser().resolve()
    output = output.expanduser().resolve(strict=False)
    summary_json = summary_json.expanduser().resolve(strict=False)
    base_cache = base_cache.expanduser().resolve() if base_cache is not None else None
    mineru_work_dir = (
        mineru_work_dir.expanduser().resolve(strict=False)
        if mineru_work_dir is not None
        else None
    )
    if output.suffix.casefold() != ".jsonl" or summary_json.suffix.casefold() != ".json":
        raise MineruAdapterError("Output must be JSONL and summary must be JSON")
    input_paths = {identity_manifest, pdf_dir}
    if base_cache is not None:
        input_paths.add(base_cache)
    if output in input_paths or summary_json in input_paths or output == summary_json:
        raise MineruAdapterError("Inputs, output cache, and summary paths must be distinct")
    if not pdf_dir.is_dir():
        raise MineruAdapterError(f"PDF directory does not exist: {pdf_dir}")
    if on_unavailable not in {"fail", "skip"}:
        raise MineruAdapterError("on_unavailable must be fail or skip")
    if on_error not in {"fail", "skip"}:
        raise MineruAdapterError("on_error must be fail or skip")
    if on_missing_pdf not in {"fail", "skip"}:
        raise MineruAdapterError("on_missing_pdf must be fail or skip")
    if page_fusion and base_cache is None:
        raise MineruAdapterError("page_fusion requires --base-cache patch mode")

    refnos = load_identity_refnos(identity_manifest)
    refno_set = set(refnos)
    explicit_refnos = set(selected_refnos)
    unknown_selected = explicit_refnos - refno_set
    if unknown_selected:
        raise MineruAdapterError(
            "Selected REFNO values are absent from identity manifest: "
            + ", ".join(sorted(unknown_selected))
        )
    selector_count = int(bool(explicit_refnos)) + int(select_all) + int(use_quality_gate)
    if selector_count > 1:
        raise MineruAdapterError(
            "--refno, --all-refnos, and --quality-gate are mutually exclusive"
        )
    if explicit_refnos and base_cache is None:
        raise MineruAdapterError("--refno is supported only with --base-cache patch mode")
    if use_quality_gate and base_cache is None:
        raise MineruAdapterError("--quality-gate requires --base-cache")
    if quality_gate_flags and (explicit_refnos or select_all):
        raise MineruAdapterError(
            "--quality-flag applies only to quality-gate selection"
        )
    if base_cache is None and not select_all:
        raise MineruAdapterError(
            "Generate mode will not run every PDF implicitly; pass --all-refnos explicitly"
        )

    base_records: dict[str, dict[str, Any]] = {}
    if base_cache is not None:
        base_records, _ = load_page_cache(base_cache, refno_set)
    selection_reasons: dict[str, int] = {}
    quality_candidate_reasons: dict[str, int] = {}
    quality_candidate_count = 0
    quality_candidates_without_local_pdf = 0
    if explicit_refnos:
        selection_mode = "explicit_refno"
        selected = explicit_refnos
    elif select_all:
        selection_mode = "all_refnos"
        selected = refno_set
    else:
        # Patch mode defaults to quality gating, so an omitted selector is safe.
        selection_mode = "quality_gate"
        configured_quality_flags = (
            tuple(quality_gate_flags)
            if quality_gate_flags
            else tuple(sorted(DEFAULT_QUALITY_GATE_FLAGS))
        )
        quality_candidates, quality_candidate_reasons = quality_gate_selection(
            refnos,
            base_records,
            configured_quality_flags,
        )
        if quality_candidates == refno_set:
            raise MineruAdapterError(
                "Quality gate selected every identity REFNO; automatic selection will not "
                "silently become a full run. Review the gate and rerun with --all-refnos "
                "only if full processing is intended."
            )
        quality_candidate_count = len(quality_candidates)
        selected = {
            refno
            for refno in quality_candidates
            if find_pdf(pdf_dir, refno) is not None
        }
        quality_candidates_without_local_pdf = len(quality_candidates - selected)
        _, selection_reasons = quality_gate_selection(
            [refno for refno in refnos if refno in selected],
            base_records,
            configured_quality_flags,
        )
    configured_quality_flags = (
        tuple(quality_gate_flags)
        if quality_gate_flags
        else tuple(sorted(DEFAULT_QUALITY_GATE_FLAGS))
    )
    cli = mineru_cli or MineruCLI()
    available = cli.available()
    if selected and not available and on_unavailable == "fail":
        raise MineruUnavailableError(
            f"MinerU executable is unavailable: {cli.executable!r}; "
            "install/configure MinerU or rerun with --on-unavailable skip"
        )

    protected_paths = [identity_manifest]
    if base_cache is not None:
        protected_paths.append(base_cache)
    protected_hashes = {str(path): file_sha256(path) for path in protected_paths}
    counters: Counter[str] = Counter()
    arbitration_reasons: Counter[str] = Counter()
    source_counts: Counter[str] = Counter()
    flag_counts: Counter[str] = Counter()
    commands: list[list[str]] = []
    table_blocks = 0
    records: list[dict[str, Any]] = []
    page_fusion_audits: list[dict[str, Any]] = []
    page_fusion_reasons: Counter[str] = Counter()

    temporary: Optional[tempfile.TemporaryDirectory[str]] = None
    if mineru_work_dir is None:
        temporary = tempfile.TemporaryDirectory(prefix="r3-mineru-")
        work_root = Path(temporary.name)
    else:
        mineru_work_dir.mkdir(parents=True, exist_ok=True)
        work_root = mineru_work_dir
    try:
        for refno in refnos:
            if refno not in selected:
                record = base_records[refno]
                counters["preserved_unselected"] += 1
                records.append(record)
                continue
            pdf_path = find_pdf(pdf_dir, refno)
            if pdf_path is None:
                if on_missing_pdf == "fail":
                    raise MineruAdapterError(f"No local PDF found for REFNO {refno}")
                if base_cache is not None:
                    record = base_records[refno]
                    counters["preserved_missing_pdf"] += 1
                else:
                    record = make_cache_record(
                        refno,
                        None,
                        None,
                        fallback_kind="mineru_missing_pdf",
                        fallback_flags=("no_local_pdf", "no_extractable_text"),
                    )
                    counters["generated_missing_pdf"] += 1
                records.append(record)
                continue
            if not available:
                if base_cache is not None:
                    record = base_records[refno]
                    counters["preserved_mineru_unavailable"] += 1
                else:
                    record = make_cache_record(
                        refno,
                        pdf_path,
                        None,
                        fallback_kind="mineru_unavailable",
                        fallback_flags=("mineru_unavailable", "no_extractable_text"),
                    )
                    counters["generated_mineru_unavailable"] += 1
                records.append(record)
                continue

            invocation_dir = Path(
                tempfile.mkdtemp(prefix=f"{refno}-", dir=str(work_root))
            )
            try:
                command = list(cli.run(pdf_path, invocation_dir))
                commands.append(command)
                parsed = parse_mineru_directory(invocation_dir, pdf_path.stem)
                if mineru_work_dir is None:
                    parsed.parsed_path = None
            except (MineruExecutionError, MineruOutputError) as exc:
                if on_error == "fail":
                    raise
                if base_cache is not None:
                    record = base_records[refno]
                    counters["preserved_mineru_error"] += 1
                else:
                    flag = (
                        "mineru_cli_failed"
                        if isinstance(exc, MineruExecutionError)
                        else "mineru_output_missing"
                    )
                    record = make_cache_record(
                        refno,
                        pdf_path,
                        None,
                        fallback_kind="mineru_failed",
                        fallback_flags=(flag, "no_extractable_text"),
                    )
                    counters["generated_mineru_error"] += 1
                records.append(record)
                continue
            if not parsed.pages and base_cache is not None:
                record = base_records[refno]
                counters["preserved_empty_mineru_output"] += 1
            else:
                candidate = make_cache_record(
                    refno,
                    pdf_path,
                    parsed,
                    fallback_kind="mineru",
                    fallback_flags=(),
                )
                if base_cache is None:
                    record = candidate
                    if set(parsed.quality_flags) & UNRELIABLE_PAGE_BOUNDARY_FLAGS:
                        counters["generated_mineru_unreliable_page_boundaries"] += 1
                    else:
                        counters["generated"] += 1
                        table_blocks += parsed.table_blocks
                elif page_fusion and base_records[refno].get("pages"):
                    fusion = fuse_cache_pages(
                        base_records[refno],
                        candidate,
                        parsed,
                        pdf_path,
                    )
                    for decision in fusion.decisions:
                        page_fusion_reasons[str(decision["reason"])] += 1
                    page_fusion_audits.append(
                        {
                            "REFNO": refno,
                            "changed": fusion.changed,
                            "candidate_pages_selected": fusion.candidate_pages_selected,
                            "base_pages_selected": fusion.base_pages_selected,
                            "candidate_only_pages": fusion.candidate_only_pages,
                            "base_only_pages": fusion.base_only_pages,
                            "page_decisions": list(fusion.decisions),
                        }
                    )
                    if fusion.changed:
                        record = fusion.record
                        counters["fused"] += 1
                        table_blocks += fusion.selected_candidate_table_blocks
                    else:
                        record = base_records[refno]
                        counters["preserved_page_fusion"] += 1
                else:
                    arbitration = compare_cache_quality(
                        base_records[refno],
                        candidate,
                    )
                    arbitration_reasons[arbitration.reason] += 1
                    if arbitration.replace:
                        record = candidate
                        counters["replaced"] += 1
                        table_blocks += parsed.table_blocks
                    else:
                        record = base_records[refno]
                        counters["preserved_quality_arbitration"] += 1
            records.append(record)
    finally:
        if temporary is not None:
            temporary.cleanup()

    if [record["REFNO"] for record in records] != refnos:
        raise MineruAdapterError("Output record order/set diverged from identity REFNO values")
    for record in records:
        source = record["source"]
        source_counts[str(source["source_kind"])] += 1
        for flag in source["source_quality_flags"] or ["none"]:
            flag_counts[str(flag)] += 1

    current_hashes = {str(path): file_sha256(path) for path in protected_paths}
    if current_hashes != protected_hashes:
        raise MineruAdapterError("An identity/base-cache input changed during MinerU parsing")
    _atomic_write_jsonl(output, records)
    load_page_cache(output, refno_set)
    skipped = sum(
        count
        for key, count in counters.items()
        if (
            key.startswith("preserved_")
            and key
            not in {
                "preserved_unselected",
                "preserved_quality_arbitration",
                "preserved_page_fusion",
            }
        )
        or key.startswith("generated_mineru_")
        or key == "generated_missing_pdf"
    )
    processed = counters["generated"] + counters["replaced"] + counters["fused"]
    status = "PASS" if not skipped else "PARTIAL"
    if selected and not available and processed == 0:
        status = "SKIPPED"
    summary: dict[str, Any] = {
        "schema_version": "r3_mineru_page_cache_adapter_v1",
        "status": status,
        "mode": "patch" if base_cache is not None else "generate",
        "inputs": {
            "identity_manifest": {
                "path": str(identity_manifest),
                "sha256": protected_hashes[str(identity_manifest)],
                "columns_read": list(IDENTITY_COLUMNS),
                "unique_refnos": len(refnos),
            },
            "pdf_dir": str(pdf_dir),
            "base_cache": (
                {"path": str(base_cache), "sha256": protected_hashes[str(base_cache)]}
                if base_cache is not None
                else None
            ),
        },
        "mineru": {
            "executable": cli.executable,
            "available": available,
            "invocations": len(commands),
            "extra_args": list(cli.extra_args),
            "work_dir": str(mineru_work_dir) if mineru_work_dir is not None else None,
        },
        "selection": {
            "mode": selection_mode,
            "selected_refnos": len(selected),
            "all_identity_refnos": len(refnos),
            "quality_gate_candidates": (
                quality_candidate_count if selection_mode == "quality_gate" else None
            ),
            "quality_candidates_without_local_pdf": (
                quality_candidates_without_local_pdf
                if selection_mode == "quality_gate"
                else None
            ),
            "quality_gate_flags": (
                list(configured_quality_flags)
                if selection_mode == "quality_gate"
                else []
            ),
            "quality_gate_reason_counts": selection_reasons,
            "quality_gate_candidate_reason_counts": quality_candidate_reasons,
            "defaulted_to_quality_gate": (
                selection_mode == "quality_gate" and not use_quality_gate
            ),
        },
        "actions": dict(sorted(counters.items())),
        "replacement_arbitration": {
            "policy": dict(ARBITRATION_POLICY),
            "metrics": [
                "nonempty_page_count",
                "normalized_character_count",
                "risk_flag_count",
                "glyph_noise_count",
                "domain_keyword_coverage",
                "r3_critical_concept_set",
                "paragraph_boundary_count",
                "structured_page_count",
                "table_row_count",
                "flattened_layout_detection",
            ],
            "compared_records": sum(arbitration_reasons.values()),
            "reason_counts": dict(sorted(arbitration_reasons.items())),
            "uses_labels": False,
            "uses_predictions": False,
        },
        "page_fusion": {
            "enabled": page_fusion,
            "policy": dict(PAGE_FUSION_POLICY),
            "compared_records": len(page_fusion_audits),
            "fused_records": counters["fused"],
            "preserved_records": counters["preserved_page_fusion"],
            "candidate_pages_selected": sum(
                audit["candidate_pages_selected"] for audit in page_fusion_audits
            ),
            "base_pages_selected": sum(
                audit["base_pages_selected"] for audit in page_fusion_audits
            ),
            "candidate_only_pages": sum(
                audit["candidate_only_pages"] for audit in page_fusion_audits
            ),
            "base_only_pages": sum(
                audit["base_only_pages"] for audit in page_fusion_audits
            ),
            "reason_counts": dict(sorted(page_fusion_reasons.items())),
            "record_audits": page_fusion_audits,
            "uses_labels": False,
            "uses_predictions": False,
            "cross_source_page_concatenation": False,
        },
        "output": {
            "path": str(output),
            "sha256": file_sha256(output),
            "records": len(records),
            "pages": sum(len(record["pages"]) for record in records),
            "table_blocks_recovered": table_blocks,
            "source_kind_counts": dict(sorted(source_counts.items())),
            "source_quality_flag_counts": dict(sorted(flag_counts.items())),
        },
        "runtime_boundary": {
            "labels_read": False,
            "pmagresult_read": False,
            "source_metadata_tables_read": False,
            "prior_evaluations_read": False,
            "prior_predictions_read": False,
            "identity_columns_only": True,
        },
        "policies": {
            "on_unavailable": on_unavailable,
            "on_error": on_error,
            "on_missing_pdf": on_missing_pdf,
            "patch_replaces_only_successful_quality_improvements": True,
        },
    }
    _atomic_write_json(summary_json, summary)
    return summary


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity-manifest", type=Path, required=True)
    parser.add_argument("--pdf-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary-json", type=Path, required=True)
    parser.add_argument("--base-cache", type=Path)
    parser.add_argument(
        "--refno",
        action="append",
        default=[],
        help="Patch only this REFNO; repeatable and valid only with --base-cache.",
    )
    parser.add_argument(
        "--all-refnos",
        action="store_true",
        help="Explicitly process every identity REFNO; required for generate mode.",
    )
    parser.add_argument(
        "--quality-gate",
        action="store_true",
        help=(
            "Select base-cache records with risky quality flags or empty page/text; "
            "this is the safe default in patch mode when no selector is supplied."
        ),
    )
    parser.add_argument(
        "--quality-flag",
        action="append",
        default=[],
        help="Override a quality-gate flag; repeat to provide the complete flag set.",
    )
    parser.add_argument("--mineru-executable", default="mineru")
    parser.add_argument(
        "--mineru-extra-arg",
        action="append",
        default=[],
        help="Additional MinerU CLI argument; repeat for multiple arguments.",
    )
    parser.add_argument("--mineru-work-dir", type=Path)
    parser.add_argument(
        "--page-fusion",
        action="store_true",
        help=(
            "Select the better complete source independently for each physical page; "
            "valid only with --base-cache."
        ),
    )
    parser.add_argument("--timeout-seconds", type=int, default=900)
    parser.add_argument(
        "--on-unavailable", choices=("fail", "skip"), default="fail"
    )
    parser.add_argument("--on-error", choices=("fail", "skip"), default="fail")
    parser.add_argument(
        "--on-missing-pdf", choices=("fail", "skip"), default="skip"
    )
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    if args.timeout_seconds <= 0:
        print("MinerU adapter error: --timeout-seconds must be positive", file=sys.stderr)
        return 2
    cli = MineruCLI(
        args.mineru_executable,
        extra_args=args.mineru_extra_arg,
        timeout_seconds=args.timeout_seconds,
    )
    try:
        summary = build_mineru_cache(
            args.identity_manifest,
            args.pdf_dir,
            args.output,
            args.summary_json,
            base_cache=args.base_cache,
            selected_refnos=args.refno,
            select_all=args.all_refnos,
            use_quality_gate=args.quality_gate,
            quality_gate_flags=args.quality_flag,
            mineru_cli=cli,
            mineru_work_dir=args.mineru_work_dir,
            page_fusion=args.page_fusion,
            on_unavailable=args.on_unavailable,
            on_error=args.on_error,
            on_missing_pdf=args.on_missing_pdf,
        )
    except (MineruAdapterError, OSError) as exc:
        print(f"MinerU adapter error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
