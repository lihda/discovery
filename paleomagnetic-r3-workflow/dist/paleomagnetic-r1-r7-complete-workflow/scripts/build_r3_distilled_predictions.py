#!/usr/bin/env python3
"""Build source-only R3 predictions with a frozen ROCKMAG distillation artifact.

Runtime scoring is deliberately limited to an exact identity manifest, an
integrity-checked original-paper page cache, and the frozen aggregate artifact.
The builder never opens PMAGRESULT, target labels, or previous predictions.
"""

from __future__ import annotations

import argparse
from collections import Counter, OrderedDict
from collections.abc import Callable, Mapping, Sequence
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook

try:
    from .paper_text_retrieval import PaperTextRetriever, canonical_target_identity
    from .r3_rockmag_distilled_reasoner import (
        R3RockmagArtifactValidationError,
        R3RockmagDistilledReasoner,
        R3RockmagRuntimeBoundaryError,
        validate_r3_rockmag_artifact,
    )
except ImportError:
    from paper_text_retrieval import PaperTextRetriever, canonical_target_identity
    from r3_rockmag_distilled_reasoner import (
        R3RockmagArtifactValidationError,
        R3RockmagDistilledReasoner,
        R3RockmagRuntimeBoundaryError,
        validate_r3_rockmag_artifact,
    )


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ARTIFACT = ROOT / "data/calibration/r3_rockmag_text_distillation_v1.json"
DEFAULT_RETRIEVAL_LIMIT = 30
DEFAULT_PER_CATEGORY_LIMIT = 12
TARGET_SCOPE_POLICIES = (
    "sentence_only",
    "exclusive_scope",
    "exclusive_scope_plus_discourse",
)
IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
PAGE_CACHE_ROW_KEYS = frozenset({"REFNO", "source", "pages"})
PAGE_CACHE_SOURCE_KEYS = frozenset(
    {
        "refno",
        "source_kind",
        "parsed_path",
        "pdf_path",
        "page_count",
        "text_length",
        "cache_sha256",
        "source_quality_flags",
    }
)
PAGE_CACHE_PAGE_KEYS = frozenset(
    {"page", "source", "text", "text_sha256", "section_hints"}
)
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
QUALITY_RISK_FLAGS = frozenset(
    {
        "short_text",
        "glyph_or_cid_text",
        "possibly_incomplete_extraction",
        "no_extractable_text",
        "poor_ocr",
        "low_ocr_quality",
        "missing_pages",
        "truncated_text",
    }
)
RUNTIME_SCORING_SOURCES = (
    "identity manifest: REFNO/RESULTNO/ROCKNAME",
    "original paper page text",
    "frozen historically supervised aggregate taxonomy",
)
FORBIDDEN_RUNTIME_SCORING_SOURCES = (
    "PMAGRESULT",
    "ROCKUNIT",
    "R3/R5/R7 labels",
    "evaluation results",
    "prior predictions",
    "record-level answer mappings",
)


class R3DistilledPredictionError(ValueError):
    """Raised when a runtime input or reasoner result violates the contract."""


class R3DistilledInputIntegrityError(R3DistilledPredictionError):
    """Raised when an input changes after its initial fingerprint is fixed."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_json_sha256(value: Any) -> str:
    try:
        encoded = json.dumps(
            value,
            ensure_ascii=True,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    except (TypeError, ValueError) as exc:
        raise R3DistilledPredictionError(
            f"JSON content is not canonicalizable: {exc}"
        ) from exc
    return hashlib.sha256(encoded).hexdigest()


def _clean_identity(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _reject_json_constant(value: str) -> None:
    raise R3DistilledPredictionError(f"non-finite JSON constant is forbidden: {value}")


def _unique_json_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    output: dict[str, Any] = {}
    for key, value in pairs:
        if key in output:
            raise R3DistilledPredictionError(f"duplicate JSON object key: {key}")
        output[key] = value
    return output


def _loads_strict_json(text: str, location: str) -> Any:
    try:
        return json.loads(
            text,
            object_pairs_hook=_unique_json_object,
            parse_constant=_reject_json_constant,
        )
    except R3DistilledPredictionError as exc:
        raise R3DistilledPredictionError(f"{location}: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise R3DistilledPredictionError(f"{location}: invalid JSON: {exc}") from exc


def _require_exact_keys(
    value: Mapping[str, Any], expected: frozenset[str], location: str
) -> None:
    actual = set(value)
    if actual == expected:
        return
    details: list[str] = []
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    if missing:
        details.append("missing=" + ",".join(missing))
    if extra:
        details.append("extra=" + ",".join(extra))
    raise R3DistilledPredictionError(
        f"{location} must use the strict page-cache schema ({'; '.join(details)})"
    )


def load_identity_manifest(path: Path) -> list[dict[str, str]]:
    """Load only the three allowed identity columns from a strict XLSX."""

    if path.suffix.casefold() != ".xlsx":
        raise R3DistilledPredictionError("identity manifest must be an XLSX file")
    if not path.is_file():
        raise R3DistilledPredictionError(f"identity manifest does not exist: {path}")
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:  # openpyxl uses several format-specific exception types
        raise R3DistilledPredictionError(
            f"could not open identity manifest {path}: {exc}"
        ) from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise R3DistilledPredictionError(
                "strict identity manifest must contain exactly one worksheet"
            )
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=False)
        header_cells = next(rows, ())
        if any(cell.data_type == "f" for cell in header_cells):
            raise R3DistilledPredictionError("identity manifest header cannot contain formulas")
        header = tuple(_clean_identity(cell.value) for cell in header_cells)
        if header != IDENTITY_COLUMNS:
            raise R3DistilledPredictionError(
                "strict identity columns must be exactly REFNO, RESULTNO, ROCKNAME"
            )

        identities: list[dict[str, str]] = []
        resultnos: set[str] = set()
        identity_tuples: set[tuple[str, str, str]] = set()
        for excel_row, cells in enumerate(rows, start=2):
            if len(cells) != len(IDENTITY_COLUMNS):
                raise R3DistilledPredictionError(
                    f"identity manifest row {excel_row} has an unexpected column count"
                )
            if any(cell.data_type == "f" for cell in cells):
                raise R3DistilledPredictionError(
                    f"identity manifest row {excel_row} contains a formula"
                )
            values = tuple(_clean_identity(cell.value) for cell in cells)
            if any(not value for value in values):
                raise R3DistilledPredictionError(
                    f"identity manifest row {excel_row} has an empty or invalid identity value"
                )
            refno, resultno, rockname = values
            if resultno in resultnos:
                raise R3DistilledPredictionError(
                    f"identity manifest contains duplicate RESULTNO: {resultno}"
                )
            if values in identity_tuples:
                raise R3DistilledPredictionError(
                    "identity manifest contains a duplicate REFNO/RESULTNO/ROCKNAME tuple"
                )
            resultnos.add(resultno)
            identity_tuples.add(values)
            identities.append(
                {"REFNO": refno, "RESULTNO": resultno, "ROCKNAME": rockname}
            )
    finally:
        workbook.close()
    if not identities:
        raise R3DistilledPredictionError("identity manifest contains no data rows")
    return identities


def _text_fingerprint(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text).strip()
    return hashlib.sha256(normalized.encode("utf-8", errors="ignore")).hexdigest()


def _validate_sha256(value: Any, location: str) -> str:
    if not isinstance(value, str) or not SHA256_RE.fullmatch(value):
        raise R3DistilledPredictionError(f"{location} must be a lowercase SHA-256")
    return value


def _validate_cache_row(
    row: Any,
    *,
    line_number: int,
    raw_record_sha256: str,
) -> tuple[str, dict[str, Any], str]:
    location = f"page cache line {line_number}"
    if not isinstance(row, Mapping):
        raise R3DistilledPredictionError(f"{location} must be an object")
    _require_exact_keys(row, PAGE_CACHE_ROW_KEYS, location)
    refno = row.get("REFNO")
    if not isinstance(refno, str) or not refno.strip() or refno != refno.strip():
        raise R3DistilledPredictionError(f"{location}.REFNO must be a trimmed string")

    source = row.get("source")
    if not isinstance(source, Mapping):
        raise R3DistilledPredictionError(f"{location}.source must be an object")
    _require_exact_keys(source, PAGE_CACHE_SOURCE_KEYS, f"{location}.source")
    if source.get("refno") != refno:
        raise R3DistilledPredictionError(
            f"{location}.source.refno must equal the row REFNO"
        )
    source_kind = source.get("source_kind")
    if not isinstance(source_kind, str) or not source_kind.strip():
        raise R3DistilledPredictionError(f"{location}.source.source_kind is invalid")
    for path_key in ("parsed_path", "pdf_path"):
        path_value = source.get(path_key)
        if path_value is not None and not isinstance(path_value, str):
            raise R3DistilledPredictionError(
                f"{location}.source.{path_key} must be a string or null"
            )
    flags = source.get("source_quality_flags")
    if (
        not isinstance(flags, list)
        or any(not isinstance(flag, str) or not flag.strip() for flag in flags)
        or len(flags) != len(set(flags))
    ):
        raise R3DistilledPredictionError(
            f"{location}.source.source_quality_flags must contain unique strings"
        )

    pages = row.get("pages")
    if not isinstance(pages, list):
        raise R3DistilledPredictionError(f"{location}.pages must be a list")
    page_hashes: list[str] = []
    page_numbers: set[int] = set()
    texts: list[str] = []
    for index, page in enumerate(pages):
        page_location = f"{location}.pages[{index}]"
        if not isinstance(page, Mapping):
            raise R3DistilledPredictionError(f"{page_location} must be an object")
        _require_exact_keys(page, PAGE_CACHE_PAGE_KEYS, page_location)
        page_number = page.get("page")
        if (
            isinstance(page_number, bool)
            or not isinstance(page_number, int)
            or page_number < 1
            or page_number in page_numbers
        ):
            raise R3DistilledPredictionError(
                f"{page_location}.page must be a unique positive integer"
            )
        page_numbers.add(page_number)
        page_source = page.get("source")
        text = page.get("text")
        hints = page.get("section_hints")
        if not isinstance(page_source, str) or not page_source.strip():
            raise R3DistilledPredictionError(f"{page_location}.source is invalid")
        if not isinstance(text, str):
            raise R3DistilledPredictionError(f"{page_location}.text must be a string")
        if not isinstance(hints, list) or any(not isinstance(hint, str) for hint in hints):
            raise R3DistilledPredictionError(
                f"{page_location}.section_hints must be a list of strings"
            )
        declared_page_hash = _validate_sha256(
            page.get("text_sha256"), f"{page_location}.text_sha256"
        )
        actual_page_hash = _text_fingerprint(text)
        if declared_page_hash != actual_page_hash:
            raise R3DistilledPredictionError(
                f"{page_location}.text_sha256 does not match the cached text"
            )
        page_hashes.append(actual_page_hash)
        texts.append(text)

    page_count = source.get("page_count")
    text_length = source.get("text_length")
    if isinstance(page_count, bool) or not isinstance(page_count, int) or page_count != len(pages):
        raise R3DistilledPredictionError(
            f"{location}.source.page_count does not match pages"
        )
    expected_text_length = len("\n".join(texts))
    if (
        isinstance(text_length, bool)
        or not isinstance(text_length, int)
        or text_length != expected_text_length
    ):
        raise R3DistilledPredictionError(
            f"{location}.source.text_length does not match pages"
        )
    declared_cache_hash = _validate_sha256(
        source.get("cache_sha256"), f"{location}.source.cache_sha256"
    )
    actual_cache_hash = hashlib.sha256("".join(page_hashes).encode("ascii")).hexdigest()
    if declared_cache_hash != actual_cache_hash:
        raise R3DistilledPredictionError(
            f"{location}.source.cache_sha256 does not match page hashes"
        )
    return refno, dict(row), raw_record_sha256


def load_page_cache(
    path: Path,
    required_refnos: set[str],
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    if path.suffix.casefold() != ".jsonl":
        raise R3DistilledPredictionError("page cache must be a JSONL file")
    if not path.is_file():
        raise R3DistilledPredictionError(f"page cache does not exist: {path}")
    cache: dict[str, dict[str, Any]] = {}
    record_hashes: dict[str, str] = {}
    with path.open("rb") as handle:
        for line_number, raw_line in enumerate(handle, start=1):
            if not raw_line.strip():
                continue
            try:
                text = raw_line.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise R3DistilledPredictionError(
                    f"page cache line {line_number} is not UTF-8: {exc}"
                ) from exc
            row = _loads_strict_json(text, f"page cache line {line_number}")
            record_hash = hashlib.sha256(raw_line.rstrip(b"\r\n")).hexdigest()
            refno, validated, record_hash = _validate_cache_row(
                row,
                line_number=line_number,
                raw_record_sha256=record_hash,
            )
            if refno in cache:
                raise R3DistilledPredictionError(f"duplicate page-cache REFNO: {refno}")
            cache[refno] = validated
            record_hashes[refno] = record_hash

    found = set(cache)
    missing = sorted(required_refnos - found)
    extra = sorted(found - required_refnos)
    if missing or extra:
        details: list[str] = []
        if missing:
            details.append("missing REFNO=" + ",".join(missing[:20]))
        if extra:
            details.append("extra REFNO=" + ",".join(extra[:20]))
        raise R3DistilledPredictionError(
            "page-cache REFNO set must exactly equal the identity manifest: "
            + "; ".join(details)
        )
    return cache, record_hashes


def load_artifact(path: Path) -> tuple[dict[str, Any], str, str]:
    if path.suffix.casefold() != ".json":
        raise R3DistilledPredictionError("distillation artifact must be a JSON file")
    if not path.is_file():
        raise R3DistilledPredictionError(f"distillation artifact does not exist: {path}")
    file_hash = file_sha256(path)
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise R3DistilledPredictionError(f"artifact is not UTF-8: {exc}") from exc
    raw = _loads_strict_json(text, str(path))
    if not isinstance(raw, Mapping):
        raise R3DistilledPredictionError("distillation artifact must be a JSON object")
    validated = validate_r3_rockmag_artifact(raw)
    return validated, file_hash, canonical_json_sha256(validated)


def _validate_track(track: Any, name: str) -> dict[str, Any]:
    if not isinstance(track, Mapping):
        raise R3DistilledPredictionError(f"reasoner {name} must be an object")
    required = {
        "policy_decision",
        "value",
        "binary_value",
        "outcome",
        "decision_path",
        "evidence",
    }
    missing = sorted(required - set(track))
    if missing:
        raise R3DistilledPredictionError(
            f"reasoner {name} is missing fields: {', '.join(missing)}"
        )
    decision = track["policy_decision"]
    if isinstance(decision, bool) or decision not in (0, 1, "ABSTAIN"):
        raise R3DistilledPredictionError(
            f"reasoner {name}.policy_decision must be 0, 1, or ABSTAIN"
        )
    expected_value = decision if decision in (0, 1) else None
    if track["value"] != expected_value or isinstance(track["value"], bool):
        raise R3DistilledPredictionError(
            f"reasoner {name}.value must mirror policy_decision or be null on abstention"
        )
    binary = track["binary_value"]
    expected_binary = decision if decision in (0, 1) else 0
    if isinstance(binary, bool) or binary != expected_binary:
        raise R3DistilledPredictionError(
            f"reasoner {name}.binary_value is inconsistent with policy_decision"
        )
    expected_outcome = (
        "POSITIVE" if decision == 1 else "NEGATIVE" if decision == 0 else "UNKNOWN"
    )
    if track["outcome"] != expected_outcome:
        raise R3DistilledPredictionError(
            f"reasoner {name}.outcome is inconsistent with policy_decision"
        )
    if not isinstance(track["decision_path"], str) or not track["decision_path"].strip():
        raise R3DistilledPredictionError(f"reasoner {name}.decision_path is invalid")
    if not isinstance(track["evidence"], Mapping):
        raise R3DistilledPredictionError(f"reasoner {name}.evidence must be an object")
    return dict(track)


def validate_reasoner_decision(
    decision: Any,
    *,
    expected_artifact_content_sha256: str,
) -> dict[str, Any]:
    if not isinstance(decision, Mapping):
        raise R3DistilledPredictionError("reasoner decision must be an object")
    normative = _validate_track(decision.get("normative_r3"), "normative_r3")
    proxy = _validate_track(decision.get("rockmag_proxy_r3"), "rockmag_proxy_r3")
    conflict = decision.get("criterion_practice_conflict")
    normative_decision = normative["policy_decision"]
    proxy_decision = proxy["policy_decision"]
    if proxy_decision == "ABSTAIN":
        expected_conflict = None
        expected_status = "UNKNOWN"
    elif normative_decision == "ABSTAIN" and proxy_decision == 1:
        expected_conflict = True
        expected_status = "INCOMPLETE_NORMATIVE_CHAIN"
    elif normative_decision == "ABSTAIN" and proxy_decision == 0:
        expected_conflict = False
        expected_status = "NO_PROXY_SUPPORT"
    elif normative_decision == proxy_decision:
        expected_conflict = False
        expected_status = "AGREE"
    else:
        expected_conflict = True
        expected_status = "TRACK_INCONSISTENCY"
    if conflict is not expected_conflict:
        raise R3DistilledPredictionError(
            "reasoner criterion_practice_conflict is inconsistent with the two tracks"
        )
    comparison = decision.get("criterion_practice_comparison")
    if not isinstance(comparison, Mapping):
        raise R3DistilledPredictionError(
            "reasoner criterion_practice_comparison must be an object"
        )
    if (
        comparison.get("status") != expected_status
        or comparison.get("normative_policy_decision") != normative_decision
        or comparison.get("rockmag_proxy_policy_decision") != proxy_decision
    ):
        raise R3DistilledPredictionError(
            "reasoner criterion_practice_comparison is inconsistent with the two tracks"
        )
    if decision.get("policy_decision") != proxy["policy_decision"]:
        raise R3DistilledPredictionError("top-level policy_decision must mirror the proxy track")
    if decision.get("binary_compatibility_value") != proxy["binary_value"]:
        raise R3DistilledPredictionError(
            "top-level binary_compatibility_value must mirror the proxy track"
        )
    artifact_hash = decision.get("artifact_sha256")
    if artifact_hash != expected_artifact_content_sha256:
        raise R3DistilledPredictionError(
            "reasoner artifact_sha256 does not match the loaded frozen artifact"
        )
    return dict(decision)


def _policy_counts(rows: Sequence[Mapping[str, Any]], track: str) -> dict[str, int]:
    counts = Counter(str(row[track]["policy_decision"]) for row in rows)
    return dict(sorted(counts.items()))


def build_prediction_rows(
    identities: Sequence[Mapping[str, str]],
    cache: Mapping[str, dict[str, Any]],
    cache_record_hashes: Mapping[str, str],
    reasoner: Any,
    *,
    artifact_file_sha256: str,
    artifact_content_sha256: str,
    identity_manifest_sha256: str,
    page_cache_sha256: str,
    retrieval_limit: int = DEFAULT_RETRIEVAL_LIMIT,
    per_category_limit: int = DEFAULT_PER_CATEGORY_LIMIT,
    target_scope_policy: str = "exclusive_scope",
) -> list[dict[str, Any]]:
    if (
        isinstance(retrieval_limit, bool)
        or not isinstance(retrieval_limit, int)
        or isinstance(per_category_limit, bool)
        or not isinstance(per_category_limit, int)
        or retrieval_limit < 1
        or per_category_limit < 1
    ):
        raise R3DistilledPredictionError("retrieval limits must be positive integers")
    if per_category_limit > retrieval_limit:
        raise R3DistilledPredictionError(
            "per-category retrieval limit cannot exceed the total retrieval limit"
        )
    if target_scope_policy not in TARGET_SCOPE_POLICIES:
        raise R3DistilledPredictionError(
            "target_scope_policy must be sentence_only, exclusive_scope, or "
            "exclusive_scope_plus_discourse"
        )
    canonical_targets_by_refno: dict[str, set[str]] = {}
    target_names_by_refno: dict[str, dict[str, str]] = {}
    result_count_by_target: Counter[tuple[str, str]] = Counter()
    for identity in identities:
        refno = identity["REFNO"]
        canonical_target = canonical_target_identity(identity["ROCKNAME"])
        canonical_targets_by_refno.setdefault(refno, set()).add(canonical_target)
        target_names_by_refno.setdefault(refno, {}).setdefault(
            canonical_target,
            identity["ROCKNAME"],
        )
        result_count_by_target[(refno, canonical_target)] += 1
    retrievers: OrderedDict[str, PaperTextRetriever] = OrderedDict()
    analysis_cache: dict[tuple[str, str, bool], dict[str, Any]] = {}
    rows: list[dict[str, Any]] = []

    for identity in identities:
        refno = identity["REFNO"]
        resultno = identity["RESULTNO"]
        rockname = identity["ROCKNAME"]
        canonical_target = canonical_target_identity(rockname)
        multi_target = len(canonical_targets_by_refno[refno]) > 1
        identity_ambiguous = result_count_by_target[(refno, canonical_target)] > 1
        if refno not in retrievers:
            retrievers[refno] = PaperTextRetriever(cache[refno])
            if len(retrievers) > 8:
                retrievers.popitem(last=False)
        else:
            retrievers.move_to_end(refno)

        analysis_key = (refno, canonical_target, multi_target)
        decision = analysis_cache.get(analysis_key)
        if decision is None:
            retrieval = retrievers[refno].retrieve(
                rockname,
                multi_target,
                "R3",
                competing_targets=[
                    name
                    for sibling_target, name in target_names_by_refno[refno].items()
                    if sibling_target != canonical_target
                ],
                allow_structured_scope=target_scope_policy != "sentence_only",
                allow_discourse_scope=(
                    target_scope_policy == "exclusive_scope_plus_discourse"
                ),
                limit=retrieval_limit,
                per_category_limit=per_category_limit,
            )
            source = cache[refno]["source"]
            raw_decision = reasoner.evaluate(
                retrieval,
                source_flags=list(source["source_quality_flags"]),
                multi_target=multi_target,
                identity_ambiguous=identity_ambiguous,
            )
            decision = validate_reasoner_decision(
                raw_decision,
                expected_artifact_content_sha256=artifact_content_sha256,
            )
            analysis_cache[analysis_key] = decision

        normative = dict(decision["normative_r3"])
        proxy = dict(decision["rockmag_proxy_r3"])
        reasoning_details = {
            key: value
            for key, value in decision.items()
            if key
            not in {
                "normative_r3",
                "rockmag_proxy_r3",
                "criterion_practice_conflict",
                "criterion_practice_comparison",
                "policy_decision",
                "binary_compatibility_value",
            }
        }
        rejected_counts = reasoning_details.get("rejected_evidence_counts")
        if not isinstance(rejected_counts, Mapping):
            rejected_counts = {}
        reason_code = str(
            reasoning_details.get("reason_code") or proxy["decision_path"]
        ).strip()
        row = {
            "schema_version": "r3_distilled_prediction_v1",
            "REFNO": refno,
            "RESULTNO": resultno,
            "ROCKNAME": rockname,
            "TARGET": rockname,
            "R3": proxy["binary_value"],
            "OUTCOME": proxy["outcome"],
            "policy_decision": proxy["policy_decision"],
            "binary_compatibility_value": proxy["binary_value"],
            "normative_r3": normative,
            "rockmag_proxy_r3": proxy,
            "criterion_practice_conflict": decision["criterion_practice_conflict"],
            "criterion_practice_comparison": dict(
                decision["criterion_practice_comparison"]
            ),
            "evidence_state": {
                "normative": normative["evidence"],
                "rockmag_proxy": proxy["evidence"],
            },
            "rejection_state": {
                "abstained": proxy["policy_decision"] == "ABSTAIN",
                "reason_code": reason_code,
                "manual_review": bool(
                    reasoning_details.get("manual_review")
                    or proxy["policy_decision"] == "ABSTAIN"
                ),
                "source_quality_risk": bool(
                    set(cache[refno]["source"]["source_quality_flags"])
                    & QUALITY_RISK_FLAGS
                ),
                "rejected_evidence_counts": dict(rejected_counts),
            },
            "reasoning_details": reasoning_details,
            "input_hashes": {
                "identity_manifest_sha256": identity_manifest_sha256,
                "page_cache_sha256": page_cache_sha256,
                "page_cache_record_sha256": cache_record_hashes[refno],
                "artifact_file_sha256": artifact_file_sha256,
                "artifact_content_sha256": artifact_content_sha256,
            },
            "runtime_scoring_sources": list(RUNTIME_SCORING_SOURCES),
            "forbidden_runtime_scoring_sources": list(
                FORBIDDEN_RUNTIME_SCORING_SOURCES
            ),
            "runtime_boundary": {
                "identity_columns_read": list(IDENTITY_COLUMNS),
                "labels_read": False,
                "source_metadata_tables_read": False,
                "evaluation_results_read": False,
                "prior_predictions_read": False,
                "offline_artifact_label_supervision": True,
            },
        }
        rows.append(row)
    return rows


def _atomic_write_jsonl(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(
                json.dumps(
                    row,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                    allow_nan=False,
                )
                + "\n"
            )
    temporary.replace(path)


def _atomic_write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
        + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _resolved_output_path(path: Path) -> Path:
    return path.expanduser().resolve(strict=False)


def _check_path_boundaries(
    identity_manifest: Path,
    page_cache: Path,
    artifact: Path,
    output: Path,
    summary_json: Path | None,
) -> None:
    inputs = {identity_manifest, page_cache, artifact}
    if len(inputs) != 3:
        raise R3DistilledPredictionError("runtime input files must be distinct")
    if output in inputs:
        raise R3DistilledPredictionError("prediction output cannot overwrite a runtime input")
    if summary_json is not None:
        if summary_json in inputs or summary_json == output:
            raise R3DistilledPredictionError(
                "summary output must be distinct from all inputs and predictions"
            )
    if output.suffix.casefold() != ".jsonl":
        raise R3DistilledPredictionError("prediction output must use a .jsonl suffix")
    if summary_json is not None and summary_json.suffix.casefold() != ".json":
        raise R3DistilledPredictionError("summary output must use a .json suffix")


def build_predictions(
    identity_manifest: Path,
    page_cache: Path,
    artifact: Path,
    output: Path,
    *,
    summary_json: Path | None = None,
    expected_artifact_sha256: str | None = None,
    retrieval_limit: int = DEFAULT_RETRIEVAL_LIMIT,
    per_category_limit: int = DEFAULT_PER_CATEGORY_LIMIT,
    target_scope_policy: str = "exclusive_scope",
    reasoner_factory: Callable[[Mapping[str, Any]], Any] = R3RockmagDistilledReasoner,
) -> dict[str, Any]:
    identity_manifest = identity_manifest.expanduser().resolve()
    page_cache = page_cache.expanduser().resolve()
    artifact = artifact.expanduser().resolve()
    output = _resolved_output_path(output)
    summary_json = _resolved_output_path(summary_json) if summary_json else None
    _check_path_boundaries(identity_manifest, page_cache, artifact, output, summary_json)

    initial_hashes = {
        "identity_manifest": file_sha256(identity_manifest),
        "page_cache": file_sha256(page_cache),
        "artifact": file_sha256(artifact),
    }
    validated_artifact, artifact_file_hash, artifact_content_hash = load_artifact(artifact)
    if artifact_file_hash != initial_hashes["artifact"]:
        raise R3DistilledInputIntegrityError(
            "artifact changed while it was being loaded"
        )
    if expected_artifact_sha256 is not None:
        normalized_expected = expected_artifact_sha256.strip().casefold()
        if not SHA256_RE.fullmatch(normalized_expected):
            raise R3DistilledPredictionError(
                "expected artifact SHA-256 must contain 64 lowercase hexadecimal characters"
            )
        if artifact_file_hash != normalized_expected:
            raise R3DistilledPredictionError(
                "frozen artifact SHA-256 does not match --expected-artifact-sha256"
            )

    identities = load_identity_manifest(identity_manifest)
    required_refnos = {row["REFNO"] for row in identities}
    cache, cache_record_hashes = load_page_cache(page_cache, required_refnos)
    reasoner = reasoner_factory(validated_artifact)
    if getattr(reasoner, "artifact_sha256", artifact_content_hash) != artifact_content_hash:
        raise R3DistilledPredictionError(
            "reasoner initialized with a different artifact content hash"
        )
    rows = build_prediction_rows(
        identities,
        cache,
        cache_record_hashes,
        reasoner,
        artifact_file_sha256=artifact_file_hash,
        artifact_content_sha256=artifact_content_hash,
        identity_manifest_sha256=initial_hashes["identity_manifest"],
        page_cache_sha256=initial_hashes["page_cache"],
        retrieval_limit=retrieval_limit,
        per_category_limit=per_category_limit,
        target_scope_policy=target_scope_policy,
    )

    final_hashes = {
        "identity_manifest": file_sha256(identity_manifest),
        "page_cache": file_sha256(page_cache),
        "artifact": file_sha256(artifact),
    }
    changed = [name for name in initial_hashes if initial_hashes[name] != final_hashes[name]]
    if changed:
        raise R3DistilledInputIntegrityError(
            "runtime input changed during scoring: " + ", ".join(changed)
        )

    _atomic_write_jsonl(output, rows)
    output_hash = file_sha256(output)
    score_counts = Counter(int(row["R3"]) for row in rows)
    outcome_counts = Counter(str(row["OUTCOME"]) for row in rows)
    comparison_counts = Counter(
        str(row["criterion_practice_comparison"]["status"]) for row in rows
    )
    summary: dict[str, Any] = {
        "schema_version": "r3_distilled_prediction_summary_v1",
        "status": "PASS",
        "workflow_profile": "rockmag_text_distilled_v1",
        "inputs": {
            "identity_manifest": {
                "path": str(identity_manifest),
                "sha256": initial_hashes["identity_manifest"],
                "rows": len(identities),
                "unique_refnos": len(required_refnos),
                "columns_read": list(IDENTITY_COLUMNS),
            },
            "page_cache": {
                "path": str(page_cache),
                "sha256": initial_hashes["page_cache"],
                "unique_refnos": len(cache),
            },
            "artifact": {
                "path": str(artifact),
                "file_sha256": artifact_file_hash,
                "content_sha256": artifact_content_hash,
                "artifact_class": validated_artifact.get("artifact_class"),
                "schema_version": validated_artifact.get("schema_version"),
                "label_supervision": True,
            },
        },
        "output": {
            "path": str(output),
            "sha256": output_hash,
            "rows": len(rows),
        },
        "score_counts": {str(key): value for key, value in sorted(score_counts.items())},
        "outcome_counts": dict(sorted(outcome_counts.items())),
        "normative_policy_counts": _policy_counts(rows, "normative_r3"),
        "rockmag_proxy_policy_counts": _policy_counts(rows, "rockmag_proxy_r3"),
        "criterion_practice_comparison_counts": dict(sorted(comparison_counts.items())),
        "criterion_practice_conflict_rows": sum(
            row["criterion_practice_conflict"] is True for row in rows
        ),
        "criterion_practice_unknown_rows": sum(
            row["criterion_practice_conflict"] is None for row in rows
        ),
        "retrieval_policy": {
            "total_limit": retrieval_limit,
            "per_category_limit": per_category_limit,
            "target_scope_policy": target_scope_policy,
        },
        "runtime_scoring_sources": list(RUNTIME_SCORING_SOURCES),
        "forbidden_runtime_scoring_sources": list(FORBIDDEN_RUNTIME_SCORING_SOURCES),
        "runtime_boundary": {
            "labels_read": False,
            "source_metadata_tables_read": False,
            "evaluation_results_read": False,
            "prior_predictions_read": False,
            "identity_cache_sets_equal": True,
            "page_and_cache_hashes_verified": True,
            "input_files_unchanged_during_scoring": True,
        },
    }
    if summary_json is not None:
        _atomic_write_json(summary_json, summary)
    return summary


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--identity-manifest",
        "--targets",
        dest="identity_manifest",
        type=Path,
        required=True,
        help="Strict XLSX containing exactly REFNO, RESULTNO, and ROCKNAME.",
    )
    parser.add_argument("--page-cache", type=Path, required=True)
    parser.add_argument("--artifact", type=Path, default=DEFAULT_ARTIFACT)
    parser.add_argument(
        "--expected-artifact-sha256",
        default=None,
        help="Optional raw-file SHA-256 pin for a preregistered run.",
    )
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary-json", type=Path, default=None)
    parser.add_argument(
        "--retrieval-limit",
        type=int,
        default=DEFAULT_RETRIEVAL_LIMIT,
        help="Maximum retained R3 retrieval hits per target.",
    )
    parser.add_argument(
        "--per-category-limit",
        type=int,
        default=DEFAULT_PER_CATEGORY_LIMIT,
        help="Maximum retained hits in each R3 evidence category.",
    )
    parser.add_argument(
        "--target-scope-policy",
        choices=TARGET_SCOPE_POLICIES,
        default="exclusive_scope",
        help=(
            "sentence_only keeps strict same-sentence binding; exclusive_scope also accepts "
            "bounded paragraphs/table rows; exclusive_scope_plus_discourse additionally "
            "accepts exact-target finite forward sentence windows."
        ),
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        summary = build_predictions(
            args.identity_manifest,
            args.page_cache,
            args.artifact,
            args.output,
            summary_json=args.summary_json,
            expected_artifact_sha256=args.expected_artifact_sha256,
            retrieval_limit=args.retrieval_limit,
            per_category_limit=args.per_category_limit,
            target_scope_policy=args.target_scope_policy,
        )
    except (
        OSError,
        R3DistilledPredictionError,
        R3RockmagArtifactValidationError,
        R3RockmagRuntimeBoundaryError,
    ) as exc:
        print(f"R3 distilled prediction error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
