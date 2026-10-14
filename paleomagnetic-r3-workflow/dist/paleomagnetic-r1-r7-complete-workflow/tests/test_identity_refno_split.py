import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.filter_identity_manifest_by_refno_split import (
    DEFAULT_HOLDOUT_FRACTION,
    DEFAULT_SALT,
    IdentitySplitError,
    holdout_bucket,
    split_identity_manifest,
)


class IdentityRefnoSplitTests(unittest.TestCase):
    def _manifest(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["REFNO", "RESULTNO", "ROCKNAME"])
        sheet.append(["6", "1", "Alpha"])
        sheet.append(["6", "2", "Beta"])
        sheet.append(["1142", "3", "Gamma"])
        sheet.append(["200", "4", "Delta"])
        workbook.save(path)

    def test_same_refno_is_never_split_between_partitions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "identity.xlsx"
            training = root / "training.xlsx"
            holdout = root / "holdout.xlsx"
            self._manifest(source)
            split_identity_manifest(source, training, partition="training")
            split_identity_manifest(source, holdout, partition="holdout")

            def rows(path: Path):
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    return list(workbook.active.iter_rows(min_row=2, values_only=True))
                finally:
                    workbook.close()

            training_rows = rows(training)
            holdout_rows = rows(holdout)
            self.assertFalse(
                {str(row[0]) for row in training_rows}
                & {str(row[0]) for row in holdout_rows}
            )
            self.assertEqual(len(training_rows) + len(holdout_rows), 4)
            self.assertEqual(
                sum(str(row[0]) == "6" for row in training_rows + holdout_rows),
                2,
            )

    def test_partition_matches_documented_sha256_rule(self):
        expected = holdout_bucket("1142", DEFAULT_SALT)
        self.assertAlmostEqual(expected, 0.3008618403058447)
        self.assertFalse(expected < DEFAULT_HOLDOUT_FRACTION)

    def test_rejects_non_identity_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "bad.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["REFNO", "RESULTNO", "ROCKNAME", "R3"])
            sheet.append(["100", "1", "Alpha", 1])
            workbook.save(source)
            with self.assertRaises(IdentitySplitError):
                split_identity_manifest(source, root / "output.xlsx")


if __name__ == "__main__":
    unittest.main()
