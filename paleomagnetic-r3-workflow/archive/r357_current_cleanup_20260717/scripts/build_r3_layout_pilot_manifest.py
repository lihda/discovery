#!/usr/bin/env python3
"""Freeze a label-blind, PDF-backed manifest for the R3 layout pilot.

The selector accepts exactly REFNO, RESULTNO, and ROCKNAME from a one-sheet
identity workbook.  A paper is eligible when it has at least eight identity
rows and exactly one local PDF whose filename stem equals REFNO.  No labels,
source-table fields, prior predictions, or evaluation artifacts are inputs.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import tempfile
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_IDENTITY = (
    ROOT / "outputs/r357_strict_eval_07111758/inputs/identity_manifest.xlsx"
)
DEFAULT_PDF_DIR = ROOT / "data/raw_pdfs"
DEFAULT_OUTPUT = ROOT / "tmp/r3_layout_pilot_v1/r3_layout_pilot_manifest_v1.json"

IDENTITY_FIELDS = ("REFNO", "RESULTNO", "ROCKNAME")
MINIMUM_ROWS_PER_REFNO = 8
SCHEMA_VERSION = "r3_layout_pilot_manifest_v1"
SELECTION_RULE_ID = "identity_row_count_gte_8_and_unique_local_pdf_v1"


class LayoutPilotManifestError(ValueError):
    """Raised when a label-blind pilot manifest cannot be frozen safely."""


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_sha256(value: Any) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def load_identity_rows(path: Path) -> list[dict[str, str]]:
    """Load the complete identity while rejecting every non-identity column."""

    path = path.expanduser().resolve()
    if not path.is_file():
        raise LayoutPilotManifestError(f"identity workbook does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise LayoutPilotManifestError(
                "identity workbook must contain exactly one worksheet"
            )
        rows = workbook.worksheets[0].iter_rows(values_only=False)
        header_cells = tuple(next(rows, ()))
        if any(cell.data_type == "f" for cell in header_cells):
            raise LayoutPilotManifestError("identity header formulas are forbidden")
        header = tuple(clean(cell.value) for cell in header_cells)
        if header != IDENTITY_FIELDS:
            raise LayoutPilotManifestError(
                "identity columns must be exactly REFNO, RESULTNO, ROCKNAME; "
                "labels and source-table fields are forbidden"
            )

        records: list[dict[str, str]] = []
        for row_number, cells in enumerate(rows, start=2):
            cells = tuple(cells)
            if any(cell.data_type == "f" for cell in cells):
                raise LayoutPilotManifestError(
                    f"identity formulas are forbidden at Excel row {row_number}"
                )
            values = tuple(clean(cell.value) for cell in cells)
            if not any(values):
                continue
            if len(values) != len(IDENTITY_FIELDS) or any(not value for value in values):
                raise LayoutPilotManifestError(
                    f"identity row {row_number} must contain exactly three nonempty values"
                )
            records.append(dict(zip(IDENTITY_FIELDS, values)))
    finally:
        workbook.close()

    if not records:
        raise LayoutPilotManifestError("identity workbook contains no data rows")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise LayoutPilotManifestError(
            "identity workbook contains duplicate RESULTNO values"
        )
    return records


def _distinct_files(paths: Sequence[Path]) -> list[Path]:
    unique: list[Path] = []
    for path in paths:
        if any(path.samefile(existing) for existing in unique):
            continue
        unique.append(path.resolve())
    return unique


def index_local_pdfs(pdf_dir: Path) -> tuple[dict[str, list[Path]], list[str]]:
    """Index top-level PDFs by exact filename stem and preserve ambiguity."""

    pdf_dir = pdf_dir.expanduser().resolve()
    if not pdf_dir.is_dir():
        raise LayoutPilotManifestError(f"PDF directory does not exist: {pdf_dir}")
    paths = sorted(
        (
            path
            for path in pdf_dir.iterdir()
            if path.is_file() and path.suffix.casefold() == ".pdf"
        ),
        key=lambda path: (path.name.casefold(), path.name),
    )
    grouped: dict[str, list[Path]] = defaultdict(list)
    for path in paths:
        grouped[path.stem].append(path)
    return (
        {refno: _distinct_files(values) for refno, values in grouped.items()},
        [path.name for path in paths],
    )


def build_layout_pilot_manifest(
    identity_path: Path,
    pdf_dir: Path,
) -> dict[str, Any]:
    """Build the deterministic pilot manifest without accessing any labels."""

    identity_path = identity_path.expanduser().resolve()
    pdf_dir = pdf_dir.expanduser().resolve()
    records = load_identity_rows(identity_path)
    pdf_index, pdf_names = index_local_pdfs(pdf_dir)

    counts = Counter(record["REFNO"] for record in records)
    eligible_refnos = {
        refno for refno, count in counts.items() if count >= MINIMUM_ROWS_PER_REFNO
    }
    ambiguous = {
        refno: paths
        for refno in eligible_refnos
        if len(paths := pdf_index.get(refno, [])) > 1
    }
    if ambiguous:
        details = "; ".join(
            f"{refno}: {', '.join(str(path) for path in paths)}"
            for refno, paths in sorted(ambiguous.items())
        )
        raise LayoutPilotManifestError(
            f"eligible REFNO values have multiple local PDF mappings: {details}"
        )

    missing_pdf_refnos = sorted(
        refno for refno in eligible_refnos if not pdf_index.get(refno)
    )
    selected_refnos = {
        refno for refno in eligible_refnos if len(pdf_index.get(refno, [])) == 1
    }
    selected_rows = [
        {field: record[field] for field in IDENTITY_FIELDS}
        for record in records
        if record["REFNO"] in selected_refnos
    ]
    if not selected_rows:
        raise LayoutPilotManifestError(
            "no REFNO satisfies the minimum row count and unique local PDF rule"
        )

    paper_order = list(dict.fromkeys(row["REFNO"] for row in selected_rows))
    papers: list[dict[str, Any]] = []
    frozen_pdf_identity: list[dict[str, str]] = []
    for refno in paper_order:
        pdf_path = pdf_index[refno][0]
        pdf_sha256 = file_sha256(pdf_path)
        papers.append(
            {
                "REFNO": refno,
                "identity_row_count": counts[refno],
                "pdf": {
                    "path": str(pdf_path),
                    "name": pdf_path.name,
                    "size_bytes": pdf_path.stat().st_size,
                    "sha256": pdf_sha256,
                },
            }
        )
        frozen_pdf_identity.append(
            {"REFNO": refno, "name": pdf_path.name, "sha256": pdf_sha256}
        )

    selection_rule = {
        "id": SELECTION_RULE_ID,
        "grouping_key": "REFNO",
        "minimum_identity_rows_inclusive": MINIMUM_ROWS_PER_REFNO,
        "pdf_match": (
            "one top-level regular file whose exact filename stem equals REFNO "
            "and whose suffix case-folds to .pdf"
        ),
        "ambiguous_pdf_policy": "fail_closed",
        "missing_pdf_policy": "exclude_and_audit",
        "row_order": "original identity workbook order",
        "label_free": True,
    }
    selection_payload = {
        "selection_rule": selection_rule,
        "rows": selected_rows,
        "pdfs": frozen_pdf_identity,
    }
    below_threshold_refnos = {
        refno for refno, count in counts.items() if count < MINIMUM_ROWS_PER_REFNO
    }
    selected_counts = [counts[refno] for refno in selected_refnos]
    runtime_boundary = {
        "identity_columns_read": list(IDENTITY_FIELDS),
        "labels_read": False,
        "pmagresult_or_rockmag_read": False,
        "prior_predictions_read": False,
        "evaluation_results_read": False,
        "source_text_read": False,
        "pdf_content_used_only_for_sha256": True,
    }
    return {
        "schema_version": SCHEMA_VERSION,
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "identity_fields": list(IDENTITY_FIELDS),
        "selection_rule": selection_rule,
        "selection_rule_sha256": json_sha256(selection_rule),
        "selection_sha256": json_sha256(selection_payload),
        "selected_rows_sha256": json_sha256(selected_rows),
        "selected_pdf_set_sha256": json_sha256(frozen_pdf_identity),
        "inputs": {
            "identity": {
                "path": str(identity_path),
                "sha256": file_sha256(identity_path),
            },
            "pdf_directory": {
                "path": str(pdf_dir),
                "top_level_pdf_files": len(pdf_names),
                "membership_sha256": json_sha256(pdf_names),
            },
        },
        "runtime_boundary": runtime_boundary,
        "audit": {
            "identity_rows": len(records),
            "identity_papers": len(counts),
            "eligible_by_row_count_rows": sum(
                counts[refno] for refno in eligible_refnos
            ),
            "eligible_by_row_count_papers": len(eligible_refnos),
            "selected_rows": len(selected_rows),
            "selected_papers": len(selected_refnos),
            "selected_row_count_minimum": min(selected_counts),
            "selected_row_count_maximum": max(selected_counts),
            "excluded_below_threshold_rows": sum(
                counts[refno] for refno in below_threshold_refnos
            ),
            "excluded_below_threshold_papers": len(below_threshold_refnos),
            "excluded_missing_pdf_rows": sum(
                counts[refno] for refno in missing_pdf_refnos
            ),
            "excluded_missing_pdf_papers": len(missing_pdf_refnos),
            "excluded_missing_pdf_refnos": missing_pdf_refnos,
            "ambiguous_pdf_papers": 0,
            "output_identity_fields": list(IDENTITY_FIELDS),
            "forbidden_sources_read": [],
        },
        "papers": papers,
        "rows": selected_rows,
    }


def freeze_layout_pilot_manifest(
    identity_path: Path,
    pdf_dir: Path,
    output_path: Path,
) -> dict[str, Any]:
    identity_path = identity_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve(strict=False)
    if output_path.suffix.casefold() != ".json":
        raise LayoutPilotManifestError("pilot manifest output must be a JSON file")
    if output_path == identity_path:
        raise LayoutPilotManifestError("identity input and manifest output must differ")
    manifest = build_layout_pilot_manifest(identity_path, pdf_dir)
    selected_pdf_paths = {
        Path(paper["pdf"]["path"]).resolve() for paper in manifest["papers"]
    }
    if output_path in selected_pdf_paths:
        raise LayoutPilotManifestError("manifest output cannot overwrite a selected PDF")
    atomic_json(output_path, manifest)
    return manifest


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    manifest = freeze_layout_pilot_manifest(
        args.identity,
        args.pdf_dir,
        args.output,
    )
    print(
        "R3 layout pilot manifest frozen: "
        f"papers={manifest['audit']['selected_papers']} "
        f"rows={manifest['audit']['selected_rows']} "
        f"selection_sha256={manifest['selection_sha256']} "
        f"output={args.output.expanduser().resolve(strict=False)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "IDENTITY_FIELDS",
    "LayoutPilotManifestError",
    "MINIMUM_ROWS_PER_REFNO",
    "build_layout_pilot_manifest",
    "file_sha256",
    "freeze_layout_pilot_manifest",
    "index_local_pdfs",
    "json_sha256",
    "load_identity_rows",
    "main",
]
