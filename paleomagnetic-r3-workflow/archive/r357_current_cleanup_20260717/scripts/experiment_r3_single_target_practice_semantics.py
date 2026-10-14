#!/usr/bin/env python3
"""Evaluate ROCKMAG-practice semantic features on single-target R3 papers."""

from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path
import re
import sys
from typing import Any, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import experiment_r3_all_labelled_pdf_grouped_oof as all_oof  # noqa: E402
import experiment_r3_exact300_nested_router as nested  # noqa: E402
import experiment_r3_pdf_only_strategies as base  # noqa: E402
import experiment_r3_small_screen_curation_style as curation  # noqa: E402


OOF_DIR = ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1"
PAGE_CACHE = ROOT / "tmp/r3_all_labelled_pdf_page_cache_v1/page_cache.jsonl"
PMAGRESULT = ROOT / "data/source_metadata/PMAGRESULT.xlsx"
FEATURE_OUTPUT = OOF_DIR / "single_target_practice_semantic_features.npz"
REPORT_OUTPUT = OOF_DIR / "single_target_practice_semantic_report.json"

REFERENCE_HEADING_RE = re.compile(
    r"(?im)^\s*(?:references|bibliography|literature\s+cited|works\s+cited)\s*[:.]?\s*$"
)
EXECUTION_RE = re.compile(
    r"\b(?:we\s+(?:measur|perform|determin|obtain|observ|analys|investigat|examin|use)|"
    r"(?:was|were|is|are|has\s+been|have\s+been)\s+(?:measur|perform|determin|obtain|"
    r"observ|analys|investigat|examin|subject|heat|cool|apply)|measurements?\s+(?:were|was)|"
    r"experiments?\s+(?:were|was)|analys(?:is|es)\s+(?:was|were)|curves?\s+(?:show|indicat|"
    r"display|reveal)|results?\s+(?:show|indicat|suggest|reveal)|Figure\s+\d+\s+(?:show|"
    r"display)|Table\s+\d+\s+(?:show|list))\w*",
    re.I,
)
RESULT_RE = re.compile(
    r"\b(?:show|indicat|suggest|reveal|demonstrat|confirm|identify|consist(?:ent)?\s+with|"
    r"attribut|domin|carrier|carried\s+by|resid(?:e|es|ing)|due\s+to|saturat|Curie|"
    r"coerciv|domain|reversib|irreversib|transition|oxidation|exsolution|alteration)\w*",
    re.I,
)
QUANTIFIED_RE = re.compile(
    r"(?:\b(?:Hc|Hcr|Bc|Bcr|Ms|Mrs|Js|Jrs|Tc|Tcurie)\b\s*(?:/\s*[A-Za-z]+)?\s*[=:~<>]?\s*"
    r"\d|\b\d+(?:\.\d+)?\s*(?:mT|T\b|K\b|°?C\b)|(?:saturat|Curie|coerciv).{0,50}\d)",
    re.I,
)
CARRIER_ATTRIBUTION_RE = re.compile(
    r"(?:reman(?:ence|ent)|magneti[sz]ation|component).{0,100}(?:carried\s+by|resid|"
    r"attribut|due\s+to|associated\s+with|held\s+by)|(?:carried\s+by|resid|attribut|"
    r"due\s+to).{0,100}(?:magnetite|titanomagnetite|ha?ematite|maghemite|pyrrhotite|"
    r"greigite|goethite)",
    re.I | re.S,
)
DEMAG_ONLY_RE = re.compile(
    r"(?:unblock(?:ing)?|removed|demagneti[sz]).{0,45}\b\d{2,3}\s*°?C\b|"
    r"\bT(?:ub|b)\s*(?:=|>|<|:)\s*\d",
    re.I,
)
NEGATED_OR_PROSPECTIVE_RE = re.compile(
    r"\b(?:not|no|without|future|should\s+be|could\s+be|would\s+be|recommended|"
    r"remains?\s+to\s+be)\b.{0,70}(?:measur|perform|investigat|analys|examin)",
    re.I,
)
METHOD_SECTION_RE = re.compile(
    r"\b(?:method|experimental|laboratory|measurement|procedure|technique)\w*\b",
    re.I,
)
RESULT_SECTION_RE = re.compile(
    r"\b(?:result|observation|discussion|interpretation|magnetic\s+mineralogy|"
    r"rock\s+magnetism|magnetic\s+properties)\b",
    re.I,
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def trim_reference_section(text: str) -> tuple[str, str]:
    match = REFERENCE_HEADING_RE.search(text)
    if match is None:
        return text, ""
    return text[: match.start()], text[match.end() :]


def occurrences(
    patterns: Sequence[re.Pattern[str]], text: str, *, cap: int = 80
) -> list[tuple[int, int]]:
    found: list[tuple[int, int]] = []
    for pattern in patterns:
        for match in pattern.finditer(text):
            found.append((match.start(), match.end()))
            if len(found) >= cap:
                return found
    return sorted(found)


def category(concept: Mapping[str, Any]) -> str:
    return curation.concept_category(concept)


def add_feature(
    values: list[float], names: list[str], name: str, value: float
) -> None:
    names.append(name)
    values.append(float(value))


def semantic_vector(
    text: str, concepts: Sequence[Mapping[str, Any]]
) -> tuple[list[float], list[str], dict[str, Any]]:
    own, references = trim_reference_section(text)
    values: list[float] = []
    names: list[str] = []
    own_hits: dict[str, list[tuple[int, int]]] = {}
    reference_counts: dict[str, int] = {}
    concepts_by_category: dict[str, list[str]] = defaultdict(list)
    for concept in concepts:
        concept_id = str(concept["concept_id"])
        concept_category = category(concept)
        key = f"{concept_category}__{concept_id}"
        own_hits[key] = occurrences(concept["patterns"], own)
        reference_counts[key] = len(occurrences(concept["patterns"], references))
        concepts_by_category[concept_category].append(key)
        own_count = len(own_hits[key])
        reference_count = reference_counts[key]
        add_feature(values, names, f"own_{key}_count_log", math.log1p(own_count))
        add_feature(values, names, f"own_{key}_present", own_count > 0)
        add_feature(
            values,
            names,
            f"reference_{key}_count_log",
            math.log1p(reference_count),
        )
        add_feature(
            values,
            names,
            f"own_share_{key}",
            own_count / max(1, own_count + reference_count),
        )

    category_counts: dict[str, int] = {}
    category_distinct: dict[str, int] = {}
    for group in ("method", "diagnostic", "carrier", "other"):
        category_counts[group] = sum(len(own_hits[key]) for key in concepts_by_category[group])
        category_distinct[group] = sum(bool(own_hits[key]) for key in concepts_by_category[group])
        add_feature(
            values, names, f"own_{group}_count_log", math.log1p(category_counts[group])
        )
        add_feature(
            values,
            names,
            f"own_{group}_distinct_log",
            math.log1p(category_distinct[group]),
        )
        ref_count = sum(reference_counts[key] for key in concepts_by_category[group])
        add_feature(values, names, f"reference_{group}_count_log", math.log1p(ref_count))

    method_contexts: list[str] = []
    method_family_events: dict[str, Counter[str]] = defaultdict(Counter)
    all_diagnostic_patterns = [
        pattern
        for concept in concepts
        if category(concept) == "diagnostic"
        for pattern in concept["patterns"]
    ]
    all_carrier_patterns = [
        pattern
        for concept in concepts
        if category(concept) == "carrier"
        for pattern in concept["patterns"]
    ]
    for concept in concepts:
        if category(concept) != "method":
            continue
        concept_id = str(concept["concept_id"])
        key = f"method__{concept_id}"
        for start, end in own_hits[key]:
            context = own[max(0, start - 260) : min(len(own), end + 260)]
            method_contexts.append(context)
            flags = {
                "execution": bool(EXECUTION_RE.search(context)),
                "result": bool(RESULT_RE.search(context)),
                "quantified": bool(QUANTIFIED_RE.search(context)),
                "diagnostic": any(pattern.search(context) for pattern in all_diagnostic_patterns),
                "carrier": any(pattern.search(context) for pattern in all_carrier_patterns),
                "negated": bool(NEGATED_OR_PROSPECTIVE_RE.search(context)),
                "methods_section": bool(METHOD_SECTION_RE.search(context)),
                "results_section": bool(RESULT_SECTION_RE.search(context)),
            }
            for flag, state in flags.items():
                method_family_events[concept_id][flag] += int(state)
            method_family_events[concept_id]["event"] += 1
            method_family_events[concept_id]["strong"] += int(
                not flags["negated"]
                and (
                    flags["execution"]
                    or flags["result"]
                    or flags["quantified"]
                    or flags["diagnostic"]
                    or flags["carrier"]
                )
            )

    event_flags = (
        "event",
        "execution",
        "result",
        "quantified",
        "diagnostic",
        "carrier",
        "negated",
        "methods_section",
        "results_section",
        "strong",
    )
    totals = Counter()
    for concept in concepts:
        if category(concept) != "method":
            continue
        concept_id = str(concept["concept_id"])
        for flag in event_flags:
            count = method_family_events[concept_id][flag]
            totals[flag] += count
            add_feature(
                values,
                names,
                f"method_event_{concept_id}_{flag}_log",
                math.log1p(count),
            )
    for flag in event_flags:
        add_feature(values, names, f"method_event_total_{flag}_log", math.log1p(totals[flag]))

    contexts_with_method_carrier = 0
    contexts_with_method_diagnostic = 0
    contexts_with_triple = 0
    for context in method_contexts:
        has_carrier = any(pattern.search(context) for pattern in all_carrier_patterns)
        has_diagnostic = any(pattern.search(context) for pattern in all_diagnostic_patterns)
        contexts_with_method_carrier += int(has_carrier)
        contexts_with_method_diagnostic += int(has_diagnostic)
        contexts_with_triple += int(has_carrier and has_diagnostic)
    add_feature(
        values,
        names,
        "window_method_carrier_count_log",
        math.log1p(contexts_with_method_carrier),
    )
    add_feature(
        values,
        names,
        "window_method_diagnostic_count_log",
        math.log1p(contexts_with_method_diagnostic),
    )
    add_feature(
        values, names, "window_triple_count_log", math.log1p(contexts_with_triple)
    )

    carrier_attribution_count = len(CARRIER_ATTRIBUTION_RE.findall(own))
    quantified_count = len(QUANTIFIED_RE.findall(own))
    demag_only_count = len(DEMAG_ONLY_RE.findall(own))
    add_feature(
        values,
        names,
        "carrier_attribution_count_log",
        math.log1p(carrier_attribution_count),
    )
    add_feature(values, names, "quantified_result_count_log", math.log1p(quantified_count))
    add_feature(values, names, "demag_only_count_log", math.log1p(demag_only_count))
    add_feature(values, names, "own_text_characters_log", math.log1p(len(own)))
    add_feature(values, names, "reference_fraction", len(references) / max(1, len(text)))

    performed_method = totals["strong"] > 0
    self_authenticating = bool(
        category_counts["diagnostic"] > 0
        and (quantified_count > 0 or totals["result"] > 0)
    )
    carrier_attribution = carrier_attribution_count > 0
    method_only = category_counts["method"] > 0 and not (
        category_counts["diagnostic"] or category_counts["carrier"]
    )
    carrier_only = category_counts["carrier"] > 0 and category_counts["method"] == 0
    practice_positive = performed_method or self_authenticating or carrier_attribution
    for name, state in (
        ("criterion_performed_method", performed_method),
        ("criterion_self_authenticating_result", self_authenticating),
        ("criterion_carrier_attribution", carrier_attribution),
        ("criterion_method_only", method_only),
        ("criterion_carrier_only", carrier_only),
        ("criterion_practice_positive", practice_positive),
    ):
        add_feature(values, names, name, state)
    return values, names, {
        "performed_method": performed_method,
        "self_authenticating_result": self_authenticating,
        "carrier_attribution": carrier_attribution,
        "practice_positive": practice_positive,
        "method_events": totals["event"],
        "strong_method_events": totals["strong"],
    }


def build_feature_matrix(
    records: Sequence[Mapping[str, str]],
    paper_text: Mapping[str, str],
    concepts: Sequence[Mapping[str, Any]],
) -> tuple[np.ndarray, list[str], dict[str, Any]]:
    vectors: dict[str, list[float]] = {}
    names: list[str] | None = None
    meta: dict[str, Any] = {}
    for number, refno in enumerate(sorted(paper_text, key=lambda value: (int(value), value)), start=1):
        vector, current_names, current_meta = semantic_vector(paper_text[refno], concepts)
        if names is None:
            names = current_names
        elif names != current_names:
            raise AssertionError("practice semantic feature order changed")
        vectors[refno] = vector
        meta[refno] = current_meta
        if number % 250 == 0:
            print(f"practice features: papers={number}/{len(paper_text)}", flush=True)
    matrix = np.asarray([vectors[record["REFNO"]] for record in records], dtype=np.float32)
    if not np.isfinite(matrix).all():
        raise AssertionError("practice semantic features are non-finite")
    return matrix, names or [], meta


def load_rockmag_nonempty(records: Sequence[Mapping[str, str]]) -> np.ndarray:
    wanted = {record["RESULTNO"] for record in records}
    nonempty: set[str] = set()
    workbook = load_workbook(PMAGRESULT, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = [clean(value) for value in next(rows, ())]
        result_index = header.index("RESULTNO")
        rockmag_index = header.index("ROCKMAG")
        for row in rows:
            resultno = clean(row[result_index])
            if resultno in wanted and clean(row[rockmag_index]):
                nonempty.add(resultno)
    finally:
        workbook.close()
    return np.asarray(
        [int(record["RESULTNO"] in nonempty) for record in records], dtype=np.int8
    )


def array_sha256(value: np.ndarray) -> str:
    return hashlib.sha256(np.ascontiguousarray(value).view(np.uint8)).hexdigest()


def main() -> int:
    receipt_path = OOF_DIR / "prediction_receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))["rows"]
    records = [
        {key: str(row[key]) for key in ("REFNO", "RESULTNO", "ROCKNAME")}
        for row in receipt
    ]
    fold_ids = np.asarray([int(row["fold"]) for row in receipt], dtype=np.int8)
    refnos = {record["REFNO"] for record in records}
    paper_text = base.load_page_text(PAGE_CACHE, refnos)
    concepts, knowledge_audit = base.load_knowledge_patterns(
        base.DEFAULT_TAXONOMY, base.DEFAULT_POSITIVE_KB
    )
    practice, feature_names, paper_meta = build_feature_matrix(
        records, paper_text, concepts
    )
    combined = np.load(OOF_DIR / "combined_features.npz")["combined"].astype(
        np.float32, copy=False
    )
    combined_practice = np.concatenate((combined, practice), axis=1).astype(
        np.float32, copy=False
    )
    np.savez_compressed(
        FEATURE_OUTPUT,
        practice=practice,
        combined_practice=combined_practice,
    )
    feature_freeze = {
        "status": "FROZEN_BEFORE_LABEL_OR_ROCKMAG_ACCESS",
        "rows": len(records),
        "practice_features": practice.shape[1],
        "combined_practice_features": combined_practice.shape[1],
        "feature_names": feature_names,
        "practice_sha256": array_sha256(practice),
        "combined_practice_sha256": array_sha256(combined_practice),
        "page_cache_sha256": base.file_sha256(PAGE_CACHE),
        "knowledge": knowledge_audit,
    }
    base.atomic_json(OOF_DIR / "single_target_practice_semantic_feature_freeze.json", feature_freeze)

    labels_by_resultno = all_oof.load_labels(
        base.DEFAULT_TARGETS, {record["RESULTNO"] for record in records}
    )
    labels = np.asarray(
        [labels_by_resultno[record["RESULTNO"]] for record in records], dtype=np.int8
    )
    rockmag_nonempty = load_rockmag_nonempty(records)
    identity = all_oof.load_identity(all_oof.IDENTITY)
    multiplicity = Counter(record["REFNO"] for record in identity)
    single = np.asarray(
        [multiplicity[record["REFNO"]] == 1 for record in records], dtype=bool
    )
    canonical_targets = [
        base.canonical_target_identity(record["ROCKNAME"]) for record in records
    ]
    experts = (
        nested.ExpertSpec("hgb", "matrix", "hgb10"),
        nested.ExpertSpec("rf", "matrix", "rf8"),
        nested.ExpertSpec("extra", "matrix", "extra8"),
        nested.ExpertSpec("logreg", "matrix", "logreg001", standardized=True),
    )
    matrices = {"practice": practice, "combined_practice": combined_practice}
    targets = {"r3": labels, "rockmag_nonempty": rockmag_nonempty}
    probabilities: dict[str, np.ndarray] = {}
    for target_name, target in targets.items():
        for matrix_name, matrix in matrices.items():
            current = {
                expert.name: np.full(len(records), np.nan, dtype=np.float64)
                for expert in experts
            }
            for fold in range(1, all_oof.N_SPLITS + 1):
                train = np.flatnonzero(fold_ids != fold)
                test = np.flatnonzero(fold_ids == fold)
                weights = base.inverse_target_multiplicity_weights(
                    train, records, canonical_targets
                )
                for expert in experts:
                    fitted = nested.fit_expert(expert, matrix, train, target, weights)
                    current[expert.name][test] = fitted.predict(matrix[test])
            current["mean4"] = np.column_stack(
                [current[expert.name] for expert in experts]
            ).mean(axis=1)
            for expert_name, values in current.items():
                probabilities[f"{target_name}__{matrix_name}__{expert_name}"] = values

    baseline = np.load(OOF_DIR / "oof_probabilities.npz")["mean4"].astype(np.float64)
    results = {
        name: base.metric_record(labels[single], values[single])
        for name, values in probabilities.items()
    }
    direct = np.asarray(
        [bool(paper_meta[record["REFNO"]]["practice_positive"]) for record in records],
        dtype=np.float64,
    )
    results["direct_practice_criterion"] = base.metric_record(
        labels[single], direct[single]
    )
    results["existing_mean4"] = base.metric_record(labels[single], baseline[single])
    blends: dict[str, Any] = {}
    for name, values in probabilities.items():
        for weight in (0.25, 0.50, 0.75):
            blend_name = f"existing_mean4__{name}__w{int(weight * 100):02d}"
            blended = (1.0 - weight) * baseline + weight * values
            blends[blend_name] = base.metric_record(labels[single], blended[single])
    ranked = sorted(
        ({"name": name, **metric} for name, metric in {**results, **blends}.items()),
        key=lambda item: (
            -item["accuracy"],
            -item["macro_f1"],
            -item["balanced_accuracy"],
            item["name"],
        ),
    )
    report = {
        "schema_version": "r3_single_target_practice_semantic_experiment_v1",
        "status": "complete_exploratory",
        "scope": {
            "single_target_rows": int(single.sum()),
            "single_target_papers": int(single.sum()),
            "all_pdf_rows_used_for_fold_training": len(records),
            "all_pdf_papers": len(refnos),
        },
        "runtime_boundary": {
            "practice_features_frozen_before_labels_or_rockmag_opened": True,
            "outer_test_refno_excluded_from_model_fit": True,
            "rockmag_used_only_as_offline_fold_training_supervision": True,
            "runtime_test_features_use_only_original_text_and_frozen_knowledge": True,
        },
        "feature_freeze": feature_freeze,
        "results": results,
        "blends": blends,
        "ranked_top_20": ranked[:20],
        "rockmag_nonempty_alignment": base.metric_record(
            labels[single], rockmag_nonempty[single].astype(np.float64)
        ),
        "inputs": {
            "receipt": {
                "path": str(receipt_path),
                "sha256": base.file_sha256(receipt_path),
            },
            "page_cache": {
                "path": str(PAGE_CACHE),
                "sha256": base.file_sha256(PAGE_CACHE),
            },
            "targets": {
                "path": str(base.DEFAULT_TARGETS),
                "sha256": base.file_sha256(base.DEFAULT_TARGETS),
            },
            "pmagresult_offline_only": {
                "path": str(PMAGRESULT),
                "sha256": base.file_sha256(PMAGRESULT),
            },
        },
    }
    base.atomic_json(REPORT_OUTPUT, report)
    np.savez_compressed(
        OOF_DIR / "single_target_practice_semantic_oof.npz",
        single_mask=single,
        label=labels,
        rockmag_nonempty=rockmag_nonempty,
        existing_mean4=baseline.astype(np.float32),
        **{name: values.astype(np.float32) for name, values in probabilities.items()},
    )
    print(json.dumps({"report": str(REPORT_OUTPUT), "top": ranked[:12]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
