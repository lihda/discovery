#!/usr/bin/env python3
"""Build frozen, label-blind R3 criterion-entailment features offline."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import importlib
import json
import math
import os
from pathlib import Path
import re
import struct
import sys
import tempfile
from typing import Any, Callable, Mapping, Optional, Sequence

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import build_r3_trace_page_graph_features as trace_freeze  # noqa: E402


DEFAULT_IDENTITY = (
    ROOT / "outputs/r357_strict_eval_07111758/inputs/identity_manifest.xlsx"
)
DEFAULT_PAGE_CACHE = ROOT / "tmp/r3_all_labelled_pdf_page_cache_v1/page_cache.jsonl"
DEFAULT_FOLD_MANIFEST = (
    ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1/trace_refno_fold_manifest_v4.json"
)
DEFAULT_CONCEPT_SNAPSHOT = (
    ROOT / "tmp/r3_all_labelled_pdf_grouped_oof_v1/trace_page_graph_concepts_v4.json"
)
DEFAULT_MODEL_SNAPSHOT = (
    Path.home()
    / ".cache/huggingface/hub/models--cross-encoder--nli-deberta-v3-base/"
    "snapshots/6c749ce3425cd33b46d187e45b92bbf96ee12ec7"
)
DEFAULT_OUTPUT_DIR = ROOT / "tmp/r3_nli_single64_v1"

SELECTION_NAME = "r3_nli_single64_selection_v1.json"
HYPOTHESES_NAME = "r3_nli_criterion_hypotheses_v1.json"
PREINFERENCE_NAME = "r3_nli_preinference_freeze_v1.json"
FEATURE_NAME = "r3_nli_criterion_features_v1.npz"
MANIFEST_NAME = "r3_nli_criterion_features_v1.manifest.json"
CANDIDATE_NAME = "r3_nli_candidate_freeze_v1.json"
EVIDENCE_NAME = "r3_nli_evidence_receipt_v1.jsonl"
AUDIT_NAME = "r3_nli_criterion_features_v1.audit.json"

IDENTITY_FIELDS = ("REFNO", "RESULTNO", "ROCKNAME")
EXPECTED_ROWS = 2751
PILOT_ROWS = 64
FOLD_QUOTAS = {1: 13, 2: 13, 3: 13, 4: 13, 5: 12}
SELECTION_SALT = "r3-nli-single-target-pilot-v1"
MAX_PREMISES_PER_GROUP = 2
MAX_PREMISE_CHARACTERS = 700
MAX_LENGTH = 384
BATCH_SIZE = 8
ENTAILMENT_THRESHOLD = 0.5
ALLOWED_PREMISE_GROUPS = frozenset(
    {"METHOD_DIAGNOSTIC", "IDENTIFY", "SIGNIFICANCE_COUNTER"}
)
MODEL_FILES = (
    "config.json",
    "model.safetensors",
    "tokenizer_config.json",
    "special_tokens_map.json",
    "tokenizer.json",
    "spm.model",
    "added_tokens.json",
)
BOUNDARY_FALSE_FIELDS = (
    "labels_read",
    "pmagresult_or_rockmag_read",
    "prior_predictions_read",
    "evaluation_results_read",
    "outer_test_metrics_used_for_selection",
)
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
FORBIDDEN_FIELD_RE = re.compile(
    r"(?:^|_)(?:R3|R5|R7|PMAGRESULT|ROCKMAG|LABELS?|GOLD|GROUND_TRUTH|"
    r"PREDICTIONS?|OOF|EVALUATION|METRICS?|OUTCOMES?)(?:$|_)",
    re.I,
)


HYPOTHESES = (
    {
        "hypothesis_id": "scope",
        "premise_groups": "all",
        "polarity": "support",
        "template": (
            "This passage describes experiments, observations, or interpretations "
            "performed by the present study."
        ),
    },
    {
        "hypothesis_id": "method",
        "premise_groups": ["METHOD_DIAGNOSTIC"],
        "polarity": "support",
        "template": (
            "The study applied a rock-magnetic or magnetic-mineralogical investigation "
            "to determine the carrier of the remanence."
        ),
    },
    {
        "hypothesis_id": "identified",
        "premise_groups": ["IDENTIFY"],
        "polarity": "support",
        "template": (
            "The current study identifies the magnetic mineral or material that carries "
            "the remanence."
        ),
    },
    {
        "hypothesis_id": "inconclusive",
        "premise_groups": ["IDENTIFY"],
        "polarity": "support",
        "template": (
            "The current study performed a reasonable carrier-focused investigation but "
            "could not conclusively identify the remanence carrier."
        ),
    },
    {
        "hypothesis_id": "significance",
        "premise_groups": ["SIGNIFICANCE_COUNTER"],
        "polarity": "support",
        "template": (
            "The current study explains what the magnetic carrier evidence implies about "
            "the origin, alteration, stability, or reliability of the remanent magnetization."
        ),
    },
    {
        "hypothesis_id": "external",
        "premise_groups": "all",
        "polarity": "counter",
        "template": (
            "This passage describes only prior or external work rather than an analysis "
            "performed by the present study."
        ),
    },
    {
        "hypothesis_id": "not_performed",
        "premise_groups": "all",
        "polarity": "counter",
        "template": (
            "The study explicitly states that the carrier-focused rock-magnetic "
            "investigation was not performed."
        ),
    },
)


def _feature_names() -> list[str]:
    names = [
        "nli_pilot_selected",
        "nli_source_available",
        "nli_retrieval_available",
        "nli_method_diagnostic_premise_count_log",
        "nli_identify_premise_count_log",
        "nli_significance_counter_premise_count_log",
    ]
    for item in HYPOTHESES:
        prefix = f"nli_{item['hypothesis_id']}"
        names.extend(
            [
                f"{prefix}_max_entailment",
                f"{prefix}_top_mean_entailment",
                f"{prefix}_max_contradiction",
                f"{prefix}_max_entailment_margin",
                f"{prefix}_entailment_fraction_ge_0_5",
            ]
        )
    names.extend(
        [
            "nli_scope_support",
            "nli_negative_control_max_entailment",
            "nli_method_scope_gated_support",
            "nli_identify_scope_gated_support",
            "nli_significance_scope_gated_support",
            "nli_method_only_support",
            "nli_practice_chain_min_support",
            "nli_normative_chain_min_support",
            "nli_criterion_chain_min_support",
        ]
    )
    return names


FEATURE_NAMES = tuple(_feature_names())


class NLICriterionFeatureError(ValueError):
    """Raised when the NLI feature freeze violates its source boundary."""


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text_sha256(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def json_sha256(value: Any) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _require_sha256(value: Any, description: str) -> str:
    digest = clean(value).casefold()
    if not SHA256_RE.fullmatch(digest):
        raise NLICriterionFeatureError(f"{description} is not a SHA-256 digest")
    return digest


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def atomic_jsonl(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True))
                handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def atomic_npz(path: Path, **arrays: np.ndarray) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    try:
        with temporary.open("wb") as handle:
            np.savez_compressed(handle, **arrays)
        temporary.replace(path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def load_identity(path: Path) -> list[dict[str, str]]:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise NLICriterionFeatureError(f"identity workbook does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise NLICriterionFeatureError(
                "identity workbook must contain exactly one worksheet"
            )
        rows = workbook.worksheets[0].iter_rows(values_only=False)
        header_cells = tuple(next(rows, ()))
        if any(cell.data_type == "f" for cell in header_cells):
            raise NLICriterionFeatureError("identity header formulas are forbidden")
        if tuple(clean(cell.value) for cell in header_cells) != IDENTITY_FIELDS:
            raise NLICriterionFeatureError(
                "identity columns must be exactly REFNO, RESULTNO, ROCKNAME"
            )
        records: list[dict[str, str]] = []
        for row_number, cells in enumerate(rows, start=2):
            cells = tuple(cells)
            if any(cell.data_type == "f" for cell in cells):
                raise NLICriterionFeatureError(
                    f"identity formulas are forbidden at Excel row {row_number}"
                )
            values = tuple(clean(cell.value) for cell in cells)
            if not any(values):
                continue
            if len(values) != 3 or any(not value for value in values):
                raise NLICriterionFeatureError(
                    f"identity row {row_number} must contain three nonempty values"
                )
            records.append(dict(zip(IDENTITY_FIELDS, values)))
    finally:
        workbook.close()
    if not records:
        raise NLICriterionFeatureError("identity workbook is empty")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise NLICriterionFeatureError("identity workbook contains duplicate RESULTNO")
    return records


def load_v4_alignment(
    identity_records: Sequence[Mapping[str, str]], fold_path: Path
) -> tuple[list[dict[str, str]], np.ndarray, dict[str, Any]]:
    payload = json.loads(fold_path.read_text(encoding="utf-8"))
    if (
        clean(payload.get("schema_version")) != "r3_trace_refno_fold_manifest_v4"
        or payload.get("status") != "FROZEN_WITHOUT_LABEL_ACCESS"
        or payload.get("labels_read") is not False
    ):
        raise NLICriterionFeatureError("V4 fold manifest has an invalid schema")
    by_resultno = {clean(row["RESULTNO"]): row for row in identity_records}
    rows = payload.get("rows") or []
    records: list[dict[str, str]] = []
    fold_ids = np.zeros(len(rows), dtype=np.int8)
    document_folds: dict[str, set[int]] = defaultdict(set)
    for index, item in enumerate(rows):
        resultno = clean(item.get("RESULTNO"))
        record = by_resultno.get(resultno)
        if (
            not isinstance(item, Mapping)
            or int(item.get("row_index", -1)) != index
            or record is None
            or clean(item.get("REFNO")) != clean(record.get("REFNO"))
        ):
            raise NLICriterionFeatureError("V4 fold identity does not align")
        records.append({field: clean(record[field]) for field in IDENTITY_FIELDS})
        fold_ids[index] = int(item.get("fold"))
        document = _require_sha256(item.get("document_sha256"), "document hash")
        document_folds[document].add(int(fold_ids[index]))
    alignment = trace_freeze._record_alignment_sha256(records)
    if alignment != clean(payload.get("alignment_sha256")):
        raise NLICriterionFeatureError("V4 alignment hash mismatch")
    if hashlib.sha256(fold_ids.tobytes()).hexdigest() != clean(
        payload.get("fold_id_sha256")
    ):
        raise NLICriterionFeatureError("V4 fold-id hash mismatch")
    for refno in {record["REFNO"] for record in records}:
        if len(
            {
                int(fold_ids[index])
                for index, record in enumerate(records)
                if record["REFNO"] == refno
            }
        ) != 1:
            raise NLICriterionFeatureError("V4 folds split a REFNO")
    if any(len(values) != 1 for values in document_folds.values()):
        raise NLICriterionFeatureError("V4 folds split normalized document content")
    return records, fold_ids, payload


def select_single64(
    records: Sequence[Mapping[str, str]], fold_ids: np.ndarray
) -> tuple[np.ndarray, dict[str, Any]]:
    counts = Counter(record["REFNO"] for record in records)
    eligible = [index for index, record in enumerate(records) if counts[record["REFNO"]] == 1]
    by_fold: dict[int, list[tuple[str, int]]] = defaultdict(list)
    for index in eligible:
        record = records[index]
        fold = int(fold_ids[index])
        digest = hashlib.sha256(
            f"{SELECTION_SALT}|{record['REFNO']}|{record['RESULTNO']}".encode("utf-8")
        ).hexdigest()
        by_fold[fold].append((digest, index))
    selected: list[int] = []
    rows: list[dict[str, Any]] = []
    for fold, quota in FOLD_QUOTAS.items():
        candidates = sorted(by_fold.get(fold, []))
        if len(candidates) < quota:
            raise NLICriterionFeatureError(
                f"fold {fold} has {len(candidates)} single-target papers; needs {quota}"
            )
        for rank, (digest, index) in enumerate(candidates[:quota], start=1):
            selected.append(index)
            record = records[index]
            rows.append(
                {
                    "row_index": index,
                    **{field: record[field] for field in IDENTITY_FIELDS},
                    "fold": fold,
                    "selection_rank_within_fold": rank,
                    "selection_digest": digest,
                }
            )
    mask = np.zeros(len(records), dtype=bool)
    mask[selected] = True
    if int(mask.sum()) != PILOT_ROWS:
        raise AssertionError("single64 selection has the wrong size")
    selection_core = {
        "salt": SELECTION_SALT,
        "fold_quotas": {str(key): value for key, value in FOLD_QUOTAS.items()},
        "rows": rows,
    }
    return mask, {
        "schema_version": "r3_nli_single64_selection_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "selection_rule": (
            "within each frozen V4 fold, select the smallest SHA256 of "
            "salt|REFNO|RESULTNO among PDF-backed REFNO values with exactly one row"
        ),
        "salt": SELECTION_SALT,
        "fold_quotas": selection_core["fold_quotas"],
        "eligible_single_target_rows": len(eligible),
        "selected_rows": len(rows),
        "selection_sha256": json_sha256(selection_core),
        "rows": rows,
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
    }


def load_concepts(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if (
        clean(payload.get("schema_version")) != "r3_trace_page_graph_concepts_v4"
        or payload.get("status") != "FROZEN_WITHOUT_LABEL_ACCESS"
    ):
        raise NLICriterionFeatureError("concept snapshot has an invalid schema")
    concepts = payload.get("concepts") or []
    if not isinstance(concepts, list) or not concepts:
        raise NLICriterionFeatureError("concept snapshot is empty")
    for concept in concepts:
        if not isinstance(concept, Mapping):
            raise NLICriterionFeatureError("concept snapshot contains a non-object")
        if FORBIDDEN_FIELD_RE.search(clean(concept.get("concept_id"))):
            raise NLICriterionFeatureError("concept snapshot contains a forbidden field")
    for name in ("criterion", "normative_concepts"):
        item = payload.get(name) or {}
        source = Path(clean(item.get("path"))).resolve()
        if not source.is_file() or file_sha256(source) != clean(item.get("sha256")):
            raise NLICriterionFeatureError(f"concept source changed: {name}")
    return [dict(item) for item in concepts], payload


def _safetensors_audit(path: Path) -> dict[str, Any]:
    dtype_bytes = {
        "F64": 8,
        "F32": 4,
        "F16": 2,
        "BF16": 2,
        "I64": 8,
        "I32": 4,
        "I16": 2,
        "I8": 1,
        "U8": 1,
        "BOOL": 1,
    }
    size = path.stat().st_size
    with path.open("rb") as handle:
        raw_length = handle.read(8)
        if len(raw_length) != 8:
            raise NLICriterionFeatureError("model.safetensors lacks a header")
        header_length = struct.unpack("<Q", raw_length)[0]
        if header_length <= 0 or 8 + header_length > size:
            raise NLICriterionFeatureError("model.safetensors header length is invalid")
        try:
            header = json.loads(handle.read(header_length))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise NLICriterionFeatureError("model.safetensors header is invalid") from exc
    intervals: list[tuple[int, int]] = []
    parameter_count = 0
    for name, item in header.items():
        if name == "__metadata__":
            continue
        if not isinstance(item, Mapping) or item.get("dtype") not in dtype_bytes:
            raise NLICriterionFeatureError("model.safetensors tensor metadata is invalid")
        shape = item.get("shape")
        offsets = item.get("data_offsets")
        if (
            not isinstance(shape, list)
            or not isinstance(offsets, list)
            or len(offsets) != 2
        ):
            raise NLICriterionFeatureError("model.safetensors tensor shape is invalid")
        count = math.prod(int(value) for value in shape)
        start, end = map(int, offsets)
        if end - start != count * dtype_bytes[item["dtype"]]:
            raise NLICriterionFeatureError("model.safetensors tensor byte size is invalid")
        intervals.append((start, end))
        parameter_count += count
    intervals.sort()
    if (
        not intervals
        or intervals[0][0] != 0
        or any(left[1] != right[0] for left, right in zip(intervals, intervals[1:]))
        or intervals[-1][1] != size - 8 - header_length
    ):
        raise NLICriterionFeatureError("model.safetensors data offsets are incomplete")
    return {
        "header_bytes": header_length,
        "tensor_count": len(intervals),
        "stored_values_including_buffers": parameter_count,
        "container_exact": True,
    }


def inspect_model_snapshot(path: Path) -> dict[str, Any]:
    path = path.expanduser().resolve()
    if not path.is_dir():
        raise NLICriterionFeatureError(f"model snapshot does not exist: {path}")
    files = []
    for name in MODEL_FILES:
        item = path / name
        if not item.is_file():
            raise NLICriterionFeatureError(f"model snapshot lacks {name}")
        files.append(
            {"name": name, "size_bytes": item.stat().st_size, "sha256": file_sha256(item)}
        )
    config = json.loads((path / "config.json").read_text(encoding="utf-8"))
    expected_labels = {
        "0": "contradiction",
        "1": "entailment",
        "2": "neutral",
    }
    if config.get("id2label") != expected_labels or config.get("num_labels", 3) != 3:
        raise NLICriterionFeatureError(
            "NLI model labels must be 0=contradiction, 1=entailment, 2=neutral"
        )
    weight_audit = _safetensors_audit(path / "model.safetensors")
    snapshot_core = {"files": files, "id2label": expected_labels}
    return {
        "path": str(path),
        "snapshot_sha256": json_sha256(snapshot_core),
        "files": files,
        "id2label": expected_labels,
        "safetensors": weight_audit,
    }


def _resolve_selector(
    selector: Optional[Callable[..., Mapping[str, Any]]],
) -> tuple[Callable[..., Mapping[str, Any]], dict[str, Any]]:
    if selector is not None:
        identity = f"{getattr(selector, '__module__', '')}.{getattr(selector, '__qualname__', '')}"
        return selector, {"kind": "injected", "identity": identity}
    try:
        module = importlib.import_module("r3_nli_premise_selector")
    except ImportError as exc:
        raise NLICriterionFeatureError(
            "r3_nli_premise_selector.select_nli_premises is not available"
        ) from exc
    selected = getattr(module, "select_nli_premises", None)
    if not callable(selected):
        raise NLICriterionFeatureError("selector lacks select_nli_premises")
    module_path = Path(module.__file__).resolve()
    return selected, {
        "kind": "module",
        "path": str(module_path),
        "sha256": file_sha256(module_path),
        "callable": "select_nli_premises",
    }


def _real_predictor(
    pairs: Sequence[Mapping[str, Any]],
    model_snapshot: Path,
    inference_config: Mapping[str, Any],
) -> np.ndarray:
    try:
        import torch
        from transformers import AutoModelForSequenceClassification, AutoTokenizer
    except ImportError as exc:
        raise NLICriterionFeatureError(
            "offline NLI inference requires torch and transformers"
        ) from exc
    tokenizer = AutoTokenizer.from_pretrained(
        model_snapshot, local_files_only=True, use_fast=True
    )
    model = AutoModelForSequenceClassification.from_pretrained(
        model_snapshot, local_files_only=True
    )
    model.to("cpu")
    model.eval()
    batches: list[np.ndarray] = []
    batch_size = int(inference_config["batch_size"])
    with torch.inference_mode():
        for start in range(0, len(pairs), batch_size):
            batch = pairs[start : start + batch_size]
            encoded = tokenizer(
                [clean(item["premise"]) for item in batch],
                [clean(item["hypothesis"]) for item in batch],
                padding=True,
                truncation="only_first",
                max_length=int(inference_config["max_length"]),
                return_tensors="pt",
            )
            logits = model(**encoded).logits
            batches.append(logits.softmax(dim=-1).cpu().numpy())
    return np.concatenate(batches, axis=0) if batches else np.zeros((0, 3))


def _validate_premise_selection(
    payload: Mapping[str, Any], refno: str
) -> list[dict[str, Any]]:
    if clean(payload.get("schema_version")) != "r3_nli_premise_selection_v1":
        raise NLICriterionFeatureError(
            f"selector returned an invalid schema for REFNO {refno}"
        )
    premises = payload.get("premises") or []
    if not isinstance(premises, list):
        raise NLICriterionFeatureError("selector premises must be a list")
    group_counts: Counter[str] = Counter()
    output = []
    for index, item in enumerate(premises):
        if not isinstance(item, Mapping):
            raise NLICriterionFeatureError("selector premise must be an object")
        group = clean(item.get("group")).upper()
        text = clean(item.get("text"))
        if group not in ALLOWED_PREMISE_GROUPS or not text:
            raise NLICriterionFeatureError("selector premise group or text is invalid")
        if len(text) > MAX_PREMISE_CHARACTERS:
            raise NLICriterionFeatureError("selector premise exceeds the character limit")
        group_counts[group] += 1
        if group_counts[group] > MAX_PREMISES_PER_GROUP:
            raise NLICriterionFeatureError("selector returned too many premises for a group")
        output.append(
            {
                "premise_index": index,
                "page": int(item.get("page") or 0),
                "section": clean(item.get("section")),
                "offset_start": int(item.get("offset_start") or 0),
                "offset_end": int(item.get("offset_end") or 0),
                "text": text,
                "text_sha256": text_sha256(text),
                "concepts": [clean(value) for value in item.get("concepts") or []],
                "roles": [clean(value) for value in item.get("roles") or []],
                "group": group,
                "score": float(item.get("score") or 0.0),
                "evidence_mode": clean(item.get("evidence_mode")),
            }
        )
    return output


def _candidate_spec(
    *,
    matrix_rows: int,
    alignment_sha256: str,
    selection_sha256: str,
    hypotheses_sha256: str,
    model_snapshot_sha256: str,
    feature_names_sha256: str,
    input_hashes: Mapping[str, str],
    selector_info: Mapping[str, Any],
) -> dict[str, Any]:
    return {
        "schema_version": "r3_nli_candidate_spec_v1",
        "status": "FROZEN_BEFORE_INFERENCE_AND_LABEL_ACCESS",
        "alignment_sha256": alignment_sha256,
        "selection_sha256": selection_sha256,
        "hypotheses_sha256": hypotheses_sha256,
        "model_snapshot_sha256": model_snapshot_sha256,
        "feature_names": list(FEATURE_NAMES),
        "feature_names_sha256": feature_names_sha256,
        "expected_matrix_shape": [matrix_rows, len(FEATURE_NAMES)],
        "inference": {
            "device": "cpu",
            "model_mode": "eval",
            "torch_context": "inference_mode",
            "pair_order": ["premise", "hypothesis"],
            "truncation": "only_first",
            "max_length": MAX_LENGTH,
            "batch_size": BATCH_SIZE,
            "label_order": ["contradiction", "entailment", "neutral"],
        },
        "feature_ablation_candidates": {
            "base_v4": {"nli_columns": []},
            "flags_only": {
                "nli_columns": [0, 1, 2],
                "feature_names": list(FEATURE_NAMES[:3]),
            },
            "full_nli": {"nli_columns": "all"},
        },
        "model_candidates": {
            "random_forest": {
                "n_estimators": 600,
                "max_features": 0.5,
                "min_samples_leaf": 8,
                "n_jobs": -1,
                "random_state": 35_720_260,
            },
            "logistic": {
                "C": 1.0,
                "standardize": True,
                "solver": "liblinear",
                "random_state": 35_720_260,
            },
        },
        "zero_shot_candidates": {
            "method_only": {
                "feature": "nli_method_only_support",
                "threshold": ENTAILMENT_THRESHOLD,
            },
            "practice_chain": {
                "feature": "nli_practice_chain_min_support",
                "threshold": ENTAILMENT_THRESHOLD,
            },
            "normative_chain": {
                "feature": "nli_normative_chain_min_support",
                "threshold": ENTAILMENT_THRESHOLD,
            },
        },
        "compatibility_alias": {
            "nli_criterion_chain_min_support": "nli_practice_chain_min_support"
        },
        "fixed_decision_threshold": ENTAILMENT_THRESHOLD,
        "input_sha256": dict(input_hashes),
        "selector": dict(selector_info),
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
    }


def _hypothesis_statistics(probabilities: np.ndarray) -> dict[str, float]:
    if len(probabilities) == 0:
        return {
            "max_entailment": 0.0,
            "top_mean_entailment": 0.0,
            "max_contradiction": 0.0,
            "max_entailment_margin": 0.0,
            "entailment_fraction_ge_0_5": 0.0,
        }
    entailment = probabilities[:, 1]
    contradiction = probabilities[:, 0]
    neutral = probabilities[:, 2]
    top = np.sort(entailment)[-min(3, len(entailment)) :]
    margin = entailment - np.maximum(contradiction, neutral)
    return {
        "max_entailment": float(np.max(entailment)),
        "top_mean_entailment": float(np.mean(top)),
        "max_contradiction": float(np.max(contradiction)),
        "max_entailment_margin": float(np.max(margin)),
        "entailment_fraction_ge_0_5": float(
            np.mean(entailment >= ENTAILMENT_THRESHOLD)
        ),
    }


def _row_features(
    premise_groups: Mapping[str, int],
    hypothesis_probabilities: Mapping[str, np.ndarray],
    premise_hypothesis_probabilities: Mapping[int, Mapping[str, np.ndarray]],
    *,
    source_available: bool,
) -> dict[str, float]:
    premise_count = sum(premise_groups.values())
    values: dict[str, float] = {
        "nli_pilot_selected": 1.0,
        "nli_source_available": float(source_available),
        "nli_retrieval_available": float(premise_count > 0),
        "nli_method_diagnostic_premise_count_log": math.log1p(
            premise_groups.get("METHOD_DIAGNOSTIC", 0)
        ),
        "nli_identify_premise_count_log": math.log1p(
            premise_groups.get("IDENTIFY", 0)
        ),
        "nli_significance_counter_premise_count_log": math.log1p(
            premise_groups.get("SIGNIFICANCE_COUNTER", 0)
        ),
    }
    stats: dict[str, dict[str, float]] = {}
    for item in HYPOTHESES:
        hypothesis_id = item["hypothesis_id"]
        stats[hypothesis_id] = _hypothesis_statistics(
            hypothesis_probabilities.get(hypothesis_id, np.zeros((0, 3)))
        )
        for metric, value in stats[hypothesis_id].items():
            values[f"nli_{hypothesis_id}_{metric}"] = value
    negative = max(
        stats["external"]["max_entailment"],
        stats["not_performed"]["max_entailment"],
    )

    def entailment(scores: Mapping[str, np.ndarray], hypothesis_id: str) -> float:
        probability = scores.get(hypothesis_id)
        if probability is None:
            return 0.0
        values = np.asarray(probability, dtype=np.float64).reshape(-1)
        if values.shape != (3,):
            raise AssertionError("same-premise NLI probability must have three labels")
        return float(values[1])

    scope_supports: list[float] = []
    method_supports: list[float] = []
    identify_supports: list[float] = []
    significance_supports: list[float] = []
    for scores in premise_hypothesis_probabilities.values():
        scope_gate = max(
            0.0,
            entailment(scores, "scope") - entailment(scores, "external"),
        )
        scope_supports.append(scope_gate)
        if "method" in scores:
            method_supports.append(
                min(
                    scope_gate,
                    entailment(scores, "method"),
                    1.0 - entailment(scores, "not_performed"),
                )
            )
        if "identified" in scores or "inconclusive" in scores:
            identify_supports.append(
                min(
                    scope_gate,
                    max(
                        entailment(scores, "identified"),
                        entailment(scores, "inconclusive"),
                    ),
                )
            )
        if "significance" in scores:
            significance_supports.append(
                min(scope_gate, entailment(scores, "significance"))
            )

    scope_support = max(scope_supports, default=0.0)
    required = [
        max(method_supports, default=0.0),
        max(identify_supports, default=0.0),
        max(significance_supports, default=0.0),
    ]
    values.update(
        {
            "nli_scope_support": scope_support,
            "nli_negative_control_max_entailment": negative,
            "nli_method_scope_gated_support": required[0],
            "nli_identify_scope_gated_support": required[1],
            "nli_significance_scope_gated_support": required[2],
            "nli_method_only_support": required[0],
            "nli_practice_chain_min_support": min(required[:2]),
            "nli_normative_chain_min_support": min(required),
            "nli_criterion_chain_min_support": min(required[:2]),
        }
    )
    if set(values) != set(FEATURE_NAMES):
        raise AssertionError("NLI feature map differs from the frozen schema")
    return values


def build_nli_criterion_features(
    *,
    identity_path: Path = DEFAULT_IDENTITY,
    page_cache_path: Path = DEFAULT_PAGE_CACHE,
    fold_manifest_path: Path = DEFAULT_FOLD_MANIFEST,
    concept_snapshot_path: Path = DEFAULT_CONCEPT_SNAPSHOT,
    model_snapshot_path: Path = DEFAULT_MODEL_SNAPSHOT,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    selector: Optional[Callable[..., Mapping[str, Any]]] = None,
    predictor: Optional[
        Callable[[Sequence[Mapping[str, Any]], Path, Mapping[str, Any]], np.ndarray]
    ] = None,
    expected_rows: Optional[int] = EXPECTED_ROWS,
) -> dict[str, Path]:
    paths = {
        "identity": identity_path.expanduser().resolve(),
        "page_cache": page_cache_path.expanduser().resolve(),
        "fold_manifest": fold_manifest_path.expanduser().resolve(),
        "concept_snapshot": concept_snapshot_path.expanduser().resolve(),
        "model_snapshot": model_snapshot_path.expanduser().resolve(),
    }
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = {
        "selection": output_dir / SELECTION_NAME,
        "hypotheses": output_dir / HYPOTHESES_NAME,
        "preinference": output_dir / PREINFERENCE_NAME,
        "features": output_dir / FEATURE_NAME,
        "manifest": output_dir / MANIFEST_NAME,
        "candidate": output_dir / CANDIDATE_NAME,
        "evidence": output_dir / EVIDENCE_NAME,
        "audit": output_dir / AUDIT_NAME,
    }
    if any(path in paths.values() for path in outputs.values()):
        raise NLICriterionFeatureError("NLI outputs must not overwrite inputs")

    identity = load_identity(paths["identity"])
    records, fold_ids, fold_manifest = load_v4_alignment(
        identity, paths["fold_manifest"]
    )
    if expected_rows is not None and len(records) != expected_rows:
        raise NLICriterionFeatureError(
            f"V4 alignment has {len(records)} rows; expected {expected_rows}"
        )
    pilot_mask, selection = select_single64(records, fold_ids)
    selection.update(
        {
            "alignment_sha256": trace_freeze._record_alignment_sha256(records),
            "fold_id_sha256": clean(fold_manifest.get("fold_id_sha256")),
            "inputs": {
                "identity": {
                    "path": str(paths["identity"]),
                    "sha256": file_sha256(paths["identity"]),
                },
                "fold_manifest": {
                    "path": str(paths["fold_manifest"]),
                    "sha256": file_sha256(paths["fold_manifest"]),
                },
            },
        }
    )
    atomic_json(outputs["selection"], selection)

    hypotheses = {
        "schema_version": "r3_nli_criterion_hypotheses_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "criterion_groups": ["method", "identify", "significance"],
        "hypotheses": list(HYPOTHESES),
        "hypotheses_sha256": json_sha256(list(HYPOTHESES)),
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
    }
    atomic_json(outputs["hypotheses"], hypotheses)

    concepts, concept_payload = load_concepts(paths["concept_snapshot"])
    model_info = inspect_model_snapshot(paths["model_snapshot"])
    selector_callable, selector_info = _resolve_selector(selector)
    input_hashes = {
        "identity": file_sha256(paths["identity"]),
        "page_cache": file_sha256(paths["page_cache"]),
        "fold_manifest": file_sha256(paths["fold_manifest"]),
        "concept_snapshot": file_sha256(paths["concept_snapshot"]),
    }
    spec = _candidate_spec(
        matrix_rows=len(records),
        alignment_sha256=selection["alignment_sha256"],
        selection_sha256=selection["selection_sha256"],
        hypotheses_sha256=hypotheses["hypotheses_sha256"],
        model_snapshot_sha256=model_info["snapshot_sha256"],
        feature_names_sha256=json_sha256(list(FEATURE_NAMES)),
        input_hashes=input_hashes,
        selector_info=selector_info,
    )
    preinference = {
        "schema_version": "r3_nli_preinference_freeze_v1",
        "status": "FROZEN_BEFORE_INFERENCE_AND_LABEL_ACCESS",
        "candidate_spec": spec,
        "candidate_spec_sha256": json_sha256(spec),
        "selection": {
            "path": str(outputs["selection"]),
            "sha256": file_sha256(outputs["selection"]),
            "selection_sha256": selection["selection_sha256"],
        },
        "hypotheses": {
            "path": str(outputs["hypotheses"]),
            "sha256": file_sha256(outputs["hypotheses"]),
            "hypotheses_sha256": hypotheses["hypotheses_sha256"],
        },
        "model_snapshot": model_info,
        "builder": {
            "path": str(Path(__file__).resolve()),
            "sha256": file_sha256(Path(__file__).resolve()),
        },
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
    }
    atomic_json(outputs["preinference"], preinference)

    selected_refnos = {
        records[index]["REFNO"] for index in np.flatnonzero(pilot_mask)
    }
    pages = trace_freeze.load_sanitized_page_cache(
        paths["page_cache"], selected_refnos
    )
    pairs: list[dict[str, Any]] = []
    row_premises: dict[int, list[dict[str, Any]]] = {}
    premise_group_counts: Counter[str] = Counter()
    for row_index in np.flatnonzero(pilot_mask):
        record = records[int(row_index)]
        page_rows = pages[record["REFNO"]]["pages"]
        selected_payload = selector_callable(
            page_rows,
            concepts,
            record["ROCKNAME"],
            max_premises_per_group=MAX_PREMISES_PER_GROUP,
            max_premise_characters=MAX_PREMISE_CHARACTERS,
        )
        premises = _validate_premise_selection(selected_payload, record["REFNO"])
        row_premises[int(row_index)] = premises
        premise_group_counts.update(item["group"] for item in premises)
        for hypothesis in HYPOTHESES:
            permitted = hypothesis["premise_groups"]
            applicable = (
                premises
                if permitted == "all"
                else [item for item in premises if item["group"] in permitted]
            )
            rendered = hypothesis["template"].format(target=record["ROCKNAME"])
            for premise in applicable:
                pairs.append(
                    {
                        "row_index": int(row_index),
                        "REFNO": record["REFNO"],
                        "RESULTNO": record["RESULTNO"],
                        "premise_index": premise["premise_index"],
                        "hypothesis_id": hypothesis["hypothesis_id"],
                        "hypothesis": rendered,
                        "premise": premise["text"],
                        "premise_metadata": premise,
                    }
                )

    predictor_callable = predictor or _real_predictor
    probabilities = np.asarray(
        predictor_callable(pairs, paths["model_snapshot"], spec["inference"]),
        dtype=np.float64,
    )
    if (
        probabilities.shape != (len(pairs), 3)
        or not np.isfinite(probabilities).all()
        or np.any(probabilities < 0.0)
        or np.any(probabilities > 1.0)
        or not np.allclose(probabilities.sum(axis=1), 1.0, atol=1e-5)
    ):
        raise NLICriterionFeatureError("NLI predictor returned invalid probabilities")

    probability_groups: dict[tuple[int, int, str], np.ndarray] = {}
    evidence_rows: list[dict[str, Any]] = []
    for pair_index, (pair, values) in enumerate(zip(pairs, probabilities)):
        probability_key = (
            pair["row_index"],
            pair["premise_index"],
            pair["hypothesis_id"],
        )
        if probability_key in probability_groups:
            raise AssertionError("duplicate same-premise NLI probability")
        probability_groups[probability_key] = values
        premise = pair["premise_metadata"]
        evidence_rows.append(
            {
                "pair_index": pair_index,
                "row_index": pair["row_index"],
                "REFNO": pair["REFNO"],
                "RESULTNO": pair["RESULTNO"],
                "premise_index": pair["premise_index"],
                "hypothesis_id": pair["hypothesis_id"],
                "hypothesis_sha256": text_sha256(pair["hypothesis"]),
                "page": premise["page"],
                "section": premise["section"],
                "offset_start": premise["offset_start"],
                "offset_end": premise["offset_end"],
                "premise_text": premise["text"],
                "text_sha256": premise["text_sha256"],
                "premise_group": premise["group"],
                "selector_score": premise["score"],
                "probabilities": {
                    "contradiction": round(float(values[0]), 8),
                    "entailment": round(float(values[1]), 8),
                    "neutral": round(float(values[2]), 8),
                },
            }
        )

    matrix = np.zeros((len(records), len(FEATURE_NAMES)), dtype=np.float32)
    for row_index in np.flatnonzero(pilot_mask):
        row_index = int(row_index)
        premises = row_premises[row_index]
        group_counts = Counter(item["group"] for item in premises)
        premise_hypothesis_values: dict[int, dict[str, np.ndarray]] = {}
        for premise in premises:
            premise_index = int(premise["premise_index"])
            premise_hypothesis_values[premise_index] = {
                hypothesis["hypothesis_id"]: probability_groups[
                    (row_index, premise_index, hypothesis["hypothesis_id"])
                ]
                for hypothesis in HYPOTHESES
                if (row_index, premise_index, hypothesis["hypothesis_id"])
                in probability_groups
            }
        hypothesis_values = {}
        for hypothesis in HYPOTHESES:
            hypothesis_id = hypothesis["hypothesis_id"]
            hypothesis_values[hypothesis_id] = np.asarray(
                [
                    premise_hypothesis_values[int(premise["premise_index"])][
                        hypothesis_id
                    ]
                    for premise in premises
                    if hypothesis_id
                    in premise_hypothesis_values[int(premise["premise_index"])]
                ],
                dtype=np.float64,
            ).reshape((-1, 3))
        values = _row_features(
            group_counts,
            hypothesis_values,
            premise_hypothesis_values,
            source_available=any(clean(page.get("text")) for page in pages[records[row_index]["REFNO"]]["pages"]),
        )
        matrix[row_index] = [values[name] for name in FEATURE_NAMES]
    if (
        matrix.shape != (len(records), len(FEATURE_NAMES))
        or not np.isfinite(matrix).all()
        or np.any(matrix[~pilot_mask] != 0.0)
        or not np.all(matrix[pilot_mask, 0] == 1.0)
    ):
        raise AssertionError("aligned NLI matrix is invalid")
    atomic_npz(outputs["features"], nli=matrix)
    atomic_jsonl(outputs["evidence"], evidence_rows)

    matrix_hash = trace_freeze.array_sha256(matrix)
    candidate = {
        **spec,
        "schema_version": "r3_nli_candidate_freeze_v1",
        "status": "FROZEN_BEFORE_LABEL_ACCESS",
        "preinference_candidate_spec_sha256": preinference["candidate_spec_sha256"],
        "preinference_freeze": {
            "path": str(outputs["preinference"]),
            "sha256": file_sha256(outputs["preinference"]),
        },
        "matrix_sha256": matrix_hash,
        "evidence_sha256": file_sha256(outputs["evidence"]),
        "finalized_after_label_blind_inference": True,
    }
    comparable = dict(candidate)
    for name in (
        "schema_version",
        "status",
        "preinference_candidate_spec_sha256",
        "preinference_freeze",
        "matrix_sha256",
        "evidence_sha256",
        "finalized_after_label_blind_inference",
    ):
        comparable.pop(name, None)
    expected_comparable = dict(spec)
    expected_comparable.pop("schema_version", None)
    expected_comparable.pop("status", None)
    if comparable != expected_comparable:
        raise AssertionError("final candidate changed the predeclared candidate spec")
    atomic_json(outputs["candidate"], candidate)

    audit = {
        "schema_version": "r3_nli_criterion_feature_audit_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "rows": len(records),
        "features": len(FEATURE_NAMES),
        "pilot_rows": int(pilot_mask.sum()),
        "source_available_rows": int(np.sum(matrix[:, 1] > 0.5)),
        "retrieval_available_rows": int(np.sum(matrix[:, 2] > 0.5)),
        "premise_group_counts": dict(sorted(premise_group_counts.items())),
        "nli_pairs": len(pairs),
        "evidence_rows": len(evidence_rows),
        "nonpilot_nonzero_values": int(np.count_nonzero(matrix[~pilot_mask])),
        "candidate_frozen_before_inference": True,
        "preinference_candidate_spec_sha256": preinference["candidate_spec_sha256"],
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
    }
    atomic_json(outputs["audit"], audit)

    manifest = {
        "schema_version": "r3_nli_criterion_features_v1",
        "status": "FROZEN_WITHOUT_LABEL_ACCESS",
        "rows": len(records),
        "features": len(FEATURE_NAMES),
        "feature_names": list(FEATURE_NAMES),
        "feature_names_sha256": json_sha256(list(FEATURE_NAMES)),
        "alignment_sha256": selection["alignment_sha256"],
        "matrix_sha256": matrix_hash,
        "selection_sha256": selection["selection_sha256"],
        "hypotheses_sha256": hypotheses["hypotheses_sha256"],
        "model_snapshot_sha256": model_info["snapshot_sha256"],
        "contains_labels": False,
        "runtime_boundary": {name: False for name in BOUNDARY_FALSE_FIELDS},
        "inputs": {
            name: {"path": str(paths[name]), "sha256": input_hashes[name]}
            for name in ("identity", "page_cache", "fold_manifest", "concept_snapshot")
        }
        | {"model_snapshot": model_info},
        "outputs": {
            name: {"path": str(outputs[name]), "sha256": file_sha256(outputs[name])}
            for name in (
                "selection",
                "hypotheses",
                "preinference",
                "features",
                "candidate",
                "evidence",
                "audit",
            )
        },
    }
    atomic_json(outputs["manifest"], manifest)
    return outputs


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--fold-manifest", type=Path, default=DEFAULT_FOLD_MANIFEST)
    parser.add_argument("--concept-snapshot", type=Path, default=DEFAULT_CONCEPT_SNAPSHOT)
    parser.add_argument("--model-snapshot", type=Path, default=DEFAULT_MODEL_SNAPSHOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args(argv)
    outputs = build_nli_criterion_features(
        identity_path=args.identity,
        page_cache_path=args.page_cache,
        fold_manifest_path=args.fold_manifest,
        concept_snapshot_path=args.concept_snapshot,
        model_snapshot_path=args.model_snapshot,
        output_dir=args.output_dir,
    )
    manifest = json.loads(outputs["manifest"].read_text(encoding="utf-8"))
    audit = json.loads(outputs["audit"].read_text(encoding="utf-8"))
    print(
        "R3 NLI criterion features frozen: "
        f"rows={manifest['rows']} pilot={audit['pilot_rows']} "
        f"pairs={audit['nli_pairs']} output={outputs['manifest']}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "FEATURE_NAMES",
    "FOLD_QUOTAS",
    "HYPOTHESES",
    "NLICriterionFeatureError",
    "build_nli_criterion_features",
    "inspect_model_snapshot",
    "load_identity",
    "load_v4_alignment",
    "select_single64",
]
