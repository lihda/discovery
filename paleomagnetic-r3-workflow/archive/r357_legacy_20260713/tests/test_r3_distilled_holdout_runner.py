from __future__ import annotations

from argparse import Namespace
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import scripts.run_r3_distilled_holdout as runner


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def split_refnos(count: int = 2) -> tuple[list[str], list[str]]:
    holdout: list[str] = []
    training: list[str] = []
    candidate = 1
    while len(holdout) < count or len(training) < count:
        refno = str(candidate)
        bucket = runner.holdout_bucket(refno)
        destination = (
            holdout if bucket < runner.DEFAULT_HOLDOUT_FRACTION else training
        )
        if len(destination) < count:
            destination.append(refno)
        candidate += 1
    return holdout, training


def write_identity(path: Path, refnos: list[str]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(runner.IDENTITY_COLUMNS)
    for index, refno in enumerate(refnos, start=1):
        worksheet.append((refno, f"result-{index}", f"Rock {index}"))
    workbook.save(path)


def write_cache(
    path: Path,
    refnos: list[str],
    text: str = "IRM acquisition curves identify magnetite as the remanence carrier.",
) -> None:
    records = []
    for refno in refnos:
        text_sha256 = hashlib.sha256(text.encode("utf-8")).hexdigest()
        cache_sha256 = hashlib.sha256(text_sha256.encode("ascii")).hexdigest()
        records.append(
            {
                "REFNO": refno,
                "source": {
                    "refno": refno,
                    "source_kind": "synthetic_original_paper",
                    "parsed_path": None,
                    "pdf_path": None,
                    "page_count": 1,
                    "text_length": len(text),
                    "cache_sha256": cache_sha256,
                    "source_quality_flags": [],
                },
                "pages": [
                    {
                        "page": 1,
                        "source": "synthetic_original_paper",
                        "text": text,
                        "text_sha256": text_sha256,
                        "section_hints": ["Rock magnetic results"],
                    }
                ],
            }
        )
    path.write_text(
        "".join(json.dumps(record) + "\n" for record in records),
        encoding="utf-8",
    )


def write_artifact(path: Path, training_count: int, holdout_count: int) -> None:
    artifact = {
        "schema_version": runner.EXPECTED_ARTIFACT_SCHEMA,
        "artifact_class": runner.EXPECTED_ARTIFACT_CLASS,
        "profile": runner.EXPECTED_ARTIFACT_PROFILE,
        "provenance": {
            "phase": "offline_training_only",
            "label_supervision": True,
            "supervision_scope": "pre_registered_training_partition_only",
            "raw_field_values_retained": False,
            "source_paper_quotes_retained": False,
            "row_identity_mapping_retained": False,
        },
        "split": {
            "unit": "paper_identity",
            "salt": runner.DEFAULT_SALT,
            "holdout_fraction": runner.DEFAULT_HOLDOUT_FRACTION,
            "held_out_supervision_accessed": False,
            "identity_counts": {
                "training": {
                    "rows": training_count,
                    "paper_groups": training_count,
                },
                "holdout": {
                    "rows": holdout_count,
                    "paper_groups": holdout_count,
                },
            },
        },
        "runtime_policy": {
            "allowed_scoring_sources": [
                "identity manifest",
                "original paper page text",
                "this frozen aggregate artifact",
            ],
            "source_metadata_table_access": False,
            "supervised_field_access": False,
            "previous_prediction_access": False,
            "contains_record_answer_mapping": False,
        },
        "method_taxonomy": [{"concept_id": "irm"}],
        "diagnostic_taxonomy": [{"concept_id": "saturation"}],
        "carrier_taxonomy": [{"concept_id": "magnetite"}],
        "decision_policy": {"missing_or_ambiguous_source": "ABSTAIN"},
    }
    path.write_text(json.dumps(artifact), encoding="utf-8")


def read_identity_rows(path: Path) -> list[tuple[str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            tuple(str(value) for value in row)
            for row in workbook.active.iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()


class R3DistilledHoldoutRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.holdout, self.training = split_refnos()
        self.all_refnos = self.holdout + self.training
        self.identity = self.root / "identity.xlsx"
        self.cache = self.root / "page_cache.jsonl"
        self.artifact = self.root / "artifact.json"
        self.labels = self.root / "labels.xlsx"
        write_identity(self.identity, self.all_refnos)
        write_cache(self.cache, self.all_refnos)
        write_artifact(self.artifact, len(self.training), len(self.holdout))
        self.labels.write_bytes(b"labels are only for the mocked post-scoring step")

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def args(
        self,
        output: Path,
        *,
        evaluate: bool = False,
        target_scope_policy: str = runner.DEFAULT_TARGET_SCOPE_POLICY,
    ) -> Namespace:
        return Namespace(
            identity_manifest=self.identity,
            page_cache=self.cache,
            artifact=self.artifact,
            expected_artifact_sha256=sha256(self.artifact),
            output_dir=output,
            evaluate_labels=self.labels if evaluate else None,
            evaluation_output=None,
            target_scope_policy=target_scope_policy,
        )

    def fake_guarded_builder(
        self,
        script: Path,
        script_args: list[str],
        allowed_inputs: tuple[Path, ...],
        write_root: Path,
    ) -> None:
        self.assertEqual(script, runner.BUILDER)
        self.assertEqual(len(allowed_inputs), 3)
        holdout_identity, holdout_cache, artifact = allowed_inputs
        self.assertEqual(artifact, self.artifact.resolve())
        self.assertNotIn(self.identity.resolve(), allowed_inputs)
        self.assertNotIn(self.cache.resolve(), allowed_inputs)
        self.assertNotIn(self.labels.resolve(), allowed_inputs)
        self.assertEqual({row[0] for row in read_identity_rows(holdout_identity)}, set(self.holdout))
        cache_refnos = {
            json.loads(line)["REFNO"]
            for line in holdout_cache.read_text(encoding="utf-8").splitlines()
            if line.strip()
        }
        self.assertEqual(cache_refnos, set(self.holdout))

        arguments = dict(zip(script_args[::2], script_args[1::2]))
        self.assertIn(arguments["--target-scope-policy"], runner.TARGET_SCOPE_POLICIES)
        output = Path(arguments["--output"])
        summary = Path(arguments["--summary-json"])
        prediction_rows = []
        for refno, resultno, rockname in read_identity_rows(holdout_identity):
            prediction_rows.append(
                {
                    "REFNO": refno,
                    "RESULTNO": resultno,
                    "ROCKNAME": rockname,
                    "R3": 1,
                    "OUTCOME": "POSITIVE",
                    "normative_r3": {
                        "policy_decision": 1,
                        "binary_value": 1,
                        "outcome": "POSITIVE",
                    },
                    "rockmag_proxy_r3": {
                        "policy_decision": 1,
                        "binary_value": 1,
                        "outcome": "POSITIVE",
                    },
                    "criterion_practice_conflict": False,
                    "criterion_practice_comparison": {
                        "status": "AGREE",
                        "normative_policy_decision": 1,
                        "rockmag_proxy_policy_decision": 1,
                    },
                    "runtime_scoring_sources": [
                        "strict identity manifest",
                        "original paper page text",
                        "frozen aggregate artifact",
                    ],
                }
            )
        output.write_text(
            "".join(json.dumps(row) + "\n" for row in prediction_rows),
            encoding="utf-8",
        )
        summary.write_text(
            json.dumps(
                {
                    "rows": len(prediction_rows),
                    "retrieval_policy": {
                        "total_limit": runner.FROZEN_RETRIEVAL_LIMIT,
                        "per_category_limit": runner.FROZEN_PER_CATEGORY_LIMIT,
                        "target_scope_policy": arguments["--target-scope-policy"],
                    },
                }
            ),
            encoding="utf-8",
        )

    def test_default_run_scores_only_derived_holdout_and_freezes_hash(self) -> None:
        output = self.root / "run"
        with patch.object(runner, "_run_guarded", side_effect=self.fake_guarded_builder), patch.object(
            runner, "_run_evaluation"
        ) as evaluate:
            report = runner.run_workflow(self.args(output))

        evaluate.assert_not_called()
        predictions = output / "r3_distilled_predictions.jsonl"
        receipt_path = output / "r3_distilled_prediction_hash.json"
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        self.assertEqual(receipt["phase"], "before_optional_label_evaluation")
        self.assertFalse(receipt["labels_opened"])
        self.assertEqual(receipt["prediction"]["sha256"], sha256(predictions))
        self.assertEqual(receipt["prediction"]["rows"], len(self.holdout))
        self.assertEqual(report["evaluation"]["status"], "NOT_RUN")
        self.assertEqual(report["boundary_audit"]["status"], "PASS")
        self.assertFalse(
            report["scoring_contract"]["labels_accepted_by_scoring_stage"]
        )
        self.assertEqual(
            {row[0] for row in read_identity_rows(output / "inputs/holdout_identity_manifest.xlsx")},
            set(self.holdout),
        )
        self.assertEqual(
            report["scoring_contract"]["frozen_retrieval_policy"][
                "target_scope_policy"
            ],
            runner.DEFAULT_TARGET_SCOPE_POLICY,
        )

    def test_sentence_only_policy_is_passed_to_builder_and_frozen(self) -> None:
        output = self.root / "run-sentence-only"
        with patch.object(
            runner, "_run_guarded", side_effect=self.fake_guarded_builder
        ), patch.object(runner, "_run_evaluation") as evaluate:
            report = runner.run_workflow(
                self.args(output, target_scope_policy="sentence_only")
            )

        evaluate.assert_not_called()
        self.assertEqual(
            report["scoring_contract"]["frozen_retrieval_policy"],
            {
                "total_limit": runner.FROZEN_RETRIEVAL_LIMIT,
                "per_category_limit": runner.FROZEN_PER_CATEGORY_LIMIT,
                "target_scope_policy": "sentence_only",
            },
        )

    def test_rejects_builder_summary_with_different_target_scope_policy(self) -> None:
        output = self.root / "run-policy-mismatch"

        def fake_mismatched_builder(*args, **kwargs) -> None:
            self.fake_guarded_builder(*args, **kwargs)
            script_args = args[1]
            arguments = dict(zip(script_args[::2], script_args[1::2]))
            summary_path = Path(arguments["--summary-json"])
            summary = json.loads(summary_path.read_text(encoding="utf-8"))
            summary["retrieval_policy"]["target_scope_policy"] = "exclusive_scope"
            summary_path.write_text(json.dumps(summary), encoding="utf-8")

        with patch.object(
            runner, "_run_guarded", side_effect=fake_mismatched_builder
        ):
            with self.assertRaisesRegex(
                runner.R3DistilledRunError, "frozen R3 retrieval policy"
            ):
                runner.run_workflow(
                    self.args(output, target_scope_policy="sentence_only")
                )

    def test_rejects_unknown_target_scope_policy_before_scoring(self) -> None:
        output = self.root / "run-invalid-policy"
        with patch.object(runner, "_run_guarded") as guarded:
            with self.assertRaisesRegex(
                runner.R3DistilledRunError, "sentence_only or exclusive_scope"
            ):
                runner.run_workflow(
                    self.args(output, target_scope_policy="same_page")
                )
        guarded.assert_not_called()

    def test_optional_evaluation_starts_only_after_prediction_receipt(self) -> None:
        output = self.root / "run-with-evaluation"

        def fake_evaluation(predictions: Path, labels: Path, report: Path) -> None:
            receipt = output / "r3_distilled_prediction_hash.json"
            self.assertTrue(receipt.is_file())
            frozen = json.loads(receipt.read_text(encoding="utf-8"))
            self.assertEqual(frozen["prediction"]["sha256"], sha256(predictions))
            manifest = json.loads(
                (output / "r3_distilled_run_manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(
                manifest["evaluation"]["status"], "REQUESTED_POST_SCORING"
            )
            self.assertEqual(labels, self.labels.resolve())
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text(json.dumps({"metrics": {"accuracy": 1.0}}), encoding="utf-8")
            evaluator_receipt = report.with_suffix(
                report.suffix + ".prediction_hashes.json"
            )
            evaluator_receipt.write_text(
                json.dumps({"prediction": {"sha256": sha256(predictions)}}),
                encoding="utf-8",
            )

        with patch.object(runner, "_run_guarded", side_effect=self.fake_guarded_builder), patch.object(
            runner, "_run_evaluation", side_effect=fake_evaluation
        ):
            result = runner.run_workflow(self.args(output, evaluate=True))

        self.assertEqual(result["evaluation"]["status"], "COMPLETED_POST_SCORING")
        self.assertTrue(result["evaluation"]["prediction_hash_fixed_before_evaluation"])

    def test_rejects_artifact_that_accessed_holdout_supervision(self) -> None:
        document = json.loads(self.artifact.read_text(encoding="utf-8"))
        document["split"]["held_out_supervision_accessed"] = True
        self.artifact.write_text(json.dumps(document), encoding="utf-8")
        with self.assertRaisesRegex(runner.R3DistilledRunError, "holdout supervision"):
            runner.validate_distilled_artifact(self.artifact, sha256(self.artifact))

    def test_guarded_worker_blocks_each_forbidden_runtime_file(self) -> None:
        attacker = self.root / "attacker.py"
        attacker.write_text(
            "from pathlib import Path\n"
            "import sys\n"
            "Path(sys.argv[1]).read_bytes()\n",
            encoding="utf-8",
        )
        allowed_inputs = []
        for index, name in enumerate(
            ("holdout_identity.xlsx", "holdout_cache.jsonl", "artifact.json")
        ):
            path = self.root / name
            path.write_text(f"allowed-{index}", encoding="utf-8")
            allowed_inputs.append(path)

        for index, forbidden_name in enumerate(
            (
                "targets.xlsx",
                "PMAGRESULT.xlsx",
                "old_evaluation.json",
                "previous_predictions.jsonl",
            )
        ):
            with self.subTest(forbidden_name=forbidden_name):
                forbidden = self.root / forbidden_name
                forbidden.write_text("forbidden", encoding="utf-8")
                stage = self.root / f"stage-{index}"
                stage.mkdir()
                with self.assertRaises(subprocess.CalledProcessError):
                    runner._run_guarded(
                        attacker,
                        [str(forbidden)],
                        allowed_inputs,
                        stage,
                    )

    def test_real_builder_cli_runs_inside_guard_on_original_text_only(self) -> None:
        identity = self.root / "single_identity.xlsx"
        cache = self.root / "single_cache.jsonl"
        stage = self.root / "real-builder-stage"
        predictions = stage / "predictions.jsonl"
        summary = stage / "summary.json"
        artifact = runner.DEFAULT_ARTIFACT.resolve()
        write_identity(identity, [self.holdout[0]])
        write_cache(
            cache,
            [self.holdout[0]],
            "IRM acquisition curves did not saturate by 1 T.",
        )
        stage.mkdir()

        runner._run_guarded(
            runner.BUILDER,
            [
                "--identity-manifest",
                str(identity),
                "--page-cache",
                str(cache),
                "--artifact",
                str(artifact),
                "--expected-artifact-sha256",
                runner.DEFAULT_ARTIFACT_SHA256,
                "--output",
                str(predictions),
                "--summary-json",
                str(summary),
            ],
            [identity, cache, artifact],
            stage,
        )

        prediction_info = runner.validate_predictions(predictions, identity)
        boundary = runner.audit_scoring_boundary([predictions], [])
        self.assertEqual(prediction_info["rows"], 1)
        self.assertTrue(boundary["passed"], boundary["violations"])
        builder_summary = json.loads(summary.read_text(encoding="utf-8"))
        self.assertFalse(builder_summary["runtime_boundary"]["labels_read"])
        prediction = json.loads(predictions.read_text(encoding="utf-8"))
        self.assertEqual(prediction["normative_r3"]["policy_decision"], "ABSTAIN")
        self.assertEqual(prediction["rockmag_proxy_r3"]["policy_decision"], 1)
        self.assertIs(prediction["criterion_practice_conflict"], True)
        self.assertEqual(
            prediction["criterion_practice_comparison"]["status"],
            "INCOMPLETE_NORMATIVE_CHAIN",
        )


if __name__ == "__main__":
    unittest.main()
