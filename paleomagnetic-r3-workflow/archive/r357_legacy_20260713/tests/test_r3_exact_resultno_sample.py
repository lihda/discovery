from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from scripts.build_r3_exact_resultno_sample import build_exact_sample


class R3ExactResultnoSampleTest(unittest.TestCase):
    def test_exact_row_count_uses_only_local_pdf_refnos(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            identity = root / "identity.xlsx"
            pdf_dir = root / "pdfs"
            pdf_dir.mkdir()
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("REFNO", "RESULTNO", "ROCKNAME"))
            resultno = 1
            for refno, size in enumerate((1, 2, 3, 1, 2, 3, 1, 2), start=1):
                for target in range(size):
                    sheet.append((str(refno), str(resultno), f"rock {refno}-{target}"))
                    resultno += 1
                (pdf_dir / f"{refno}.pdf").write_bytes(b"%PDF-test")
            workbook.save(identity)
            output = root / "sample.xlsx"

            audit = build_exact_sample(
                identity, pdf_dir, output, target_rows=10, salt="fixed-test"
            )

            self.assertEqual(audit["output"]["rows"], 10)
            self.assertFalse(audit["selection"]["labels_read"])
            self.assertFalse(audit["selection"]["pdf_text_read"])
            workbook_out = load_workbook(output, read_only=True)
            try:
                values = list(
                    workbook_out.active.iter_rows(min_row=2, values_only=True)
                )
            finally:
                workbook_out.close()
            self.assertEqual(len(values), 10)
            self.assertTrue(all((pdf_dir / f"{row[0]}.pdf").is_file() for row in values))
            self.assertEqual(
                audit["selection"]["required_evaluation_grouping"],
                "all selected rows sharing REFNO must remain in the same fold",
            )

    def test_selection_is_deterministic(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            identity = root / "identity.xlsx"
            pdf_dir = root / "pdfs"
            pdf_dir.mkdir()
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("REFNO", "RESULTNO", "ROCKNAME"))
            for refno in range(1, 13):
                sheet.append((str(refno), str(refno), f"rock {refno}"))
                (pdf_dir / f"{refno}.pdf").write_bytes(b"%PDF-test")
            workbook.save(identity)

            first = build_exact_sample(
                identity, pdf_dir, root / "a.xlsx", target_rows=7, salt="same"
            )
            second = build_exact_sample(
                identity, pdf_dir, root / "b.xlsx", target_rows=7, salt="same"
            )

            self.assertEqual(
                first["output"]["selected_row_rank_digest_sha256"],
                second["output"]["selected_row_rank_digest_sha256"],
            )


if __name__ == "__main__":
    unittest.main()
