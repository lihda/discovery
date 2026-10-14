#!/usr/bin/env python3
"""Freeze and evaluate a source-derived R3 component manifest.

The ``freeze`` phase accepts only the fixed training identity, original-paper
page cache, local-PDF membership, and frozen knowledge artifacts. It extracts
target aliases, sample/site identifiers, magnetization components, and
same-scope/bridge relations from table rows, captions, and paragraphs. It
writes and hashes the candidate manifest without accepting or opening labels.

The separate ``evaluate`` phase verifies the frozen receipt and every allowed
input hash before opening training R3 cells. It then compares target-local,
article-shared, and manifest-gated propagation with five-fold REFNO GroupKFold.
No PMAGRESULT, ROCKMAG source cell, prior prediction, or prior evaluation is an
input to either phase.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import html
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

try:
    import numpy as np
    from scipy.sparse import csr_matrix, hstack
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.linear_model import LogisticRegression
    from sklearn.metrics import (
        accuracy_score,
        balanced_accuracy_score,
        confusion_matrix,
        f1_score,
    )
    from sklearn.model_selection import GroupKFold
    from sklearn.preprocessing import MaxAbsScaler
except ModuleNotFoundError as exc:  # Extraction remains dependency-light.
    np = None  # type: ignore[assignment]
    _MODEL_DEPENDENCY_ERROR: ModuleNotFoundError | None = exc
else:
    _MODEL_DEPENDENCY_ERROR = None


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_IDENTITY = ROOT / "tmp/r3_distilled_training_preflight/identity.xlsx"
DEFAULT_PAGE_CACHE = ROOT / "tmp/mineru_cache_all_hard_risk.jsonl"
DEFAULT_PDF_DIR = ROOT / "data/raw_pdfs"
DEFAULT_TAXONOMY = ROOT / "data/calibration/r3_rockmag_text_distillation_v1.json"
DEFAULT_POSITIVE_KB = ROOT / "data/calibration/r3_positive_rockmag_expert_kb_v2.json"
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "tmp/r3_component_manifest"

SPLIT_SALT = "r3-rockmag-distillation-v1"
HOLDOUT_FRACTION = 0.20
N_SPLITS = 5
RANDOM_STATE = 357
PROPAGATION_WEIGHT = 0.50
MAX_SCOPE_CHARS = 3500
MAX_TARGET_DOCUMENT_CHARS = 24000
MAX_ARTICLE_DOCUMENT_CHARS = 36000

SPACE_RE = re.compile(r"\s+")
TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9'-]{1,}")
HTML_ROW_RE = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.I | re.S)
HTML_TAG_RE = re.compile(r"<[^>]+>")
CAPTION_RE = re.compile(
    r"(?im)^\s*((?:fig(?:ure)?|table)\s*\.?\s*[A-Z]?\d+[A-Za-z]?[.:]?[^\n]{0,1200})"
)
REFERENCE_HEADING_RE = re.compile(
    r"(?im)^\s*(?:references|bibliography|literature\s+cited)\s*$"
)
COLLECTIVE_RE = re.compile(
    r"\b(?:all|both|these|those|the)\s+(?:samples?|specimens?|sites?|units?|"
    r"rocks?|formations?|lavas?|d[yi]kes?|intrusions?|components?)\b|"
    r"\b(?:samples?|specimens?|sites?)\s+from\s+(?:all|both|each)\b",
    re.I,
)
LABELED_SAMPLE_RE = re.compile(
    r"\b(?:samples?|specimens?|sites?|stations?|localities|locality|sections?|units?)"
    r"\s*(?:nos?\.?|numbers?|codes?)?\s*[:#]?\s*"
    r"([A-Z0-9][A-Z0-9._/-]*(?:\s*(?:,|and|to|through)\s*"
    r"[A-Z0-9][A-Z0-9._/-]*){0,12})",
    re.I,
)
SAMPLE_TOKEN_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:[A-Z]{1,6}[-/]?\d{1,4}[A-Za-z]?|\d{1,3}[A-Za-z]?)(?![A-Za-z0-9])"
)
DESCRIPTIVE_COMPONENT_AFTER_RE = re.compile(
    r"\b(?:component|magneti[sz]ation)\s+(?:designated|labelled|labeled|named|called)?\s*"
    r"(high[- ]temperature|low[- ]temperature|intermediate[- ]temperature|"
    r"high[- ]coercivity|low[- ]coercivity|primary|secondary|normal|reverse[sd]?|"
    r"characteristic)\b",
    re.I,
)
DESCRIPTIVE_COMPONENT_BEFORE_RE = re.compile(
    r"\b(high[- ]temperature|low[- ]temperature|intermediate[- ]temperature|"
    r"high[- ]coercivity|low[- ]coercivity|primary|secondary|normal|reverse[sd]?|"
    r"characteristic)[- ]component\b",
    re.I,
)
CODE_COMPONENT_AFTER_RE = re.compile(
    r"\b(?:component|magnetization|magnetisation)\s+"
    r"(?:designated|labelled|labeled|named|called)?\s*([A-Z][A-Z0-9-]{0,5})\b"
)
CODE_COMPONENT_BEFORE_RE = re.compile(r"\b([A-Z])[- ]component\b")
NAMED_COMPONENT_RE = re.compile(r"\b(?:ChRM|NRM|HTC|LTC|HRC|LRC)\b", re.I)

GENERIC_TARGET_TERMS = {
    "and", "area", "basalt", "basalts", "bed", "beds", "body", "complex",
    "dike", "dikes", "dyke", "dykes", "formation", "formations", "group",
    "intrusion", "intrusions", "lava", "lavas", "member", "members", "of",
    "province", "rock", "rocks", "sample", "samples", "section", "sections",
    "series", "site", "sites", "suite", "the", "tuff", "tuffs", "unit",
    "units", "volcanics",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: Any) -> str:
    return SPACE_RE.sub(" ", html.unescape(clean(value))).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text_sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def holdout_bucket(refno: str) -> float:
    digest = hashlib.sha256(f"{SPLIT_SALT}\0{refno}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def is_training_refno(refno: str) -> bool:
    return holdout_bucket(refno) >= HOLDOUT_FRACTION


def canonical_target(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", clean(value).casefold()).strip()


def _header_indices(header: Sequence[Any], required: Sequence[str]) -> dict[str, int]:
    names = [clean(cell.value if hasattr(cell, "value") else cell) for cell in header]
    output: dict[str, int] = {}
    for name in required:
        matches = [index for index, value in enumerate(names) if value == name]
        if len(matches) != 1:
            raise ValueError(f"expected exactly one {name} column; found {len(matches)}")
        output[name] = matches[0]
    return output


def load_training_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("training identity workbook must contain one worksheet")
        rows = workbook.worksheets[0].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "ROCKNAME"))
        records: list[dict[str, str]] = []
        for row in rows:
            record = {name: clean(row[index].value) for name, index in columns.items()}
            if not any(record.values()):
                continue
            if any(not value for value in record.values()):
                raise ValueError("training identity contains an empty value")
            if not is_training_refno(record["REFNO"]):
                raise ValueError("identity contains a held-out paper assignment")
            records.append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError("training identity is empty")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise ValueError("training identity contains duplicate record identities")
    return records


def local_pdf_refnos(pdf_dir: Path) -> set[str]:
    if not pdf_dir.is_dir():
        raise ValueError(f"PDF directory does not exist: {pdf_dir}")
    return {
        path.stem
        for path in pdf_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".pdf"
    }


def pdf_membership_digest(refnos: Iterable[str]) -> str:
    return text_sha256("\n".join(sorted(set(refnos))) + "\n")


def select_pdf_identity(
    identity: Sequence[Mapping[str, str]], pdf_dir: Path
) -> tuple[list[dict[str, str]], dict[str, Any]]:
    available = local_pdf_refnos(pdf_dir)
    selected = [dict(row) for row in identity if row["REFNO"] in available]
    if not selected:
        raise ValueError("no fixed-training identity has a local PDF")
    selected_refnos = {row["REFNO"] for row in selected}
    skipped_refnos = {row["REFNO"] for row in identity} - selected_refnos
    return selected, {
        "fixed_training_rows": len(identity),
        "fixed_training_paper_groups": len({row["REFNO"] for row in identity}),
        "pdf_only_rows": len(selected),
        "pdf_only_paper_groups": len(selected_refnos),
        "skipped_rows_without_pdf": len(identity) - len(selected),
        "skipped_paper_groups_without_pdf": len(skipped_refnos),
        "pdf_membership_sha256": pdf_membership_digest(selected_refnos),
    }


def load_page_cache(path: Path, refnos: set[str]) -> dict[str, list[dict[str, str]]]:
    papers: dict[str, list[dict[str, str]]] = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            row = json.loads(line)
            refno = clean(row.get("REFNO"))
            if refno not in refnos:
                continue
            if refno in papers:
                raise ValueError(f"duplicate page-cache paper at line {line_number}")
            pages: list[dict[str, str]] = []
            for index, page in enumerate(row.get("pages", []), start=1):
                if not isinstance(page, Mapping):
                    continue
                text = clean(page.get("text"))
                if not text:
                    continue
                pages.append({"page": clean(page.get("page")) or str(index), "text": text})
            papers[refno] = pages
    missing = sorted(refnos - papers.keys())
    empty = sorted(refno for refno in refnos if not papers.get(refno))
    if missing or empty:
        raise ValueError(
            "every PDF-only paper must have nonempty cached source text: "
            f"missing={missing[:10]}, empty={empty[:10]}"
        )
    return papers


def _iter_runtime_patterns(
    value: Any, path: tuple[str, ...] = ()
) -> Iterable[tuple[str, str, str]]:
    if isinstance(value, Mapping):
        patterns = value.get("runtime_original_text_patterns")
        if isinstance(patterns, list) and all(isinstance(item, str) for item in patterns):
            concept_id = clean(value.get("concept_id") or ":".join(path))
            path_text = ":".join(path).casefold()
            category = (
                "method" if "method" in path_text else
                "diagnostic" if "diagnostic" in path_text else
                "carrier" if "carrier" in path_text else
                "other"
            )
            for pattern in patterns:
                yield category, concept_id, pattern
        for key, child in value.items():
            if key != "runtime_original_text_patterns":
                yield from _iter_runtime_patterns(child, (*path, str(key)))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _iter_runtime_patterns(child, (*path, str(index)))


def load_knowledge_patterns(paths: Sequence[Path]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    compiled: list[dict[str, Any]] = []
    seen: set[str] = set()
    per_source: Counter[str] = Counter()
    for path in paths:
        artifact = json.loads(path.read_text(encoding="utf-8"))
        if path.name == "r3_positive_rockmag_expert_kb_v2.json":
            if artifact.get("schema_version") != "r3_positive_rockmag_expert_kb_v2":
                raise ValueError("positive expert KB has an unexpected schema")
            provenance = artifact.get("provenance", {})
            runtime = artifact.get("runtime_policy", {})
            if provenance.get("held_out_labels_opened") is not False:
                raise ValueError("positive expert KB reports held-out label access")
            if provenance.get("held_out_field_cells_opened") is not False:
                raise ValueError("positive expert KB reports held-out field access")
            if runtime.get("contains_record_answer_mapping") is not False:
                raise ValueError("positive expert KB contains an answer mapping")
        for category, concept_id, expression in _iter_runtime_patterns(artifact):
            if expression in seen:
                continue
            seen.add(expression)
            compiled.append(
                {
                    "category": category,
                    "concept_id": concept_id,
                    "pattern_text": expression,
                    "pattern": re.compile(expression, re.I | re.S),
                    "source": path.name,
                }
            )
            per_source[path.name] += 1
    if not compiled:
        raise ValueError("knowledge artifacts contain no runtime original-text patterns")
    return compiled, {
        "concept_pattern_count": len(compiled),
        "patterns_by_source": dict(sorted(per_source.items())),
        "inputs": {
            path.name: {"path": str(path.resolve()), "sha256": file_sha256(path)}
            for path in paths
        },
    }


def _chunk_text(text: str, maximum: int = MAX_SCOPE_CHARS) -> list[str]:
    value = clean(text)
    if len(value) <= maximum:
        return [value] if len(normalized_text(value)) >= 35 else []
    pieces: list[str] = []
    cursor = 0
    while cursor < len(value):
        end = min(len(value), cursor + maximum)
        if end < len(value):
            boundary = max(value.rfind(". ", cursor, end), value.rfind("\n", cursor, end))
            if boundary > cursor + maximum // 2:
                end = boundary + 1
        piece = value[cursor:end].strip()
        if len(normalized_text(piece)) >= 35:
            pieces.append(piece)
        cursor = max(end, cursor + 1)
    return pieces


def extract_scopes(pages: Sequence[Mapping[str, str]]) -> list[dict[str, Any]]:
    """Create deterministic table-row, caption, and paragraph scopes."""

    scopes: list[dict[str, Any]] = []
    seen: set[str] = set()
    references_started = False
    for page_index, page in enumerate(pages):
        raw = clean(page.get("text"))
        if not raw:
            continue
        if page_index >= max(1, len(pages) // 2) and REFERENCE_HEADING_RE.search(raw):
            raw = raw[: REFERENCE_HEADING_RE.search(raw).start()]
            references_started = True
        elif references_started:
            continue

        candidates: list[tuple[str, str]] = []
        for match in HTML_ROW_RE.finditer(raw):
            row_text = HTML_TAG_RE.sub(" ", match.group(1))
            candidates.append(("table_row", row_text))
        for line in raw.splitlines():
            if line.count("|") >= 2 and not re.fullmatch(r"[\s|:-]+", line):
                candidates.append(("table_row", line.replace("|", " ")))
        candidates.extend(("figure_or_table_caption", match.group(1)) for match in CAPTION_RE.finditer(raw))

        without_html = HTML_ROW_RE.sub("\n", raw)
        paragraphs = re.split(r"\n\s*\n+", without_html)
        if len(paragraphs) <= 2:
            lines = [line.strip() for line in without_html.splitlines() if line.strip()]
            paragraphs = []
            buffer: list[str] = []
            length = 0
            for line in lines:
                buffer.append(line)
                length += len(line)
                if length >= 900 or re.search(r"[.!?][\"')\]]?$", line):
                    paragraphs.append(" ".join(buffer))
                    buffer = []
                    length = 0
            if buffer:
                paragraphs.append(" ".join(buffer))
        for paragraph in paragraphs:
            candidates.extend(("paragraph", piece) for piece in _chunk_text(paragraph))

        for kind, candidate in candidates:
            text = normalized_text(candidate)
            minimum_length = 8 if kind == "table_row" else 35
            if len(text) < minimum_length:
                continue
            digest = text_sha256(text)
            if digest in seen:
                continue
            seen.add(digest)
            scopes.append(
                {
                    "scope_id": digest[:24],
                    "text_sha256": digest,
                    "page": clean(page.get("page")) or str(page_index + 1),
                    "kind": kind,
                    "text": text[:MAX_SCOPE_CHARS],
                }
            )
    return scopes


def target_alias_candidates(rockname: str) -> list[str]:
    """Generate identity-derived aliases; source confirmation is stored later."""

    raw = normalized_text(rockname)
    candidates = {raw}
    candidates.update(normalized_text(part) for part in re.split(r"[,;/]", raw))
    candidates.update(
        normalized_text(match)
        for match in re.findall(r"\(([^()]{2,60})\)", raw)
    )
    stripped = re.sub(
        r"\b(?:formation|group|member|series|suite|complex|basalts?|volcanics?|"
        r"d[yi]kes?|tuffs?|granites?)\b",
        " ",
        raw,
        flags=re.I,
    )
    candidates.add(normalized_text(stripped))
    output: list[str] = []
    for candidate in candidates:
        tokens = [token.casefold() for token in TOKEN_RE.findall(candidate)]
        informative = [token for token in tokens if token not in GENERIC_TARGET_TERMS]
        if not candidate or len(candidate) < 3:
            continue
        if not informative and len(tokens) == 1 and raw.casefold() != candidate.casefold():
            continue
        output.append(candidate)
    return sorted(set(output), key=lambda value: (-len(value), value.casefold()))


def alias_pattern(alias: str) -> re.Pattern[str]:
    tokens = TOKEN_RE.findall(alias)
    if not tokens:
        return re.compile(r"(?!x)x")
    return re.compile(
        r"(?<![A-Za-z0-9])" + r"\W+".join(re.escape(token) for token in tokens) + r"(?![A-Za-z0-9])",
        re.I,
    )


def source_acronym_aliases(
    scopes: Sequence[Mapping[str, Any]], base_aliases: Sequence[str]
) -> list[dict[str, str]]:
    output: dict[str, dict[str, str]] = {}
    base_patterns = [(alias, alias_pattern(alias)) for alias in base_aliases]
    acronym_re = re.compile(
        r"\b([A-Z][A-Za-z'-]+(?:\s+[A-Z][A-Za-z'-]+){0,7})\s*\(([A-Z][A-Z0-9-]{1,8})\)"
    )
    for scope in scopes:
        text = str(scope["text"])
        for phrase, acronym in acronym_re.findall(text):
            anchored = any(pattern.search(phrase) or pattern.search(text[max(0, text.find(phrase)-80):text.find(phrase)+len(phrase)+80]) for _, pattern in base_patterns)
            if not anchored:
                continue
            output[acronym.casefold()] = {
                "value": acronym,
                "origin": "source_parenthetical_acronym",
                "scope_id": str(scope["scope_id"]),
            }
    return [output[key] for key in sorted(output)]


def target_hits(text: str, aliases: Sequence[str]) -> list[str]:
    hits = [alias for alias in aliases if alias_pattern(alias).search(text)]
    return sorted(set(hits), key=lambda value: (-len(value), value.casefold()))[:12]


def extract_sample_ids(text: str, kind: str) -> list[str]:
    values: set[str] = set()
    for match in LABELED_SAMPLE_RE.finditer(text):
        segment = match.group(1)[:240]
        values.update(token.upper().replace(" ", "") for token in SAMPLE_TOKEN_RE.findall(segment))
    if kind == "table_row" and re.search(r"\b(?:site|sample|specimen|station|unit)\b", text, re.I):
        values.update(token.upper().replace(" ", "") for token in SAMPLE_TOKEN_RE.findall(text[:300]))
    return sorted(value for value in values if 1 < len(value) <= 16)[:40]


def normalize_component(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")


def extract_components(text: str) -> list[str]:
    values = {
        normalize_component(match)
        for match in DESCRIPTIVE_COMPONENT_AFTER_RE.findall(text)
    }
    values.update(
        normalize_component(match)
        for match in DESCRIPTIVE_COMPONENT_BEFORE_RE.findall(text)
    )
    values.update(
        normalize_component(match) for match in CODE_COMPONENT_AFTER_RE.findall(text)
    )
    values.update(
        normalize_component(match) for match in CODE_COMPONENT_BEFORE_RE.findall(text)
    )
    values.update(normalize_component(match) for match in NAMED_COMPONENT_RE.findall(text))
    return sorted(value for value in values if value)[:30]


def concept_hits(text: str, patterns: Sequence[Mapping[str, Any]]) -> list[str]:
    values = {
        f"{item['category']}:{item['concept_id']}"
        for item in patterns
        if item["pattern"].search(text)
    }
    return sorted(values)


def _bounded_document(scopes: Iterable[Mapping[str, Any]], cap: int) -> str:
    pieces: list[str] = []
    seen: set[str] = set()
    characters = 0
    for scope in scopes:
        scope_id = str(scope["scope_id"])
        if scope_id in seen:
            continue
        seen.add(scope_id)
        text = str(scope["text"])
        remaining = cap - characters
        if remaining <= 0:
            break
        pieces.append(text[:remaining])
        characters += min(len(text), remaining)
    return "\n".join(pieces)


def build_paper_manifest(
    refno: str,
    identity_rows: Sequence[Mapping[str, str]],
    pages: Sequence[Mapping[str, str]],
    knowledge_patterns: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    scopes = extract_scopes(pages)
    targets: dict[str, dict[str, Any]] = {}
    for row in identity_rows:
        canonical = canonical_target(row["ROCKNAME"])
        target = targets.setdefault(
            canonical,
            {
                "canonical_target": canonical,
                "rocknames": set(),
                "resultnos": set(),
                "identity_aliases": set(),
            },
        )
        target["rocknames"].add(row["ROCKNAME"])
        target["resultnos"].add(row["RESULTNO"])
        target["identity_aliases"].update(target_alias_candidates(row["ROCKNAME"]))

    for target in targets.values():
        source_aliases = source_acronym_aliases(scopes, sorted(target["identity_aliases"]))
        target["source_aliases"] = source_aliases
        target["all_aliases"] = sorted(
            target["identity_aliases"] | {item["value"] for item in source_aliases},
            key=lambda value: (-len(value), value.casefold()),
        )

    relevant_scopes: list[dict[str, Any]] = []
    scope_lookup: dict[str, dict[str, Any]] = {}
    for scope in scopes:
        text = str(scope["text"])
        hits = {
            canonical: target_hits(text, target["all_aliases"])
            for canonical, target in targets.items()
        }
        hits = {key: value for key, value in hits.items() if value}
        samples = extract_sample_ids(text, str(scope["kind"]))
        components = extract_components(text)
        concepts = concept_hits(text, knowledge_patterns)
        if not (hits or samples or components or concepts):
            continue
        record = {
            "scope_id": scope["scope_id"],
            "text_sha256": scope["text_sha256"],
            "page": scope["page"],
            "kind": scope["kind"],
            "target_hits": hits,
            "sample_ids": samples,
            "components": components,
            "concepts": concepts,
            "collective_language": bool(COLLECTIVE_RE.search(text)),
            "text": text,
        }
        relevant_scopes.append(record)
        scope_lookup[str(scope["scope_id"])] = record

    sample_scopes: dict[str, set[str]] = defaultdict(set)
    component_scopes: dict[str, set[str]] = defaultdict(set)
    for scope in relevant_scopes:
        for sample in scope["sample_ids"]:
            sample_scopes[sample].add(scope["scope_id"])
        for component in scope["components"]:
            component_scopes[component].add(scope["scope_id"])

    target_output: list[dict[str, Any]] = []
    for canonical, target in sorted(targets.items()):
        direct = [scope for scope in relevant_scopes if canonical in scope["target_hits"]]
        direct_ids = {scope["scope_id"] for scope in direct}
        samples = sorted({value for scope in direct for value in scope["sample_ids"]})
        components = sorted({value for scope in direct for value in scope["components"]})
        direct_concepts = sorted({value for scope in direct for value in scope["concepts"]})
        sample_bridge_ids = {
            scope_id
            for sample in samples
            for scope_id in sample_scopes[sample]
            if scope_lookup[scope_id]["concepts"]
        }
        component_bridge_ids = {
            scope_id
            for component in components
            for scope_id in component_scopes[component]
            if scope_lookup[scope_id]["concepts"]
        }
        bridge_ids = (sample_bridge_ids | component_bridge_ids) - direct_ids
        bridge_scopes = [scope_lookup[scope_id] for scope_id in sorted(bridge_ids)]
        bridged_concepts = sorted(
            {value for scope in bridge_scopes for value in scope["concepts"]}
        )
        collective_direct = any(
            scope["collective_language"] and scope["concepts"] for scope in direct
        )
        eligible = bool(sample_bridge_ids or component_bridge_ids or collective_direct)
        linked_scopes = direct + bridge_scopes
        source_confirmed_identity_aliases = sorted(
            {
                alias
                for scope in direct
                for alias in scope["target_hits"].get(canonical, [])
            }
        )
        target_output.append(
            {
                "canonical_target": canonical,
                "target_id": text_sha256(f"{refno}\0{canonical}")[:20],
                "rocknames": sorted(target["rocknames"]),
                "resultnos": sorted(target["resultnos"]),
                "aliases": {
                    "identity_derived": sorted(target["identity_aliases"]),
                    "source_confirmed_identity": source_confirmed_identity_aliases,
                    "source_derived": target["source_aliases"],
                },
                "sample_ids": samples,
                "components": components,
                "direct_concepts": direct_concepts,
                "bridged_concepts": bridged_concepts,
                "scope_links": {
                    "direct_scope_ids": sorted(direct_ids),
                    "sample_bridge_scope_ids": sorted(sample_bridge_ids),
                    "component_bridge_scope_ids": sorted(component_bridge_ids),
                    "bridged_scope_ids": sorted(bridge_ids),
                },
                "gate": {
                    "sample_bridge": bool(sample_bridge_ids),
                    "component_bridge": bool(component_bridge_ids),
                    "collective_target_scope": collective_direct,
                    "eligible_for_article_propagation": eligible,
                },
                "features": {
                    "direct_scope_count": len(direct_ids),
                    "linked_scope_count": len({scope["scope_id"] for scope in linked_scopes}),
                    "table_scope_count": sum(scope["kind"] == "table_row" for scope in linked_scopes),
                    "caption_scope_count": sum(scope["kind"] == "figure_or_table_caption" for scope in linked_scopes),
                    "sample_id_count": len(samples),
                    "component_count": len(components),
                    "direct_concept_count": len(direct_concepts),
                    "bridged_concept_count": len(bridged_concepts),
                },
                "target_local_document": _bounded_document(direct, MAX_TARGET_DOCUMENT_CHARS),
                "manifest_linked_document": _bounded_document(linked_scopes, MAX_TARGET_DOCUMENT_CHARS),
            }
        )

    concept_scopes = [scope for scope in relevant_scopes if scope["concepts"]]
    shared_scopes = [
        scope
        for scope in concept_scopes
        if len(scope["target_hits"]) != 1 or scope["collective_language"]
    ]
    auditable_scopes = [
        {
            key: value
            for key, value in scope.items()
            if key != "text"
        }
        for scope in relevant_scopes
    ]
    return {
        "refno": refno,
        "paper_text_sha256": text_sha256("\n".join(page["text"] for page in pages)),
        "page_count": len(pages),
        "target_count": len(target_output),
        "article": {
            "relevant_scope_count": len(relevant_scopes),
            "knowledge_scope_count": len(concept_scopes),
            "shared_knowledge_scope_count": len(shared_scopes),
            "collective_knowledge_scope_count": sum(
                scope["collective_language"] for scope in concept_scopes
            ),
            "has_shared_evidence": bool(shared_scopes),
            "concepts": sorted({value for scope in concept_scopes for value in scope["concepts"]}),
            "knowledge_document": _bounded_document(concept_scopes, MAX_ARTICLE_DOCUMENT_CHARS),
        },
        "targets": target_output,
        "source_scopes": auditable_scopes,
    }


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def atomic_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
                handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def atomic_strategy_npz(path: Path, strategies: Mapping[str, Any]) -> None:
    """Persist aligned OOF outputs without labels or identity columns."""

    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        arrays: dict[str, Any] = {}
        for name, probability in sorted(strategies.items()):
            arrays[f"{name}__probability"] = np.asarray(
                probability, dtype=np.float32
            )
            arrays[f"{name}__decision"] = np.asarray(
                probability >= 0.5, dtype=np.uint8
            )
        with os.fdopen(descriptor, "wb") as handle:
            np.savez_compressed(handle, **arrays)
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def freeze_manifest(
    identity_path: Path,
    page_cache_path: Path,
    pdf_dir: Path,
    taxonomy_path: Path,
    positive_kb_path: Path,
    manifest_path: Path,
    receipt_path: Path,
) -> dict[str, Any]:
    """Freeze source relations. This function has no label-source argument."""

    inputs = {
        "identity": identity_path.resolve(),
        "page_cache": page_cache_path.resolve(),
        "taxonomy": taxonomy_path.resolve(),
        "positive_kb": positive_kb_path.resolve(),
    }
    for name, path in inputs.items():
        if not path.is_file():
            raise ValueError(f"missing {name}: {path}")
    identity = load_training_identity(inputs["identity"])
    selected, membership = select_pdf_identity(identity, pdf_dir.resolve())
    by_paper: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in selected:
        by_paper[row["REFNO"]].append(row)
    pages = load_page_cache(inputs["page_cache"], set(by_paper))
    patterns, knowledge_audit = load_knowledge_patterns(
        [inputs["taxonomy"], inputs["positive_kb"]]
    )
    manifests = [
        build_paper_manifest(refno, by_paper[refno], pages[refno], patterns)
        for refno in sorted(by_paper)
    ]
    atomic_jsonl(manifest_path.resolve(), manifests)
    manifest_hash = file_sha256(manifest_path.resolve())
    target_manifests = [target for paper in manifests for target in paper["targets"]]
    receipt = {
        "schema_version": "r3_source_component_manifest_freeze_v1",
        "status": "FROZEN_BEFORE_LABEL_ACCESS",
        "labels_opened": False,
        "manifest": {
            "path": str(manifest_path.resolve()),
            "sha256": manifest_hash,
            "paper_records": len(manifests),
            "target_records": len(target_manifests),
            "identity_rows": sum(len(target["resultnos"]) for target in target_manifests),
        },
        "inputs": {
            name: {"path": str(path), "sha256": file_sha256(path)}
            for name, path in inputs.items()
        },
        "pdf_dir": {
            "path": str(pdf_dir.resolve()),
            "selected_membership_sha256": membership["pdf_membership_sha256"],
        },
        "membership": membership,
        "knowledge": knowledge_audit,
        "extraction": {
            "scope_types": ["table_row", "figure_or_table_caption", "paragraph"],
            "target_aliases": "identity-derived aliases plus source-anchored parenthetical acronyms",
            "bridges": ["target-sample-concept", "target-component-concept"],
            "manifest_gate": "sample bridge, component bridge, or collective target-bound knowledge scope",
            "aggregate": {
                "papers_with_shared_evidence": sum(paper["article"]["has_shared_evidence"] for paper in manifests),
                "targets_with_direct_scope": sum(target["features"]["direct_scope_count"] > 0 for target in target_manifests),
                "targets_with_sample_ids": sum(bool(target["sample_ids"]) for target in target_manifests),
                "targets_with_components": sum(bool(target["components"]) for target in target_manifests),
                "targets_manifest_gate_open": sum(target["gate"]["eligible_for_article_propagation"] for target in target_manifests),
                "source_scope_records": sum(len(paper["source_scopes"]) for paper in manifests),
            },
        },
        "guardrails": {
            "allowed_inputs": [
                "fixed training identity fields",
                "original-paper page cache",
                "local PDF membership",
                "frozen taxonomy",
                "frozen positive expert KB",
            ],
            "forbidden_sources_not_read": [
                "held-out labels",
                "source metadata tables",
                "historical ROCKMAG cells",
                "prior predictions",
                "prior evaluations",
            ],
            "manifest_contains_labels": False,
            "manifest_contains_predictions": False,
        },
        "script": {
            "path": str(Path(__file__).resolve()),
            "sha256": file_sha256(Path(__file__).resolve()),
        },
    }
    atomic_json(receipt_path.resolve(), receipt)
    return receipt


def load_manifest(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            row = json.loads(line)
            if not isinstance(row, dict) or not clean(row.get("refno")):
                raise ValueError(f"invalid manifest row {line_number}")
            rows.append(row)
    if not rows:
        raise ValueError("component manifest is empty")
    if len({row["refno"] for row in rows}) != len(rows):
        raise ValueError("component manifest contains duplicate papers")
    return rows


def verify_frozen_manifest(
    receipt_path: Path,
    manifest_path: Path,
    identity_path: Path,
    page_cache_path: Path,
    pdf_dir: Path,
    taxonomy_path: Path,
    positive_kb_path: Path,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    if receipt.get("status") != "FROZEN_BEFORE_LABEL_ACCESS":
        raise ValueError("manifest receipt is not frozen-before-label-access")
    if receipt.get("labels_opened") is not False:
        raise ValueError("manifest receipt reports label access")
    expected_manifest_hash = receipt.get("manifest", {}).get("sha256")
    if file_sha256(manifest_path) != expected_manifest_hash:
        raise ValueError("component manifest hash does not match frozen receipt")
    current_inputs = {
        "identity": identity_path.resolve(),
        "page_cache": page_cache_path.resolve(),
        "taxonomy": taxonomy_path.resolve(),
        "positive_kb": positive_kb_path.resolve(),
    }
    for name, path in current_inputs.items():
        declared = receipt.get("inputs", {}).get(name, {}).get("sha256")
        if not path.is_file() or file_sha256(path) != declared:
            raise ValueError(f"frozen input changed before evaluation: {name}")
    if file_sha256(Path(__file__).resolve()) != receipt.get("script", {}).get("sha256"):
        raise ValueError("component manifest script changed after freeze")
    identity = load_training_identity(current_inputs["identity"])
    selected, membership = select_pdf_identity(identity, pdf_dir.resolve())
    if membership["pdf_membership_sha256"] != receipt.get("pdf_dir", {}).get("selected_membership_sha256"):
        raise ValueError("local PDF membership changed after freeze")
    manifests = load_manifest(manifest_path)
    manifest_resultnos = {
        resultno
        for paper in manifests
        for target in paper.get("targets", [])
        for resultno in target.get("resultnos", [])
    }
    selected_resultnos = {row["RESULTNO"] for row in selected}
    if manifest_resultnos != selected_resultnos:
        raise ValueError("frozen manifest membership differs from PDF-only identity")
    return receipt, manifests


def load_training_labels(path: Path, permitted_resultnos: set[str]) -> dict[str, int]:
    """Open R3 only for fixed-training, PDF-only identities after freeze verification."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[workbook.sheetnames[0]].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "R3"))
        labels: dict[str, int] = {}
        for row in rows:
            refno = clean(row[columns["REFNO"]].value)
            if not refno or not is_training_refno(refno):
                continue
            resultno = clean(row[columns["RESULTNO"]].value)
            if resultno not in permitted_resultnos:
                continue
            raw = row[columns["R3"]].value
            try:
                label = int(float(raw))
            except (TypeError, ValueError) as exc:
                raise ValueError("non-binary permitted training label") from exc
            if label not in {0, 1} or resultno in labels:
                raise ValueError("invalid or duplicate permitted training label")
            labels[resultno] = label
    finally:
        workbook.close()
    if set(labels) != permitted_resultnos:
        raise ValueError("training label membership does not equal frozen manifest membership")
    return labels


def flatten_manifest(manifests: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for paper in manifests:
        article = paper["article"]
        for target in paper["targets"]:
            aliases = target["aliases"]
            alias_text = " ".join(
                aliases.get("identity_derived", [])
                + aliases.get("source_confirmed_identity", [])
                + [item["value"] for item in aliases.get("source_derived", [])]
            )
            concept_text = " ".join(target["direct_concepts"] + target["bridged_concepts"])
            component_text = " ".join(target["components"] + target["sample_ids"])
            local_document = normalized_text(
                " ".join(target["rocknames"]) + " " + alias_text + " "
                + target["target_local_document"] + " " + concept_text + " " + component_text
            )
            linked_document = normalized_text(
                local_document + " " + target["manifest_linked_document"]
            )
            numeric = target["features"]
            for resultno in target["resultnos"]:
                rows.append(
                    {
                        "refno": paper["refno"],
                        "resultno": resultno,
                        "canonical_target": target["canonical_target"],
                        "local_document": local_document,
                        "linked_document": linked_document,
                        "article_document": article["knowledge_document"],
                        "shared_evidence": bool(article["has_shared_evidence"]),
                        "manifest_gate": bool(target["gate"]["eligible_for_article_propagation"]),
                        "numeric": [
                            numeric["direct_scope_count"],
                            numeric["linked_scope_count"],
                            numeric["table_scope_count"],
                            numeric["caption_scope_count"],
                            numeric["sample_id_count"],
                            numeric["component_count"],
                            numeric["direct_concept_count"],
                            numeric["bridged_concept_count"],
                            article["knowledge_scope_count"],
                            article["shared_knowledge_scope_count"],
                            article["collective_knowledge_scope_count"],
                            paper["target_count"],
                            int(article["has_shared_evidence"]),
                            int(target["gate"]["sample_bridge"]),
                            int(target["gate"]["component_bridge"]),
                            int(target["gate"]["collective_target_scope"]),
                        ],
                    }
                )
    return rows


def require_model_dependencies() -> None:
    if _MODEL_DEPENDENCY_ERROR is not None:
        raise SystemExit(
            "Component-manifest evaluation requires numpy, scipy, and scikit-learn. "
            f"Missing import: {_MODEL_DEPENDENCY_ERROR.name}"
        )


def _fit_text_numeric_model(
    train_documents: Sequence[str],
    test_documents: Sequence[str],
    train_numeric: Any,
    test_numeric: Any,
    train_labels: Any,
    sample_weight: Any,
) -> Any:
    vectorizer = TfidfVectorizer(
        analyzer="word",
        ngram_range=(1, 2),
        min_df=2,
        max_df=0.995,
        max_features=50000,
        sublinear_tf=True,
        strip_accents="unicode",
    )
    text_train = vectorizer.fit_transform(train_documents)
    text_test = vectorizer.transform(test_documents)
    scaler = MaxAbsScaler()
    numeric_train = scaler.fit_transform(csr_matrix(np.log1p(train_numeric)))
    numeric_test = scaler.transform(csr_matrix(np.log1p(test_numeric)))
    features_train = hstack([text_train, numeric_train], format="csr")
    features_test = hstack([text_test, numeric_test], format="csr")
    model = LogisticRegression(
        C=2.0,
        max_iter=1200,
        solver="liblinear",
        class_weight="balanced",
        random_state=RANDOM_STATE,
    )
    model.fit(features_train, train_labels, sample_weight=sample_weight)
    return model.predict_proba(features_test)[:, 1]


def inverse_group_weights(indices: Any, groups: Any) -> Any:
    counts = Counter(str(groups[index]) for index in indices)
    return np.asarray([1.0 / counts[str(groups[index])] for index in indices], dtype=float)


def grouped_oof_probabilities(rows: Sequence[Mapping[str, Any]], labels: Any) -> tuple[dict[str, Any], list[Any]]:
    groups = np.asarray([row["refno"] for row in rows], dtype=object)
    local_docs = [str(row["local_document"]) for row in rows]
    linked_docs = [str(row["linked_document"]) for row in rows]
    article_docs = [str(row["article_document"]) for row in rows]
    numeric = np.asarray([row["numeric"] for row in rows], dtype=float)
    outputs = {
        "target_local": np.full(len(rows), np.nan, dtype=float),
        "manifest_linked_local": np.full(len(rows), np.nan, dtype=float),
        "article": np.full(len(rows), np.nan, dtype=float),
    }
    fold_indices: list[Any] = []
    splitter = GroupKFold(n_splits=N_SPLITS)
    for fold, (train_index, test_index) in enumerate(splitter.split(numeric, labels, groups), start=1):
        fold_indices.append(test_index)
        weights = inverse_group_weights(train_index, groups)
        outputs["target_local"][test_index] = _fit_text_numeric_model(
            [local_docs[index] for index in train_index],
            [local_docs[index] for index in test_index],
            numeric[train_index], numeric[test_index], labels[train_index], weights,
        )
        outputs["manifest_linked_local"][test_index] = _fit_text_numeric_model(
            [linked_docs[index] for index in train_index],
            [linked_docs[index] for index in test_index],
            numeric[train_index], numeric[test_index], labels[train_index], weights,
        )
        outputs["article"][test_index] = _fit_text_numeric_model(
            [article_docs[index] for index in train_index],
            [article_docs[index] for index in test_index],
            numeric[train_index], numeric[test_index], labels[train_index], weights,
        )
        print(f"grouped fold {fold}/{N_SPLITS} complete", flush=True)
    for name, values in outputs.items():
        if not np.isfinite(values).all():
            raise AssertionError(f"incomplete grouped OOF probabilities: {name}")
    return outputs, fold_indices


def build_strategy_probabilities(
    rows: Sequence[Mapping[str, Any]], base: Mapping[str, Any]
) -> dict[str, Any]:
    local = base["target_local"]
    article = base["article"]
    linked = base["manifest_linked_local"]
    shared = np.asarray([bool(row["shared_evidence"]) for row in rows])
    gate = np.asarray([bool(row["manifest_gate"]) for row in rows])
    shared_probability = local.copy()
    shared_probability[shared] = (
        (1.0 - PROPAGATION_WEIGHT) * local[shared]
        + PROPAGATION_WEIGHT * article[shared]
    )
    manifest_probability = local.copy()
    manifest_probability[gate] = (
        0.50 * linked[gate]
        + 0.25 * local[gate]
        + 0.25 * article[gate]
    )
    return {
        "target_local_current_like": local,
        "shared_propagation_50": shared_probability,
        "manifest_gated_propagation": manifest_probability,
    }


def metric_record(labels: Any, probabilities: Any) -> dict[str, Any]:
    decisions = (probabilities >= 0.5).astype(np.int8)
    matrix = confusion_matrix(labels, decisions, labels=[0, 1])
    tn, fp, fn, tp = (int(value) for value in matrix.ravel())
    return {
        "rows": int(len(labels)),
        "coverage": 1.0,
        "accuracy": round(float(accuracy_score(labels, decisions)), 8),
        "balanced_accuracy": round(float(balanced_accuracy_score(labels, decisions)), 8),
        "macro_f1": round(float(f1_score(labels, decisions, average="macro", zero_division=0)), 8),
        "f1_0": round(float(f1_score(labels, decisions, labels=[0], average="macro", zero_division=0)), 8),
        "f1_1": round(float(f1_score(labels, decisions, labels=[1], average="macro", zero_division=0)), 8),
        "confusion": {"tn": tn, "fp": fp, "fn": fn, "tp": tp},
        "predicted_positive": int(decisions.sum()),
        "threshold": 0.5,
    }


def fold_metrics(labels: Any, probabilities: Any, fold_indices: Sequence[Any]) -> list[dict[str, Any]]:
    return [
        {"fold": fold, **metric_record(labels[index], probabilities[index])}
        for fold, index in enumerate(fold_indices, start=1)
    ]


def subset_metric(labels: Any, probabilities: Any, mask: Any) -> dict[str, Any] | None:
    if not bool(mask.any()):
        return None
    return metric_record(labels[mask], probabilities[mask])


def evaluate_manifest(
    receipt_path: Path,
    manifest_path: Path,
    identity_path: Path,
    page_cache_path: Path,
    pdf_dir: Path,
    taxonomy_path: Path,
    positive_kb_path: Path,
    targets_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    require_model_dependencies()
    # This verification is intentionally complete before targets_path is opened.
    receipt, manifests = verify_frozen_manifest(
        receipt_path, manifest_path, identity_path, page_cache_path, pdf_dir,
        taxonomy_path, positive_kb_path,
    )
    rows = flatten_manifest(manifests)
    permitted_resultnos = {row["resultno"] for row in rows}
    labels_by_resultno = load_training_labels(targets_path.resolve(), permitted_resultnos)
    labels = np.asarray([labels_by_resultno[row["resultno"]] for row in rows], dtype=np.int8)
    base, folds = grouped_oof_probabilities(rows, labels)
    strategies = build_strategy_probabilities(rows, base)
    probability_path = output_path.resolve().parent / "oof_probabilities.npz"
    atomic_strategy_npz(probability_path, strategies)
    shared_mask = np.asarray([bool(row["shared_evidence"]) for row in rows])
    gate_mask = np.asarray([bool(row["manifest_gate"]) for row in rows])
    strategy_results = {
        name: {
            "aggregate": metric_record(labels, probability),
            "folds": fold_metrics(labels, probability, folds),
            "subsets": {
                "shared_evidence": subset_metric(labels, probability, shared_mask),
                "not_shared_evidence": subset_metric(labels, probability, ~shared_mask),
                "manifest_gate_open": subset_metric(labels, probability, gate_mask),
                "manifest_gate_closed": subset_metric(labels, probability, ~gate_mask),
            },
        }
        for name, probability in strategies.items()
    }
    baseline_accuracy = strategy_results["target_local_current_like"]["aggregate"]["accuracy"]
    for result in strategy_results.values():
        result["accuracy_delta_vs_target_local"] = round(
            result["aggregate"]["accuracy"] - baseline_accuracy, 8
        )
    label_counts = Counter(int(value) for value in labels)
    report = {
        "schema_version": "r3_component_manifest_grouped_oof_v1",
        "status": "complete",
        "scope": "development-only fixed PDF-only training membership",
        "freeze_boundary": {
            "manifest_frozen_before_label_access": True,
            "receipt_status": receipt["status"],
            "verified_manifest_sha256": receipt["manifest"]["sha256"],
            "all_allowed_input_hashes_verified_before_label_access": True,
            "label_file_hash_omitted_to_avoid_reading_heldout_bytes": True,
        },
        "guardrails": {
            "cross_validation": "5-fold GroupKFold by REFNO",
            "fold_text_vocabulary": "fit on outer-training papers only",
            "full_coverage": True,
            "allowed_feature_sources": [
                "frozen identity ROCKNAME",
                "frozen source-derived component manifest",
                "frozen taxonomy and positive expert KB patterns",
            ],
            "forbidden_sources_not_read": receipt["guardrails"]["forbidden_sources_not_read"],
            "prior_result_selection": False,
            "formal_scorer_modified": False,
        },
        "data": {
            **receipt["membership"],
            "evaluated_rows": len(rows),
            "evaluated_paper_groups": len({row["refno"] for row in rows}),
            "label_counts": {str(key): value for key, value in sorted(label_counts.items())},
            "shared_evidence_rows": int(shared_mask.sum()),
            "manifest_gate_open_rows": int(gate_mask.sum()),
        },
        "strategy_definitions": {
            "target_local_current_like": "target-local source scopes and target-local expert concepts",
            "shared_propagation_50": "50% article probability only where the article has a target-unbound, multi-target, or collective knowledge scope",
            "manifest_gated_propagation": "article/local/linked evidence fusion only where a source sample, component, or collective target scope opens the frozen manifest gate",
        },
        "strategies": strategy_results,
        "comparison": {
            name: {
                "accuracy": result["aggregate"]["accuracy"],
                "accuracy_delta_vs_target_local": result["accuracy_delta_vs_target_local"],
                "balanced_accuracy": result["aggregate"]["balanced_accuracy"],
                "macro_f1": result["aggregate"]["macro_f1"],
                "confusion": result["aggregate"]["confusion"],
                "fold_accuracies": [fold["accuracy"] for fold in result["folds"]],
            }
            for name, result in strategy_results.items()
        },
        "interpretation": {
            "manifest_improves_full_coverage_accuracy": (
                strategy_results["manifest_gated_propagation"]["aggregate"]["accuracy"]
                > baseline_accuracy
            ),
            "full_coverage_target_accuracy": 0.90,
            "any_strategy_reaches_target": any(
                result["aggregate"]["accuracy"] >= 0.90
                for result in strategy_results.values()
            ),
            "external_paper_group_validation_required": True,
        },
        "artifacts": {
            "manifest": receipt["manifest"],
            "freeze_receipt": {"path": str(receipt_path.resolve()), "sha256": file_sha256(receipt_path)},
            "oof_probabilities": {
                "path": str(probability_path),
                "sha256": file_sha256(probability_path),
                "row_count": len(rows),
                "array_keys": sorted(
                    key
                    for name in strategies
                    for key in (
                        f"{name}__probability",
                        f"{name}__decision",
                    )
                ),
                "contains_labels": False,
                "contains_identifiers": False,
                "alignment": (
                    "component manifest JSONL paper order, then canonical-target "
                    "order, then sorted result identity order"
                ),
                "manifest_sha256": receipt["manifest"]["sha256"],
            },
            "script": {"path": str(Path(__file__).resolve()), "sha256": file_sha256(Path(__file__).resolve())},
            "output": str(output_path.resolve()),
        },
    }
    atomic_json(output_path.resolve(), report)
    summary_path = output_path.resolve().parent / "comparison_summary.json"
    atomic_json(
        summary_path,
        {
            "schema_version": "r3_component_manifest_comparison_summary_v1",
            "manifest_sha256": receipt["manifest"]["sha256"],
            "oof_probabilities_sha256": file_sha256(probability_path),
            "rows": len(rows),
            "paper_groups": len({row["refno"] for row in rows}),
            "comparison": report["comparison"],
            "interpretation": report["interpretation"],
        },
    )
    return report


def _freeze_parser(subparsers: Any) -> None:
    parser = subparsers.add_parser("freeze", help="extract and freeze the label-free manifest")
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--positive-kb", type=Path, default=DEFAULT_POSITIVE_KB)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_OUTPUT_DIR / "component_manifest.jsonl")
    parser.add_argument("--receipt", type=Path, default=DEFAULT_OUTPUT_DIR / "freeze_receipt.json")


def _evaluate_parser(subparsers: Any) -> None:
    parser = subparsers.add_parser("evaluate", help="verify freeze, then run grouped OOF")
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--pdf-dir", type=Path, default=DEFAULT_PDF_DIR)
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--positive-kb", type=Path, default=DEFAULT_POSITIVE_KB)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_OUTPUT_DIR / "component_manifest.jsonl")
    parser.add_argument("--receipt", type=Path, default=DEFAULT_OUTPUT_DIR / "freeze_receipt.json")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_DIR / "evaluation.json")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="phase", required=True)
    _freeze_parser(subparsers)
    _evaluate_parser(subparsers)
    args = parser.parse_args(argv)
    if args.phase == "freeze":
        receipt = freeze_manifest(
            args.identity, args.page_cache, args.pdf_dir, args.taxonomy,
            args.positive_kb, args.manifest, args.receipt,
        )
        print(json.dumps({
            "status": receipt["status"],
            "manifest": receipt["manifest"],
            "membership": receipt["membership"],
            "extraction": receipt["extraction"]["aggregate"],
        }, ensure_ascii=False, indent=2))
        return 0
    report = evaluate_manifest(
        args.receipt, args.manifest, args.identity, args.page_cache, args.pdf_dir,
        args.taxonomy, args.positive_kb, args.targets, args.output,
    )
    print(json.dumps(report["comparison"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
