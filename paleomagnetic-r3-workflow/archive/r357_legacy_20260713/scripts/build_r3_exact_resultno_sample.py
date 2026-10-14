#!/usr/bin/env python3
"""Freeze exactly N RESULTNO rows from records whose REFNO has a local PDF."""

from __future__ import annotations

import argparse
from collections import defaultdict
import hashlib
import json
import os
from pathlib import Path
import tempfile
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook


IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
DEFAULT_SALT = "r3-exact-300-resultno-v1"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rank_digest(refno: str, resultno: str, salt: str) -> str:
    return hashlib.sha256(
        f"{salt}\0{refno}\0{resultno}".encode("utf-8")
    ).hexdigest()


def load_identity(path: Path) -> list[tuple[str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("identity manifest must contain one worksheet")
        iterator = workbook.worksheets[0].iter_rows()
        header = tuple(clean(cell.value) for cell in next(iterator, ()))
        if header != IDENTITY_COLUMNS:
            raise ValueError("identity columns must be REFNO, RESULTNO, ROCKNAME")
        rows: list[tuple[str, str, str]] = []
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"identity formula at Excel row {row_number}")
            values = tuple(clean(cell.value) for cell in cells[:3])
            if not any(values):
                continue
            if len(values) != 3 or any(not value for value in values):
                raise ValueError(f"empty identity value at Excel row {row_number}")
            rows.append(values)  # type: ignore[arg-type]
    finally:
        workbook.close()
    if not rows:
        raise ValueError("identity manifest is empty")
    if len({row[1] for row in rows}) != len(rows):
        raise ValueError("identity manifest contains duplicate RESULTNO values")
    return rows


def local_pdf_refnos(pdf_dir: Path) -> set[str]:
    if not pdf_dir.is_dir():
        raise ValueError(f"PDF directory does not exist: {pdf_dir}")
    return {
        path.stem
        for path in pdf_dir.iterdir()
        if path.is_file() and path.suffix.casefold() == ".pdf"
    }


def write_identity(path: Path, rows: Sequence[tuple[str, str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("identity")
    sheet.append(IDENTITY_COLUMNS)
    for row in rows:
        sheet.append(row)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".xlsx", dir=path.parent
    )
    os.close(descriptor)
    Path(temporary_name).unlink(missing_ok=True)
    try:
        workbook.save(temporary_name)
        Path(temporary_name).replace(path)
    finally:
        Path(temporary_name).unlink(missing_ok=True)


def build_exact_sample(
    identity_path: Path,
    pdf_dir: Path,
    output_path: Path,
    *,
    target_rows: int,
    salt: str,
) -> dict[str, Any]:
    if target_rows < 5:
        raise ValueError("target_rows must be at least five")
    if not salt:
        raise ValueError("salt must be nonempty")
    identity_path = identity_path.resolve()
    pdf_dir = pdf_dir.resolve()
    output_path = output_path.resolve(strict=False)
    if identity_path == output_path:
        raise ValueError("input and output identities must differ")
    rows = load_identity(identity_path)
    pdf_refnos = local_pdf_refnos(pdf_dir)
    eligible_rows = [row for row in rows if row[0] in pdf_refnos]
    grouped: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    for row in eligible_rows:
        grouped[row[0]].append(row)
    if len(eligible_rows) < target_rows:
        raise ValueError(
            f"requested {target_rows} rows but only "
            f"{len(eligible_rows)} belong to local-PDF papers"
        )
    ranked = sorted(
        eligible_rows,
        key=lambda row: (rank_digest(row[0], row[1], salt), row[0], row[1]),
    )
    selected_rows = ranked[:target_rows]
    if len(selected_rows) != target_rows:
        raise AssertionError("exact row selection produced the wrong row count")
    selected_refnos = {row[0] for row in selected_rows}
    selected_counts: dict[str, int] = defaultdict(int)
    for row in selected_rows:
        selected_counts[row[0]] += 1
    write_identity(output_path, selected_rows)
    return {
        "schema_version": "r3_exact_resultno_local_pdf_sample_v1",
        "status": "PASS",
        "selection": {
            "requested_resultno_rows": target_rows,
            "actual_resultno_rows": len(selected_rows),
            "sampling_unit": "RESULTNO identity row",
            "method": "first N rows in deterministic hash rank",
            "salt": salt,
            "rank_expression": "SHA256(UTF8(salt + NUL + REFNO + NUL + RESULTNO))",
            "required_evaluation_grouping": (
                "all selected rows sharing REFNO must remain in the same fold"
            ),
            "labels_read": False,
            "source_tables_read": False,
            "pdf_text_read": False,
            "local_pdf_membership_used": True,
        },
        "input": {
            "identity": str(identity_path),
            "identity_sha256": file_sha256(identity_path),
            "rows": len(rows),
            "paper_groups": len({row[0] for row in rows}),
            "local_pdf_rows": len(eligible_rows),
            "local_pdf_paper_groups": len(grouped),
            "pdf_dir": str(pdf_dir),
        },
        "output": {
            "identity": str(output_path),
            "identity_sha256": file_sha256(output_path),
            "rows": len(selected_rows),
            "paper_groups": len(selected_refnos),
            "columns": list(IDENTITY_COLUMNS),
            "selected_rows_per_refno_counts": {
                str(size): sum(value == size for value in selected_counts.values())
                for size in sorted(set(selected_counts.values()))
            },
            "maximum_selected_rows_for_one_refno": max(selected_counts.values()),
            "selected_row_rank_digest_sha256": hashlib.sha256(
                "\n".join(rank_digest(row[0], row[1], salt) for row in selected_rows).encode(
                    "ascii"
                )
            ).hexdigest(),
            "eligible_row_rank_digest_sha256": hashlib.sha256(
                "\n".join(rank_digest(row[0], row[1], salt) for row in ranked).encode(
                    "ascii"
                )
            ).hexdigest(),
        },
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--identity", type=Path, required=True)
    parser.add_argument("--pdf-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--target-rows", type=int, default=300)
    parser.add_argument("--salt", default=DEFAULT_SALT)
    args = parser.parse_args(argv)
    try:
        audit = build_exact_sample(
            args.identity,
            args.pdf_dir,
            args.output,
            target_rows=args.target_rows,
            salt=args.salt,
        )
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}")
        return 2
    audit_path = args.audit.resolve(strict=False)
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
