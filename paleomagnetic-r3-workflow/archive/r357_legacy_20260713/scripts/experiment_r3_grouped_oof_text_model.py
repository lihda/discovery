#!/usr/bin/env python3
"""Grouped OOF experiment for transparent R3 original-paper text models.

This development-only experiment is deliberately isolated from the runtime
scorer. It reads R3 labels only after a row is proven to belong to the frozen
training identity partition. REFNO and RESULTNO are alignment/grouping keys;
neither is inserted into model text.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Iterable, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    confusion_matrix,
    f1_score,
)
from sklearn.model_selection import GroupKFold
from sklearn.preprocessing import OneHotEncoder


ROOT = Path(__file__).resolve().parents[1]
EXPERIMENT_DIR = ROOT / "tmp/r3_text_model_oof"
DEFAULT_TARGETS = ROOT / "data/targets.xlsx"
DEFAULT_IDENTITY = ROOT / "tmp/r3_distilled_training_preflight/identity.xlsx"
DEFAULT_PAGE_CACHE = ROOT / "tmp/mineru_cache_all_hard_risk.jsonl"
DEFAULT_TAXONOMY = ROOT / "data/calibration/r3_rockmag_text_distillation_v1.json"
DEFAULT_OUTPUT = EXPERIMENT_DIR / "aggregate_result.json"
DEFAULT_CHECKPOINT = EXPERIMENT_DIR / "oof_probability_checkpoint.npz"

SPLIT_SALT = "r3-rockmag-distillation-v1"
HOLDOUT_FRACTION = 0.20
N_SPLITS = 5
RANDOM_STATE = 357

TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z'-]{1,}|\d+(?:\.\d+)?")
SPACE_RE = re.compile(r"\s+")
GENERIC_TARGET_TERMS = {
    "and",
    "area",
    "basalt",
    "basalts",
    "bed",
    "beds",
    "body",
    "complex",
    "dike",
    "dikes",
    "dyke",
    "dykes",
    "formation",
    "formations",
    "group",
    "intrusion",
    "intrusions",
    "lava",
    "lavas",
    "member",
    "members",
    "of",
    "province",
    "rock",
    "rocks",
    "sample",
    "samples",
    "section",
    "sections",
    "site",
    "sites",
    "suite",
    "the",
    "unit",
    "units",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def holdout_bucket(refno: str) -> float:
    digest = hashlib.sha256(f"{SPLIT_SALT}\0{refno}".encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") / float(1 << 64)


def is_training_refno(refno: str) -> bool:
    return holdout_bucket(refno) >= HOLDOUT_FRACTION


def _header_indices(header: Sequence[Any], required: Sequence[str]) -> dict[str, int]:
    names = [clean(cell.value if hasattr(cell, "value") else cell) for cell in header]
    result: dict[str, int] = {}
    for name in required:
        matches = [index for index, value in enumerate(names) if value == name]
        if len(matches) != 1:
            raise ValueError(f"expected exactly one {name} column; found {len(matches)}")
        result[name] = matches[0]
    return result


def load_training_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("training identity workbook must contain one worksheet")
        rows = workbook.worksheets[0].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "ROCKNAME"))
        records: list[dict[str, str]] = []
        for row in rows:
            record = {
                name: clean(row[index].value)
                for name, index in columns.items()
            }
            if not any(record.values()):
                continue
            if any(not value for value in record.values()):
                raise ValueError("training identity contains an empty value")
            if not is_training_refno(record["REFNO"]):
                raise ValueError(
                    f"identity contains held-out REFNO {record['REFNO']}; refusing label access"
                )
            records.append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError("training identity is empty")
    resultnos = [record["RESULTNO"] for record in records]
    if len(resultnos) != len(set(resultnos)):
        raise ValueError("training identity contains duplicate RESULTNO values")
    return records


def load_training_labels(path: Path, training_resultnos: set[str]) -> dict[str, int]:
    """Dereference R3 only after RESULTNO membership is known to be training."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[workbook.sheetnames[0]].iter_rows()
        columns = _header_indices(next(rows, ()), ("REFNO", "RESULTNO", "R3"))
        labels: dict[str, int] = {}
        for row in rows:
            refno = clean(row[columns["REFNO"]].value)
            if not refno or not is_training_refno(refno):
                # The R3 and RESULTNO cells of held-out rows are not dereferenced.
                continue
            resultno = clean(row[columns["RESULTNO"]].value)
            if resultno not in training_resultnos:
                continue
            raw_label = row[columns["R3"]].value
            try:
                label = int(float(raw_label))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"non-binary training R3 for RESULTNO {resultno}") from exc
            if label not in {0, 1}:
                raise ValueError(f"non-binary training R3 for RESULTNO {resultno}")
            if resultno in labels:
                raise ValueError(f"duplicate target row for training RESULTNO {resultno}")
            labels[resultno] = label
    finally:
        workbook.close()
    missing = sorted(training_resultnos - labels.keys())
    extra = sorted(labels.keys() - training_resultnos)
    if missing or extra:
        raise ValueError(
            f"training label/identity mismatch: missing={missing[:10]}, extra={extra[:10]}"
        )
    return labels


def load_page_text(path: Path, required_refnos: set[str]) -> dict[str, str]:
    papers: dict[str, str] = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            row = json.loads(line)
            refno = clean(row.get("REFNO"))
            if refno not in required_refnos:
                continue
            if refno in papers:
                raise ValueError(f"duplicate cache REFNO {refno} at line {line_number}")
            page_text = [
                clean(page.get("text"))
                for page in row.get("pages", [])
                if isinstance(page, Mapping)
            ]
            papers[refno] = "\n\n".join(text for text in page_text if text)
    missing = sorted(required_refnos - papers.keys())
    if missing:
        raise ValueError(f"page cache is missing training REFNO values: {missing[:20]}")
    return papers


def load_taxonomy_patterns(
    path: Path,
) -> tuple[
    dict[str, list[tuple[str, re.Pattern[str]]]],
    list[re.Pattern[str]],
    re.Pattern[str],
]:
    artifact = json.loads(path.read_text(encoding="utf-8"))
    concepts: dict[str, list[tuple[str, re.Pattern[str]]]] = {}
    patterns: list[re.Pattern[str]] = []
    for category, key in (
        ("method", "method_taxonomy"),
        ("diagnostic", "diagnostic_taxonomy"),
        ("carrier", "carrier_taxonomy"),
    ):
        entries = artifact.get(key)
        if not isinstance(entries, list):
            raise ValueError(f"taxonomy artifact is missing {key}")
        concepts[category] = []
        for entry in entries:
            expressions = [
                str(expression)
                for expression in entry.get("runtime_original_text_patterns", [])
            ]
            if not expressions:
                continue
            compiled_patterns = [
                re.compile(expression, re.I | re.S) for expression in expressions
            ]
            patterns.extend(compiled_patterns)
            concept_pattern = re.compile(
                "|".join(f"(?:{expression})" for expression in expressions),
                re.I | re.S,
            )
            concepts[category].append((clean(entry.get("concept_id")), concept_pattern))
    if not patterns:
        raise ValueError("taxonomy artifact contains no runtime patterns")
    combined = re.compile(
        "|".join(f"(?:{pattern.pattern})" for pattern in patterns),
        re.I | re.S,
    )
    return concepts, patterns, combined


def _merge_spans(spans: Iterable[tuple[int, int]], text_length: int) -> list[tuple[int, int]]:
    normalized = sorted(
        (max(0, start), min(text_length, end))
        for start, end in spans
        if end > start
    )
    merged: list[list[int]] = []
    for start, end in normalized:
        if not merged or start > merged[-1][1] + 80:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    return [(start, end) for start, end in merged]


def _bounded_matches(
    text: str,
    patterns: Iterable[re.Pattern[str]],
    radius: int,
    cap: int,
) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    for pattern in patterns:
        for match in pattern.finditer(text):
            spans.append((match.start() - radius, match.end() + radius))
            if len(spans) >= cap:
                return spans
    return spans


def _target_patterns(rockname: str) -> list[re.Pattern[str]]:
    tokens = [token.casefold() for token in TOKEN_RE.findall(rockname)]
    informative = [
        token
        for token in tokens
        if token not in GENERIC_TARGET_TERMS
        and (len(token) >= 3 or token.replace(".", "").isdigit())
    ]
    patterns: list[re.Pattern[str]] = []
    normalized_name = " ".join(tokens)
    if len(tokens) >= 2 and len(normalized_name) >= 7:
        expression = r"\W+".join(re.escape(token) for token in tokens)
        patterns.append(re.compile(expression, re.I))
    # Longest names and localities are more discriminative and are tried first.
    for token in sorted(set(informative), key=lambda value: (-len(value), value))[:8]:
        if token.replace(".", "").isdigit():
            patterns.append(re.compile(rf"(?<!\d){re.escape(token)}(?!\d)", re.I))
        else:
            patterns.append(re.compile(rf"\b{re.escape(token)}\b", re.I))
    return patterns


def target_neighborhood(
    text: str,
    rockname: str,
    knowledge_spans: Sequence[tuple[int, int]],
) -> str:
    if not text:
        return ""
    target_spans = _bounded_matches(
        text,
        _target_patterns(rockname),
        radius=700,
        cap=80,
    )
    # Keep taxonomy evidence paper-wide and target passages target-specific.
    # Their overlap naturally yields the strongest target-bound descriptions.
    spans = _merge_spans([*target_spans, *knowledge_spans], len(text))
    if not spans:
        # Title/abstract still provides a bounded paper signal when neither the
        # target name nor frozen terminology survives PDF extraction.
        return SPACE_RE.sub(" ", text[:2500]).strip()
    excerpts = [SPACE_RE.sub(" ", text[start:end]).strip() for start, end in spans]
    return "\n".join(excerpt for excerpt in excerpts if excerpt)


def _count_matches(pattern: re.Pattern[str], text: str, cap: int = 100) -> int:
    count = 0
    for _ in pattern.finditer(text):
        count += 1
        if count >= cap:
            break
    return count


def taxonomy_count_features(
    text: str,
    concepts: Mapping[str, Sequence[tuple[str, re.Pattern[str]]]],
) -> dict[str, int]:
    output: dict[str, int] = {}
    for category in ("method", "diagnostic", "carrier"):
        distinct = 0
        hits = 0
        for _, pattern in concepts[category]:
            concept_hits = _count_matches(pattern, text)
            distinct += int(concept_hits > 0)
            hits = min(100, hits + concept_hits)
        output[f"{category}_distinct"] = distinct
        output[f"{category}_hits"] = hits
    return output


def target_hit_count(text: str, rockname: str) -> int:
    hits = 0
    for pattern in _target_patterns(rockname):
        hits = min(100, hits + _count_matches(pattern, text))
        if hits >= 100:
            break
    return hits


def count_bin(value: int, *, distinct: bool) -> int:
    """Return a stable ordinal bin; the model one-hot encodes this value."""

    limits = (0, 1, 2, 3) if distinct else (0, 1, 2, 5, 10, 25)
    for index, upper in enumerate(limits):
        if value <= upper:
            return index
    return len(limits)


ENGINEERED_FEATURE_NAMES = (
    "method_distinct_bin",
    "method_hits_bin",
    "diagnostic_distinct_bin",
    "diagnostic_hits_bin",
    "carrier_distinct_bin",
    "carrier_hits_bin",
    "target_hits_bin",
)

CONTINUOUS_PAPER_FEATURE_NAMES = (
    "method_distinct",
    "method_hits",
    "diagnostic_distinct",
    "diagnostic_hits",
    "carrier_distinct",
    "carrier_hits",
    "paper_character_count",
)

CONTINUOUS_ROW_FEATURE_NAMES = (
    *CONTINUOUS_PAPER_FEATURE_NAMES[:-1],
    "target_hits",
    "paper_character_count",
)


def binned_feature_row(paper_counts: Mapping[str, int], target_hits: int) -> list[int]:
    return [
        count_bin(paper_counts["method_distinct"], distinct=True),
        count_bin(paper_counts["method_hits"], distinct=False),
        count_bin(paper_counts["diagnostic_distinct"], distinct=True),
        count_bin(paper_counts["diagnostic_hits"], distinct=False),
        count_bin(paper_counts["carrier_distinct"], distinct=True),
        count_bin(paper_counts["carrier_hits"], distinct=False),
        count_bin(target_hits, distinct=False),
    ]


def metric_record(y_true: np.ndarray, probabilities: np.ndarray) -> dict[str, Any]:
    prediction = (probabilities >= 0.5).astype(np.int8)
    tn, fp, fn, tp = confusion_matrix(y_true, prediction, labels=[0, 1]).ravel()
    return {
        "accuracy": round(float(accuracy_score(y_true, prediction)), 8),
        "macro_f1": round(float(f1_score(y_true, prediction, average="macro")), 8),
        "balanced_accuracy": round(float(balanced_accuracy_score(y_true, prediction)), 8),
        "confusion": {
            "tn": int(tn),
            "fp": int(fp),
            "fn": int(fn),
            "tp": int(tp),
        },
        "predicted_positive": int(prediction.sum()),
        "threshold": 0.5,
    }


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.name + ".", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        Path(temporary_name).replace(path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def save_checkpoint(
    path: Path,
    probabilities: Mapping[str, np.ndarray],
    completed: Sequence[str],
) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp.npz")
    np.savez_compressed(
        temporary,
        completed=np.asarray(completed),
        **probabilities,
    )
    temporary.replace(path)


def model_specs() -> dict[str, dict[str, Any]]:
    return {
        "target_word": {
            "document": "target_neighborhood",
            "vectorizer": {
                "analyzer": "word",
                "ngram_range": (1, 2),
                "min_df": 2,
                "max_df": 0.995,
                "max_features": 80000,
                "sublinear_tf": True,
                "strip_accents": "unicode",
            },
            "logistic_c": 2.0,
        },
        "target_char": {
            "document": "target_neighborhood",
            "vectorizer": {
                "analyzer": "char_wb",
                "ngram_range": (3, 5),
                "min_df": 2,
                "max_features": 100000,
                "sublinear_tf": True,
            },
            "logistic_c": 2.0,
        },
        "paper_word": {
            "document": "full_paper",
            "vectorizer": {
                "analyzer": "word",
                "ngram_range": (1, 2),
                "min_df": 2,
                "max_df": 0.995,
                "max_features": 80000,
                "sublinear_tf": True,
                "strip_accents": "unicode",
            },
            "logistic_c": 2.0,
        },
    }


def run_oof_model(
    name: str,
    spec: Mapping[str, Any],
    documents: Sequence[str],
    labels: np.ndarray,
    groups: np.ndarray,
    checkpoint_path: Path,
    accumulated: dict[str, np.ndarray],
    completed: list[str],
) -> tuple[np.ndarray, list[dict[str, Any]]]:
    probabilities = np.full(len(labels), np.nan, dtype=np.float64)
    folds: list[dict[str, Any]] = []
    splitter = GroupKFold(n_splits=N_SPLITS)
    for fold_index, (train_index, test_index) in enumerate(
        splitter.split(np.zeros(len(labels)), labels, groups), start=1
    ):
        vectorizer = TfidfVectorizer(dtype=np.float32, **dict(spec["vectorizer"]))
        train_text = [documents[index] for index in train_index]
        test_text = [documents[index] for index in test_index]
        train_matrix = vectorizer.fit_transform(train_text)
        test_matrix = vectorizer.transform(test_text)
        classifier = LogisticRegression(
            C=float(spec["logistic_c"]),
            solver="liblinear",
            max_iter=1000,
            random_state=RANDOM_STATE,
        )
        classifier.fit(train_matrix, labels[train_index])
        fold_probabilities = classifier.predict_proba(test_matrix)[:, 1]
        probabilities[test_index] = fold_probabilities
        fold_metric = metric_record(labels[test_index], fold_probabilities)
        fold_metric.update(
            {
                "fold": fold_index,
                "train_rows": int(len(train_index)),
                "test_rows": int(len(test_index)),
                "train_groups": int(len(set(groups[train_index]))),
                "test_groups": int(len(set(groups[test_index]))),
                "vocabulary_size": int(len(vectorizer.vocabulary_)),
            }
        )
        folds.append(fold_metric)
        accumulated[name] = probabilities
        save_checkpoint(checkpoint_path, accumulated, [*completed, f"{name}:fold{fold_index}"])
        print(
            f"{name} fold {fold_index}/{N_SPLITS}: "
            f"accuracy={fold_metric['accuracy']:.4f}, vocab={len(vectorizer.vocabulary_)}",
            flush=True,
        )
    if np.isnan(probabilities).any():
        raise AssertionError(f"{name} has incomplete OOF probabilities")
    completed.append(name)
    accumulated[name] = probabilities
    save_checkpoint(checkpoint_path, accumulated, completed)
    return probabilities, folds


def run_oof_binned_logistic(
    name: str,
    features: np.ndarray,
    labels: np.ndarray,
    groups: np.ndarray,
    checkpoint_path: Path,
    accumulated: dict[str, np.ndarray],
    completed: list[str],
) -> tuple[np.ndarray, list[dict[str, Any]]]:
    probabilities = np.full(len(labels), np.nan, dtype=np.float64)
    folds: list[dict[str, Any]] = []
    splitter = GroupKFold(n_splits=N_SPLITS)
    for fold_index, (train_index, test_index) in enumerate(
        splitter.split(features, labels, groups), start=1
    ):
        encoder = OneHotEncoder(handle_unknown="ignore", sparse_output=True)
        train_matrix = encoder.fit_transform(features[train_index])
        test_matrix = encoder.transform(features[test_index])
        classifier = LogisticRegression(
            C=1.0,
            solver="liblinear",
            max_iter=1000,
            random_state=RANDOM_STATE,
        )
        classifier.fit(train_matrix, labels[train_index])
        fold_probabilities = classifier.predict_proba(test_matrix)[:, 1]
        probabilities[test_index] = fold_probabilities
        fold_metric = metric_record(labels[test_index], fold_probabilities)
        fold_metric.update(
            {
                "fold": fold_index,
                "train_rows": int(len(train_index)),
                "test_rows": int(len(test_index)),
                "train_groups": int(len(set(groups[train_index]))),
                "test_groups": int(len(set(groups[test_index]))),
                "encoded_feature_count": int(train_matrix.shape[1]),
            }
        )
        folds.append(fold_metric)
        accumulated[name] = probabilities
        save_checkpoint(checkpoint_path, accumulated, [*completed, f"{name}:fold{fold_index}"])
        print(
            f"{name} fold {fold_index}/{N_SPLITS}: "
            f"accuracy={fold_metric['accuracy']:.4f}, features={train_matrix.shape[1]}",
            flush=True,
        )
    if np.isnan(probabilities).any():
        raise AssertionError(f"{name} has incomplete OOF probabilities")
    completed.append(name)
    accumulated[name] = probabilities
    save_checkpoint(checkpoint_path, accumulated, completed)
    return probabilities, folds


def run_oof_gradient_boosting(
    name: str,
    features: np.ndarray,
    labels: np.ndarray,
    groups: np.ndarray,
    checkpoint_path: Path,
    accumulated: dict[str, np.ndarray],
    completed: list[str],
    *,
    paper_majority_training: bool,
    min_samples_leaf: int,
) -> tuple[np.ndarray, list[dict[str, Any]]]:
    """Estimate an interpretable nonlinear knowledge-count model by paper group."""

    probabilities = np.full(len(labels), np.nan, dtype=np.float64)
    folds: list[dict[str, Any]] = []
    splitter = GroupKFold(n_splits=N_SPLITS)
    for fold_index, (train_index, test_index) in enumerate(
        splitter.split(features, labels, groups), start=1
    ):
        fit_features = features[train_index]
        fit_labels = labels[train_index]
        fit_rows = len(train_index)
        if paper_majority_training:
            train_groups = np.unique(groups[train_index])
            representative_indices = np.asarray(
                [
                    train_index[np.flatnonzero(groups[train_index] == group)[0]]
                    for group in train_groups
                ]
            )
            fit_features = features[representative_indices]
            fit_labels = np.asarray(
                [
                    int(labels[train_index][groups[train_index] == group].mean() >= 0.5)
                    for group in train_groups
                ],
                dtype=np.int8,
            )
            fit_rows = len(representative_indices)
        classifier = GradientBoostingClassifier(
            n_estimators=150,
            learning_rate=0.04,
            max_depth=1,
            min_samples_leaf=min_samples_leaf,
            random_state=RANDOM_STATE,
        )
        classifier.fit(fit_features, fit_labels)
        fold_probabilities = classifier.predict_proba(features[test_index])[:, 1]
        probabilities[test_index] = fold_probabilities
        fold_metric = metric_record(labels[test_index], fold_probabilities)
        fold_metric.update(
            {
                "fold": fold_index,
                "train_rows": int(len(train_index)),
                "fit_rows": int(fit_rows),
                "test_rows": int(len(test_index)),
                "train_groups": int(len(set(groups[train_index]))),
                "test_groups": int(len(set(groups[test_index]))),
            }
        )
        folds.append(fold_metric)
        accumulated[name] = probabilities
        save_checkpoint(checkpoint_path, accumulated, [*completed, f"{name}:fold{fold_index}"])
        print(
            f"{name} fold {fold_index}/{N_SPLITS}: "
            f"accuracy={fold_metric['accuracy']:.4f}, fit_rows={fit_rows}",
            flush=True,
        )
    if np.isnan(probabilities).any():
        raise AssertionError(f"{name} has incomplete OOF probabilities")
    completed.append(name)
    accumulated[name] = probabilities
    save_checkpoint(checkpoint_path, accumulated, completed)
    return probabilities, folds


def fusion_results(labels: np.ndarray, probabilities: Mapping[str, np.ndarray]) -> dict[str, Any]:
    names = ("target_word", "target_char", "paper_word")
    equal_probability = sum(probabilities[name] for name in names) / 3.0
    equal = metric_record(labels, equal_probability)
    equal["weights"] = dict.fromkeys(names, round(1.0 / 3.0, 8))

    candidates: list[dict[str, Any]] = []
    for target_word_tenths in range(11):
        for target_char_tenths in range(11 - target_word_tenths):
            paper_word_tenths = 10 - target_word_tenths - target_char_tenths
            weights = np.asarray(
                [target_word_tenths, target_char_tenths, paper_word_tenths],
                dtype=np.float64,
            ) / 10.0
            fused = sum(
                weights[index] * probabilities[name]
                for index, name in enumerate(names)
            )
            metric = metric_record(labels, fused)
            metric["weights"] = {
                name: round(float(weights[index]), 1)
                for index, name in enumerate(names)
            }
            candidates.append(metric)
    best_accuracy = max(
        candidates,
        key=lambda item: (
            item["accuracy"],
            item["macro_f1"],
            item["balanced_accuracy"],
        ),
    )
    best_macro_f1 = max(
        candidates,
        key=lambda item: (
            item["macro_f1"],
            item["accuracy"],
            item["balanced_accuracy"],
        ),
    )
    return {
        "equal_weight": equal,
        "grid": {
            "step": 0.1,
            "candidate_count": len(candidates),
            "selection_scope": "post-hoc on training-partition OOF probabilities only",
            "best_accuracy": best_accuracy,
            "best_macro_f1": best_macro_f1,
        },
    }


def engineered_char_fusion(
    labels: np.ndarray,
    target_char: np.ndarray,
    taxonomy_counts: np.ndarray,
) -> dict[str, Any]:
    candidates: list[dict[str, Any]] = []
    for char_tenths in range(11):
        char_weight = char_tenths / 10.0
        count_weight = 1.0 - char_weight
        fused = char_weight * target_char + count_weight * taxonomy_counts
        metric = metric_record(labels, fused)
        metric["weights"] = {
            "target_char": round(char_weight, 1),
            "taxonomy_count_logistic": round(count_weight, 1),
        }
        candidates.append(metric)
    best_accuracy = max(
        candidates,
        key=lambda item: (
            item["accuracy"],
            item["macro_f1"],
            item["balanced_accuracy"],
        ),
    )
    best_macro_f1 = max(
        candidates,
        key=lambda item: (
            item["macro_f1"],
            item["accuracy"],
            item["balanced_accuracy"],
        ),
    )
    return {
        "step": 0.1,
        "candidate_count": len(candidates),
        "selection_scope": "post-hoc on training-partition OOF probabilities only",
        "best_accuracy": best_accuracy,
        "best_macro_f1": best_macro_f1,
    }


def nonlinear_knowledge_fusion(
    labels: np.ndarray,
    probabilities: Mapping[str, np.ndarray],
) -> dict[str, Any]:
    names = (
        "target_char",
        "taxonomy_count_logistic",
        "row_knowledge_gradient_boosting",
        "paper_majority_knowledge_gradient_boosting",
    )
    candidates: list[dict[str, Any]] = []
    for target_char_tenths in range(11):
        for count_tenths in range(11 - target_char_tenths):
            for row_tenths in range(11 - target_char_tenths - count_tenths):
                paper_tenths = 10 - target_char_tenths - count_tenths - row_tenths
                weights = np.asarray(
                    [target_char_tenths, count_tenths, row_tenths, paper_tenths],
                    dtype=np.float64,
                ) / 10.0
                fused = sum(
                    weights[index] * probabilities[name]
                    for index, name in enumerate(names)
                )
                metric = metric_record(labels, fused)
                metric["weights"] = {
                    name: round(float(weights[index]), 1)
                    for index, name in enumerate(names)
                }
                candidates.append(metric)
    best_accuracy = max(
        candidates,
        key=lambda item: (
            item["accuracy"],
            item["macro_f1"],
            item["balanced_accuracy"],
        ),
    )
    best_macro_f1 = max(
        candidates,
        key=lambda item: (
            item["macro_f1"],
            item["accuracy"],
            item["balanced_accuracy"],
        ),
    )
    return {
        "step": 0.1,
        "candidate_count": len(candidates),
        "selection_scope": "post-hoc on training-partition OOF probabilities only",
        "best_accuracy": best_accuracy,
        "best_macro_f1": best_macro_f1,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--targets", type=Path, default=DEFAULT_TARGETS)
    parser.add_argument("--identity", type=Path, default=DEFAULT_IDENTITY)
    parser.add_argument("--page-cache", type=Path, default=DEFAULT_PAGE_CACHE)
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--checkpoint", type=Path, default=DEFAULT_CHECKPOINT)
    args = parser.parse_args()

    input_paths = {
        "targets": args.targets.resolve(),
        "training_identity": args.identity.resolve(),
        "page_cache": args.page_cache.resolve(),
        "frozen_taxonomy": args.taxonomy.resolve(),
    }
    for label, path in input_paths.items():
        if not path.is_file():
            raise SystemExit(f"missing {label}: {path}")

    identity = load_training_identity(input_paths["training_identity"])
    training_resultnos = {record["RESULTNO"] for record in identity}
    labels_by_resultno = load_training_labels(input_paths["targets"], training_resultnos)
    required_refnos = {record["REFNO"] for record in identity}
    paper_text = load_page_text(input_paths["page_cache"], required_refnos)
    taxonomy_concepts, taxonomy_patterns, combined_taxonomy_pattern = (
        load_taxonomy_patterns(input_paths["frozen_taxonomy"])
    )

    labels = np.asarray(
        [labels_by_resultno[record["RESULTNO"]] for record in identity], dtype=np.int8
    )
    groups = np.asarray([record["REFNO"] for record in identity], dtype=object)
    knowledge_spans_by_refno = {
        refno: _bounded_matches(
            text,
            (combined_taxonomy_pattern,),
            radius=320,
            cap=160,
        )
        for refno, text in paper_text.items()
    }
    target_document_cache: dict[tuple[str, str], str] = {}
    target_documents: list[str] = []
    for record in identity:
        target_key = (record["REFNO"], record["ROCKNAME"])
        if target_key not in target_document_cache:
            target_document_cache[target_key] = target_neighborhood(
                paper_text[record["REFNO"]],
                record["ROCKNAME"],
                knowledge_spans_by_refno[record["REFNO"]],
            )
        target_documents.append(target_document_cache[target_key])
    paper_documents = [paper_text[record["REFNO"]] for record in identity]
    paper_count_features = {
        refno: taxonomy_count_features(text, taxonomy_concepts)
        for refno, text in paper_text.items()
    }
    target_hit_cache: dict[tuple[str, str], int] = {}
    engineered_rows: list[list[int]] = []
    continuous_row_features: list[list[float]] = []
    continuous_paper_features: list[list[float]] = []
    for record in identity:
        target_key = (record["REFNO"], record["ROCKNAME"])
        if target_key not in target_hit_cache:
            target_hit_cache[target_key] = target_hit_count(
                paper_text[record["REFNO"]], record["ROCKNAME"]
            )
        engineered_rows.append(
            binned_feature_row(
                paper_count_features[record["REFNO"]],
                target_hit_cache[target_key],
            )
        )
        paper_counts = paper_count_features[record["REFNO"]]
        paper_length = len(paper_text[record["REFNO"]])
        continuous_paper_features.append(
            [
                paper_counts["method_distinct"],
                paper_counts["method_hits"],
                paper_counts["diagnostic_distinct"],
                paper_counts["diagnostic_hits"],
                paper_counts["carrier_distinct"],
                paper_counts["carrier_hits"],
                paper_length,
            ]
        )
        continuous_row_features.append(
            [
                paper_counts["method_distinct"],
                paper_counts["method_hits"],
                paper_counts["diagnostic_distinct"],
                paper_counts["diagnostic_hits"],
                paper_counts["carrier_distinct"],
                paper_counts["carrier_hits"],
                target_hit_cache[target_key],
                paper_length,
            ]
        )
    engineered_features = np.asarray(engineered_rows, dtype=np.int8)
    continuous_row_feature_array = np.asarray(
        continuous_row_features,
        dtype=np.float64,
    )
    continuous_paper_feature_array = np.asarray(
        continuous_paper_features,
        dtype=np.float64,
    )
    documents_by_view = {
        "target_neighborhood": target_documents,
        "full_paper": paper_documents,
    }

    specs = model_specs()
    all_probabilities: dict[str, np.ndarray] = {}
    completed: list[str] = []
    model_results: dict[str, Any] = {}
    for name, spec in specs.items():
        probabilities, folds = run_oof_model(
            name,
            spec,
            documents_by_view[str(spec["document"])],
            labels,
            groups,
            args.checkpoint.resolve(),
            all_probabilities,
            completed,
        )
        model_results[name] = {
            "feature_view": spec["document"],
            "vectorizer": {
                key: list(value) if isinstance(value, tuple) else value
                for key, value in spec["vectorizer"].items()
            },
            "classifier": {
                "type": "LogisticRegression",
                "C": spec["logistic_c"],
                "solver": "liblinear",
                "class_weight": None,
            },
            "aggregate": metric_record(labels, probabilities),
            "folds": folds,
        }

    taxonomy_count_probabilities, taxonomy_count_folds = run_oof_binned_logistic(
        "taxonomy_count_logistic",
        engineered_features,
        labels,
        groups,
        args.checkpoint.resolve(),
        all_probabilities,
        completed,
    )
    taxonomy_count_metric = metric_record(labels, taxonomy_count_probabilities)
    taxonomy_count_fusion = engineered_char_fusion(
        labels,
        all_probabilities["target_char"],
        taxonomy_count_probabilities,
    )
    row_gradient_probabilities, row_gradient_folds = run_oof_gradient_boosting(
        "row_knowledge_gradient_boosting",
        continuous_row_feature_array,
        labels,
        groups,
        args.checkpoint.resolve(),
        all_probabilities,
        completed,
        paper_majority_training=False,
        min_samples_leaf=10,
    )
    paper_gradient_probabilities, paper_gradient_folds = run_oof_gradient_boosting(
        "paper_majority_knowledge_gradient_boosting",
        continuous_paper_feature_array,
        labels,
        groups,
        args.checkpoint.resolve(),
        all_probabilities,
        completed,
        paper_majority_training=True,
        min_samples_leaf=8,
    )
    row_gradient_metric = metric_record(labels, row_gradient_probabilities)
    paper_gradient_metric = metric_record(labels, paper_gradient_probabilities)
    nonlinear_fusion = nonlinear_knowledge_fusion(labels, all_probabilities)

    fusions = fusion_results(labels, all_probabilities)
    candidate_metrics: list[tuple[str, Mapping[str, Any]]] = [
        (name, result["aggregate"]) for name, result in model_results.items()
    ]
    candidate_metrics.extend(
        [
            ("fusion_equal_weight", fusions["equal_weight"]),
            ("fusion_grid_best_accuracy", fusions["grid"]["best_accuracy"]),
            ("taxonomy_count_logistic", taxonomy_count_metric),
            ("row_knowledge_gradient_boosting", row_gradient_metric),
            ("paper_majority_knowledge_gradient_boosting", paper_gradient_metric),
            (
                "taxonomy_count_target_char_fusion",
                taxonomy_count_fusion["best_accuracy"],
            ),
            (
                "nonlinear_knowledge_fusion",
                nonlinear_fusion["best_accuracy"],
            ),
        ]
    )
    best_name, best_metric = max(
        candidate_metrics,
        key=lambda pair: (
            pair[1]["accuracy"],
            pair[1]["macro_f1"],
            pair[1]["balanced_accuracy"],
        ),
    )

    label_counts = Counter(int(value) for value in labels)
    empty_refnos = {refno for refno, text in paper_text.items() if not text.strip()}
    result: dict[str, Any] = {
        "schema_version": "r3_training_grouped_oof_text_experiment_v1",
        "status": "complete",
        "scope": "development-only fixed training partition",
        "guardrails": {
            "grouping": "REFNO GroupKFold 5 folds",
            "labels": "R3 dereferenced only for fixed training identity rows",
            "runtime_feature_sources": [
                "training identity ROCKNAME used only to select target neighborhoods",
                "original-paper text from MinerU-assisted page cache",
                "frozen R3 taxonomy runtime patterns",
            ],
            "identifier_feature_exclusion": [
                "REFNO is grouping/alignment only and is never inserted into text",
                "RESULTNO is label alignment only and is never inserted into text",
            ],
            "forbidden_sources_not_read": [
                "held-out R3 labels",
                "PMAGRESULT",
                "ROCKMAG",
                "old predictions/evaluations",
            ],
        },
        "inputs": {
            label: {"path": str(path), "sha256": file_sha256(path)}
            for label, path in input_paths.items()
        },
        "data": {
            "rows": len(identity),
            "paper_groups": len(required_refnos),
            "label_counts": {str(key): value for key, value in sorted(label_counts.items())},
            "papers_with_empty_text": len(empty_refnos),
            "rows_with_empty_paper_text": sum(
                record["REFNO"] in empty_refnos for record in identity
            ),
            "rows_with_empty_target_neighborhood": sum(
                not document.strip() for document in target_documents
            ),
            "taxonomy_runtime_pattern_count": len(taxonomy_patterns),
        },
        "cross_validation": {
            "type": "GroupKFold",
            "n_splits": N_SPLITS,
            "group": "REFNO",
            "paper_overlap_between_train_and_test": False,
        },
        "models": model_results,
        "fusion": fusions,
        "taxonomy_count_experiment": {
            "feature_scope": "label-free, deterministically recomputable from frozen taxonomy, ROCKNAME, and original-paper text",
            "raw_paper_features": [
                "method distinct concept count",
                "method total hit count capped at 100",
                "diagnostic distinct concept count",
                "diagnostic total hit count capped at 100",
                "carrier distinct concept count",
                "carrier total hit count capped at 100",
            ],
            "raw_target_features": [
                "count of ROCKNAME-derived target-pattern hits capped at 100"
            ],
            "binning": {
                "distinct_count_upper_bounds": [0, 1, 2, 3, "4+"],
                "hit_count_upper_bounds": [0, 1, 2, 5, 10, 25, "26+"],
                "encoding": "fold-local one-hot encoding",
            },
            "feature_names": list(ENGINEERED_FEATURE_NAMES),
            "observed_bin_counts": {
                name: {
                    str(bin_id): int(count)
                    for bin_id, count in sorted(
                        Counter(engineered_features[:, index].tolist()).items()
                    )
                }
                for index, name in enumerate(ENGINEERED_FEATURE_NAMES)
            },
            "model": {
                "type": "OneHotEncoder + LogisticRegression",
                "C": 1.0,
                "solver": "liblinear",
                "class_weight": None,
                "aggregate": taxonomy_count_metric,
                "folds": taxonomy_count_folds,
            },
            "target_char_weight_grid_fusion": taxonomy_count_fusion,
        },
        "nonlinear_knowledge_experiment": {
            "feature_scope": "label-free counts recomputable from ROCKNAME, original-paper text, and the frozen taxonomy",
            "row_feature_names": list(CONTINUOUS_ROW_FEATURE_NAMES),
            "paper_feature_names": list(CONTINUOUS_PAPER_FEATURE_NAMES),
            "model": {
                "type": "GradientBoostingClassifier",
                "n_estimators": 150,
                "learning_rate": 0.04,
                "max_depth": 1,
                "row_min_samples_leaf": 10,
                "paper_min_samples_leaf": 8,
            },
            "row_model": {
                "aggregate": row_gradient_metric,
                "folds": row_gradient_folds,
            },
            "paper_majority_model": {
                "training_unit": "one majority-labelled row per training REFNO; test predictions remain row-aligned",
                "aggregate": paper_gradient_metric,
                "folds": paper_gradient_folds,
            },
            "four_model_weight_grid_fusion": nonlinear_fusion,
        },
        "highest_accuracy": {"candidate": best_name, **dict(best_metric)},
        "interpretation": {
            "target_accuracy": 0.9,
            "target_reached": bool(best_metric["accuracy"] >= 0.9),
            "fusion_grid_is_training_partition_post_hoc": True,
            "formal_scorer_modified": False,
        },
        "artifacts": {
            "experiment_script": str(Path(__file__).resolve()),
            "experiment_script_sha256": file_sha256(Path(__file__).resolve()),
            "probability_checkpoint": str(args.checkpoint.resolve()),
        },
    }
    atomic_json(args.output.resolve(), result)
    print(json.dumps(result["highest_accuracy"], ensure_ascii=False, indent=2))
    print(f"wrote {args.output.resolve()}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
