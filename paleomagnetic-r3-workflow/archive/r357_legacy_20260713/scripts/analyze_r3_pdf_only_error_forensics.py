#!/usr/bin/env python3
"""Post-freeze PDF-only R3 error forensics and constrained propagation trials.

Candidate decisions are built before labels or prior predictions are opened.
Runtime candidate inputs are restricted to the strict identity manifest, the
original-paper page cache, and the frozen aggregate ROCKMAG distillation.
Labels and frozen predictions are used only in the second, evaluation phase.
"""

from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.evaluate_r3_distilled_predictions import is_holdout_refno
from scripts.paper_text_retrieval import PaperTextRetriever, canonical_target_identity
from scripts.r3_rockmag_distilled_reasoner import R3RockmagDistilledReasoner


IDENTITY = ROOT / "tmp/r3_distilled_training_preflight/identity.xlsx"
PAGE_CACHE = ROOT / "tmp/mineru_cache_all_hard_risk.jsonl"
ARTIFACT = ROOT / "data/calibration/r3_rockmag_text_distillation_v1.json"
LABELS = ROOT / "data/targets.xlsx"
LATEST = ROOT / "tmp/r3_pdf_v2/training_target_section.jsonl"
PREVIOUS = ROOT / "tmp/r3_proxy_v2/training_exclusive_scope_plus_discourse.jsonl"
OUTPUT_DIR = ROOT / "tmp/r3_error_forensics"
CANDIDATES = OUTPUT_DIR / "constrained_candidates.jsonl"
RECEIPT = OUTPUT_DIR / "candidate_freeze_receipt.json"
REPORT = OUTPUT_DIR / "report.json"
REPORT_MD = OUTPUT_DIR / "report.md"

IDENTITY_COLUMNS = ("REFNO", "RESULTNO", "ROCKNAME")
SHARED_SCOPE_PATTERNS = (
    re.compile(
        r"\b(?:all|both|each)\s+(?:of\s+)?(?:the\s+)?"
        r"(?:samples?|specimens?|sites?|localities|formations?|units?|rocks?)\b",
        re.I,
    ),
    re.compile(
        r"\b(?:samples?|specimens?|sites?|localities|formations?|units?|rocks?)"
        r"\s+(?:from|of|at)\s+(?:all|both|each)\b",
        re.I,
    ),
    re.compile(
        r"\b(?:representative\s+)?samples?\s+from\s+each\s+of\b",
        re.I,
    ),
    re.compile(
        r"\b(?:for|on)\s+(?:all|both|each)\s+(?:of\s+)?(?:the\s+)?"
        r"(?:samples?|specimens?|sites?|localities|formations?|units?|rocks?)\b",
        re.I,
    ),
)
EXECUTION_RE = re.compile(
    r"\b(?:we|our|this\s+(?:study|work)|present\s+(?:study|work)|"
    r"measure(?:d|ments?)|determin(?:ed|ation)|perform(?:ed)?|conduct(?:ed)?|"
    r"subjected|acquir(?:ed|ing)|record(?:ed)?|analys(?:ed|ed|is|es)|"
    r"experiment(?:s|al)?|curves?|loops?|observ(?:ed|ations?)|show(?:ed|s)|"
    r"indicat(?:ed|es)|reveal(?:ed|s)|yield(?:ed|s))\b",
    re.I,
)
REFERENCE_HEADING_RE = re.compile(
    r"(?im)^\s*(?:references|bibliography|literature\s+cited)\s*$"
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def atomic_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(
                json.dumps(
                    row,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                + "\n"
            )
    temporary.replace(path)


def load_identity(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = tuple(clean(value) for value in next(rows))
        if header != IDENTITY_COLUMNS:
            raise ValueError(f"strict identity schema mismatch: {header!r}")
        records = [dict(zip(header, map(clean, row))) for row in rows if any(row)]
    finally:
        workbook.close()
    if any(not all(record.values()) for record in records):
        raise ValueError("strict identity contains an empty value")
    if len({record["RESULTNO"] for record in records}) != len(records):
        raise ValueError("strict identity RESULTNO values are not unique")
    return records


def load_pdf_cache(path: Path, required_refnos: set[str]) -> dict[str, dict[str, Any]]:
    cache: dict[str, dict[str, Any]] = {}
    found: set[str] = set()
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            row = json.loads(line)
            refno = clean(row.get("REFNO"))
            if refno in found:
                raise ValueError(f"duplicate page-cache REFNO: {refno}")
            found.add(refno)
            source = row.get("source") or {}
            flags = set(map(clean, source.get("source_quality_flags") or []))
            if clean(source.get("pdf_path")) and "no_local_pdf" not in flags:
                cache[refno] = row
    if found != required_refnos:
        raise ValueError(
            f"page-cache identity mismatch: missing={len(required_refnos-found)}, "
            f"extra={len(found-required_refnos)}"
        )
    return cache


def body_text(cache_row: Mapping[str, Any]) -> str:
    text = "\n".join(clean(page.get("text")) for page in cache_row.get("pages") or [])
    headings = list(REFERENCE_HEADING_RE.finditer(text))
    return text[: headings[-1].start()] if headings else text


def method_patterns(artifact: Mapping[str, Any]) -> list[tuple[str, re.Pattern[str]]]:
    output: list[tuple[str, re.Pattern[str]]] = []
    for concept in artifact["method_taxonomy"]:
        for pattern in concept["runtime_original_text_patterns"]:
            output.append((concept["concept_id"], re.compile(pattern, re.I | re.S)))
    return output


def explicit_shared_episode(
    cache_row: Mapping[str, Any], patterns: Sequence[tuple[str, re.Pattern[str]]]
) -> dict[str, Any]:
    """Find a method execution explicitly quantified over a shared sample set."""

    text = body_text(cache_row)
    best: dict[str, Any] | None = None
    for concept_id, pattern in patterns:
        for match in pattern.finditer(text):
            start = max(0, match.start() - 650)
            end = min(len(text), match.end() + 650)
            window = text[start:end]
            shared = next((item.search(window) for item in SHARED_SCOPE_PATTERNS if item.search(window)), None)
            if shared is None or EXECUTION_RE.search(window) is None:
                continue
            candidate = {
                "found": True,
                "method_concept": concept_id,
                "method_text_sha256": hashlib.sha256(
                    match.group(0).encode("utf-8")
                ).hexdigest(),
                "scope_text_sha256": hashlib.sha256(
                    shared.group(0).encode("utf-8")
                ).hexdigest(),
                "window_sha256": hashlib.sha256(window.encode("utf-8")).hexdigest(),
            }
            if best is None:
                best = candidate
    return best or {"found": False}


def proxy_summary(decision: Mapping[str, Any]) -> dict[str, Any]:
    proxy = decision["rockmag_proxy_r3"]
    conditions = decision.get("conditions") or {}
    counts = decision.get("evidence_counts") or {}
    rejected = decision.get("rejected_evidence_counts") or {}
    selected = decision.get("selected_evidence") or {}
    selected_anchors = sorted(
        {
            clean(hit.get("target_anchor"))
            for hit in selected.values()
            if isinstance(hit, Mapping) and clean(hit.get("target_anchor"))
        }
    )
    return {
        "score": int(proxy["binary_value"]),
        "outcome": clean(proxy["outcome"]),
        "reason_code": clean(decision.get("reason_code")),
        "conditions": {
            key: conditions.get(key)
            for key in (
                "accepted_method",
                "carrier_identification",
                "diagnostic_result",
                "alteration_result",
                "connected_positive_path",
                "result_component_identity_ambiguous",
                "source_quality_sufficient",
            )
        },
        "evidence_counts": {
            key: int(counts.get(key) or 0)
            for key in (
                "retrieved",
                "eligible",
                "accepted_methods",
                "carrier_identifications",
                "diagnostic_results",
                "alteration_results",
                "microscopy_methods",
                "weak_only",
            )
        },
        "rejected_counts": {
            clean(key): int(value) for key, value in rejected.items()
        },
        "selected_target_anchors": selected_anchors,
    }


def target_is_included(anchor: Mapping[str, Any]) -> bool:
    return bool(
        anchor.get("found")
        and clean(anchor.get("strength")) in {"strong", "moderate"}
        and clean(anchor.get("status")) != "generic"
    )


def build_candidates() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Phase 1: build and freeze candidates without labels or predictions."""

    input_hashes_before = {
        "identity": sha256_file(IDENTITY),
        "page_cache": sha256_file(PAGE_CACHE),
        "artifact": sha256_file(ARTIFACT),
    }
    identities = load_identity(IDENTITY)
    required_refnos = {row["REFNO"] for row in identities}
    cache = load_pdf_cache(PAGE_CACHE, required_refnos)
    artifact = json.loads(ARTIFACT.read_text(encoding="utf-8"))
    reasoner = R3RockmagDistilledReasoner(artifact)
    patterns = method_patterns(artifact)

    rows_by_refno: dict[str, list[dict[str, str]]] = defaultdict(list)
    for identity in identities:
        if identity["REFNO"] in cache:
            rows_by_refno[identity["REFNO"]].append(identity)

    candidates: list[dict[str, Any]] = []
    for paper_index, (refno, paper_rows) in enumerate(rows_by_refno.items(), start=1):
        if paper_index % 100 == 0:
            print(
                f"candidate phase: {paper_index}/{len(rows_by_refno)} PDF papers",
                flush=True,
            )
        target_names: dict[str, str] = {}
        target_row_counts: Counter[str] = Counter()
        for row in paper_rows:
            canonical = canonical_target_identity(row["ROCKNAME"])
            target_names.setdefault(canonical, row["ROCKNAME"])
            target_row_counts[canonical] += 1
        target_count = len(target_names)
        multi_target = target_count > 1
        shared_episode = explicit_shared_episode(cache[refno], patterns)
        retriever = PaperTextRetriever(cache[refno])
        target_facts: dict[str, dict[str, Any]] = {}

        for canonical, rockname in target_names.items():
            competitors = [
                name for sibling, name in target_names.items() if sibling != canonical
            ]
            strict_retrieval = retriever.retrieve(
                rockname,
                multi_target,
                "R3",
                competing_targets=competitors,
                allow_structured_scope=True,
                allow_discourse_scope=True,
                limit=30,
                per_category_limit=12,
            )
            strict = reasoner.evaluate(
                strict_retrieval,
                source_flags=(cache[refno].get("source") or {}).get(
                    "source_quality_flags", []
                ),
                multi_target=multi_target,
                identity_ambiguous=target_row_counts[canonical] > 1,
            )
            if target_row_counts[canonical] > 1:
                canonical_target = reasoner.evaluate(
                    strict_retrieval,
                    source_flags=(cache[refno].get("source") or {}).get(
                        "source_quality_flags", []
                    ),
                    multi_target=multi_target,
                    identity_ambiguous=False,
                )
            else:
                canonical_target = strict
            if multi_target:
                paper_retrieval = retriever.retrieve(
                    rockname,
                    False,
                    "R3",
                    competing_targets=(),
                    allow_structured_scope=True,
                    allow_discourse_scope=True,
                    limit=30,
                    per_category_limit=12,
                )
            else:
                paper_retrieval = strict_retrieval
            paper = reasoner.evaluate(
                paper_retrieval,
                source_flags=(cache[refno].get("source") or {}).get(
                    "source_quality_flags", []
                ),
                multi_target=False,
                identity_ambiguous=False,
            )
            anchor = strict_retrieval.get("target_anchor") or {}
            target_facts[canonical] = {
                "strict": proxy_summary(strict),
                "canonical_target": proxy_summary(canonical_target),
                "paper": proxy_summary(paper),
                "target_anchor": {
                    "found": bool(anchor.get("found")),
                    "status": clean(anchor.get("status")),
                    "strength": clean(anchor.get("strength")),
                    "specificity": anchor.get("specificity"),
                },
                "target_included": target_is_included(anchor),
            }

        any_strict_positive = any(
            facts["strict"]["score"] == 1 for facts in target_facts.values()
        )
        any_paper_positive = any(
            facts["paper"]["score"] == 1 for facts in target_facts.values()
        )
        all_targets_included = all(
            facts["target_included"] for facts in target_facts.values()
        )
        shared_gate = bool(shared_episode["found"] and all_targets_included)

        for row in paper_rows:
            canonical = canonical_target_identity(row["ROCKNAME"])
            facts = target_facts[canonical]
            direct = facts["strict"]["score"]
            canonical_target_score = facts["canonical_target"]["score"]
            paper_score = facts["paper"]["score"]
            single_target_score = paper_score if not multi_target else direct
            explicit_pair_score = (
                paper_score if multi_target and shared_gate else single_target_score
            )
            explicit_episode_score = (
                1 if multi_target and shared_gate else single_target_score
            )
            strict_trigger_shared_score = (
                1
                if multi_target and shared_gate and any_strict_positive
                else single_target_score
            )
            unsafe_paper_score = (
                1 if multi_target and any_paper_positive else single_target_score
            )
            ordered_pair_score = canonical_target_score
            ordered_episode_score = canonical_target_score
            if not multi_target:
                ordered_pair_score = paper_score
                ordered_episode_score = paper_score
            elif shared_gate:
                ordered_pair_score = paper_score
                ordered_episode_score = 1
            candidates.append(
                {
                    "REFNO": refno,
                    "RESULTNO": row["RESULTNO"],
                    "ROCKNAME": row["ROCKNAME"],
                    "canonical_target": canonical,
                    "paper_target_count": target_count,
                    "same_target_result_count": target_row_counts[canonical],
                    "direct_strict": direct,
                    "canonical_target_propagation": canonical_target_score,
                    "single_target_propagation": single_target_score,
                    "explicit_shared_pair_propagation": explicit_pair_score,
                    "explicit_shared_episode_propagation": explicit_episode_score,
                    "explicit_shared_strict_trigger_propagation": strict_trigger_shared_score,
                    "criterion_ordered_pair_propagation": ordered_pair_score,
                    "criterion_ordered_episode_propagation": ordered_episode_score,
                    "unsafe_any_paper_propagation": unsafe_paper_score,
                    "paper_scope": {
                        "explicit_shared_episode": shared_episode,
                        "shared_gate_all_targets_included": shared_gate,
                        "all_targets_included": all_targets_included,
                        "any_strict_positive": any_strict_positive,
                        "any_paper_positive": any_paper_positive,
                    },
                    "target_scope": facts["target_anchor"],
                    "strict_facts": facts["strict"],
                    "canonical_target_facts": facts["canonical_target"],
                    "paper_facts": facts["paper"],
                    "runtime_boundary": {
                        "identity_columns": list(IDENTITY_COLUMNS),
                        "original_paper_text": True,
                        "frozen_aggregate_expert_knowledge": True,
                        "labels_read": False,
                        "pmagresult_read": False,
                        "rockmag_row_value_read": False,
                        "prior_predictions_read": False,
                    },
                }
            )

    input_hashes_after = {
        "identity": sha256_file(IDENTITY),
        "page_cache": sha256_file(PAGE_CACHE),
        "artifact": sha256_file(ARTIFACT),
    }
    if input_hashes_before != input_hashes_after:
        raise RuntimeError("candidate runtime input changed during generation")
    atomic_jsonl(CANDIDATES, candidates)
    receipt = {
        "phase": "candidates_frozen_before_labels_or_predictions_opened",
        "candidate_sha256": sha256_file(CANDIDATES),
        "candidate_rows": len(candidates),
        "pdf_eligible_papers": len(rows_by_refno),
        "input_hashes": input_hashes_before,
        "runtime_sources": [
            "strict REFNO/RESULTNO/ROCKNAME identity manifest",
            "original-paper PDF page cache",
            "frozen aggregate ROCKMAG expert taxonomy",
        ],
        "forbidden_sources_not_read": [
            "PMAGRESULT",
            "row-level ROCKMAG values",
            "R3 labels",
            "prior predictions",
            "old evaluations",
        ],
    }
    atomic_json(RECEIPT, receipt)
    return candidates, receipt


def metrics(truth: Sequence[int], predicted: Sequence[int]) -> dict[str, Any]:
    pairs = Counter(zip(truth, predicted))
    tn, fp, fn, tp = pairs[(0, 0)], pairs[(0, 1)], pairs[(1, 0)], pairs[(1, 1)]
    total = tn + fp + fn + tp
    tnr = tn / (tn + fp) if tn + fp else 0.0
    tpr = tp / (tp + fn) if tp + fn else 0.0
    return {
        "rows": total,
        "accuracy": (tn + tp) / total if total else 0.0,
        "balanced_accuracy": (tnr + tpr) / 2,
        "confusion_matrix": {"TN": tn, "FP": fp, "FN": fn, "TP": tp},
    }


def aggregate_counter(rows: Sequence[Mapping[str, Any]], key) -> dict[str, int]:
    return dict(sorted(Counter(clean(key(row)) for row in rows).items()))


def oracle_ceiling(rows: Sequence[Mapping[str, Any]], group_key) -> dict[str, Any]:
    groups: dict[Any, Counter[int]] = defaultdict(Counter)
    for row in rows:
        groups[group_key(row)][int(row["label"])] += 1
    errors = sum(sum(counts.values()) - max(counts.values()) for counts in groups.values())
    return {
        "groups": len(groups),
        "irreducible_rows_under_grouping": errors,
        "accuracy_ceiling": 1 - errors / len(rows),
    }


def load_frozen_predictions(path: Path) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                row = json.loads(line)
                output[clean(row["RESULTNO"])] = row
    return output


def evaluate_candidates(
    candidates: list[dict[str, Any]], receipt: Mapping[str, Any]
) -> dict[str, Any]:
    """Phase 2: labels and frozen predictions are now allowed for analysis."""

    if sha256_file(CANDIDATES) != receipt["candidate_sha256"]:
        raise RuntimeError("candidate file changed before labels were opened")
    candidate_by_result = {row["RESULTNO"]: row for row in candidates}

    import pandas as pd

    labels = pd.read_excel(
        LABELS,
        usecols=["REFNO", "RESULTNO", "ROCKNAME", "R3"],
        dtype={"REFNO": str, "RESULTNO": str, "ROCKNAME": str},
    )
    labels["REFNO"] = labels["REFNO"].map(clean)
    labels["RESULTNO"] = labels["RESULTNO"].map(clean)
    labels["ROCKNAME"] = labels["ROCKNAME"].map(clean)
    labels = labels.loc[
        ~labels["REFNO"].map(is_holdout_refno)
        & labels["RESULTNO"].isin(candidate_by_result)
    ].copy()
    labels["R3"] = pd.to_numeric(labels["R3"], errors="raise").astype(int)
    latest = load_frozen_predictions(LATEST)
    previous = load_frozen_predictions(PREVIOUS)

    rows: list[dict[str, Any]] = []
    for record in labels.to_dict(orient="records"):
        resultno = record["RESULTNO"]
        candidate = candidate_by_result[resultno]
        frozen = latest[resultno]
        details = frozen.get("reasoning_details") or {}
        rows.append(
            {
                **candidate,
                "label": int(record["R3"]),
                "latest": int(frozen["R3"]),
                "previous": int(previous[resultno]["R3"]),
                "latest_reason": clean(details.get("reason_code")),
                "latest_counts": details.get("evidence_counts") or {},
                "latest_rejected": details.get("rejected_evidence_counts") or {},
                "latest_conditions": details.get("conditions") or {},
                "latest_selected": details.get("selected_evidence") or {},
            }
        )
    if len(rows) != len(candidates):
        raise RuntimeError(f"label/candidate mismatch: {len(rows)} != {len(candidates)}")

    strategy_names = (
        "direct_strict",
        "canonical_target_propagation",
        "single_target_propagation",
        "explicit_shared_pair_propagation",
        "explicit_shared_episode_propagation",
        "explicit_shared_strict_trigger_propagation",
        "criterion_ordered_pair_propagation",
        "criterion_ordered_episode_propagation",
        "unsafe_any_paper_propagation",
    )
    strategy_metrics = {
        name: metrics(
            [row["label"] for row in rows], [int(row[name]) for row in rows]
        )
        for name in strategy_names
    }
    strategy_metrics["frozen_latest_target_section"] = metrics(
        [row["label"] for row in rows], [row["latest"] for row in rows]
    )
    latest_metrics = strategy_metrics["frozen_latest_target_section"]
    strategy_deltas = {}
    for name in strategy_names:
        result = strategy_metrics[name]
        strategy_deltas[name] = {
            "accuracy_delta": result["accuracy"] - latest_metrics["accuracy"],
            "net_true_positive": (
                result["confusion_matrix"]["TP"]
                - latest_metrics["confusion_matrix"]["TP"]
            ),
            "net_false_positive": (
                result["confusion_matrix"]["FP"]
                - latest_metrics["confusion_matrix"]["FP"]
            ),
            "rows_changed": sum(
                int(row[name]) != row["latest"] for row in rows
            ),
        }

    paper_labels: dict[str, set[int]] = defaultdict(set)
    paper_targets: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        paper_labels[row["REFNO"]].add(row["label"])
        paper_targets[row["REFNO"]].add(row["canonical_target"])
    for row in rows:
        row["paper_label_structure"] = (
            "homogeneous" if len(paper_labels[row["REFNO"]]) == 1 else "mixed"
        )
        row["paper_target_structure"] = (
            "single_target" if len(paper_targets[row["REFNO"]]) == 1 else "multi_target"
        )

    def error_metrics(subset: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
        return metrics(
            [row["label"] for row in subset], [row["latest"] for row in subset]
        )

    strata: dict[str, Any] = {}
    for field in ("paper_label_structure", "paper_target_structure"):
        strata[field] = {
            value: error_metrics([row for row in rows if row[field] == value])
            for value in sorted({row[field] for row in rows})
        }
    strata["joint_paper_structure"] = {
        f"{label_structure}|{target_structure}": error_metrics(
            [
                row
                for row in rows
                if row["paper_label_structure"] == label_structure
                and row["paper_target_structure"] == target_structure
            ]
        )
        for label_structure in ("homogeneous", "mixed")
        for target_structure in ("single_target", "multi_target")
        if any(
            row["paper_label_structure"] == label_structure
            and row["paper_target_structure"] == target_structure
            for row in rows
        )
    }

    target_group_labels: dict[tuple[str, str], set[int]] = defaultdict(set)
    for row in rows:
        target_group_labels[(row["REFNO"], row["canonical_target"])].add(row["label"])
    duplicate_rows = [row for row in rows if row["same_target_result_count"] > 1]
    conflicting_duplicate_rows = [
        row
        for row in duplicate_rows
        if len(target_group_labels[(row["REFNO"], row["canonical_target"])]) > 1
    ]
    duplicate_analysis = {
        "duplicate_canonical_target_groups": len(
            {
                (row["REFNO"], row["canonical_target"])
                for row in duplicate_rows
            }
        ),
        "duplicate_rows": len(duplicate_rows),
        "mixed_label_duplicate_groups": len(
            {
                (row["REFNO"], row["canonical_target"])
                for row in conflicting_duplicate_rows
            }
        ),
        "mixed_label_duplicate_rows": len(conflicting_duplicate_rows),
        "canonical_target_propagation_on_duplicate_rows": metrics(
            [row["label"] for row in duplicate_rows],
            [row["canonical_target_propagation"] for row in duplicate_rows],
        ),
        "latest_on_duplicate_rows": error_metrics(duplicate_rows),
        "latest_on_mixed_label_duplicate_rows": error_metrics(
            conflicting_duplicate_rows
        ),
    }

    false_negatives = [row for row in rows if row["label"] == 1 and row["latest"] == 0]
    false_positives = [row for row in rows if row["label"] == 0 and row["latest"] == 1]

    def missing_facet(row: Mapping[str, Any]) -> str:
        counts = row["latest_counts"]
        rejected = row["latest_rejected"]
        if row["latest_reason"] == "SOURCE_QUALITY_RISK":
            return "source_quality_risk"
        if int(rejected.get("cross_target_or_unbound") or 0) > 0:
            return "cross_target_or_unbound_rejection"
        if int(counts.get("eligible") or 0) == 0:
            return "no_eligible_evidence"
        if int(counts.get("accepted_methods") or 0) == 0:
            return "accepted_method_missing"
        companions = sum(
            int(counts.get(key) or 0)
            for key in (
                "carrier_identifications",
                "diagnostic_results",
                "alteration_results",
            )
        )
        if companions == 0:
            return "carrier_or_diagnostic_missing"
        return "connection_or_scope_gate"

    def selected_scope(row: Mapping[str, Any]) -> str:
        selected = row["latest_selected"]
        anchors = sorted(
            {
                clean(hit.get("target_anchor"))
                for hit in selected.values()
                if isinstance(hit, Mapping) and clean(hit.get("target_anchor"))
            }
        )
        return "+".join(anchors) or "no_selected_scope"

    forensic = {
        "false_negatives": {
            "rows": len(false_negatives),
            "by_paper_target_structure": aggregate_counter(
                false_negatives,
                lambda row: f"{row['paper_label_structure']}|{row['paper_target_structure']}",
            ),
            "by_reason_code": aggregate_counter(
                false_negatives, lambda row: row["latest_reason"]
            ),
            "by_primary_missing_or_rejected_facet": aggregate_counter(
                false_negatives, missing_facet
            ),
        },
        "false_positives": {
            "rows": len(false_positives),
            "by_paper_target_structure": aggregate_counter(
                false_positives,
                lambda row: f"{row['paper_label_structure']}|{row['paper_target_structure']}",
            ),
            "by_reason_code": aggregate_counter(
                false_positives, lambda row: row["latest_reason"]
            ),
            "by_selected_evidence_scope": aggregate_counter(
                false_positives, selected_scope
            ),
        },
    }

    changed = [row for row in rows if row["previous"] != row["latest"]]
    target_section_delta = {
        "changed_rows": len(changed),
        "zero_to_one": sum(row["previous"] == 0 and row["latest"] == 1 for row in changed),
        "true_positive_gains": sum(
            row["previous"] == 0 and row["latest"] == 1 and row["label"] == 1
            for row in changed
        ),
        "false_positive_costs": sum(
            row["previous"] == 0 and row["latest"] == 1 and row["label"] == 0
            for row in changed
        ),
    }

    shared_papers = {
        row["REFNO"]
        for row in rows
        if row["paper_scope"]["shared_gate_all_targets_included"]
    }
    ceilings = {
        "canonical_target_identity_only": oracle_ceiling(
            rows, lambda row: (row["REFNO"], row["canonical_target"])
        ),
        "paper_uniform_prediction": oracle_ceiling(rows, lambda row: row["REFNO"]),
        "single_target_and_explicit_shared_forced_uniform": oracle_ceiling(
            rows,
            lambda row: (
                ("paper", row["REFNO"])
                if row["paper_target_count"] == 1 or row["REFNO"] in shared_papers
                else ("target", row["REFNO"], row["canonical_target"])
            ),
        ),
    }

    direct_mismatches = sum(row["direct_strict"] != row["latest"] for row in rows)
    report = {
        "schema_version": "r3_pdf_only_error_forensics_v1",
        "scope": {
            "partition": "fixed training partition, development-only",
            "pdf_only": True,
            "rows": len(rows),
            "papers": len({row["REFNO"] for row in rows}),
            "no_pdf_rows_skipped": 2473 - len(rows),
        },
        "integrity": {
            "candidate_receipt": str(RECEIPT),
            "candidate_sha256": receipt["candidate_sha256"],
            "labels_opened_only_after_candidate_freeze": True,
            "candidate_direct_recompute_vs_frozen_latest_mismatches": direct_mismatches,
            "group_leakage": {
                "status": "PASS",
                "fitting_performed": False,
                "unit_of_propagation": "REFNO paper group",
                "labels_used_as_runtime_features": False,
                "note": "Deterministic rules require no CV; all paper operations are REFNO-grouped.",
            },
            "input_hashes": {
                **receipt["input_hashes"],
                "labels_postfreeze": sha256_file(LABELS),
                "latest_prediction_postfreeze": sha256_file(LATEST),
                "previous_prediction_postfreeze": sha256_file(PREVIOUS),
            },
        },
        "full_coverage_metrics": strategy_metrics,
        "strategy_deltas_vs_frozen_latest": strategy_deltas,
        "latest_error_strata": strata,
        "same_name_duplicate_resultno": duplicate_analysis,
        "error_forensics": forensic,
        "target_section_scope_delta": target_section_delta,
        "shared_scope_detection": {
            "eligible_papers": len(shared_papers),
            "eligible_rows": sum(row["REFNO"] in shared_papers for row in rows),
        },
        "oracle_ceilings": ceilings,
        "interpretation": {
            "target_accuracy": 0.9,
            "target_reached_by_constrained_full_coverage_strategy": any(
                strategy_metrics[name]["accuracy"] >= 0.9
                for name in (
                    "single_target_propagation",
                    "explicit_shared_pair_propagation",
                    "explicit_shared_episode_propagation",
                    "explicit_shared_strict_trigger_propagation",
                    "criterion_ordered_pair_propagation",
                    "criterion_ordered_episode_propagation",
                )
            ),
            "unsafe_any_paper_is_policy_violating_diagnostic_only": True,
            "independent_test_claim": False,
        },
    }
    return report


def markdown_report(report: Mapping[str, Any]) -> str:
    metrics_rows = []
    for name, result in report["full_coverage_metrics"].items():
        cm = result["confusion_matrix"]
        metrics_rows.append(
            f"| {name} | {result['accuracy']:.2%} | {cm['TN']} | {cm['FP']} | "
            f"{cm['FN']} | {cm['TP']} |"
        )
    fn = report["error_forensics"]["false_negatives"]
    fp = report["error_forensics"]["false_positives"]
    ceilings = report["oracle_ceilings"]
    return "\n".join(
        [
            "# R3 PDF-only Error Forensics",
            "",
            "Development-only analysis. Candidate decisions were frozen before labels or prior predictions were opened.",
            "",
            "## Full-coverage results",
            "",
            "| Strategy | Accuracy | TN | FP | FN | TP |",
            "|---|---:|---:|---:|---:|---:|",
            *metrics_rows,
            "",
            "## Main bottleneck",
            "",
            f"Latest false negatives: {fn['rows']}; latest false positives: {fp['rows']}.",
            "",
            f"FN structure: `{json.dumps(fn['by_paper_target_structure'], ensure_ascii=False)}`",
            "",
            f"FN primary gate: `{json.dumps(fn['by_primary_missing_or_rejected_facet'], ensure_ascii=False)}`",
            "",
            f"FP decision paths: `{json.dumps(fp['by_reason_code'], ensure_ascii=False)}`",
            "",
            "## Oracle ceilings",
            "",
            f"Canonical target-equivalence ceiling: {ceilings['canonical_target_identity_only']['accuracy_ceiling']:.2%}.",
            f"Paper-uniform ceiling: {ceilings['paper_uniform_prediction']['accuracy_ceiling']:.2%}.",
            f"Permitted single/shared grouping ceiling: {ceilings['single_target_and_explicit_shared_forced_uniform']['accuracy_ceiling']:.2%}.",
            "",
            "The ceilings use labels only after candidate freeze and are not deployable classifiers.",
            "",
        ]
    )


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidates, receipt = build_candidates()
    report = evaluate_candidates(candidates, receipt)
    atomic_json(REPORT, report)
    REPORT_MD.write_text(markdown_report(report), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
